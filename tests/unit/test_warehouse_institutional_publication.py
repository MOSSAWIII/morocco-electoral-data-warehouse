from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from fractions import Fraction
from pathlib import Path

import pytest
import duckdb
import openpyxl

from morocco_elections.warehouse.history import affiliation_at_date, geographies_comparable, validate_geo_history, validate_party_history
from morocco_elections.warehouse.institutional import reproduce_largest_remainder_allocation, validate_candidacies_and_seats
from morocco_elections.warehouse.publication import (
    PublicationContext,
    REQUIRED_GATES,
    _as_of,
    _boundary,
    _immutability,
    _result_status,
    evaluate_publication_gates,
    validate_coverage_matrix,
    validate_publication,
    write_evidence_bundle,
)
from morocco_elections.warehouse.contracts import TABLE_CONTRACTS
from morocco_elections.warehouse.schema import ddl
from morocco_elections.warehouse.validation import validate_legal_regimes, validate_rows


def test_covered_election_requires_applicable_sourced_legal_regime() -> None:
    election = {"election_id": "E", "election_date": "2021-09-08", "covered": True}
    regime = {"legal_regime_id": "L", "valid_from": "2021-01-01", "legal_article": "art. 1", "legal_basis": "law", "official_source_id": "S", "official_source_url": "https://official.example/law", "allocation_formula": "D_HONDT", "electoral_quotient_denominator": "VALID_VOTES", "threshold_rule": "NONE", "district_magnitude_rule": "CONTEST_SEATS", "remainder_rule": "LARGEST_REMAINDER", "tie_break_rule": "SOURCE_RULE", "list_type": "LOCAL"}
    link = {"election_id": "E", "legal_regime_id": "L", "valid_from": "2021-01-01", "source_id": "S"}
    assert validate_legal_regimes([election], [regime], [link]) == []
    assert "NO_APPLICABLE_LEGAL_REGIME" in {i.code for i in validate_legal_regimes([election], [], [])}
    contest = {"contest_id": "C", "election_id": "E", "list_type": "LOCAL"}
    contest_link = {"contest_id": "C", "election_id": "E", "legal_regime_id": "L", "valid_from": "2021-01-01", "source_id": "S"}
    assert validate_legal_regimes([election], [regime], [link], [contest], [contest_link]) == []
    assert "NO_APPLICABLE_CONTEST_LEGAL_REGIME" in {i.code for i in validate_legal_regimes([election], [regime], [link], [contest], [])}


def test_malformed_legal_link_fails_closed_without_crashing() -> None:
    election = {"election_id": "E", "election_date": "2021-01-02", "covered": True}
    regime = {"legal_regime_id": "L", "valid_from": "2021-01-01", "legal_article": "art. 1", "legal_basis": "law", "official_source_id": "S", "official_source_url": "https://official.example/law", "allocation_formula": "D_HONDT", "electoral_quotient_denominator": "VALID_VOTES", "threshold_rule": "NONE", "district_magnitude_rule": "CONTEST_SEATS", "remainder_rule": "LARGEST_REMAINDER", "tie_break_rule": "SOURCE_RULE", "list_type": "LOCAL"}
    codes = {issue.code for issue in validate_legal_regimes([election], [regime], [{}])}
    assert {"REQUIRED_FIELD_MISSING", "NO_APPLICABLE_LEGAL_REGIME"} <= codes


def test_party_affiliations_are_temporal_and_name_only_matching_fails() -> None:
    affiliations = [{"affiliation_id": "A", "person_id": "P", "party_id": "OLD", "party_version_id": "OLD-V1", "valid_from": "2010-01-01", "valid_to": "2015-12-31", "matching_method": "OFFICIAL_IDENTIFIER", "source_id": "S"}]
    assert affiliation_at_date(affiliations, "P", "2012-01-01")[0]["party_id"] == "OLD"
    assert affiliation_at_date(affiliations, "P", "2021-01-01") == []
    bad = [{**affiliations[0], "matching_method": "NAME_ONLY"}]
    assert "NAME_ONLY_IDENTITY_FORBIDDEN" in {i.code for i in validate_party_history([], [], bad)}


def test_party_history_enforces_versions_intervals_and_acyclic_lineages() -> None:
    versions = [
        {"party_version_id": "A1", "party_id": "A", "name": "A", "valid_from": "2010-01-01", "valid_to": "2015-12-31", "source_id": "S"},
        {"party_version_id": "A2", "party_id": "A", "name": "A2", "valid_from": "2016-01-01", "source_id": "S"},
        {"party_version_id": "B1", "party_id": "B", "name": "B", "valid_from": "2010-01-01", "source_id": "S"},
    ]
    affiliation = {"affiliation_id": "AF", "person_id": "P", "party_id": "A", "party_version_id": "A1", "valid_from": "2011-01-01", "valid_to": "2014-01-01", "matching_method": "OFFICIAL_IDENTIFIER", "source_id": "S"}
    lineage = {"lineage_id": "L1", "predecessor_party_id": "A", "successor_party_id": "B", "party_lineage_type": "MERGER", "effective_date": "2016-01-01", "source_id": "S"}
    assert validate_party_history(versions, [lineage], [affiliation]) == []

    unknown = {**affiliation, "party_version_id": "MISSING"}
    outside = {**affiliation, "valid_to": "2017-01-01"}
    overlapping = {**versions[1], "party_version_id": "A3", "valid_from": "2015-01-01"}
    reverse = {**lineage, "lineage_id": "L2", "predecessor_party_id": "B", "successor_party_id": "A"}
    assert "UNKNOWN_AFFILIATION_PARTY_VERSION" in {i.code for i in validate_party_history(versions, [], [unknown])}
    assert "AFFILIATION_OUTSIDE_PARTY_VERSION" in {i.code for i in validate_party_history(versions, [], [outside])}
    assert "OVERLAPPING_PARTY_VERSIONS" in {i.code for i in validate_party_history([*versions, overlapping], [], [])}
    assert "CYCLIC_PARTY_LINEAGE" in {i.code for i in validate_party_history(versions, [lineage, reverse], [])}


def test_geography_crosswalk_requires_method_source_confidence_and_direct_compatibility() -> None:
    versions = [
        {"geo_version_id": "G1", "geo_id": "G", "boundary_version": "2015", "valid_from": "2015-01-01", "valid_to": "2020-12-31", "source_id": "S"},
        {"geo_version_id": "G2", "geo_id": "G", "boundary_version": "2021", "valid_from": "2021-01-01", "source_id": "S"},
    ]
    link = {"geo_lineage_id": "X", "from_geo_version_id": "G1", "to_geo_version_id": "G2", "geo_lineage_type": "REDISTRICTED", "method": "AREAL_INTERPOLATION", "source_id": "S", "confidence": 0.8}
    assert validate_geo_history(versions, [link]) == []
    assert not geographies_comparable("G1", "G2", [link], direct=True)
    assert geographies_comparable("G1", "G2", [link], direct=False)
    overlapping = [{**versions[0], "valid_to": None}, versions[1]]
    assert "OVERLAPPING_GEO_VERSIONS" in {issue.code for issue in validate_geo_history(overlapping, [link])}


def test_direct_boundary_comparison_requires_matching_package_geometries(tmp_path: Path) -> None:
    geometry = {"type": "Polygon", "coordinates": [[[0, 0], [1, 0], [1, 1], [0, 0]]]}
    (tmp_path / "first.geojson").write_text(json.dumps(geometry), encoding="utf-8")
    (tmp_path / "second.geojson").write_text(json.dumps(geometry), encoding="utf-8")
    versions = [
        {"geo_version_id": "G1", "geo_id": "G", "boundary_version": "2015", "valid_from": "2015-01-01", "valid_to": "2020-12-31", "source_id": "S", "geometry_path": "first.geojson"},
        {"geo_version_id": "G2", "geo_id": "G", "boundary_version": "2021", "valid_from": "2021-01-01", "source_id": "S", "geometry_path": "second.geojson"},
    ]
    link = {"geo_lineage_id": "X", "from_geo_version_id": "G1", "to_geo_version_id": "G2", "geo_lineage_type": "SAME_BOUNDARY", "method": "GEOMETRY_IDENTITY", "source_id": "S", "confidence": 1.0}
    evidence_root = tmp_path / "source-evidence"
    evidence_root.mkdir()
    source_bytes = b"official geography order"
    (evidence_root / "order.pdf").write_bytes(source_bytes)
    official_source = {
        "source_id": "S", "title": "Official boundary order", "authority": "Official authority",
        "source_url": "https://official.example/order", "verification_status": "VERIFIED_AUTHORITY_AND_BYTES",
        "license_status": "UNKNOWN", "public_package_disposition": "METADATA_ONLY",
        "raw_path": "order.pdf", "bytes": len(source_bytes), "sha256": hashlib.sha256(source_bytes).hexdigest(),
    }
    database = tmp_path / "geography.duckdb"
    connection = duckdb.connect(str(database))
    try:
        connection.execute(ddl("dim_geo_version"))
        connection.execute(ddl("bridge_geo_lineage"))
        for version in versions:
            connection.execute(
                "INSERT INTO dim_geo_version (geo_version_id, geo_id, boundary_version, valid_from, valid_to, source_id, geometry_path) VALUES (?, ?, ?, ?, ?, ?, ?)",
                [version.get(field) for field in ("geo_version_id", "geo_id", "boundary_version", "valid_from", "valid_to", "source_id", "geometry_path")],
            )
        connection.execute(
            "INSERT INTO bridge_geo_lineage (geo_lineage_id, from_geo_version_id, to_geo_version_id, geo_lineage_type, method, source_id, confidence) VALUES (?, ?, ?, ?, ?, ?, ?)",
            [link[field] for field in TABLE_CONTRACTS["bridge_geo_lineage"]["required"]],
        )
    finally:
        connection.close()

    def context() -> PublicationContext:
        return PublicationContext(
            release_id="TEST-SNAPSHOT", as_of_date="2021-09-08", package_root=tmp_path,
            files=[
                {"relative_path": name, "sha256": hashlib.sha256((tmp_path / name).read_bytes()).hexdigest()}
                for name in ("geography.duckdb", "first.geojson", "second.geojson")
            ], coverage_matrix=[], datasets={"geo_versions": versions, "geo_lineages": [link], "official_sources": [official_source]},
            checks={"boundary_checks": [{"check_id": "BC", "from_geo_version_id": "G1", "to_geo_version_id": "G2", "comparison_status": "DIRECTLY_COMPARABLE"}]},
            package_database_path="geography.duckdb", evidence_root=evidence_root,
        )

    assert _boundary(context()).status == "PASS", _boundary(context()).as_dict()
    changed = {"type": "Polygon", "coordinates": [[[0, 0], [2, 0], [2, 2], [0, 0]]]}
    (tmp_path / "second.geojson").write_text(json.dumps(changed), encoding="utf-8")
    dishonest = _boundary(context())
    assert dishonest.status == "FAIL"
    assert "X" in dishonest.affected_records
    assert "BC" in dishonest.affected_records
    (tmp_path / "second.geojson").write_text(json.dumps(geometry), encoding="utf-8")
    undeclared = replace(context(), files=context().files[:2])
    assert _boundary(undeclared).status == "FAIL"
    (tmp_path / "second.geojson").write_text(json.dumps({**geometry, "crs": "Lambert"}), encoding="utf-8")
    assert _boundary(context()).status == "FAIL"
    (tmp_path / "second.geojson").write_text(json.dumps({"type": "Polygon", "coordinates": [[[0, 0], [1, 0], [1, 1], [0, 1]]]}), encoding="utf-8")
    assert _boundary(context()).status == "FAIL"
    (tmp_path / "second.geojson").write_text(json.dumps(geometry), encoding="utf-8")
    unbound = _dataset(context(), "geo_versions", [{**versions[0], "boundary_version": "OTHER"}, versions[1]])
    assert "G1" in _boundary(unbound).affected_records
    (evidence_root / "order.pdf").write_bytes(b"altered geography order")
    assert "S" in _boundary(context()).affected_records


def test_candidacies_are_separate_and_unresolved_seat_gap_is_documented() -> None:
    lists = [{"candidacy_list_id": "L", "contest_id": "C", "list_type": "LOCAL", "source_id": "S"}]
    candidates = [{"candidacy_id": "K", "candidacy_list_id": "L", "person_id": "P", "position": 1, "identity_match_method": "HUMAN_REVIEW", "source_id": "S"}]
    allocation = {"allocation_id": "A", "contest_id": "C", "candidacy_list_id": "L", "seat_category_id": "SC", "legal_regime_id": "R", "official_seats": 2, "recomputed_seats": 1, "validation_status": "FAIL", "source_id": "S"}
    codes = {i.code for i in validate_candidacies_and_seats(lists, candidates, [allocation], [{"contest_id": "C"}], [{"legal_regime_id": "R"}])}
    assert "UNEXPLAINED_SEAT_DIFFERENCE" in codes
    wrong_status = [{**allocation, "validation_status": "PASS", "difference_explanation": "official source differs"}]
    codes = {i.code for i in validate_candidacies_and_seats(lists, candidates, wrong_status, [{"contest_id": "C"}], [{"legal_regime_id": "R"}])}
    assert "ALLOCATION_STATUS_MISMATCH" in codes


def test_largest_remainder_allocation_is_exact_and_fails_closed_on_ties() -> None:
    lists = [
        {"candidacy_list_id": "A", "votes": 400, "candidate_count": 3},
        {"candidacy_list_id": "B", "votes": 300, "candidate_count": 3},
        {"candidacy_list_id": "C", "votes": 200, "candidate_count": 3},
    ]
    result = reproduce_largest_remainder_allocation("C", lists, seats_to_fill=5, quotient_base=1000)
    assert result["metric_status"] == "COMPUTED"
    assert result["quotient"] == "200"
    assert {row["candidacy_list_id"]: row["recomputed_seats"] for row in result["allocations"]} == {"A": 2, "B": 2, "C": 1}

    tied = [
        {"candidacy_list_id": "A", "votes": 300, "candidate_count": 2},
        {"candidacy_list_id": "B", "votes": 300, "candidate_count": 2},
    ]
    ambiguous = reproduce_largest_remainder_allocation("T", tied, seats_to_fill=1, quotient_base=1000)
    assert ambiguous["metric_status"] == "NOT_COMPUTED"
    resolved = reproduce_largest_remainder_allocation(
        "T", [{**tied[0], "tie_break_rank": 2}, {**tied[1], "tie_break_rank": 1}], seats_to_fill=1, quotient_base=1000
    )
    assert resolved["metric_status"] == "COMPUTED"
    assert resolved["allocations"][1]["recomputed_seats"] == 1


def test_largest_remainder_checks_unique_list_threshold_and_candidate_capacity() -> None:
    below = reproduce_largest_remainder_allocation(
        "U", [{"candidacy_list_id": "A", "votes": 19, "candidate_count": 2}], seats_to_fill=2,
        quotient_base=100, unique_list_minimum=Fraction(1, 5), registered_voters=100,
    )
    assert below["metric_status"] == "COMPUTED"
    assert below["allocations"][0]["recomputed_seats"] == 0
    capacity = reproduce_largest_remainder_allocation(
        "K", [{"candidacy_list_id": "A", "votes": 100, "candidate_count": 1}], seats_to_fill=2, quotient_base=100,
    )
    assert capacity["metric_status"] == "NOT_COMPUTED"


def _file(sha256: str = "0" * 64) -> dict:
    return {"release_id": "TEST-SNAPSHOT", "relative_path": "data.duckdb", "sha256": sha256, "source_id": "S", "acquired_at": "2026-09-14", "license_status": "REDISTRIBUTABLE", "claim_class": "OBSERVED_FACT", "fact_status": "OBSERVED", "privacy_review_required": True}


def _matrix(scope_id: str = "ACQUIRED") -> dict:
    return {"release_id": "TEST-SNAPSHOT", "scope_id": scope_id, "universe_ids_json": json.dumps([f"U-{scope_id}"], separators=(",", ":")), "acquired": 1, "expected": 1, "covered": 1, "missing": 0, "non_comparable": 0, "redistribution_forbidden": 0, "status": "COMPLETE"}


def _matrices() -> list[dict]:
    return [_matrix(dimension) for dimension in ("ACQUIRED", "OFFICIAL", "TERRITORIAL", "TEMPORAL", "DOCUMENTARY", "FIELD")]


def _publication_context(package_root: Path) -> PublicationContext:
    package_root.mkdir(parents=True, exist_ok=True)
    evidence_root = package_root.parent / f"{package_root.name}-source-evidence"
    evidence_root.mkdir(parents=True, exist_ok=True)
    regime = {"legal_regime_id": "L", "valid_from": "2021-01-01", "legal_article": "art. 1", "legal_basis": "law", "official_source_id": "S", "official_source_url": "https://official.example/law", "allocation_formula": "D_HONDT", "electoral_quotient_denominator": "VALID_VOTES", "threshold_rule": "NONE", "district_magnitude_rule": "CONTEST_SEATS", "remainder_rule": "LARGEST_REMAINDER", "tie_break_rule": "SOURCE_RULE", "list_type": "LOCAL"}
    legal_election_link = {"election_id": "E", "legal_regime_id": "L", "valid_from": "2021-01-01", "source_id": "S"}
    legal_contest_link = {"contest_id": "C", "election_id": "E", "legal_regime_id": "L", "valid_from": "2021-01-01", "source_id": "S"}
    revision = {
        "revision_id": "R", "result_id": "X", "contest_id": "C", "election_id": "E",
        "geo_id": "G", "geo_version_id": "GV", "result_status": "FINAL",
        "valid_from": "2021-09-10", "source_id": "S", "published_at": "2021-09-10",
        "known_at": "2021-09-10", "verification_method": "STRUCTURED_SOURCE_CLAIM",
        "verification_status": "VERIFIED",
    }
    package_database = package_root / "data.duckdb"
    connection = duckdb.connect(str(package_database))
    try:
        for table in ("fact_election_result", "fact_communal_election_result"):
            connection.execute(
                f'CREATE TABLE "{table}" (result_id VARCHAR, contest_id VARCHAR, election_id VARCHAR, geo_id VARCHAR, year INTEGER)'
            )
        connection.execute(
            "INSERT INTO fact_election_result VALUES ('X', 'C', 'E', 'G', 2021)"
        )
        connection.execute(
            "INSERT INTO fact_election_result VALUES ('Y', 'C', 'E', 'G', 2021)"
        )
        connection.execute(
            "CREATE TABLE fact_electoral_mobilization (result_id VARCHAR, contest_id VARCHAR, "
            "election_id VARCHAR, geo_id VARCHAR, year INTEGER, contest_seats BIGINT, "
            "registered_voters BIGINT, voters BIGINT, valid_votes BIGINT, invalid_votes BIGINT, "
            "blank_votes BIGINT, turnout_rate DOUBLE, source_id VARCHAR)"
        )
        connection.execute(
            "INSERT INTO fact_electoral_mobilization VALUES "
            "('M', 'C', 'E', 'G', 2021, NULL, 200, 120, NULL, NULL, NULL, 0.6, 'S')"
        )
        connection.execute("CREATE TABLE dim_election (election_id VARCHAR, election_date DATE)")
        connection.execute("INSERT INTO dim_election VALUES ('E', '2021-09-08')")
        connection.execute("CREATE TABLE dim_electoral_contest (contest_id VARCHAR, election_id VARCHAR, geo_id VARCHAR)")
        connection.execute("INSERT INTO dim_electoral_contest VALUES ('C', 'E', 'G')")
        connection.execute("CREATE TABLE dim_geo (geo_id VARCHAR, geo_type VARCHAR, geo_name VARCHAR, parent_geo_id VARCHAR)")
        connection.execute("INSERT INTO dim_geo VALUES ('G', NULL, NULL, NULL), ('G2', NULL, NULL, NULL)")
        connection.execute(ddl("bridge_geo_parent"))
        for table, row in (
            ("dim_legal_regime", regime),
            ("bridge_election_legal_regime", legal_election_link),
            ("bridge_contest_legal_regime", legal_contest_link),
        ):
            connection.execute(ddl(table))
            fields = TABLE_CONTRACTS[table]["required"]
            selected = ", ".join(f'"{field}"' for field in fields)
            placeholders = ", ".join("?" for _ in fields)
            connection.execute(
                f'INSERT INTO "{table}" ({selected}) VALUES ({placeholders})',
                [row[field] for field in fields],
            )
        for table in ("fact_geo_population", "bridge_geo_official_identifier"):
            connection.execute(ddl(table))
        connection.execute(ddl("dim_geo_version"))
        connection.execute("INSERT INTO dim_geo_version (geo_version_id, geo_id, boundary_version, valid_from, source_id) VALUES ('GV', 'G', '2021', '2021-01-01', 'S')")
        connection.execute(ddl("bridge_geo_lineage"))
        connection.execute(ddl("fact_result_reconciliation"))
        connection.execute(
            "INSERT INTO fact_result_reconciliation "
            "(reconciliation_id, contest_id, metric, tolerance, validation_status, source_id, "
            "official_value, recomputed_value, difference) "
            "VALUES ('RC', 'C', 'turnout_rate', 0, 'PASS', 'S', 0.6, 0.6, 0)"
        )
        connection.execute(
            "INSERT INTO fact_result_reconciliation "
            "(reconciliation_id, contest_id, metric, tolerance, validation_status, source_id, "
            "official_value, explanation, not_computable_reason) VALUES "
            "('RC-B', 'C', 'ballot_categories_vs_voters', 0, 'NOT_COMPUTABLE', 'S', 120, "
            "'Sourced ballot categories unavailable', 'BALLOT_COMPONENTS_MISSING'), "
            "('RC-P', 'C', 'party_votes_vs_valid_votes', 0, 'NOT_COMPUTABLE', 'S', NULL, "
            "'Official expected party identifiers unavailable', 'PARTY_UNIVERSE_UNVERIFIED'), "
            "('RC-S', 'C', 'allocated_seats_vs_contest_seats', 0, 'NOT_COMPUTABLE', 'S', NULL, "
            "'Official expected seat identifiers unavailable', 'SEAT_UNIVERSE_UNVERIFIED')"
        )
        connection.execute(ddl("fact_result_revision"))
        revision_fields = [
            row[0] for row in connection.execute("DESCRIBE fact_result_revision").fetchall()
            if row[0] in revision
        ]
        connection.execute(
            "INSERT INTO fact_result_revision (" + ", ".join(f'\"{field}\"' for field in revision_fields) +
            ") VALUES (" + ", ".join("?" for _ in revision_fields) + ")",
            [revision[field] for field in revision_fields],
        )
        connection.execute(ddl("fact_legal_decision"))
        connection.execute(
            "CREATE TABLE warehouse_metadata (release VARCHAR, schema_version BIGINT, source_release VARCHAR, as_of_date DATE)"
        )
        connection.execute("INSERT INTO warehouse_metadata VALUES ('TEST-SNAPSHOT', 16, 'V15', '2026-09-14')")
    finally:
        connection.close()
    published = package_database.read_bytes()
    database_sha = hashlib.sha256(published).hexdigest()
    manifest = json.dumps({"files": [{"relative_path": "data.duckdb", "sha256": database_sha}]}, sort_keys=True)
    (package_root / "package-manifest.json").write_text(manifest, encoding="utf-8")
    official_source = json.dumps({
        "result_history_claims": [{
            field: revision.get(field) for field in
            ("revision_id", "result_id", "contest_id", "result_status", "valid_from", "published_at")
        }]
    }, sort_keys=True).encode("utf-8")
    (evidence_root / "official-law.pdf").write_bytes(official_source)
    universe_source = json.dumps({
        "universes": {f"U-{dimension}": [f"ID-{dimension}"] for dimension in
        ("ACQUIRED", "OFFICIAL", "TERRITORIAL", "TEMPORAL", "DOCUMENTARY", "FIELD")}
    }, sort_keys=True).encode("utf-8")
    (evidence_root / "official-universe.json").write_bytes(universe_source)
    manifest_sha = hashlib.sha256(manifest.encode("utf-8")).hexdigest()
    clean_shas = []
    for build_id, report_name in (("B1", "clean-build-1.json"), ("B2", "clean-build-2.json")):
        clean_report = json.dumps({
            "build_id": build_id,
            "clean_environment": True,
            "environment": {
                "workspace_id": f"WORKSPACE-{build_id}",
                "source_materialization": "SOURCE_ARCHIVE_EXTRACTION",
                "initial_entries": [],
            },
            "source_revision": "test-revision",
            "dependency_lock_sha256": "a" * 64,
            "manifest_path": "package-manifest.json",
            "manifest_sha256": manifest_sha,
            "commands": [
                {"name": "tests", "command": "python -m pytest", "exit_code": 0},
                {"name": "validate_warehouse", "command": "python tools/validate_warehouse.py", "exit_code": 0},
            ],
        }, sort_keys=True)
        clean_shas.append(hashlib.sha256(clean_report.encode("utf-8")).hexdigest())
        (package_root / report_name).write_text(clean_report, encoding="utf-8")
    dimensions = ("ACQUIRED", "OFFICIAL", "TERRITORIAL", "TEMPORAL", "DOCUMENTARY", "FIELD")
    universes = [
        {"universe_id": f"U-{dimension}", "election_id": "E", "coverage_dimension": dimension, "universe_type": "OFFICIAL_RECORDS", "denominator": 1, "source_id": "S-U", "source_url": "https://official.example/universe", "acquired_at": "2026-09-14", "verification_status": "VERIFIED", "is_external": True, "member_extraction_method": "JSON_UNIVERSES_OBJECT"}
        for dimension in dimensions
    ]
    reconciliation = {"reconciliation_id": "RC", "contest_id": "C", "metric": "turnout_rate", "official_value": 0.6, "recomputed_value": 0.6, "difference": 0, "tolerance": 0, "validation_status": "PASS", "source_id": "S"}
    incomplete_reconciliations = [
        {"reconciliation_id": "RC-B", "contest_id": "C", "metric": "ballot_categories_vs_voters", "official_value": 120, "recomputed_value": None, "difference": None, "tolerance": 0, "validation_status": "NOT_COMPUTABLE", "source_id": "S", "explanation": "Sourced ballot categories unavailable", "not_computable_reason": "BALLOT_COMPONENTS_MISSING"},
        {"reconciliation_id": "RC-P", "contest_id": "C", "metric": "party_votes_vs_valid_votes", "official_value": None, "recomputed_value": None, "difference": None, "tolerance": 0, "validation_status": "NOT_COMPUTABLE", "source_id": "S", "explanation": "Official expected party identifiers unavailable", "not_computable_reason": "PARTY_UNIVERSE_UNVERIFIED"},
        {"reconciliation_id": "RC-S", "contest_id": "C", "metric": "allocated_seats_vs_contest_seats", "official_value": None, "recomputed_value": None, "difference": None, "tolerance": 0, "validation_status": "NOT_COMPUTABLE", "source_id": "S", "explanation": "Official expected seat identifiers unavailable", "not_computable_reason": "SEAT_UNIVERSE_UNVERIFIED"},
    ]
    context = PublicationContext(
        release_id="TEST-SNAPSHOT",
        as_of_date="2026-09-14",
        files=[_file(database_sha)],
        coverage_matrix=_matrices(),
        datasets={
            "elections": [{"election_id": "E", "election_date": "2021-09-08", "covered": True}],
            "contests": [{"contest_id": "C", "election_id": "E", "geo_id": "G", "list_type": "LOCAL"}],
            "geographies": [{"geo_id": "G", "parent_geo_id": None}, {"geo_id": "G2", "parent_geo_id": None}],
            "geo_parent_relations": [],
            "semantic_facts": [
                {"result_id": "X", "contest_id": "C", "election_id": "E", "geo_id": "G", "year": 2021},
                {"result_id": "Y", "contest_id": "C", "election_id": "E", "geo_id": "G", "year": 2021},
                {"result_id": "M", "contest_id": "C", "election_id": "E", "geo_id": "G", "year": 2021, "contest_seats": None, "registered_voters": 200, "voters": 120, "valid_votes": None, "invalid_votes": None, "blank_votes": None, "turnout_rate": 0.6, "source_id": "S"},
            ],
            "legal_regimes": [regime],
            "election_legal_regimes": [legal_election_link],
            "contest_legal_regimes": [legal_contest_link],
            "official_sources": [{
                "source_id": "S", "title": "Official law", "authority": "Official authority",
                "source_url": "https://official.example/law", "verification_status": "VERIFIED_AUTHORITY_AND_BYTES",
                "license_status": "UNKNOWN", "public_package_disposition": "METADATA_ONLY",
                "raw_path": "official-law.pdf", "bytes": len(official_source),
                "sha256": hashlib.sha256(official_source).hexdigest(),
                "acquired_at": "2021-09-10",
            }, {
                "source_id": "S-U", "title": "Official universe", "authority": "Official authority",
                "source_url": "https://official.example/universe", "verification_status": "VERIFIED_AUTHORITY_AND_BYTES",
                "license_status": "UNKNOWN", "public_package_disposition": "METADATA_ONLY",
                "raw_path": "official-universe.json", "bytes": len(universe_source),
                "sha256": hashlib.sha256(universe_source).hexdigest(),
            }],
            "result_revisions": [revision],
            "legal_decisions": [],
            "warehouse_metadata": [{"release": "TEST-SNAPSHOT", "schema_version": 16, "source_release": "V15", "as_of_date": "2026-09-14"}],
            "geo_versions": [{"geo_version_id": "GV", "geo_id": "G", "boundary_version": "2021", "valid_from": "2021-01-01", "source_id": "S"}],
            "coverage_universes": universes,
            "coverage_universe_members": [{"universe_id": f"U-{dimension}", "expected_id": f"ID-{dimension}", "source_id": "S-U"} for dimension in dimensions],
            "coverage_observations": [{"coverage_dimension": dimension, "observation_id": f"ID-{dimension}"} for dimension in dimensions],
            "reconciliations": [reconciliation, *incomplete_reconciliations],
            "candidacy_lists": [{"candidacy_list_id": "CL", "contest_id": "C", "list_type": "LOCAL", "source_id": "S"}],
            "candidates": [{"candidacy_id": "CA", "candidacy_list_id": "CL", "person_id": "PERSON", "position": 1, "identity_match_method": "OFFICIAL_IDENTIFIER", "source_id": "S"}],
            "seat_categories": [{"seat_category_id": "SC", "label": "General", "legal_regime_id": "L"}],
            "seat_allocations": [{"allocation_id": "A", "contest_id": "C", "candidacy_list_id": "CL", "seat_category_id": "SC", "legal_regime_id": "L", "official_seats": 1, "recomputed_seats": 1, "difference": 0, "validation_status": "PASS", "source_id": "S"}],
            "metric_validations": [{"metric_validation_id": "MV", "metric_name": "HHI", "formula": "SUM(p_i^2)", "denominator": "complete shares", "scope": "C", "metric_status": "COMPUTED", "coverage_status": "COMPLETE", "limitations": ["descriptive"], "failed_preconditions": []}],
            "analytic_output": [{"contest_id": "C", "party_id": "P", "vote_share": 1.0}],
            "source_observations": [{"observation_id": "SO1", "subject_id": "C", "field": "valid_votes", "normalized_value": 10, "source_id": "S"}],
            "party_versions": [
                {"party_version_id": "P0-V1", "party_id": "P0", "name": "Predecessor", "valid_from": "2000-01-01", "valid_to": "2020-12-31", "source_id": "S"},
                {"party_version_id": "P-V1", "party_id": "P", "name": "Party", "valid_from": "2021-01-01", "source_id": "S"},
            ],
            "party_lineages": [{"lineage_id": "PL", "predecessor_party_id": "P0", "successor_party_id": "P", "party_lineage_type": "SUCCESSION", "effective_date": "2021-01-01", "source_id": "S"}],
            "party_affiliations": [{"affiliation_id": "PA", "person_id": "PERSON", "party_id": "P", "party_version_id": "P-V1", "valid_from": "2021-01-01", "matching_method": "OFFICIAL_IDENTIFIER", "source_id": "S"}],
        },
        checks={
            "grain_checks": [{"check_id": "GR", "dataset": "analytic_output", "grain_keys": ["contest_id", "party_id"], "expected_row_count": 1}],
            "boundary_checks": [{"check_id": "GE", "from_geo_version_id": "GV", "to_geo_version_id": "GV", "comparison_status": "DIRECTLY_COMPARABLE"}],
            "party_lineage_reviews": [{"lineage_id": "PL", "reviewed_by": "reviewer", "reviewed_at": "2026-09-13", "source_id": "S"}],
            "source_conflict_checks": [],
            "claim_reviews": [{"relative_path": "data.duckdb", "file_sha256": database_sha, "claim_class": "OBSERVED_FACT", "uncertainty_applicable": False, "reviewed_by": "reviewer", "reviewed_at": "2026-09-13", "evidence_id": "E"}],
            "privacy_reviews": [{"relative_path": "data.duckdb", "file_sha256": database_sha, "decision": "PASS", "reviewed_by": "reviewer", "reviewed_at": "2026-09-13", "evidence_id": "E"}],
            "license_reviews": [{
                "relative_path": "data.duckdb", "file_sha256": database_sha, "decision": "REDISTRIBUTABLE",
                "legal_basis": "Official open-data terms permit redistribution",
                "evidence_url": "https://official.example/license",
                "reviewed_by": "reviewer", "reviewed_at": "2026-09-13", "evidence_id": "E", "proof_sha256": "1" * 64,
            }],
            "clean_builds": [
                {"build_id": "B1", "report_path": "clean-build-1.json", "report_sha256": clean_shas[0]},
                {"build_id": "B2", "report_path": "clean-build-2.json", "report_sha256": clean_shas[1]},
            ],
        },
        package_root=package_root,
        package_database_path="data.duckdb",
        evidence_root=evidence_root,
    )
    return write_evidence_bundle(context)


def _dataset(context: PublicationContext, name: str, rows: list[dict]) -> PublicationContext:
    return replace(context, datasets={**context.datasets, name: rows})


def _materialize_classified_link(context: PublicationContext, link: dict) -> PublicationContext:
    """Reauthor a test-owned package and all reports after a real DuckDB change."""
    root = context.package_root
    assert root is not None
    database = root / "data.duckdb"
    connection = duckdb.connect(str(database))
    try:
        for row in context.datasets["geographies"]:
            if row.get("geo_type") is not None:
                connection.execute(
                    "INSERT INTO dim_geo (geo_id, geo_type, geo_name, parent_geo_id) VALUES (?, ?, ?, ?)",
                    [row["geo_id"], row.get("geo_type"), row.get("geo_name"), row.get("parent_geo_id")],
                )
        connection.execute(
            "UPDATE bridge_contest_legal_regime SET classification_basis = ?, "
            "classification_source_id = ?, classification_value = ?, classification_rule = ? WHERE contest_id = ?",
            [link["classification_basis"], link["classification_source_id"],
             link["classification_value"], link["classification_rule"], link["contest_id"]],
        )
        for table, dataset in (
            ("fact_geo_population", "geo_populations"),
            ("bridge_geo_official_identifier", "geo_official_identifier_crosswalks"),
        ):
            fields = [*TABLE_CONTRACTS[table]["required"]]
            if table == "bridge_geo_official_identifier":
                fields.append("geo_type")
            columns = ", ".join(f'"{field}"' for field in fields)
            placeholders = ", ".join("?" for _ in fields)
            for row in context.datasets[dataset]:
                connection.execute(
                    f'INSERT INTO "{table}" ({columns}) VALUES ({placeholders})',
                    [row.get(field) for field in fields],
                )
        connection.execute("CHECKPOINT")
    finally:
        connection.close()
    return _reauthor_test_package(context)


def _reauthor_test_package(context: PublicationContext) -> PublicationContext:
    root = context.package_root
    assert root is not None
    database = root / "data.duckdb"
    database_sha = hashlib.sha256(database.read_bytes()).hexdigest()
    manifest = json.dumps({"files": [{"relative_path": "data.duckdb", "sha256": database_sha}]}, sort_keys=True)
    (root / "package-manifest.json").write_text(manifest, encoding="utf-8")
    manifest_sha = hashlib.sha256(manifest.encode("utf-8")).hexdigest()
    clean_builds = []
    for row in context.checks["clean_builds"]:
        report_path = root / row["report_path"]
        report = json.loads(report_path.read_text(encoding="utf-8"))
        report["manifest_sha256"] = manifest_sha
        payload = json.dumps(report, sort_keys=True)
        report_path.write_text(payload, encoding="utf-8")
        clean_builds.append({**row, "report_sha256": hashlib.sha256(payload.encode("utf-8")).hexdigest()})
    bundle = root / "validation-evidence.json"
    bundle.unlink()  # This helper only changes temporary files created by this test.
    updated = replace(
        context,
        files=[{**context.files[0], "sha256": database_sha}],
        checks={**context.checks, "clean_builds": clean_builds},
        evidence_bundle_path=None, evidence_bundle_sha256=None,
    )
    return write_evidence_bundle(updated)


def _check(context: PublicationContext, name: str, rows: list[dict]) -> PublicationContext:
    return replace(context, checks={**context.checks, name: rows})


def test_release_matrix_reconciles_all_disposition_counts() -> None:
    assert validate_coverage_matrix([_matrix()]) == []
    assert "COVERAGE_MATRIX_NOT_RECONCILED" in {i.code for i in validate_coverage_matrix([{**_matrix(), "expected": 2}])}
    assert "COVERAGE_UNIVERSE_REFERENCE_REQUIRED" in {
        i.code for i in validate_coverage_matrix([{**_matrix(), "universe_ids_json": "[]"}])
    }


def test_coverage_gate_rejects_false_universe_identity_with_identical_counts(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    original = {row.gate_id: row for row in evaluate_publication_gates(context)}["COVERAGE_DISCLOSED"]
    assert original.status == "PASS"
    claimed = [
        {**context.coverage_matrix[0], "universe_ids_json": '["U-UNRELATED"]'},
        *context.coverage_matrix[1:],
    ]
    result = {
        row.gate_id: row
        for row in evaluate_publication_gates(replace(context, coverage_matrix=claimed))
    }["COVERAGE_DISCLOSED"]
    assert result.status == "FAIL"
    assert "universe_ids_json" in result.justification
    assert result.evidence_id != original.evidence_id


def test_metric_gate_recomputes_status_and_difference(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    dishonest = [{**context.datasets["reconciliations"][0], "recomputed_value": 11, "difference": 0, "validation_status": "PASS"}, *context.datasets["reconciliations"][1:]]
    results = {row.gate_id: row for row in evaluate_publication_gates(_dataset(context, "reconciliations", dishonest))}
    assert results["METRIC_RECONCILED"].status == "FAIL"
    assert "recomputed" in results["METRIC_RECONCILED"].justification


def test_metric_gate_rejects_omitted_materialized_reconciliation_with_valid_bundle(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    omitted = _reauthor_test_package(_dataset(context, "reconciliations", []))
    gates = {row.gate_id: row for row in evaluate_publication_gates(omitted)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["METRIC_RECONCILED"].status == "FAIL"
    assert "materialized fact_result_reconciliation row" in gates["METRIC_RECONCILED"].justification


def test_metric_gate_requires_each_applicable_contest_check(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    missing = _dataset(context, "reconciliations", [
        row for row in context.datasets["reconciliations"] if row["metric"] != "party_votes_vs_valid_votes"
    ])
    missing = _reauthor_test_package(missing)
    gate = {row.gate_id: row for row in evaluate_publication_gates(missing)}["METRIC_RECONCILED"]
    assert gate.status == "FAIL"
    assert "C:party_votes_vs_valid_votes" in gate.affected_records
    assert "every contest requires exactly one applicable reconciliation status" in gate.justification


def test_metric_gate_rederives_ballot_not_computable_reason_from_complete_package_facts(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    assert context.package_root is not None
    connection = duckdb.connect(str(context.package_root / "data.duckdb"))
    try:
        connection.execute(
            "UPDATE fact_electoral_mobilization SET valid_votes = 100, invalid_votes = 15, "
            "blank_votes = 5 WHERE contest_id = 'C'"
        )
        connection.execute("CHECKPOINT")
    finally:
        connection.close()
    coordinated_facts = [
        {**row, "valid_votes": 100, "invalid_votes": 15, "blank_votes": 5}
        if row.get("result_id") == "M" else row
        for row in context.datasets["semantic_facts"]
    ]
    coordinated = _reauthor_test_package(_dataset(context, "semantic_facts", coordinated_facts))
    gates = {row.gate_id: row for row in evaluate_publication_gates(coordinated)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["SEMANTIC_FACTS_VALIDATED"].status == "PASS"
    assert gates["METRIC_RECONCILED"].status == "FAIL"
    assert "source-derived reason BALLOT_TAXONOMY_UNVERIFIED" in gates["METRIC_RECONCILED"].justification


def test_metric_gate_rejects_convenient_not_computable_for_complete_turnout(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    assert context.package_root is not None
    connection = duckdb.connect(str(context.package_root / "data.duckdb"))
    try:
        connection.execute(
            "UPDATE fact_result_reconciliation SET official_value = NULL, recomputed_value = NULL, "
            "difference = NULL, validation_status = 'NOT_COMPUTABLE', explanation = 'Missing', "
            "not_computable_reason = 'MOBILIZATION_FIELDS_MISSING' WHERE reconciliation_id = 'RC'"
        )
        connection.execute("CHECKPOINT")
    finally:
        connection.close()
    hidden = _dataset(context, "reconciliations", [{
        **context.datasets["reconciliations"][0], "official_value": None, "recomputed_value": None,
        "difference": None, "validation_status": "NOT_COMPUTABLE", "explanation": "Missing",
        "not_computable_reason": "MOBILIZATION_FIELDS_MISSING",
    }, *context.datasets["reconciliations"][1:]])
    hidden = _reauthor_test_package(hidden)
    gate = {row.gate_id: row for row in evaluate_publication_gates(hidden)}["METRIC_RECONCILED"]
    assert gate.status == "FAIL"
    assert "package fact derivation" in gate.justification
    assert "OFFICIAL_TURNOUT_MISSING" in gate.justification


def test_metric_gate_rejects_mutated_materialized_reconciliation_with_valid_manifest(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    assert context.package_root is not None
    connection = duckdb.connect(str(context.package_root / "data.duckdb"))
    try:
        connection.execute("UPDATE fact_result_reconciliation SET recomputed_value = 11 WHERE reconciliation_id = 'RC'")
        connection.execute("CHECKPOINT")
    finally:
        connection.close()
    mutated = _reauthor_test_package(context)
    gates = {row.gate_id: row for row in evaluate_publication_gates(mutated)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["METRIC_RECONCILED"].status == "FAIL"
    assert "materialized fact_result_reconciliation" in gates["METRIC_RECONCILED"].justification


def test_metric_gate_rejects_coordinated_invented_values_against_package_facts(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    assert context.package_root is not None
    connection = duckdb.connect(str(context.package_root / "data.duckdb"))
    try:
        connection.execute(
            "UPDATE fact_result_reconciliation SET official_value = 0.5, "
            "recomputed_value = 0.5 WHERE reconciliation_id = 'RC'"
        )
        connection.execute("CHECKPOINT")
    finally:
        connection.close()
    invented = _dataset(context, "reconciliations", [{
        **context.datasets["reconciliations"][0], "official_value": 0.5, "recomputed_value": 0.5,
    }, *context.datasets["reconciliations"][1:]])
    invented = _reauthor_test_package(invented)
    gates = {row.gate_id: row for row in evaluate_publication_gates(invented)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["METRIC_RECONCILED"].status == "FAIL"
    assert "materialized fact_result_reconciliation" not in gates["METRIC_RECONCILED"].justification
    assert "package fact derivation" in gates["METRIC_RECONCILED"].justification


def test_metric_gate_rejects_unsupported_pass_even_when_table_and_bundle_agree(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    assert context.package_root is not None
    connection = duckdb.connect(str(context.package_root / "data.duckdb"))
    try:
        connection.execute(
            "UPDATE fact_result_reconciliation SET metric = 'fabricated_metric' WHERE reconciliation_id = 'RC'"
        )
        connection.execute("CHECKPOINT")
    finally:
        connection.close()
    unsupported = _dataset(context, "reconciliations", [{
        **context.datasets["reconciliations"][0], "metric": "fabricated_metric",
    }, *context.datasets["reconciliations"][1:]])
    unsupported = _reauthor_test_package(unsupported)
    gates = {row.gate_id: row for row in evaluate_publication_gates(unsupported)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["METRIC_RECONCILED"].status == "FAIL"
    assert "no supported package-fact derivation" in gates["METRIC_RECONCILED"].justification


def test_metric_gate_requires_verified_official_source_bytes(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    unsourced = _dataset(
        context,
        "reconciliations",
        [{**context.datasets["reconciliations"][0], "source_id": "MISSING"}, *context.datasets["reconciliations"][1:]],
    )
    result = {row.gate_id: row for row in evaluate_publication_gates(unsourced)}["METRIC_RECONCILED"]
    assert result.status == "FAIL"
    assert "verified local bytes" in result.justification


def test_legal_gate_recomputes_the_official_source_checksum(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    original = {row.gate_id: row for row in evaluate_publication_gates(context)}["LEGAL_REGIME_PINNED"]
    assert context.evidence_root is not None
    (context.evidence_root / "official-law.pdf").write_bytes(b"mutated legal source\n")
    result = {row.gate_id: row for row in evaluate_publication_gates(context)}["LEGAL_REGIME_PINNED"]
    assert result.status == "FAIL"
    assert "SOURCE_CHECKSUM_MISMATCH" in result.justification
    assert result.evidence_id != original.evidence_id


def test_official_universe_gate_recomputes_source_checksum_and_url(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    original = {row.gate_id: row for row in evaluate_publication_gates(context)}["OFFICIAL_UNIVERSE_DECLARED"]
    assert original.status == "PASS"
    assert context.evidence_root is not None
    (context.evidence_root / "official-universe.json").write_bytes(b"mutated universe\n")
    result = {row.gate_id: row for row in evaluate_publication_gates(context)}["OFFICIAL_UNIVERSE_DECLARED"]
    assert result.status == "FAIL"
    assert "SOURCE_CHECKSUM_MISMATCH" in result.justification
    assert result.evidence_id != original.evidence_id

    fresh = _publication_context(tmp_path / "fresh")
    bad_url = _dataset(
        fresh,
        "coverage_universes",
        [{**fresh.datasets["coverage_universes"][0], "source_url": "https://unregistered.example"}],
    )
    result = {row.gate_id: row for row in evaluate_publication_gates(bad_url)}["OFFICIAL_UNIVERSE_DECLARED"]
    assert result.status == "FAIL"
    assert "URL differs" in result.justification


def test_official_universe_gate_rejects_members_not_present_in_source(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    original = {row.gate_id: row for row in evaluate_publication_gates(context)}["OFFICIAL_UNIVERSE_DECLARED"]
    assert original.status == "PASS"
    members = [row for row in context.datasets["coverage_universe_members"] if row["universe_id"] != "U-TERRITORIAL"]
    result = {
        row.gate_id: row
        for row in evaluate_publication_gates(_dataset(context, "coverage_universe_members", members))
    }["OFFICIAL_UNIVERSE_DECLARED"]
    assert result.status == "FAIL"
    assert "source-derived expected identifier is absent" in result.justification
    assert "ID-TERRITORIAL" in result.affected_records
    assert result.evidence_id != original.evidence_id

    fabricated = [
        {**row, "expected_id": "NOT-IN-SOURCE"} if row["universe_id"] == "U-TERRITORIAL" else row
        for row in context.datasets["coverage_universe_members"]
    ]
    result = {
        row.gate_id: row
        for row in evaluate_publication_gates(_dataset(context, "coverage_universe_members", fabricated))
    }["OFFICIAL_UNIVERSE_DECLARED"]
    assert result.status == "FAIL"
    assert "declared universe member is absent" in result.justification
    assert "NOT-IN-SOURCE" in result.affected_records


def test_result_status_gate_requires_sourced_revision_and_linked_legal_decision(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    unsourced = _dataset(
        context,
        "result_revisions",
        [{**context.datasets["result_revisions"][0], "source_id": "MISSING"}],
    )
    result = {row.gate_id: row for row in evaluate_publication_gates(unsourced)}["RESULT_STATUS_KNOWN"]
    assert result.status == "FAIL"
    assert "verified local bytes" in result.justification

    base = {**context.datasets["result_revisions"][0], "result_status": "FINAL", "valid_to": "2021-09-10"}
    rectified = {
        **base,
        "revision_id": "R2",
        "result_status": "RECTIFIED",
        "valid_from": "2021-09-11",
        "published_at": "2021-09-12",
        "known_at": "2021-09-12",
        "supersedes_revision_id": "R",
        "decision_id": "D",
    }
    rectified.pop("valid_to")
    decision = {
        "decision_id": "D", "decision_type": "RECTIFICATION", "decision_date": "2021-09-11",
        "legal_basis": "Judgment", "source_id": "S", "source_url": "https://official.example/law",
        "affected_revision_id": "R2",
    }
    linked = _dataset(context, "result_revisions", [base, rectified])
    linked = _dataset(linked, "legal_decisions", [decision])
    claim_fields = ("revision_id", "result_id", "contest_id", "result_status", "valid_from", "published_at")
    source_bytes = json.dumps({
        "result_history_claims": [
            {field: row.get(field) for field in claim_fields} for row in (base, rectified)
        ]
    }, sort_keys=True).encode("utf-8")
    assert linked.evidence_root is not None
    (linked.evidence_root / "official-law.pdf").write_bytes(source_bytes)
    linked_sources = [
        {**row, "bytes": len(source_bytes), "sha256": hashlib.sha256(source_bytes).hexdigest()}
        if row["source_id"] == "S" else row
        for row in linked.datasets["official_sources"]
    ]
    linked = _dataset(linked, "official_sources", linked_sources)
    assert linked.package_root is not None and linked.package_database_path is not None
    connection = duckdb.connect(str(linked.package_root / linked.package_database_path))
    try:
        connection.execute("DELETE FROM fact_result_revision")
        for row in (base, rectified):
            fields = [item[0] for item in connection.execute("DESCRIBE fact_result_revision").fetchall() if item[0] in row]
            connection.execute(
                "INSERT INTO fact_result_revision (" + ", ".join(f'\"{field}\"' for field in fields) +
                ") VALUES (" + ", ".join("?" for _ in fields) + ")",
                [row[field] for field in fields],
            )
        fields = [item[0] for item in connection.execute("DESCRIBE fact_legal_decision").fetchall() if item[0] in decision]
        connection.execute(
            "INSERT INTO fact_legal_decision (" + ", ".join(f'\"{field}\"' for field in fields) +
            ") VALUES (" + ", ".join("?" for _ in fields) + ")",
            [decision[field] for field in fields],
        )
    finally:
        connection.close()
    database_sha = hashlib.sha256((linked.package_root / linked.package_database_path).read_bytes()).hexdigest()
    linked = replace(linked, files=[
        {**row, "sha256": database_sha} if row["relative_path"] == linked.package_database_path else row
        for row in linked.files
    ])
    result = {row.gate_id: row for row in evaluate_publication_gates(linked)}["RESULT_STATUS_KNOWN"]
    assert result.status == "PASS", result.as_dict()

    broken = _dataset(linked, "legal_decisions", [{**decision, "affected_revision_id": "OTHER"}])
    result = {row.gate_id: row for row in evaluate_publication_gates(broken)}["RESULT_STATUS_KNOWN"]
    assert result.status == "FAIL"
    assert "affected revision" in result.justification


def test_temporal_gate_rejects_future_and_reversed_dates(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    revision = context.datasets["result_revisions"][0]
    known_future = _dataset(context, "result_revisions", [{**revision, "known_at": "2027-01-01"}])
    assert "known_at is later" in _as_of(known_future).justification
    published_future = _dataset(context, "result_revisions", [{**revision, "published_at": "2027-01-01"}])
    assert "published_at is later" in _as_of(published_future).justification
    reversed_dates = _dataset(
        context, "result_revisions",
        [{**revision, "published_at": "2021-09-12", "known_at": "2021-09-11"}],
    )
    assert "known_at cannot precede published_at" in _result_status(reversed_dates).justification


def test_result_status_gate_rejects_unproved_final_date_substitution_and_omission(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    revision = context.datasets["result_revisions"][0]
    substituted = _dataset(context, "result_revisions", [{**revision, "published_at": "2021-09-08"}])
    result = _result_status(substituted)
    assert result.status == "FAIL"
    assert "source bytes do not prove" in result.justification

    claim_fields = ("revision_id", "result_id", "contest_id", "result_status", "valid_from", "published_at")
    provisional_claim = {field: revision.get(field) for field in claim_fields}
    provisional_claim["result_status"] = "PROVISIONAL"
    source_bytes = json.dumps({"result_history_claims": [provisional_claim]}, sort_keys=True).encode("utf-8")
    assert context.evidence_root is not None
    (context.evidence_root / "official-law.pdf").write_bytes(source_bytes)
    sources = [
        {**row, "bytes": len(source_bytes), "sha256": hashlib.sha256(source_bytes).hexdigest()}
        if row["source_id"] == "S" else row
        for row in context.datasets["official_sources"]
    ]
    false_final = _dataset(context, "official_sources", sources)
    result = _result_status(false_final)
    assert result.status == "FAIL"
    assert "source bytes do not prove" in result.justification

    fabricated = _dataset(context, "result_revisions", [{**revision, "result_status": "PROCLAIMED"}])
    result = _result_status(fabricated)
    assert result.status == "FAIL"
    assert "source bytes do not prove" in result.justification

    omitted = _dataset(context, "result_revisions", [])
    result = _result_status(omitted)
    assert result.status == "FAIL"
    assert "materialized fact_result_revision" in result.justification


def test_result_status_gate_recomputes_source_bytes(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    assert _result_status(context).status == "PASS"
    assert context.evidence_root is not None
    (context.evidence_root / "official-law.pdf").write_bytes(b"mutated source")
    result = _result_status(context)
    assert result.status == "FAIL"
    assert "SOURCE_CHECKSUM_MISMATCH" in result.justification


def test_coordinated_revision_and_bundle_mutation_cannot_replace_source_proof(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    revision = {**context.datasets["result_revisions"][0], "result_status": "PROCLAIMED"}
    assert context.package_root is not None and context.package_database_path is not None
    connection = duckdb.connect(str(context.package_root / context.package_database_path))
    try:
        connection.execute("UPDATE fact_result_revision SET result_status = 'PROCLAIMED' WHERE revision_id = 'R'")
        connection.execute("CHECKPOINT")
    finally:
        connection.close()
    coordinated = _reauthor_test_package(_dataset(context, "result_revisions", [revision]))
    result = _result_status(coordinated)
    assert result.status == "FAIL"
    assert "source bytes do not prove" in result.justification


def test_v15_manifest_identity_is_location_independent_and_byte_exact(tmp_path: Path) -> None:
    context = _publication_context(tmp_path / "package")
    moved_manifest = tmp_path / "moved" / "v15-immutable-checksums.json"
    moved_manifest.parent.mkdir()
    moved_manifest.write_bytes(context.v15_manifest_path.read_bytes())
    moved = replace(context, v15_manifest_path=moved_manifest, v15_root=None)
    assert _immutability(moved).status == "PASS"
    moved_manifest.write_bytes(moved_manifest.read_bytes() + b"\n")
    mutated = _immutability(moved)
    assert mutated.status == "FAIL"
    assert "byte length differs" in mutated.justification
    assert "SHA-256 differs" in mutated.justification


def test_legal_gate_recomputes_population_based_classification(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    source_geo_id, official_code = "MA-01-051-0101", "01.051.01.01."
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Indic.Ensemble"
    for _ in range(3):
        sheet.append([])
    sheet.append([1, 51, 1, None, 1, None, None, "Commune A", 10])
    assert context.evidence_root is not None
    workbook_path = context.evidence_root / "hcp-individus.xlsx"
    workbook.save(workbook_path)
    workbook.close()
    workbook_bytes = workbook_path.read_bytes()
    population_source = {
        "source_id": "S-HCP", "title": "Test HCP individuals", "authority": "Test authority",
        "source_url": "https://official.example/hcp", "verification_status": "VERIFIED_AUTHORITY_AND_BYTES",
        "license_status": "UNKNOWN", "public_package_disposition": "METADATA_ONLY",
        "raw_path": workbook_path.name, "bytes": len(workbook_bytes),
        "sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        "extraction_method": "HCP_RGPH2014_INDIVIDUS",
    }
    context = _dataset(context, "official_sources", [*context.datasets["official_sources"], population_source])
    context = _dataset(context, "contests", [{**context.datasets["contests"][0], "geo_id": source_geo_id}])
    context = _dataset(context, "geographies", [
        *context.datasets["geographies"],
        {"geo_id": source_geo_id, "geo_type": "commune", "geo_name": "Commune A", "parent_geo_id": "MA-01-051"},
    ])
    classified = [{
        **context.datasets["contest_legal_regimes"][0],
        "classification_basis": "POPULATION_THRESHOLD", "classification_source_id": "S-HCP",
        "classification_value": 10, "classification_rule": "population <= 20",
    }]
    context = _dataset(context, "contest_legal_regimes", classified)
    context = _dataset(context, "population_legal_rules", [{
        "election_id": "E", "geo_type": "commune", "threshold": 20,
        "at_or_below_regime_id": "L", "above_regime_id": "OTHER", "population_source_id": "S-HCP",
    }])
    context = _dataset(context, "geo_populations", [{
        "geo_population_id": f"RGPH2014-INDIVIDUS:{official_code}", "official_geo_code": official_code,
        "population": 10, "census_date": "2014-09-01", "source_id": "S-HCP",
        "source_url": population_source["source_url"],
    }])
    context = _dataset(context, "geo_official_identifier_crosswalks", [{
        "crosswalk_id": f"HCP-RGPH2014-INDIVIDUS:{official_code}",
        "geo_id": source_geo_id, "official_geo_code": official_code,
        "matching_method": "OFFICIAL_IDENTIFIER", "source_id": "S-HCP",
        "geo_type": "commune", "confidence": 1.0,
    }])
    context = _materialize_classified_link(context, classified[0])
    assert {row.gate_id: row for row in evaluate_publication_gates(context)}["LEGAL_REGIME_PINNED"].status == "PASS"
    rule_changed = _dataset(context, "contest_legal_regimes", [{**classified[0], "classification_rule": "unsourced alternative"}])
    (tmp_path / "validation-evidence.json").unlink()  # Test-owned temporary artifact only.
    rule_changed = write_evidence_bundle(replace(rule_changed, evidence_bundle_path=None, evidence_bundle_sha256=None))
    gates = {row.gate_id: row for row in evaluate_publication_gates(rule_changed)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["LEGAL_REGIME_PINNED"].status == "FAIL"
    assert "C" in gates["LEGAL_REGIME_PINNED"].affected_records
    changed_population = _dataset(context, "geo_populations", [
        {**context.datasets["geo_populations"][0], "population": 11}
    ])
    (tmp_path / "validation-evidence.json").unlink()  # Test-owned temporary artifact only.
    changed_population = write_evidence_bundle(replace(changed_population, evidence_bundle_path=None, evidence_bundle_sha256=None))
    gates = {row.gate_id: row for row in evaluate_publication_gates(changed_population)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["LEGAL_REGIME_PINNED"].status == "FAIL"
    assert "materialized fact_geo_population row(s) were omitted or changed" in gates["LEGAL_REGIME_PINNED"].justification
    assert f"RGPH2014-INDIVIDUS:{official_code}" in gates["LEGAL_REGIME_PINNED"].affected_records
    dishonest = _dataset(context, "contest_legal_regimes", [{**classified[0], "classification_value": 11}])
    result = {row.gate_id: row for row in evaluate_publication_gates(dishonest)}["LEGAL_REGIME_PINNED"]
    assert result.status == "FAIL"
    assert "does not match" in result.justification
    uncertain = _dataset(context, "geo_official_identifier_crosswalks", [{
        **context.datasets["geo_official_identifier_crosswalks"][0], "confidence": 0.9,
    }])
    assert {row.gate_id: row for row in evaluate_publication_gates(uncertain)}["LEGAL_REGIME_PINNED"].status == "FAIL"
    # A coordinated falsification of package rows, manifest, reports, and bundle
    # must still fail because the independently pinned workbook has population 10.
    coordinated = _dataset(context, "geo_populations", [
        {**context.datasets["geo_populations"][0], "population": 11}
    ])
    coordinated = _dataset(coordinated, "contest_legal_regimes", [
        {**classified[0], "classification_value": 11}
    ])
    connection = duckdb.connect(str(tmp_path / "data.duckdb"))
    try:
        connection.execute("UPDATE fact_geo_population SET population = 11 WHERE geo_population_id = ?", [
            f"RGPH2014-INDIVIDUS:{official_code}",
        ])
        connection.execute("UPDATE bridge_contest_legal_regime SET classification_value = 11 WHERE contest_id = 'C'")
        connection.execute("CHECKPOINT")
    finally:
        connection.close()
    coordinated = _reauthor_test_package(coordinated)
    gates = {row.gate_id: row for row in evaluate_publication_gates(coordinated)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert "materialized fact_geo_population" not in gates["LEGAL_REGIME_PINNED"].justification
    assert "source-derived fact_geo_population" in gates["LEGAL_REGIME_PINNED"].justification
    assert gates["LEGAL_REGIME_PINNED"].status == "FAIL"
    assert validate_publication(coordinated)["publication_status"] == "NOT_PUBLICATION_READY"


def test_publication_gates_are_derived_and_reject_caller_booleans(tmp_path: Path) -> None:
    ready = validate_publication(_publication_context(tmp_path))
    assert ready["publication_status"] == "PUBLICATION_READY"
    assert tuple(row["gate_id"] for row in ready["gates"]) == REQUIRED_GATES
    assert all(row["status"] == "PASS" and row["justification"] and row["evidence_id"].startswith("sha256:") for row in ready["gates"])
    assert all(not validate_rows("publication_gate_result", [row]) for row in ready["gate_results"])
    with pytest.raises(TypeError, match="boolean gate maps are forbidden"):
        validate_publication({gate: True for gate in REQUIRED_GATES})  # type: ignore[arg-type]


def test_semantic_gate_rejects_omitted_package_fact_even_with_rewritten_valid_bundle(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    assert {row.gate_id: row for row in evaluate_publication_gates(context)}["SEMANTIC_FACTS_VALIDATED"].status == "PASS"
    omitted = _dataset(context, "semantic_facts", [dict(context.datasets["semantic_facts"][0])])
    (tmp_path / "validation-evidence.json").unlink()  # Test-owned temporary artifact only.
    omitted = write_evidence_bundle(replace(omitted, evidence_bundle_path=None, evidence_bundle_sha256=None))
    gates = {row.gate_id: row for row in evaluate_publication_gates(omitted)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["SEMANTIC_FACTS_VALIDATED"].status == "FAIL"
    assert "package fact row(s) were omitted" in gates["SEMANTIC_FACTS_VALIDATED"].justification
    assert "Y" in gates["SEMANTIC_FACTS_VALIDATED"].affected_records
    assert validate_publication(omitted)["publication_status"] == "NOT_PUBLICATION_READY"


def test_semantic_gate_rejects_unobserved_geography_omission_even_with_valid_bundle(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    assert {row.gate_id: row for row in evaluate_publication_gates(context)}["SEMANTIC_FACTS_VALIDATED"].status == "PASS"
    omitted = _dataset(context, "geographies", [dict(context.datasets["geographies"][0])])
    (tmp_path / "validation-evidence.json").unlink()  # Test-owned temporary artifact only.
    omitted = write_evidence_bundle(replace(omitted, evidence_bundle_path=None, evidence_bundle_sha256=None))
    gates = {row.gate_id: row for row in evaluate_publication_gates(omitted)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["SEMANTIC_FACTS_VALIDATED"].status == "FAIL"
    assert "package geographies row(s) were omitted" in gates["SEMANTIC_FACTS_VALIDATED"].justification
    assert "G2" in gates["SEMANTIC_FACTS_VALIDATED"].affected_records


def test_semantic_gate_reads_materialized_geo_parent_table(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    connection = duckdb.connect(str(tmp_path / "data.duckdb"))
    try:
        connection.execute(
            "INSERT INTO bridge_geo_parent "
            "(relation_id, child_geo_id, parent_geo_id, relationship_type, election_id, valid_from, "
            "source_id, matching_method, confidence, review_status) VALUES "
            "('P', 'G', 'G2', 'ADMINISTRATIVE_PARENT', 'E', '2021-09-08', "
            "'S', 'OFFICIAL_IDENTIFIER', 1, 'VERIFIED_SOURCE_DERIVED')"
        )
        connection.execute("CHECKPOINT")
    finally:
        connection.close()
    context = _reauthor_test_package(context)
    gate = {row.gate_id: row for row in evaluate_publication_gates(context)}["SEMANTIC_FACTS_VALIDATED"]
    assert gate.status == "FAIL"
    assert "package geo_parent_relations row(s) were omitted" in gate.justification
    assert "P" in gate.affected_records


def test_legal_gate_rejects_changed_formula_even_with_rewritten_valid_bundle(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    assert {row.gate_id: row for row in evaluate_publication_gates(context)}["LEGAL_REGIME_PINNED"].status == "PASS"
    changed = _dataset(
        context, "legal_regimes",
        [{**context.datasets["legal_regimes"][0], "allocation_formula": "LARGEST_REMAINDER"}],
    )
    (tmp_path / "validation-evidence.json").unlink()  # Test-owned temporary artifact only.
    changed = write_evidence_bundle(replace(changed, evidence_bundle_path=None, evidence_bundle_sha256=None))
    gates = {row.gate_id: row for row in evaluate_publication_gates(changed)}
    assert gates["EVIDENCE_BUNDLE_VERIFIED"].status == "PASS"
    assert gates["LEGAL_REGIME_PINNED"].status == "FAIL"
    assert "materialized dim_legal_regime row(s) were omitted or changed" in gates["LEGAL_REGIME_PINNED"].justification
    assert "L" in gates["LEGAL_REGIME_PINNED"].affected_records
    assert validate_publication(changed)["publication_status"] == "NOT_PUBLICATION_READY"


def test_evidence_bundle_writer_is_confined_and_immutable(tmp_path: Path) -> None:
    context = _publication_context(tmp_path / "package")
    with pytest.raises(FileExistsError, match="refusing to overwrite"):
        write_evidence_bundle(context)

    fresh = replace(context, package_root=tmp_path / "fresh", evidence_bundle_path=None, evidence_bundle_sha256=None)
    with pytest.raises(ValueError, match="unsafe evidence bundle path"):
        write_evidence_bundle(fresh, "../outside.json")
    assert not (tmp_path / "outside.json").exists()


def test_evidence_bundle_gate_detects_byte_mutation(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    original = {row.gate_id: row for row in evaluate_publication_gates(context)}["EVIDENCE_BUNDLE_VERIFIED"]
    assert original.status == "PASS"
    (tmp_path / "validation-evidence.json").write_bytes(b"{}")
    mutated = {row.gate_id: row for row in evaluate_publication_gates(context)}["EVIDENCE_BUNDLE_VERIFIED"]
    assert mutated.status == "FAIL"
    assert "checksum mismatch" in mutated.justification
    assert mutated.evidence_id != original.evidence_id


def test_evidence_bundle_gate_rejects_unmanaged_package_file(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    original = {row.gate_id: row for row in evaluate_publication_gates(context)}["EVIDENCE_BUNDLE_VERIFIED"]
    assert original.status == "PASS"
    (tmp_path / "undeclared.csv").write_text("hidden payload\n", encoding="utf-8")
    result = {row.gate_id: row for row in evaluate_publication_gates(context)}["EVIDENCE_BUNDLE_VERIFIED"]
    assert result.status == "FAIL"
    assert "unmanaged file" in result.justification
    assert result.affected_records == ("undeclared.csv",)
    assert result.evidence_id != original.evidence_id


def test_evidence_bundle_gate_rejects_package_symlink(tmp_path: Path) -> None:
    context = _publication_context(tmp_path / "package")
    assert context.evidence_root is not None
    assert context.package_root is not None
    link = context.package_root / "linked-source.pdf"
    try:
        link.symlink_to(context.evidence_root / "official-law.pdf")
    except OSError as error:
        pytest.skip(f"symbolic links are unavailable: {error}")
    result = {row.gate_id: row for row in evaluate_publication_gates(context)}["EVIDENCE_BUNDLE_VERIFIED"]
    assert result.status == "FAIL"
    assert "regular files physically contained" in result.justification
    assert result.affected_records == ("linked-source.pdf",)


def test_immutability_gate_rejects_noncanonical_or_invalid_manifest(tmp_path: Path) -> None:
    context = _publication_context(tmp_path / "package")
    alternate = tmp_path / "alternate-v15-manifest.json"
    alternate.write_text("{}", encoding="utf-8")
    result = {
        row.gate_id: row
        for row in evaluate_publication_gates(replace(context, v15_manifest_path=alternate))
    }["V15_IMMUTABILITY_VERIFIED"]
    assert result.status == "FAIL"
    assert "canonical" in result.justification
    assert "invalid" in result.justification


def test_privacy_gate_rejects_caller_exemption_duplicate_or_future_review(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    exempted = replace(context, files=[{**context.files[0], "privacy_review_required": False}])
    result = {row.gate_id: row for row in evaluate_publication_gates(exempted)}["PRIVACY_REVIEW_PASSED"]
    assert result.status == "FAIL"
    assert "exemptions are forbidden" in result.justification

    review = dict(context.checks["privacy_reviews"][0])
    duplicated = _check(context, "privacy_reviews", [review, review])
    result = {row.gate_id: row for row in evaluate_publication_gates(duplicated)}["PRIVACY_REVIEW_PASSED"]
    assert result.status == "FAIL"
    assert "exactly one" in result.justification

    future = _check(context, "privacy_reviews", [{**review, "reviewed_at": "2026-09-15"}])
    result = {row.gate_id: row for row in evaluate_publication_gates(future)}["PRIVACY_REVIEW_PASSED"]
    assert result.status == "FAIL"
    assert "after release as_of_date" in result.justification


def test_redistribution_gate_requires_one_sourced_license_decision(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    metadata_only = replace(context, files=[{**context.files[0], "license_status": "METADATA_ONLY"}])
    result = {row.gate_id: row for row in evaluate_publication_gates(metadata_only)}["REDISTRIBUTION_PERMITTED"]
    assert result.status == "FAIL"
    assert "only REDISTRIBUTABLE" in result.justification
    assert "differs from its reviewed decision" in result.justification

    missing = _check(context, "license_reviews", [])
    result = {row.gate_id: row for row in evaluate_publication_gates(missing)}["REDISTRIBUTION_PERMITTED"]
    assert result.status == "FAIL"
    assert "exactly one license review" in result.justification

    review = dict(context.checks["license_reviews"][0])
    future = _check(context, "license_reviews", [{**review, "reviewed_at": "2026-09-15"}])
    result = {row.gate_id: row for row in evaluate_publication_gates(future)}["REDISTRIBUTION_PERMITTED"]
    assert result.status == "FAIL"
    assert "after release as_of_date" in result.justification


def test_reviews_reject_duplicate_license_missing_proof_and_wrong_file_sha(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    review = dict(context.checks["license_reviews"][0])
    duplicate = _check(context, "license_reviews", [review, review])
    result = {row.gate_id: row for row in evaluate_publication_gates(duplicate)}["REDISTRIBUTION_PERMITTED"]
    assert result.status == "FAIL"
    assert "exactly one" in result.justification

    unproved = _check(context, "license_reviews", [{**review, "proof_sha256": None}])
    result = {row.gate_id: row for row in evaluate_publication_gates(unproved)}["REDISTRIBUTION_PERMITTED"]
    assert result.status == "FAIL"
    assert "proof linkage" in result.justification

    wrong_sha = _check(context, "license_reviews", [{**review, "file_sha256": "f" * 64}])
    result = {row.gate_id: row for row in evaluate_publication_gates(wrong_sha)}["REDISTRIBUTION_PERMITTED"]
    assert result.status == "FAIL"
    assert "wrong file SHA-256" in result.justification


def test_claim_gate_rejects_missing_class_and_inference_as_observed_fact(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    missing = replace(context, files=[{**context.files[0], "claim_class": None}])
    result = {row.gate_id: row for row in evaluate_publication_gates(missing)}["CLAIM_CLASS_DECLARED"]
    assert result.status == "FAIL"

    inferred = replace(context, files=[{**context.files[0], "fact_status": "INFERRED"}])
    result = {row.gate_id: row for row in evaluate_publication_gates(inferred)}["CLAIM_CLASS_DECLARED"]
    assert result.status == "FAIL"
    assert "cannot be an observed fact" in result.justification


def test_clean_build_gate_derives_isolation_instead_of_trusting_boolean(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    report_path = tmp_path / "clean-build-1.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["clean_environment"] is True
    report["environment"]["initial_entries"] = ["preexisting-file"]
    encoded = json.dumps(report, sort_keys=True)
    report_path.write_text(encoded, encoding="utf-8")
    clean_builds = [
        {**context.checks["clean_builds"][0], "report_sha256": hashlib.sha256(encoded.encode("utf-8")).hexdigest()},
        dict(context.checks["clean_builds"][1]),
    ]
    result = {
        row.gate_id: row
        for row in evaluate_publication_gates(_check(context, "clean_builds", clean_builds))
    }["REPRODUCIBLE_FROM_CLEAN_ENVIRONMENT"]
    assert result.status == "FAIL"
    assert "was not empty" in result.justification


def _mutate_gate(context: PublicationContext, gate_id: str) -> PublicationContext:
    if gate_id == "EVIDENCE_BUNDLE_VERIFIED":
        return replace(context, evidence_bundle_sha256="0" * 64)
    if gate_id == "SEMANTIC_FACTS_VALIDATED":
        return _dataset(context, "semantic_facts", [{**context.datasets["semantic_facts"][0], "election_id": "OTHER"}])
    if gate_id == "LEGAL_REGIME_PINNED":
        return _dataset(context, "legal_regimes", [])
    if gate_id == "RESULT_STATUS_KNOWN":
        rows = [{**context.datasets["result_revisions"][0], "result_status": "MADE_UP"}]
        return _dataset(context, "result_revisions", rows)
    if gate_id == "AS_OF_DATE_VALID":
        rows = [{**context.datasets["result_revisions"][0], "known_at": "2027-01-01"}]
        return _dataset(context, "result_revisions", rows)
    if gate_id == "OFFICIAL_UNIVERSE_DECLARED":
        rows = [{**context.datasets["coverage_universes"][0], "is_external": False}]
        return _dataset(context, "coverage_universes", rows)
    if gate_id == "DENOMINATOR_TYPED":
        rows = [{**context.datasets["coverage_universes"][0], "denominator": 2}]
        return _dataset(context, "coverage_universes", rows)
    if gate_id == "GRAIN_COMPATIBLE":
        rows = [*context.datasets["analytic_output"], dict(context.datasets["analytic_output"][0])]
        return _dataset(context, "analytic_output", rows)
    if gate_id == "BOUNDARY_COMPATIBLE":
        return _check(context, "boundary_checks", [{"check_id": "GE", "from_geo_version_id": "GV", "to_geo_version_id": "MISSING", "comparison_status": "NOT_COMPARABLE"}])
    if gate_id == "PARTY_LINEAGE_REVIEWED":
        rows = [{**context.datasets["party_affiliations"][0], "party_version_id": "MISSING"}]
        return _dataset(context, "party_affiliations", rows)
    if gate_id == "SOURCE_CONFLICTS_RESOLVED_OR_EXPOSED":
        rows = [*context.datasets["source_observations"], {"observation_id": "SO2", "subject_id": "C", "field": "valid_votes", "normalized_value": 11, "source_id": "S2"}]
        return _dataset(context, "source_observations", rows)
    if gate_id == "CANDIDACY_AND_SEATS_VALIDATED":
        return _dataset(context, "candidates", [])
    if gate_id == "METRIC_RECONCILED":
        rows = [{**context.datasets["reconciliations"][0], "validation_status": "FAIL", "difference": 1}]
        return _dataset(context, "reconciliations", rows)
    if gate_id == "ANALYTIC_METRICS_ADMISSIBLE":
        rows = [{**context.datasets["metric_validations"][0], "failed_preconditions": ["INCOMPLETE_COVERAGE"]}]
        return _dataset(context, "metric_validations", rows)
    if gate_id == "COVERAGE_DISCLOSED":
        return replace(context, coverage_matrix=context.coverage_matrix[:-1])
    if gate_id == "UNCERTAINTY_DISCLOSED_WHEN_APPLICABLE":
        return _check(context, "claim_reviews", [{"relative_path": "data.duckdb", "claim_class": "OBSERVED_FACT", "uncertainty_applicable": True}])
    if gate_id == "PRIVACY_REVIEW_PASSED":
        return _check(context, "privacy_reviews", [])
    if gate_id == "CLAIM_CLASS_DECLARED":
        return replace(context, files=[{**context.files[0], "fact_status": "INFERRED"}])
    if gate_id == "REDISTRIBUTION_PERMITTED":
        return replace(context, files=[{**context.files[0], "license_status": "UNKNOWN"}])
    if gate_id == "REPRODUCIBLE_FROM_CLEAN_ENVIRONMENT":
        rows = [dict(context.checks["clean_builds"][0]), {**context.checks["clean_builds"][1], "report_sha256": "0" * 64}]
        return _check(context, "clean_builds", rows)
    if gate_id == "V15_IMMUTABILITY_VERIFIED":
        return replace(context, v15_root=Path(__file__).parent / "missing-v15-root")
    raise AssertionError(gate_id)


@pytest.mark.parametrize("gate_id", REQUIRED_GATES)
def test_each_mandatory_gate_has_a_negative_mutation(gate_id: str, tmp_path: Path) -> None:
    context = _mutate_gate(_publication_context(tmp_path), gate_id)
    results = {row.gate_id: row for row in evaluate_publication_gates(context)}
    assert results[gate_id].status == "FAIL"
    assert results[gate_id].justification
    assert results[gate_id].evidence_id.startswith("sha256:")
    assert results[gate_id].affected_records


def test_association_causal_and_fraud_claims_fail_closed(tmp_path: Path) -> None:
    context = _publication_context(tmp_path)
    association = replace(context, files=[{**context.files[0], "claim_class": "ASSOCIATION", "fact_status": "RECOMPUTED"}])
    association = _check(association, "claim_reviews", [{"relative_path": "data.duckdb", "claim_class": "ASSOCIATION", "uncertainty_applicable": False}])
    claims = {row.gate_id: row for row in evaluate_publication_gates(association)}
    assert claims["CLAIM_CLASS_DECLARED"].status == "FAIL"
    causal = replace(context, files=[{**context.files[0], "claim_class": "CAUSAL", "fact_status": "RECOMPUTED"}])
    causal = _check(causal, "claim_reviews", [{"relative_path": "data.duckdb", "claim_class": "CAUSAL", "uncertainty_applicable": True, "uncertainty_disclosure": "95% CI", "automatic_fraud_inference": True}])
    claims = {row.gate_id: row for row in evaluate_publication_gates(causal)}
    assert claims["CLAIM_CLASS_DECLARED"].status == "FAIL"
    assert "fraud" in claims["CLAIM_CLASS_DECLARED"].justification
