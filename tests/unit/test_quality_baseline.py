from __future__ import annotations

from pathlib import Path

import openpyxl

from morocco_elections.quality.baseline import build_baseline, render_report, validate_baseline


def _sheet(workbook: openpyxl.Workbook, name: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append([name])
    sheet.append(["fixture synthétique"])
    sheet.append([])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)


def _synthetic_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "RAW_COMM2015_FULL", ["idCommune", "nInscrits"], [[1, None]])
    _sheet(workbook, "RAW_COMM2021_FULL", ["idCommune", "nInscrits"], [[1, None]])
    _sheet(workbook, "COMMUNE_ELECTION_PANEL", ["geo_id", "election_id"], [["MA-1", "COMM2015"], ["MA-1", "COMM2021"]])
    _sheet(workbook, "LOCAL_MANDATES", ["local_mandate_id"], [["M1"]])
    _sheet(workbook, "PARLIAMENTARY_MANDATES", ["mandate_id"], [["PM1"]])
    _sheet(workbook, "POPULATION", ["record_id"], [["POP1"]])
    _sheet(workbook, "RESULTS", ["record_id"], [["R1"]])
    _sheet(
        workbook,
        "COUNCIL_SEAT_STATUS_V10",
        ["legal_seat_count", "reconciled_seat_count"],
        [[1, 1]],
    )
    _sheet(
        workbook,
        "LOCAL_COUNCIL_CONTROL",
        ["geo_id", "president_person_id"],
        [["MA-1", "P1"]],
    )
    _sheet(
        workbook,
        "DATA_COVERAGE",
        ["domain_sheet", "next_granularity_target"],
        [["RESULTS", "offre partisane exhaustive"]],
    )
    _sheet(
        workbook,
        "QUALITY_CONTROL",
        ["issue_id", "sheet_name", "severity", "description", "resolution", "resolved_flag", "source_id"],
        [["QA_V10_INSCRITS_2015", "RAW_COMM2015_FULL", "high", "absent", "trouver la source", 0, "SRC"]],
    )
    audit_headers = ["sheet", "classification", "quality_status", "known_limitations"]
    audit_rows = [[name, "FACT", None, None] for name in workbook.sheetnames]
    _sheet(workbook, "WORKBOOK_AUDIT_V10", audit_headers, audit_rows)
    workbook.save(path)


def test_baseline_is_deterministic_and_rates_have_denominators(tmp_path: Path) -> None:
    workbook = tmp_path / "v10.xlsx"
    _synthetic_workbook(workbook)
    first = build_baseline(workbook, "2026-09-08")
    second = build_baseline(workbook, "2026-09-08")
    assert first == second
    assert validate_baseline(first) == []
    for item in first["coverage"]:
        if item["coverage_percent"] is not None:
            assert item["denominator_defined"] is True
            assert item["expected_count"] is not None


def test_report_is_a_pure_view_of_json(tmp_path: Path) -> None:
    workbook = tmp_path / "v10.xlsx"
    _synthetic_workbook(workbook)
    data = build_baseline(workbook, "2026-09-08")
    assert render_report(data) == render_report(data)
    assert "Aucune note globale" in render_report(data)
    assert "ANA_REGISTERED_VOTERS" not in render_report(data)


def test_blocker_without_analysis_restriction_is_rejected(tmp_path: Path) -> None:
    workbook = tmp_path / "v10.xlsx"
    _synthetic_workbook(workbook)
    data = build_baseline(workbook, "2026-09-08")
    data["anomalies"][0]["status"] = "BLOQUÉ"
    data["anomalies"][0]["restricts_analyses"] = []
    assert any("blocker sans restriction" in error for error in validate_baseline(data))
