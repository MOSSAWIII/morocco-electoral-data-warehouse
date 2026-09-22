from __future__ import annotations

import openpyxl
import pytest

from morocco_elections.v16.coverage import coverage_report
from morocco_elections.v16.demography import (
    exact_code_crosswalks,
    load_hcp_rgph2014_arrondissement_identifiers,
    load_hcp_rgph2014_territorial_universe,
    normalize_hcp_communal_code,
    validate_population_crosswalks,
)


def test_hcp_arrondissement_identifiers_require_explicit_label_and_complete_code(tmp_path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Communes"
    for _ in range(5):
        sheet.append([])
    sheet.append(["01.511.01.03.", "Bni Makada (Arrond.)"])
    sheet.append(["01.511.01.09.", "Gueznaia (Mun.)"])
    path = tmp_path / "hcp.xlsx"
    workbook.save(path)
    rows = load_hcp_rgph2014_arrondissement_identifiers(path, source_id="HCP")
    assert rows == [{
        "geo_id": "MA-01-511-0103", "official_geo_code": "01.511.01.03.",
        "official_name": "Bni Makada (Arrond.)", "prefecture_code": "01.511.", "source_id": "HCP",
    }]
    sheet.append(["01.511.01.03.", "Duplicate (Arrond.)"])
    workbook.save(path)
    with pytest.raises(ValueError, match="duplicated"):
        load_hcp_rgph2014_arrondissement_identifiers(path, source_id="HCP")
    sheet.cell(sheet.max_row, 1).value = "01.511.01."
    workbook.save(path)
    with pytest.raises(ValueError, match="complete official code"):
        load_hcp_rgph2014_arrondissement_identifiers(path, source_id="HCP")


def test_hcp_territorial_universe_retains_codes_with_suppressed_population(tmp_path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Communes"
    sheet.append(["header"])
    for _ in range(4):
        sheet.append([])
    sheet.append(["01.051.01.01.", "Commune A", None, 56716])
    sheet.append(["01.051.01.03.", "Commune B", None, "pm"])
    sheet.append(["not-an-official-code", "Unknown", None, 1])
    path = tmp_path / "hcp.xlsx"
    workbook.save(path)
    universe, members = load_hcp_rgph2014_territorial_universe(
        path, election_id="COMM2015", source_id="HCP",
        source_url="https://www.hcp.ma/file/230057/", acquired_at="2026-09-14",
    )
    assert universe["denominator"] == 2
    assert universe["universe_type"] == "OFFICIAL_TERRITORIES"
    assert [row["expected_id"] for row in members] == ["MA-01-051-0101", "MA-01-051-0103"]
    observations = [{"election_id": "COMM2015", "coverage_dimension": "TERRITORIAL", "geo_id": "MA-01-051-0101"}]
    report = coverage_report(observations, [universe], members)[2]
    assert report["status"] == "PARTIAL"
    assert report["missing_ids"] == ["MA-01-051-0103"]
    sheet.append(["01.051.01.01.", "Duplicate", None, 1])
    workbook.save(path)
    with pytest.raises(ValueError, match="duplicate territorial identifiers"):
        load_hcp_rgph2014_territorial_universe(
            path, election_id="COMM2015", source_id="HCP",
            source_url="https://www.hcp.ma/file/230057/", acquired_at="2026-09-14",
        )


def test_hcp_codes_are_crosswalked_exactly_without_name_matching() -> None:
    assert normalize_hcp_communal_code("01.051.01.01.") == "MA-01-051-0101"
    assert normalize_hcp_communal_code("Al Hoceima") is None
    populations = [{
        "geo_population_id": "RGPH2014:MA-01-051-0101",
        "official_geo_code": "01.051.01.01.",
        "normalized_geo_id": "MA-01-051-0101",
        "population": 56716,
        "census_date": "2014-09-01",
        "source_id": "HCP",
        "source_url": "https://www.hcp.ma/file/230057/",
    }]
    geographies = [
        {"geo_id": "MA-01-051-0101", "geo_type": "commune", "geo_name": "Different spelling"},
        {"geo_id": "MA-01-051-9999", "geo_type": "commune", "geo_name": "Al Hoceima"},
    ]
    crosswalks = exact_code_crosswalks(populations, geographies)
    assert [row["geo_id"] for row in crosswalks] == ["MA-01-051-0101"]
    assert crosswalks[0]["confidence"] == 1.0
    assert validate_population_crosswalks(populations, crosswalks, geographies) == []


def test_name_only_geo_crosswalk_is_rejected() -> None:
    population = {
        "geo_population_id": "P", "official_geo_code": "01.051.01.01.", "population": 1,
        "census_date": "2014-09-01", "source_id": "HCP", "source_url": "https://www.hcp.ma/file/230057/",
    }
    crosswalk = {
        "crosswalk_id": "X", "geo_id": "G", "official_geo_code": "01.051.01.01.",
        "matching_method": "NAME_ONLY", "source_id": "HCP", "confidence": 0.5,
    }
    codes = {issue.code for issue in validate_population_crosswalks([population], [crosswalk], [{"geo_id": "G"}])}
    assert "NAME_ONLY_GEO_CROSSWALK_FORBIDDEN" in codes
