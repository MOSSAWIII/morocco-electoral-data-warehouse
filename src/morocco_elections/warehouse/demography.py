from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any, Iterable, Mapping

import openpyxl

from morocco_elections.warehouse.validation import ValidationIssue, validate_rows


HCP_COMMUNAL_CODE = re.compile(r"^(\d{2})\.(\d{3})\.(\d{2})\.(\d{2})\.$")


def normalize_hcp_communal_code(value: Any) -> str | None:
    """Translate a complete HCP communal/arrondissement code without fuzzy matching."""
    match = HCP_COMMUNAL_CODE.fullmatch(str(value).strip()) if value is not None else None
    return f"MA-{match[1]}-{match[2]}-{match[3]}{match[4]}" if match else None


def _normalized_name(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).casefold()
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = re.sub(r"\s*\((?:mun\.|arrond\.)\)\s*$", "", text)
    text = re.sub(r"[^a-z0-9\u0600-\u06ff]+", " ", text)
    return " ".join(text.split())


def load_hcp_rgph2014_individuals(
    workbook_path: Path,
    geographies: Iterable[Mapping[str, Any]],
    *,
    source_id: str,
    source_url: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Extract population and code crosswalks, preserving lower-confidence name/parent matches."""
    geographies = list(geographies)
    by_id = {str(row.get("geo_id")): row for row in geographies}
    by_name_parent: dict[tuple[str, str], list[Mapping[str, Any]]] = {}
    for row in geographies:
        if row.get("geo_type") in {"commune", "arrondissement"}:
            by_name_parent.setdefault(
                (str(row.get("parent_geo_id")), _normalized_name(row.get("geo_name"))), []
            ).append(row)
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    populations: list[dict[str, Any]] = []
    crosswalks: list[dict[str, Any]] = []
    try:
        for values in workbook["Indic.Ensemble"].iter_rows(min_row=4, values_only=True):
            try:
                if not (all(values[index] not in (None, "") for index in (0, 1, 2, 4)) and values[5] in (None, "")):
                    continue
                region, province, group, unit = (int(values[index]) for index in (0, 1, 2, 4))
                source_geo_id = f"MA-{region:02d}-{province:03d}-{group:02d}{unit:02d}"
                official_code = f"{region:02d}.{province:03d}.{group:02d}.{unit:02d}."
                population = float(values[8])
            except (TypeError, ValueError, IndexError):
                continue
            if population < 0 or not population.is_integer():
                raise ValueError(f"invalid population for official code {official_code}")
            geography = by_id.get(source_geo_id)
            method, confidence = "OFFICIAL_IDENTIFIER", 1.0
            if geography is None:
                parent = f"MA-{region:02d}-{province:03d}"
                candidates = by_name_parent.get((parent, _normalized_name(values[7])), [])
                if len(candidates) != 1:
                    continue
                geography = candidates[0]
                method, confidence = "DOCUMENTED_CROSSWALK", 0.9
            geo_id = str(geography["geo_id"])
            populations.append({
                "geo_population_id": f"RGPH2014-INDIVIDUS:{official_code}",
                "official_geo_code": official_code,
                "normalized_geo_id": geo_id,
                "population": int(population),
                "census_date": "2014-09-01",
                "source_id": source_id,
                "source_url": source_url,
            })
            crosswalks.append({
                "crosswalk_id": f"HCP-RGPH2014-INDIVIDUS:{official_code}",
                "geo_id": geo_id,
                "geo_type": geography.get("geo_type"),
                "official_geo_code": official_code,
                "matching_method": method,
                "source_id": source_id,
                "confidence": confidence,
            })
    finally:
        workbook.close()
    issues = validate_population_crosswalks(populations, crosswalks, geographies)
    if issues:
        raise ValueError("invalid HCP population extraction: " + ", ".join(sorted({issue.code for issue in issues})))
    return populations, crosswalks


def load_hcp_rgph2014_population(
    workbook_path: Path,
    *,
    source_id: str,
    source_url: str,
) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Communes"]
        rows: list[dict[str, Any]] = []
        for values in sheet.iter_rows(min_row=6, values_only=True):
            official_code, population = values[0], values[3]
            geo_id = normalize_hcp_communal_code(official_code)
            if geo_id is None:
                continue
            if population in {None, "pm", "p.m."}:
                continue
            if (
                not isinstance(population, (int, float))
                or isinstance(population, bool)
                or population < 0
                or int(population) != population
            ):
                raise ValueError(f"invalid population for official code {official_code!r}")
            rows.append(
                {
                    "geo_population_id": f"RGPH2014:{geo_id}",
                    "official_geo_code": str(official_code),
                    "normalized_geo_id": geo_id,
                    "population": int(population),
                    "census_date": "2014-09-01",
                    "source_id": source_id,
                    "source_url": source_url,
                }
            )
    finally:
        workbook.close()
    issues = validate_rows("fact_geo_population", rows)
    if issues:
        raise ValueError("invalid HCP population rows: " + ", ".join(sorted({issue.code for issue in issues})))
    return rows


def load_hcp_rgph2014_territorial_universe(
    workbook_path: Path,
    *,
    election_id: str,
    source_id: str,
    source_url: str,
    acquired_at: str,
) -> tuple[dict[str, Any], list[dict[str, str]]]:
    """Use every complete official code, including rows with suppressed population."""
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        extracted_codes = [
            normalized
            for values in workbook["Communes"].iter_rows(min_row=6, values_only=True)
            if (normalized := normalize_hcp_communal_code(values[0])) is not None
        ]
    finally:
        workbook.close()
    codes = set(extracted_codes)
    if len(codes) != len(extracted_codes):
        raise ValueError("HCP workbook contains duplicate territorial identifiers")
    if not codes:
        raise ValueError("HCP workbook has no inspectable territorial identifiers")
    universe_id = f"HCP-RGPH2014-TERRITORIES:{election_id}"
    universe = {
        "universe_id": universe_id,
        "election_id": election_id,
        "coverage_dimension": "TERRITORIAL",
        "universe_type": "OFFICIAL_TERRITORIES",
        "denominator": len(codes),
        "source_id": source_id,
        "source_url": source_url,
        "acquired_at": acquired_at,
        "verification_status": "VERIFIED",
        "is_external": True,
        "member_extraction_method": "HCP_RGPH2014_COMMUNES",
    }
    members = [
        {"universe_id": universe_id, "expected_id": code, "source_id": source_id}
        for code in sorted(codes)
    ]
    return universe, members


def load_hcp_rgph2014_arrondissement_identifiers(
    workbook_path: Path,
    *,
    source_id: str,
) -> list[dict[str, str]]:
    """Read explicitly labelled arrondissement codes, without inferring council parents."""
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    rows: list[dict[str, str]] = []
    try:
        for values in workbook["Communes"].iter_rows(min_row=6, values_only=True):
            name = str(values[1] or "")
            if "(Arrond.)" not in name:
                continue
            code = str(values[0] or "")
            geo_id = normalize_hcp_communal_code(code)
            if geo_id is None:
                raise ValueError(f"labelled HCP arrondissement lacks a complete official code: {code!r}")
            rows.append({
                "geo_id": geo_id,
                "official_geo_code": code,
                "official_name": name,
                "prefecture_code": code[:7],
                "source_id": source_id,
            })
    finally:
        workbook.close()
    identifiers = [row["geo_id"] for row in rows]
    if not rows or len(identifiers) != len(set(identifiers)):
        raise ValueError("HCP arrondissement identifiers are absent or duplicated")
    return sorted(rows, key=lambda row: row["geo_id"])


def exact_code_crosswalks(
    populations: Iterable[Mapping[str, Any]],
    geographies: Iterable[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Crosswalk only byte-equivalent normalized codes; never infer identity from names."""
    geographies_by_id = {str(row.get("geo_id")): row for row in geographies}
    rows = []
    for population in populations:
        geo_id = str(population.get("normalized_geo_id"))
        geography = geographies_by_id.get(geo_id)
        if geography is None or geography.get("geo_type") not in {"commune", "arrondissement"}:
            continue
        rows.append(
            {
                "crosswalk_id": f"HCP-RGPH2014:{geo_id}",
                "geo_id": geo_id,
                "geo_type": geography.get("geo_type"),
                "official_geo_code": population.get("official_geo_code"),
                "matching_method": "DOCUMENTED_CROSSWALK",
                "source_id": population.get("source_id"),
                "confidence": 1.0,
            }
        )
    return rows


def validate_population_crosswalks(
    populations: Iterable[Mapping[str, Any]],
    crosswalks: Iterable[Mapping[str, Any]],
    geographies: Iterable[Mapping[str, Any]],
) -> list[ValidationIssue]:
    populations, crosswalks, geographies = list(populations), list(crosswalks), list(geographies)
    issues = validate_rows("fact_geo_population", populations)
    issues += validate_rows("bridge_geo_official_identifier", crosswalks)
    population_codes = {row.get("official_geo_code") for row in populations}
    geography_ids = {row.get("geo_id") for row in geographies}
    for row in crosswalks:
        rid = str(row.get("crosswalk_id", "<unknown>"))
        if row.get("official_geo_code") not in population_codes:
            issues.append(ValidationIssue("UNKNOWN_OFFICIAL_GEO_CODE", "bridge_geo_official_identifier", rid, str(row.get("official_geo_code"))))
        if row.get("geo_id") not in geography_ids:
            issues.append(ValidationIssue("UNKNOWN_CROSSWALK_GEOGRAPHY", "bridge_geo_official_identifier", rid, str(row.get("geo_id"))))
        if row.get("matching_method") == "NAME_ONLY":
            issues.append(ValidationIssue("NAME_ONLY_GEO_CROSSWALK_FORBIDDEN", "bridge_geo_official_identifier", rid, "geography identity cannot rely only on a name"))
    return issues
