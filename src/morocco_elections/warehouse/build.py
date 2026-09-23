from __future__ import annotations

import hashlib
import json
import os
import tempfile
from pathlib import Path
from typing import Any

import duckdb
import openpyxl

from morocco_elections.warehouse import DEFAULT_SNAPSHOT_ID, SCHEMA_REVISION
from morocco_elections.warehouse.bo6374_reconciliation import normalize_latin
from morocco_elections.warehouse.contracts import TABLE_CONTRACTS
from morocco_elections.warehouse.demography import load_hcp_rgph2014_individuals
from morocco_elections.warehouse.evidence import (
    derive_contest_legal_regime_links,
    load_json,
    validate_legal_regime_seed,
    validate_official_source_registry,
)
from morocco_elections.warehouse.geo_parents import build_geo_parent_evidence, geo_parent_report
from morocco_elections.warehouse.reconciliation import derive_reconciliation_matrix
from morocco_elections.warehouse.result_history import build_result_history_diagnostic
from morocco_elections.warehouse.schema import create_schema
from morocco_elections.warehouse.semantic import create_analytical_views
from morocco_elections.warehouse.sources import canonical_source_id, source_registry_issues


def _quoted(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def _insert_rows(connection: duckdb.DuckDBPyConnection, table: str, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    columns = [row[0] for row in connection.execute(f"DESCRIBE {_quoted(table)}").fetchall()]
    selected = [column for column in columns if any(column in row for row in rows)]
    placeholders = ", ".join("?" for _ in selected)
    quoted_columns = ", ".join(_quoted(column) for column in selected)
    values = []
    for row in rows:
        values.append(
            [json.dumps(row.get(column), ensure_ascii=False) if isinstance(row.get(column), (list, dict)) else row.get(column) for column in selected]
        )
    connection.executemany(
        f"INSERT INTO {_quoted(table)} ({quoted_columns}) VALUES ({placeholders})",
        values,
    )


def _normalize_historical_source_ids(connection: duckdb.DuckDBPyConnection) -> int:
    """Remove legacy release labels from every copied source-reference column."""
    if not connection.execute(
        "SELECT count(*) FROM information_schema.tables "
        "WHERE table_schema='main' AND table_name='sources'"
    ).fetchone()[0]:
        return 0
    source_ids = [row[0] for row in connection.execute("SELECT source_id FROM sources").fetchall()]
    aliases = {source_id: canonical_source_id(source_id) for source_id in source_ids}
    if len(set(aliases.values())) != len(aliases):
        raise RuntimeError("historical source aliases are not unique")
    columns = connection.execute(
        "SELECT table_name, column_name FROM information_schema.columns "
        "WHERE table_schema='main' AND (column_name='source_id' OR column_name LIKE '%_source_id') "
        "ORDER BY (table_name='sources'), table_name, column_name"
    ).fetchall()
    changed = 0
    for table, column in columns:
        for old, new in aliases.items():
            if old == new:
                continue
            result = connection.execute(
                f"UPDATE {_quoted(table)} SET {_quoted(column)} = ? WHERE {_quoted(column)} = ?",
                [new, old],
            )
            changed += result.fetchone()[0] if result.description else 0
    return changed


def _normalize_seed_identifiers(connection: duckdb.DuckDBPyConnection) -> int:
    """Remove any release marker from copied identifiers at the seed boundary."""
    columns = connection.execute(
        "SELECT table_name, column_name FROM information_schema.columns "
        "WHERE table_schema='main' AND data_type='VARCHAR' "
        "AND (column_name LIKE '%_id' OR column_name='candidate_sources') "
        "ORDER BY table_name, column_name"
    ).fetchall()
    changed = 0
    for table, column in columns:
        # Normalize list members first. Otherwise the embedded-marker pattern
        # can partially consume a release marker with a numeric sub-suffix.
        delimited = connection.execute(
            f"UPDATE {_quoted(table)} SET {_quoted(column)} = "
            f"regexp_replace({_quoted(column)}, "
            "'_V[0-9]+(_[0-9]+)*([;, ]+)', '\\2', 'gi') "
            f"WHERE {_quoted(column)} IS NOT NULL AND "
            f"regexp_matches({_quoted(column)}, '_V[0-9]+(_[0-9]+)*[;, ]', 'i')"
        )
        changed += delimited.fetchone()[0] if delimited.description else 0
        result = connection.execute(
            f"UPDATE {_quoted(table)} SET {_quoted(column)} = "
            f"regexp_replace(regexp_replace({_quoted(column)}, "
            "'_V[0-9]+(_[0-9]+)*_', '_', 'gi'), "
            "'_V[0-9]+(_[0-9]+)*$', '', 'gi') "
            f"WHERE {_quoted(column)} IS NOT NULL AND "
            f"regexp_matches({_quoted(column)}, '_V[0-9]+(_[0-9]+)*(_|$)', 'i')"
        )
        changed += result.fetchone()[0] if result.description else 0
    return changed


def _normalize_historical_release_text(connection: duckdb.DuckDBPyConnection) -> int:
    """Describe copied release-labelled notes as historical seed metadata."""
    note_columns = connection.execute(
        "SELECT table_name, column_name FROM information_schema.columns "
        "WHERE table_schema='main' AND data_type='VARCHAR' AND column_name='notes' "
        "ORDER BY table_name"
    ).fetchall()
    changed = 0
    pattern = r"(^|[^a-z0-9])V(9|10|11|12|13|14|15|16)(-QA)?([^a-z0-9]|$)"
    for table, column in note_columns:
        result = connection.execute(
            f"UPDATE {_quoted(table)} SET {_quoted(column)} = "
            f"regexp_replace({_quoted(column)}, ?, '\\1historical seed\\4', 'gi') "
            f"WHERE {_quoted(column)} IS NOT NULL AND regexp_matches({_quoted(column)}, ?, 'i')",
            [pattern, pattern],
        )
        changed += result.fetchone()[0] if result.description else 0
    return changed


_COMMUNAL_RESULT_SOURCES = {
    "COMM2015": "SRC_TAFRA_COMM2015",
    "COMM2021": "SRC_TAFRA_COMM2021",
}


def _communal_source_index(repository_root: Path, descriptor: dict[str, Any]) -> dict[tuple[str, str], tuple[str, str]]:
    """Read pinned TAFRA commune IDs without treating them as official identities."""
    relative = Path(str(descriptor["raw_path"]))
    path = (repository_root / relative).resolve()
    root = repository_root.resolve()
    if relative.is_absolute() or ".." in relative.parts or root not in path.parents:
        raise ValueError(f"unsafe communal-result source path: {relative}")
    observed_sha = None
    if path.is_file():
        with path.open("rb") as stream:
            observed_sha = hashlib.file_digest(stream, "sha256").hexdigest()
    if not path.is_file() or path.stat().st_size != int(descriptor["bytes"]) or observed_sha != descriptor["sha256"]:
        raise ValueError(f"communal-result source bytes differ: {descriptor['source_id']}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        records = workbook.worksheets[0].iter_rows(values_only=True)
        headers = list(next(records))
        required = {"idCommune", "prefProv", "commune"}
        if not required <= set(headers):
            raise ValueError(f"communal-result source lacks required columns: {descriptor['source_id']}")
        index: dict[tuple[str, str], tuple[str, str]] = {}
        source_ids: set[str] = set()
        for values in records:
            row = dict(zip(headers, values, strict=False))
            if row["idCommune"] is None:
                raise ValueError(f"communal-result source has a missing idCommune: {descriptor['source_id']}")
            source_id = str(int(row["idCommune"]))
            source_label = str(row["commune"])
            key = (normalize_latin(row["prefProv"]), normalize_latin(source_label))
            if key in index and index[key] != (source_id, source_label):
                raise ValueError(f"communal-result source has an ambiguous place label: {descriptor['source_id']}")
            index[key] = (source_id, source_label)
            source_ids.add(source_id)
    finally:
        workbook.close()
    if len(index) != 1538 or len(source_ids) != 1538:
        raise ValueError(f"communal-result source does not contain 1,538 unique communes: {descriptor['source_id']}")
    return index


def _materialize_communal_source_identifiers(
    connection: duckdb.DuckDBPyConnection,
    repository_root: Path,
    source_registry: dict[str, Any],
) -> dict[str, Any]:
    """Preserve secondary source IDs while keeping name-only identity links explicit."""
    descriptors = {row["source_id"]: row for row in source_registry["sources"]}
    indexes = {
        election_id: _communal_source_index(repository_root, descriptors[source_id])
        for election_id, source_id in _COMMUNAL_RESULT_SOURCES.items()
    }
    counts: dict[str, dict[str, int]] = {}
    for election_id, index in indexes.items():
        rows = connection.execute(
            "SELECT c.contest_id, g.geo_name, p.geo_name "
            "FROM dim_electoral_contest c JOIN dim_geo g USING(geo_id) "
            "LEFT JOIN dim_geo p ON p.geo_id=g.parent_geo_id WHERE c.election_id=? "
            "ORDER BY c.contest_id",
            [election_id],
        ).fetchall()
        matched = 0
        for contest_id, commune_name, province_name in rows:
            source = index.get((normalize_latin(province_name), normalize_latin(commune_name)))
            if source is None:
                connection.execute(
                    "UPDATE dim_electoral_contest SET source_contest_id=NULL, source_label=NULL, "
                    "normalized_label=NULL, identity_review_status='SECONDARY_SOURCE_IDENTIFIER_UNRESOLVED' "
                    "WHERE contest_id=?",
                    [contest_id],
                )
                continue
            source_id, source_label = source
            connection.execute(
                "UPDATE dim_electoral_contest SET source_contest_id=?, source_label=?, normalized_label=?, "
                "identity_review_status='SECONDARY_SOURCE_IDENTIFIER_NAME_MATCHED' WHERE contest_id=?",
                [source_id, source_label, normalize_latin(source_label), contest_id],
            )
            matched += 1
        counts[election_id] = {"contests": len(rows), "matched": matched, "unmatched": len(rows) - matched}
    uncertain = connection.execute(
        "SELECT x.geo_id, count(DISTINCT c.source_contest_id), count(*) "
        "FROM bridge_geo_official_identifier x JOIN dim_electoral_contest c USING(geo_id) "
        "WHERE x.confidence < 1 AND c.election_id IN ('COMM2015', 'COMM2021') "
        "GROUP BY x.geo_id"
    ).fetchall()
    stable_uncertain = sum(distinct_ids == 1 and row_count == 2 for _, distinct_ids, row_count in uncertain)
    return {
        "by_election": counts,
        "uncertain_crosswalks": len(uncertain),
        "stable_secondary_identifiers": stable_uncertain,
    }


def _rename_parliamentary_source_keys(connection: duckdb.DuckDBPyConnection) -> None:
    """Reserve source_id for canonical provenance, not parliamentary file identity."""
    for table in ("dim_parliamentary_source", "bridge_question_source"):
        present = connection.execute(
            "SELECT count(*) FROM information_schema.tables "
            "WHERE table_schema='main' AND table_name=?", [table]
        ).fetchone()[0]
        if not present:
            continue
        columns = {row[0] for row in connection.execute(f"DESCRIBE {_quoted(table)}").fetchall()}
        if "source_id" in columns and "parliamentary_source_id" not in columns:
            connection.execute(
                f"ALTER TABLE {_quoted(table)} RENAME COLUMN source_id TO parliamentary_source_id"
            )


def _materialize_canonical_sources(
    connection: duckdb.DuckDBPyConnection, source_registry: dict[str, Any]
) -> None:
    """Make the package source dimension an exact projection of the one registry."""
    connection.execute("DELETE FROM sources")
    rows = []
    checked_at = source_registry.get("as_of_date")
    for source in source_registry.get("sources", []):
        rows.append({
            "source_id": source["source_id"],
            "source_name": source["title"],
            "publisher": source["authority"],
            "source_type": ",".join(source.get("usages", [])),
            "domain": "canonical_source_registry",
            "url": source["source_url"],
            "geo_granularity": source["grain"],
            "format": source.get("format"),
            "access_method": source["acquisition_status"],
            "license": source["license_status"],
            "reliability_score": source.get("reliability_score"),
            "last_checked": checked_at,
            "status": source["status"],
            "notes": source.get("notes"),
        })
    _insert_rows(connection, "sources", rows)


def _validate_provenance_references(connection: duckdb.DuckDBPyConnection) -> None:
    """Require every column actually named source_id to use the canonical registry."""
    present = connection.execute(
        "SELECT count(*) FROM information_schema.tables "
        "WHERE table_schema='main' AND table_name='sources'"
    ).fetchone()[0]
    if not present:
        return
    registered = {row[0] for row in connection.execute("SELECT source_id FROM sources").fetchall()}
    failures: list[str] = []
    tables = connection.execute(
        "SELECT table_name FROM information_schema.columns "
        "WHERE table_schema='main' AND column_name='source_id' ORDER BY table_name"
    ).fetchall()
    for (table,) in tables:
        unknown = connection.execute(
            f"SELECT DISTINCT source_id FROM {_quoted(table)} "
            "WHERE source_id IS NOT NULL AND source_id NOT IN (SELECT source_id FROM sources) LIMIT 1"
        ).fetchone()
        if unknown:
            failures.append(f"{table}.source_id={unknown[0]}")
    if failures:
        raise RuntimeError("non-canonical provenance reference: " + failures[0])


def build_development_database(
    seed_database: Path,
    output_database: Path,
    *,
    legal_seed_path: Path | None = None,
    source_registry_path: Path | None = None,
    evidence_root: Path | None = None,
    replace_existing: bool = False,
    as_of_date: str = "2026-09-21",
) -> dict[str, Any]:
    """Create the canonical warehouse from an immutable historical seed."""
    seed_database = seed_database.resolve()
    output_database = output_database.resolve()
    if not seed_database.is_file():
        raise FileNotFoundError(seed_database)
    if output_database.exists() and not replace_existing:
        raise FileExistsError(f"refusing to overwrite existing warehouse database: {output_database}")
    legal_payload: dict[str, Any] | None = None
    source_registry: dict[str, Any] | None = None
    population_rows: list[dict[str, Any]] = []
    population_source: dict[str, Any] | None = None
    geo_parent_rows: list[dict[str, Any]] = []
    geo_parent_details: dict[str, Any] = {}
    parent_report: dict[str, Any] = {}
    reconciliation_rows: list[dict[str, Any]] = []
    reconciliation_report: dict[str, Any] = {}
    result_history_report: dict[str, Any] = {}
    communal_source_identifiers: dict[str, Any] = {}
    if legal_seed_path is not None or source_registry_path is not None:
        if legal_seed_path is None or source_registry_path is None or evidence_root is None:
            raise ValueError("legal_seed_path, source_registry_path and evidence_root must be supplied together")
        legal_payload = load_json(legal_seed_path)
        source_registry = load_json(source_registry_path)
        registry_issues = source_registry_issues(source_registry)
        if registry_issues:
            raise RuntimeError(f"invalid canonical source registry: {registry_issues[0]}")
        official_sources = [
            row for row in source_registry.get("sources", [])
            if "official_evidence" in row.get("usages", [])
        ]
        evidence_issues = validate_official_source_registry(
            evidence_root, {"sources": official_sources}
        )
        evidence_issues += validate_legal_regime_seed(legal_payload, source_registry)
        if evidence_issues:
            codes = ", ".join(sorted({issue.code for issue in evidence_issues}))
            raise RuntimeError(f"refusing to seed unverified legal evidence: {codes}")
        population_source_id = legal_payload["population_link_rules"][0]["population_source_id"]
        population_source = next(row for row in source_registry["sources"] if row["source_id"] == population_source_id)
    output_database.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="warehouse-build-", dir=output_database.parent) as temporary_name:
        temporary = Path(temporary_name) / output_database.name
        contest_count = 0
        crosswalks: list[dict[str, Any]] = []
        connection = duckdb.connect(str(temporary))
        try:
            seed_sql = str(seed_database).replace("'", "''")
            connection.execute(f"ATTACH '{seed_sql}' AS historical_seed (READ_ONLY)")
            source_tables = [
                row[0]
                for row in connection.execute(
                    "SELECT table_name FROM information_schema.tables "
                    "WHERE table_catalog='historical_seed' AND table_type='BASE TABLE' ORDER BY table_name"
                ).fetchall()
            ]
            collisions = set(source_tables) & TABLE_CONTRACTS.keys()
            if collisions:
                raise RuntimeError("seed/canonical table collision: " + ", ".join(sorted(collisions)))
            for table in source_tables:
                quoted = _quoted(table)
                connection.execute(f"CREATE TABLE main.{quoted} AS SELECT * FROM historical_seed.{quoted}")
            connection.execute("DETACH historical_seed")
            normalized_source_references = _normalize_historical_source_ids(connection)
            normalized_seed_identifiers = _normalize_seed_identifiers(connection)
            normalized_historical_text = _normalize_historical_release_text(connection)
            _rename_parliamentary_source_keys(connection)
            create_schema(connection)
            if "warehouse_metadata" not in source_tables:
                connection.execute(
                    "CREATE TABLE warehouse_metadata (release VARCHAR, schema_version BIGINT, seed_identity VARCHAR)"
                )
                connection.execute(
                    "INSERT INTO warehouse_metadata VALUES (?, ?, 'historical_seed')",
                    [DEFAULT_SNAPSHOT_ID, SCHEMA_REVISION],
                )
            metadata_columns = {
                row[0] for row in connection.execute("DESCRIBE warehouse_metadata").fetchall()
            }
            legacy_identity_columns = metadata_columns - {
                "release", "schema_version", "as_of_date", "seed_identity",
            }
            if "seed_identity" not in metadata_columns and len(legacy_identity_columns) == 1:
                legacy_identity = next(iter(legacy_identity_columns))
                connection.execute(
                    f"ALTER TABLE warehouse_metadata RENAME COLUMN {_quoted(legacy_identity)} TO seed_identity"
                )
                metadata_columns.remove(legacy_identity)
                metadata_columns.add("seed_identity")
            if "seed_identity" not in metadata_columns:
                connection.execute("ALTER TABLE warehouse_metadata ADD COLUMN seed_identity VARCHAR")
            if "as_of_date" not in metadata_columns:
                connection.execute("ALTER TABLE warehouse_metadata ADD COLUMN as_of_date DATE")
            connection.execute(
                "UPDATE warehouse_metadata SET release = ?, schema_version = ?, "
                "seed_identity = 'historical_seed', as_of_date = ?",
                [DEFAULT_SNAPSHOT_ID, SCHEMA_REVISION, as_of_date],
            )
            if source_registry is not None:
                _materialize_canonical_sources(connection, source_registry)
            if evidence_root is not None:
                geo_parent_rows, geo_parent_details = build_geo_parent_evidence(connection, evidence_root)
                _insert_rows(connection, "bridge_geo_parent", geo_parent_rows)
                parent_report = geo_parent_report(connection, geo_parent_rows)
                reconciliation_rows, reconciliation_report = derive_reconciliation_matrix(connection, evidence_root)
                _insert_rows(connection, "fact_result_reconciliation", reconciliation_rows)
            contest_links: list[dict[str, Any]] = []
            if legal_payload is not None:
                _insert_rows(connection, "dim_legal_regime", list(legal_payload["legal_regimes"]))
                _insert_rows(connection, "bridge_election_legal_regime", list(legal_payload["election_links"]))
                if "dim_electoral_contest" in source_tables:
                    columns = {row[0] for row in connection.execute("DESCRIBE dim_electoral_contest").fetchall()}
                    if {"contest_id", "election_id", "list_type"} <= columns:
                        contests = [
                            {"contest_id": row[0], "election_id": row[1], "source_list_type": row[2], "geo_id": row[3]}
                            for row in connection.execute(
                                "SELECT contest_id, election_id, list_type, geo_id FROM dim_electoral_contest"
                            ).fetchall()
                        ]
                        contest_count = len(contests)
                        geographies = [
                            {"geo_id": row[0], "geo_type": row[1], "geo_name": row[2], "parent_geo_id": row[3]}
                            for row in connection.execute(
                                "SELECT geo_id, geo_type, geo_name, parent_geo_id FROM dim_geo"
                            ).fetchall()
                        ]
                        assert population_source is not None and evidence_root is not None
                        population_rows, crosswalks = load_hcp_rgph2014_individuals(
                            evidence_root / population_source["raw_path"],
                            geographies,
                            source_id=population_source["source_id"],
                            source_url=population_source["source_url"],
                        )
                        _insert_rows(connection, "fact_geo_population", population_rows)
                        trusted_ids = {row["geo_id"] for row in crosswalks if row["confidence"] == 1.0}
                        populations_by_geo = {
                            row["normalized_geo_id"]: row for row in population_rows if row["normalized_geo_id"] in trusted_ids
                        }
                        geography_types = {row["geo_id"]: row["geo_type"] for row in geographies}
                        contest_links = derive_contest_legal_regime_links(
                            contests,
                            legal_payload,
                            populations_by_geo=populations_by_geo,
                            geography_types=geography_types,
                        )
                        _insert_rows(connection, "bridge_geo_official_identifier", crosswalks)
                        _insert_rows(connection, "bridge_contest_legal_regime", contest_links)
            if evidence_root is not None and source_registry is not None and crosswalks:
                communal_source_identifiers = _materialize_communal_source_identifiers(
                    connection, evidence_root, source_registry,
                )
            if evidence_root is not None:
                result_history_report = build_result_history_diagnostic(
                    connection, evidence_root, as_of_date,
                )
            create_analytical_views(connection)
            _validate_provenance_references(connection)
            connection.execute("CHECKPOINT")
        finally:
            connection.close()
        os.replace(temporary, output_database)
    return {
        "source_database": str(seed_database),
        "output_database": str(output_database),
        "copied_historical_tables": len(source_tables),
        "normalized_historical_source_references": normalized_source_references,
        "normalized_seed_identifiers": normalized_seed_identifiers,
        "normalized_historical_text": normalized_historical_text,
        "created_canonical_tables": len(TABLE_CONTRACTS),
        "seeded_legal_regimes": len(legal_payload["legal_regimes"]) if legal_payload is not None else 0,
        "seeded_election_legal_links": len(legal_payload["election_links"]) if legal_payload is not None else 0,
        "seeded_contest_legal_links": len(contest_links),
        "seeded_geo_populations": len(population_rows),
        "seeded_geo_crosswalks": len(crosswalks),
        "seeded_geo_parent_relations": len(geo_parent_rows),
        "geo_parent_evidence": geo_parent_details,
        "geo_parent_report": parent_report,
        "seeded_result_reconciliations": len(reconciliation_rows),
        "reconciliation_report": reconciliation_report,
        "result_history_report": result_history_report,
        "communal_source_identifiers": communal_source_identifiers,
        "as_of_date": as_of_date,
        "unmapped_contests": contest_count - len(contest_links),
        "publication_status": "NOT_PUBLICATION_READY",
    }
