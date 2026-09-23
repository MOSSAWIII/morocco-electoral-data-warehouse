"""Reconcile reviewed BO 6374 candidates with TAFRA's 2015 workbooks.

The HCP RGPH 2014 workbook is used only as a bilingual, coded bridge.  TAFRA
remains a secondary source and no candidate transcription becomes an official
universe through this module.
"""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable, Mapping

import openpyxl

from morocco_elections.warehouse.bo6374 import validate_visual_candidate
from morocco_elections.warehouse.sources import load_source_registry
from morocco_elections.warehouse.validation import ValidationIssue


BO_SOURCE_ID = "MA_SGG_BO_6374_DECREE_2_15_402_COMMUNES_2015"
HCP_SOURCE_ID = "MA_HCP_RGPH2014_POPULATION_LEGALE_COMMUNES_12_REGIONS"
LAW_SOURCE_ID = "MA_SGG_LO_34_15"
ARRONDISSEMENT_ANNEX_SOURCE_ID = "MA_SGG_BO_6381_DECREE_2_15_577_ARRONDISSEMENTS_2015"
TAFRA_RESULTS_SOURCE_ID = "SRC_TAFRA_COMM2015"
TAFRA_ELECTED_SOURCE_ID = "TAFRA_COUNCILS_2015_V1_0_0"

RESULTS_PATH = Path("data/raw/tafra/communal_results/2015/communes-elections-2015-1-0.xlsx")
ELECTED_PATH = Path("data/staging/elected-2015/communes-elus-2015-1-0.xlsx")
HCP_PATH = Path("data/raw/elections/warehouse/demography/hcp_rgph2014_population_legale_12_regions.xlsx")
DEFAULT_ARTIFACT_PATH = Path("data/exports/open/warehouse/bo6374_tafra_2015_reconciliation.json")
DEFAULT_REPORT_PATH = Path("data/exports/open/warehouse/bo6374_tafra_2015_reconciliation_report.json")

EXPECTED_RESULTS_SHA256 = "2b63ad67ab4817bbd86a7faa06d29f3ff6a0fe5d1d609d343121e2b680722c35"
EXPECTED_ELECTED_SHA256 = "8b5e6d23756087409a77fbaa00c0e25428c699245dbc3241d280b914600fd341"

_HCP_PROVINCE = re.compile(r"^\d{2}\.\d{3}\.$")
_HCP_UNIT = re.compile(r"^\d{2}\.\d{3}\.\d{2}\.\d{2}\.$")
_ARABIC_TRANSLATION = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه", "ؤ": "و", "ئ": "ي"})


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def repair_mojibake(value: Any) -> str:
    text = str(value or "")
    try:
        return text.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return text


def normalize_arabic(value: Any) -> str:
    text = repair_mojibake(value).replace("ـ", "")
    text = re.sub(r"\s*\([^)]*\)\s*$", "", text)
    if ":" in text:
        text = text.split(":", 1)[1]
    text = text.translate(_ARABIC_TRANSLATION)
    return re.sub(r"[^\u0600-\u06ff0-9]", "", text)


def normalize_latin(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).casefold()
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = re.sub(r"\s*\((?:mun\.|arrond\.)\)\s*$", "", text)
    text = re.sub(r"^(?:province|prefecture)\s*:\s*", "", text)
    return re.sub(r"[^a-z0-9]", "", text)


def _similarity(left: Any, right: Any) -> float:
    return SequenceMatcher(None, normalize_arabic(left), normalize_arabic(right)).ratio()


def _official_source(registry: Mapping[str, Any], source_id: str) -> Mapping[str, Any]:
    return next((row for row in registry.get("sources", []) if row.get("source_id") == source_id), {})


def _verified_descriptor(root: Path, source: Mapping[str, Any]) -> dict[str, Any]:
    relative = Path(str(source["raw_path"]))
    path = root / relative
    observed = sha256_file(path)
    if observed != source.get("sha256") or path.stat().st_size != source.get("bytes"):
        raise ValueError(f"source bytes differ from registry: {source.get('source_id')}")
    return {
        "source_id": source["source_id"],
        "role": "PRIMARY_OFFICIAL" if source["source_id"] != HCP_SOURCE_ID else "BILINGUAL_IDENTIFIER_BRIDGE",
        "path": relative.as_posix(),
        "source_url": source["source_url"],
        "bytes": path.stat().st_size,
        "sha256": observed,
    }


def _load_candidates(root: Path, bo_source: Mapping[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for path in sorted((root / "metadata/warehouse").glob("bo6374_page*_visual_transcription.candidate.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        issues = validate_visual_candidate(payload, bo_source)
        if issues:
            raise ValueError(f"invalid visual candidate {path.name}: {'; '.join(issues)}")
        for group in payload["groups"]:
            for row in group["rows"]:
                rows.append({
                    "candidate_id": f"BO6374:P{payload['printed_page']}:R{row['row']:03d}",
                    "printed_page": payload["printed_page"],
                    "pdf_page_index": payload["pdf_page_index"],
                    "page_row": row["row"],
                    "bo_province_prefecture_ar": repair_mojibake(group["prefecture_province_ar"]),
                    "bo_commune_name_ar": repair_mojibake(row["commune_name_ar"]),
                    "bo_council_members": row["council_members"],
                })
    if len(rows) != 1255:
        raise ValueError(f"expected 1255 BO candidate rows, observed {len(rows)}")
    return rows


def _load_hcp_bridge(path: Path) -> tuple[dict[str, dict[str, str]], list[dict[str, str]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    provinces: dict[str, dict[str, str]] = {}
    units: list[dict[str, str]] = []
    try:
        for values in workbook["Communes"].iter_rows(min_row=6, values_only=True):
            code = str(values[0] or "")
            if _HCP_PROVINCE.fullmatch(code):
                provinces[code[:6]] = {"province_code": code, "name_latin": str(values[1]), "name_ar": str(values[6])}
            elif _HCP_UNIT.fullmatch(code):
                units.append({
                    "official_geo_code": code,
                    "province_key": code[:6],
                    "name_latin": str(values[1]),
                    "name_ar": str(values[6]),
                })
    finally:
        workbook.close()
    if len(units) != 1538:
        raise ValueError(f"expected 1538 HCP coded units, observed {len(units)}")
    return provinces, units


def _sheet_records(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        values = workbook.worksheets[0].iter_rows(values_only=True)
        headers = next(values)
        return [dict(zip(headers, row)) for row in values]
    finally:
        workbook.close()


def _load_tafra(
    results_path: Path, elected_path: Path
) -> tuple[list[dict[str, Any]], Counter[int], dict[int, dict[str, str]]]:
    results = _sheet_records(results_path)
    elected = _sheet_records(elected_path)
    if len(results) != 1538 or len(elected) != 31482:
        raise ValueError(f"unexpected TAFRA shapes: results={len(results)}, elected={len(elected)}")
    ids = [int(row["idCommune"]) for row in results]
    if len(ids) != len(set(ids)):
        raise ValueError("TAFRA result idCommune is not unique")
    elected_counts = Counter(int(row["idCommune"]) for row in elected)
    if set(elected_counts) != set(ids):
        raise ValueError("TAFRA elected and result commune universes differ")
    elected_names: dict[int, dict[str, str]] = {}
    for row in elected:
        source_id = int(row["idCommune"])
        names = {"province": str(row["prefProv"]), "commune": str(row["commune"])}
        if source_id in elected_names and elected_names[source_id] != names:
            raise ValueError(f"TAFRA elected names vary within idCommune {source_id}")
        elected_names[source_id] = names
    return results, elected_counts, elected_names


def _unit_type(name: str) -> str:
    if "(Arrond.)" in name:
        return "ARRONDISSEMENT"
    if "(Mun.)" in name:
        return "MUNICIPALITY"
    return "COMMUNE"


def _source_descriptors(root: Path, registry: Mapping[str, Any]) -> list[dict[str, Any]]:
    descriptors = [
        _verified_descriptor(root, _official_source(registry, BO_SOURCE_ID)),
        _verified_descriptor(root, _official_source(registry, HCP_SOURCE_ID)),
        _verified_descriptor(root, _official_source(registry, LAW_SOURCE_ID)),
        _verified_descriptor(root, _official_source(registry, ARRONDISSEMENT_ANNEX_SOURCE_ID)),
    ]
    secondary = (
        (TAFRA_RESULTS_SOURCE_ID, RESULTS_PATH, EXPECTED_RESULTS_SHA256, "SECONDARY_RESULTS", "https://open.africa/dataset/cecd6468-53f8-4fdb-8269-f1a51c029ac3/resource/6d21d289-41d2-4f62-8719-e2a4b3aae055/download/communes-elections-2015-1-0.xlsx"),
        (TAFRA_ELECTED_SOURCE_ID, ELECTED_PATH, EXPECTED_ELECTED_SHA256, "SECONDARY_ELECTED_MEMBERS_CANDIDATE", "https://open.africa/dataset/07a04224-c0ad-4861-9705-0518f5d49dbd/resource/7ae81ece-1b3d-4cdc-ac49-acd6ba37f6ea/download/communes-elus-2015-1-0.xlsx"),
    )
    for source_id, relative, expected_hash, role, url in secondary:
        path = root / relative
        observed = sha256_file(path)
        if observed != expected_hash:
            raise ValueError(f"source bytes differ from pinned SHA-256: {source_id}")
        descriptors.append({
            "source_id": source_id,
            "role": role,
            "path": relative.as_posix(),
            "source_url": url,
            "bytes": path.stat().st_size,
            "sha256": observed,
        })
    return descriptors


def _match_rows(
    candidates: list[dict[str, Any]],
    provinces: Mapping[str, Mapping[str, str]],
    units: list[dict[str, str]],
) -> tuple[dict[int, tuple[dict[str, str], float, str]], dict[int, list[dict[str, Any]]]]:
    units_by_province: dict[str, list[dict[str, str]]] = defaultdict(list)
    for unit in units:
        units_by_province[unit["province_key"]].append(unit)

    for candidate in candidates:
        ranked = sorted(
            ((_similarity(candidate["bo_province_prefecture_ar"], province["name_ar"]), key) for key, province in provinces.items()),
            reverse=True,
        )
        if ranked[0][0] < 0.80 or ranked[0][0] - ranked[1][0] < 0.08:
            raise ValueError(f"ambiguous BO province: {candidate['bo_province_prefecture_ar']}")
        candidate["_province_key"] = ranked[0][1]

    matched: dict[int, tuple[dict[str, str], float, str]] = {}
    used: dict[str, set[str]] = defaultdict(set)
    for index, candidate in enumerate(candidates):
        exact = [
            unit for unit in units_by_province[candidate["_province_key"]]
            if normalize_arabic(unit["name_ar"]) == normalize_arabic(candidate["bo_commune_name_ar"])
        ]
        if len(exact) == 1 and exact[0]["official_geo_code"] not in used[candidate["_province_key"]]:
            matched[index] = (exact[0], 1.0, "HCP_BILINGUAL_EXACT")
            used[candidate["_province_key"]].add(exact[0]["official_geo_code"])

    # Recompute after each pass: exact assignments eliminate false fuzzy rivals.
    for minimum_score in (0.70, 0.55):
        while True:
            options: list[tuple[float, float, int, dict[str, str]]] = []
            for index, candidate in enumerate(candidates):
                if index in matched:
                    continue
                available = [
                    unit for unit in units_by_province[candidate["_province_key"]]
                    if unit["official_geo_code"] not in used[candidate["_province_key"]]
                ]
                ranked = sorted(
                    ((_similarity(candidate["bo_commune_name_ar"], unit["name_ar"]), unit) for unit in available),
                    key=lambda item: item[0],
                    reverse=True,
                )
                if not ranked:
                    continue
                if (
                    normalize_arabic(candidate["bo_commune_name_ar"])
                    == normalize_arabic(candidate["bo_province_prefecture_ar"])
                    and ranked[0][0] < 0.70
                ):
                    continue  # A city council cannot be forced onto one arrondissement.
                margin = ranked[0][0] - (ranked[1][0] if len(ranked) > 1 else 0.0)
                if ranked[0][0] >= minimum_score and margin >= 0.08:
                    options.append((ranked[0][0], margin, index, ranked[0][1]))
            if not options:
                break
            progress = False
            for score, _margin, index, unit in sorted(options, reverse=True):
                province_key = candidates[index]["_province_key"]
                if index in matched or unit["official_geo_code"] in used[province_key]:
                    continue
                matched[index] = (unit, score, "HCP_BILINGUAL_FUZZY_UNIQUE")
                used[province_key].add(unit["official_geo_code"])
                progress = True
            if not progress:
                break

    alternatives: dict[int, list[dict[str, Any]]] = {}
    for index, candidate in enumerate(candidates):
        if index in matched:
            continue
        ranked = sorted(
            ((_similarity(candidate["bo_commune_name_ar"], unit["name_ar"]), unit) for unit in units_by_province[candidate["_province_key"]]),
            key=lambda item: item[0],
            reverse=True,
        )
        alternatives[index] = [
            {"official_geo_code": unit["official_geo_code"], "hcp_name_ar": unit["name_ar"], "hcp_name_latin": unit["name_latin"], "name_score": round(score, 6)}
            for score, unit in ranked[:5]
        ]
    return matched, alternatives


def _distribution(rows: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    counts: Counter[tuple[Any, ...]] = Counter()
    for row in rows:
        counts[(row["commune_type"], row["tafra_province_prefecture"], row["seat_difference"])] += 1
    return [
        {"commune_type": key[0], "province_prefecture": key[1], "seat_difference": key[2], "count": count}
        for key, count in sorted(counts.items(), key=lambda item: (str(item[0][0]), str(item[0][1]), str(item[0][2])))
    ]


def build_reconciliation(root: Path) -> dict[str, Any]:
    root = root.resolve()
    registry = load_source_registry(root)
    bo_source = _official_source(registry, BO_SOURCE_ID)
    candidates = _load_candidates(root, bo_source)
    provinces, units = _load_hcp_bridge(root / HCP_PATH)
    tafra_rows, elected_counts, elected_names = _load_tafra(root / RESULTS_PATH, root / ELECTED_PATH)
    descriptors = _source_descriptors(root, registry)
    parent_seed = json.loads((root / "metadata/warehouse/comm2015_geo_parents.seed.json").read_text(encoding="utf-8"))
    groups_by_city = {normalize_latin(row["parent_name"]): row for row in parent_seed["groups"]}
    city_key_by_arabic = {
        normalize_arabic("طنجة"): "tanger",
        normalize_arabic("فاس"): "fes",
        normalize_arabic("الرباط"): "rabat",
        normalize_arabic("سلا"): "sale",
        normalize_arabic("الدار البيضاء"): "casablanca",
        normalize_arabic("مراكش"): "marrakech",
    }

    tafra_by_name = {
        (normalize_latin(row["prefProv"]), normalize_latin(row["commune"])): row
        for row in tafra_rows
    }
    for unit in units:
        province = provinces[unit["province_key"]]
        key = (normalize_latin(province["name_latin"]), normalize_latin(unit["name_latin"]))
        if key not in tafra_by_name:
            raise ValueError(f"HCP unit does not map exactly to TAFRA: {unit['official_geo_code']}")
        unit["_tafra"] = tafra_by_name[key]  # type: ignore[assignment]

    matched, alternatives = _match_rows(candidates, provinces, units)
    rows: list[dict[str, Any]] = []
    for index, candidate in enumerate(candidates):
        base = {key: value for key, value in candidate.items() if not key.startswith("_")}
        if index not in matched:
            group = groups_by_city.get(city_key_by_arabic.get(normalize_arabic(candidate["bo_commune_name_ar"]), ""))
            if group is None:
                raise ValueError(f"unmatched large-city council has no verified WAREHOUSE parent: {candidate['candidate_id']}")
            rows.append({
                **base,
                "tafra_commune_id": None,
                "tafra_province_prefecture": None,
                "tafra_commune_name": None,
                "hcp_official_geo_code": None,
                "commune_type": "COMMUNE_DIVIDED_INTO_ARRONDISSEMENTS",
                "tafra_nSieges": None,
                "observed_elected_members": None,
                "seat_difference": None,
                "matching_method": "NO_EQUIVALENT_TAFRA_GRAIN",
                "confidence": 1.0,
                "status": "UNMATCHED",
                "explanation_code": "CITY_COUNCIL_VS_TAFRA_ARRONDISSEMENTS",
                "warehouse_communal_parent_id": group["parent_geo_id"],
                "parent_relation_status": "VERIFIED_PARENT_ONLY",
                "verified_arrondissement_count": group["annex_count"],
                "remaining_blocker": "TAFRA_HAS_ARRONDISSEMENT_RESULTS_BUT_NO_EQUIVALENT_CITY_COUNCIL_RESULT_ROW",
                "candidate_alternatives": alternatives[index],
                "source_ids": [BO_SOURCE_ID, ARRONDISSEMENT_ANNEX_SOURCE_ID, HCP_SOURCE_ID, TAFRA_RESULTS_SOURCE_ID],
            })
            continue
        unit, confidence, method = matched[index]
        tafra = unit["_tafra"]
        tafra_id = int(tafra["idCommune"])
        tafra_names = elected_names[tafra_id]
        tafra_seats = int(tafra["nSieges"])
        observed = elected_counts[tafra_id]
        difference = tafra_seats - candidate["bo_council_members"]
        if difference == 0 and observed == tafra_seats:
            status, explanation = "MATCH", "ALL_COUNTS_EQUAL"
        elif difference == 4 and observed == tafra_seats:
            status, explanation = "EXPLAINED_DIFFERENCE", "LO_34_15_ART_128_BIS_FOUR_ADDITIONAL_WOMEN_SEATS"
        else:
            status, explanation = "AMBIGUOUS", "UNEXPLAINED_SOURCE_COUNT_CONFLICT"
        rows.append({
            **base,
            "tafra_commune_id": tafra_id,
            "tafra_province_prefecture": tafra_names["province"],
            "tafra_commune_name": tafra_names["commune"],
            "hcp_official_geo_code": unit["official_geo_code"],
            "commune_type": _unit_type(unit["name_latin"]),
            "tafra_nSieges": tafra_seats,
            "observed_elected_members": observed,
            "seat_difference": difference,
            "matching_method": method,
            "confidence": round(confidence, 6),
            "status": status,
            "explanation_code": explanation,
            "candidate_alternatives": [],
            "source_ids": [BO_SOURCE_ID, HCP_SOURCE_ID, TAFRA_RESULTS_SOURCE_ID, TAFRA_ELECTED_SOURCE_ID]
            + ([LAW_SOURCE_ID] if status == "EXPLAINED_DIFFERENCE" else []),
        })

    status_counts = Counter(row["status"] for row in rows)
    method_counts = Counter(row["matching_method"] for row in rows)
    report = {
        "candidate_rows": len(rows),
        "identity_matches": sum(row["tafra_commune_id"] is not None for row in rows),
        "exact_name_matches": method_counts["HCP_BILINGUAL_EXACT"],
        "fuzzy_unique_name_matches": method_counts["HCP_BILINGUAL_FUZZY_UNIQUE"],
        "exact_count_matches": status_counts["MATCH"],
        "explained_differences": status_counts["EXPLAINED_DIFFERENCE"],
        "ambiguous": status_counts["AMBIGUOUS"],
        "unmatched": status_counts["UNMATCHED"],
        "unmatched_with_verified_parent": sum(
            row.get("parent_relation_status") == "VERIFIED_PARENT_ONLY" for row in rows
        ),
        "status_distribution": dict(sorted(status_counts.items())),
        "matching_method_distribution": dict(sorted(method_counts.items())),
        "seat_difference_distribution": _distribution(row for row in rows if row["tafra_commune_id"] is not None),
        "legal_explanation": {
            "rule": "For COMM2015, Organic Law 34-15 added Article 128 bis to Organic Law 59-11; the four additional women seats explain only rows where TAFRA and the elected-member count both equal BO + 4.",
            "source_id": LAW_SOURCE_ID,
            "applicability": "2015_ONLY",
            "non_rule_differences_remain_conflicts": True,
        },
        "limitations": [
            "BO rows remain single-review visual candidates and are not promoted to an official universe.",
            "TAFRA data are secondary copies attributed to elections.ma and are not treated as the legal source.",
            "Large city council rows are not forced onto TAFRA arrondissement rows.",
        ],
    }
    return {"schema_version": 1, "scope": "BO6374_PAGES_6105_6130_CANDIDATE_ONLY", "sources": descriptors, "rows": rows, "report": report}


def validate_reconciliation(payload: Mapping[str, Any], root: Path) -> list[ValidationIssue]:
    """Recompute the artifact and report field-level mutations."""
    issues: list[ValidationIssue] = []
    try:
        expected = build_reconciliation(root)
    except (OSError, KeyError, TypeError, ValueError) as exc:
        return [ValidationIssue("BO_TAFRA_RECONCILIATION_INPUT_INVALID", "bo_tafra_reconciliation", "WAREHOUSE", str(exc))]
    if payload.get("schema_version") != expected["schema_version"] or payload.get("scope") != expected["scope"]:
        issues.append(ValidationIssue("BO_TAFRA_RECONCILIATION_HEADER_MUTATION", "bo_tafra_reconciliation", "WAREHOUSE", "schema_version or scope differs"))
    if payload.get("sources") != expected["sources"]:
        issues.append(ValidationIssue("BO_TAFRA_RECONCILIATION_SOURCE_MUTATION", "bo_tafra_reconciliation", "WAREHOUSE", "source fingerprints or identifiers differ"))
    actual_rows = {row.get("candidate_id"): row for row in payload.get("rows", []) if isinstance(row, Mapping)}
    expected_rows = {row["candidate_id"]: row for row in expected["rows"]}
    if set(actual_rows) != set(expected_rows):
        issues.append(ValidationIssue("BO_TAFRA_RECONCILIATION_ROW_SET_MUTATION", "bo_tafra_reconciliation", "WAREHOUSE", "candidate row identifiers differ"))
    protected = ("tafra_commune_id", "bo_council_members", "tafra_nSieges", "observed_elected_members", "source_ids")
    for candidate_id in sorted(set(actual_rows) & set(expected_rows)):
        actual, wanted = actual_rows[candidate_id], expected_rows[candidate_id]
        changed = [field for field in protected if actual.get(field) != wanted.get(field)]
        if changed:
            issues.append(ValidationIssue("BO_TAFRA_RECONCILIATION_PROTECTED_VALUE_MUTATION", "bo_tafra_reconciliation", candidate_id, ", ".join(changed)))
        elif actual != wanted:
            issues.append(ValidationIssue("BO_TAFRA_RECONCILIATION_ROW_MUTATION", "bo_tafra_reconciliation", candidate_id, "derived row differs from recomputation"))
    if payload.get("report") != expected["report"]:
        issues.append(ValidationIssue("BO_TAFRA_RECONCILIATION_REPORT_MUTATION", "bo_tafra_reconciliation", "WAREHOUSE", "report differs from recomputation"))
    return issues


def write_reconciliation(root: Path, artifact_path: Path, report_path: Path) -> dict[str, Any]:
    payload = build_reconciliation(root)
    artifact_path.parent.mkdir(parents=True, exist_ok=True)
    artifact_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(payload["report"], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return payload
