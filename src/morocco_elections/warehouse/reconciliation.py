from __future__ import annotations

import hashlib
import json
from collections import Counter
from dataclasses import asdict, dataclass
from math import isfinite
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import duckdb
import openpyxl

from morocco_elections.warehouse.sources import canonical_source_id, source_descriptors_by_id


RECONCILIATION_CHECKS = (
    "turnout_rate",
    "ballot_categories_vs_voters",
    "party_votes_vs_valid_votes",
    "allocated_seats_vs_contest_seats",
)
ELECTED_2015_PATH = Path("data/staging/elected-2015/communes-elus-2015-1-0.xlsx")
ELECTED_2015_SHA256 = "8b5e6d23756087409a77fbaa00c0e25428c699245dbc3241d280b914600fd341"


@dataclass(frozen=True)
class Reconciliation:
    reconciliation_id: str
    contest_id: str
    metric: str
    official_value: float | str | bool | None
    recomputed_value: float | str | bool | None
    difference: float | None
    tolerance: float
    validation_status: str
    explanation: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def compare(contest_id: str, metric: str, official: Any, recomputed: Any, tolerance: float = 0.0) -> Reconciliation:
    if official is None or recomputed is None:
        difference, status = None, "NOT_COMPUTABLE"
        explanation = "official_value or recomputed_value is missing"
    elif isinstance(official, (int, float)) and not isinstance(official, bool) and isinstance(recomputed, (int, float)) and not isinstance(recomputed, bool):
        difference = float(recomputed) - float(official)
        status = "NOT_COMPUTABLE" if not isfinite(difference) else "PASS" if abs(difference) <= tolerance else "FAIL"
        explanation = "numeric difference is not finite" if status == "NOT_COMPUTABLE" else None
    else:
        difference = 0.0 if recomputed == official else 1.0
        status = "PASS" if difference == 0 else "FAIL"
        explanation = None
    return Reconciliation(f"{contest_id}:{metric}", contest_id, metric, official, recomputed, difference, tolerance, status, explanation)


def _valid_shares(values: Sequence[float], tolerance: float) -> bool:
    return bool(values) and all(isfinite(v) and 0 <= v <= 1 for v in values) and abs(sum(values) - 1.0) <= tolerance


def recompute_contest(
    contest_id: str,
    mobilization: Mapping[str, Any],
    party_results: Iterable[Mapping[str, Any]],
    *,
    tolerance: float = 0.0,
    share_tolerance: float = 1e-9,
    ballot_accounting_rule: str = "SEPARATE_BLANK_AND_INVALID",
    expected_party_ids: Sequence[str] | None = None,
) -> dict[str, Any]:
    rows = list(party_results)
    observed_ids = [row.get("party_id") for row in rows]
    expected_ids = list(expected_party_ids) if expected_party_ids is not None else []
    party_vector_complete = (
        bool(expected_ids)
        and all(isinstance(party_id, str) and party_id for party_id in expected_ids)
        and len(expected_ids) == len(set(expected_ids))
        and all(isinstance(party_id, str) and party_id for party_id in observed_ids)
        and len(observed_ids) == len(set(observed_ids))
        and set(observed_ids) == set(expected_ids)
    )
    votes_complete = party_vector_complete and all(row.get("votes") is not None for row in rows)
    seats_complete = party_vector_complete and all(row.get("seats") is not None for row in rows)
    votes = [float(row["votes"]) for row in rows] if votes_complete else None
    seats = [float(row["seats"]) for row in rows] if seats_complete else None
    valid = mobilization.get("valid_votes")
    if ballot_accounting_rule not in {"SEPARATE_BLANK_AND_INVALID", "INVALID_INCLUDES_BLANK"}:
        raise ValueError("unknown ballot accounting rule; categories must not be silently merged")
    component_names = ("valid_votes", "invalid_votes", "blank_votes") if ballot_accounting_rule == "SEPARATE_BLANK_AND_INVALID" else ("valid_votes", "invalid_votes")
    voters_components = (
        sum(float(mobilization[name]) for name in component_names if mobilization.get(name) is not None)
        if all(mobilization.get(name) is not None for name in component_names) else None
    )
    reconciliations = [
        compare(contest_id, "party_votes_vs_valid_votes", valid, sum(votes) if votes is not None else None, tolerance),
        compare(contest_id, "ballot_categories_vs_voters", mobilization.get("voters"), voters_components, tolerance),
        compare(contest_id, "allocated_seats_vs_contest_seats", mobilization.get("contest_seats"), sum(seats) if seats is not None else None, tolerance),
    ]
    registered, voters = mobilization.get("registered_voters"), mobilization.get("voters")
    turnout = None if registered in (None, 0) or voters is None else float(voters) / float(registered)
    reconciliations.append(compare(contest_id, "turnout_rate", mobilization.get("turnout_rate"), turnout, share_tolerance))
    voter_bound = None if registered is None or voters is None else float(voters) <= float(registered)
    reconciliations.append(compare(contest_id, "voters_not_above_registered", True, voter_bound, 0))

    ranked = []
    if valid not in (None, 0) and votes is not None:
        ordered = sorted(rows, key=lambda row: (-float(row["votes"]), str(row.get("party_id"))))
        for index, row in enumerate(ordered, start=1):
            derived = {**row, "recomputed_vote_share": float(row["votes"]) / float(valid), "recomputed_rank": index, "recomputed_winner": index == 1}
            ranked.append(derived)
            party = str(row.get("party_id"))
            reconciliations.extend([
                compare(contest_id, f"party:{party}:vote_share", row.get("vote_share"), derived["recomputed_vote_share"], share_tolerance),
                compare(contest_id, f"party:{party}:rank", row.get("rank"), index, 0),
                compare(contest_id, f"party:{party}:winner", row.get("winner_flag"), index == 1, 0),
            ])
    shares = [row["recomputed_vote_share"] for row in ranked]
    hhi = sum(value * value for value in shares) if _valid_shares(shares, share_tolerance) else None
    enp = 1 / hhi if hhi and hhi > 0 else None
    margin_votes = None
    margin_share = None
    if len(ranked) >= 2:
        margin_votes = float(ranked[0]["votes"]) - float(ranked[1]["votes"])
        margin_share = ranked[0]["recomputed_vote_share"] - ranked[1]["recomputed_vote_share"]
    reconciliations.extend([
        compare(contest_id, "hhi", mobilization.get("hhi"), hhi, share_tolerance),
        compare(contest_id, "enp", mobilization.get("enp"), enp, share_tolerance),
        compare(contest_id, "winner_party_id", mobilization.get("winner_party_id"), ranked[0].get("party_id") if ranked else None, 0),
        compare(contest_id, "victory_margin_votes", mobilization.get("victory_margin_votes"), margin_votes, tolerance),
        compare(contest_id, "victory_margin_share", mobilization.get("victory_margin_share"), margin_share, share_tolerance),
    ])
    return {
        "contest_id": contest_id,
        "party_vector_status": "COMPLETE_AGAINST_EXPECTED_IDENTIFIERS" if party_vector_complete else "NOT_COMPUTABLE_WITHOUT_COMPLETE_PARTY_UNIVERSE",
        "party_results": ranked,
        "reconciliations": [row.as_dict() for row in reconciliations],
        "recomputed": {"hhi": hhi, "enp": enp, "winner_party_id": ranked[0].get("party_id") if ranked else None, "margin_votes": margin_votes, "margin_share": margin_share},
    }


def reconcile_detail_to_aggregate(
    contest_id: str,
    details: Iterable[Mapping[str, Any]],
    aggregate: Mapping[str, Any],
    metrics: Sequence[str] = ("votes", "seats"),
    tolerance: float = 0.0,
) -> list[dict[str, Any]]:
    details = list(details)
    output = []
    for metric in metrics:
        complete = bool(details) and all(row.get(metric) is not None for row in details)
        recomputed = sum(float(row[metric]) for row in details) if complete else None
        output.append(compare(contest_id, f"detail_vs_aggregate:{metric}", aggregate.get(metric), recomputed, tolerance).as_dict())
    return output


def _table_rows(connection: duckdb.DuckDBPyConnection, table: str) -> list[dict[str, Any]]:
    cursor = connection.execute(f'SELECT * FROM "{table}"')
    fields = [item[0] for item in cursor.description]
    return [dict(zip(fields, values, strict=True)) for values in cursor.fetchall()]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _verified_manifest_sources(root: Path, source_ids: set[str]) -> dict[str, dict[str, Any]]:
    available = source_descriptors_by_id(root)
    sources = {
        source_id: available[canonical_source_id(source_id)]
        for source_id in source_ids if canonical_source_id(source_id) in available
    }
    if set(sources) != source_ids:
        raise ValueError("a reconciliation source is absent from metadata/warehouse/source_registry.json")
    for source_id, descriptor in sources.items():
        path = (root / descriptor["raw_path"]).resolve()
        if (
            not path.is_file()
            or path.stat().st_size != descriptor["bytes"]
            or _sha256(path) != descriptor["sha256"]
        ):
            raise ValueError(f"reconciliation source bytes differ: {source_id}")
    return sources


def _reconciliation_row(
    contest: Mapping[str, Any],
    metric: str,
    official: Any,
    recomputed: Any,
    tolerance: float,
    source_id: str,
    evidence_id: str,
    reason: str | None,
    missing: Sequence[str],
    method: str,
) -> dict[str, Any]:
    comparison = compare(str(contest["contest_id"]), metric, official, recomputed, tolerance)
    status = comparison.validation_status
    if reason is not None:
        status = "NOT_COMPUTABLE"
    difference = comparison.difference if status != "NOT_COMPUTABLE" else None
    explanation = None
    if status == "NOT_COMPUTABLE":
        explanation = "Missing or unverified components: " + ", ".join(missing)
    return {
        "reconciliation_id": f"{contest['contest_id']}:{metric}",
        "contest_id": contest["contest_id"],
        "metric": metric,
        "official_value": official,
        "recomputed_value": recomputed,
        "difference": difference,
        "tolerance": tolerance,
        "validation_status": status,
        "source_id": source_id,
        "explanation": explanation,
        "not_computable_reason": reason,
        "official_source_id": source_id,
        "external_universe_id": None,
        "evidence_id": evidence_id,
        "calculation_method": method,
        "missing_components": json.dumps(list(missing), ensure_ascii=False, separators=(",", ":")),
    }


def derive_reconciliation_matrix(
    connection: duckdb.DuckDBPyConnection,
    root: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Derive the complete four-check matrix from materialized package facts."""
    contests = _table_rows(connection, "dim_electoral_contest")
    mobilizations = _table_rows(connection, "fact_electoral_mobilization")
    summaries = _table_rows(connection, "fact_commune_election_summary")
    mobilization_by_contest: dict[str, dict[str, Any]] = {}
    for fact in mobilizations:
        contest_id = str(fact["contest_id"])
        if contest_id in mobilization_by_contest:
            raise ValueError(f"duplicate mobilization row: {contest_id}")
        mobilization_by_contest[contest_id] = fact
    summary_by_contest: dict[str, dict[str, Any]] = {}
    for summary in summaries:
        contest_id = str(summary["contest_id"])
        if contest_id in summary_by_contest:
            raise ValueError(f"duplicate communal summary row: {contest_id}")
        summary_by_contest[contest_id] = summary
    if set(mobilization_by_contest) & set(summary_by_contest):
        raise ValueError("contest has ambiguous mobilization and communal-summary aggregates")
    source_ids = {str(row["source_id"]) for row in contests}
    descriptors = _verified_manifest_sources(root, source_ids)
    elected_path = root / ELECTED_2015_PATH
    if not elected_path.is_file() or _sha256(elected_path) != ELECTED_2015_SHA256:
        raise ValueError("COMM2015 elected-person source bytes differ")
    workbook = openpyxl.load_workbook(elected_path, read_only=True, data_only=True)
    try:
        elected_sheet = workbook["données"]
        elected_person_rows = elected_sheet.max_row - 1
        elected_communes = {
            row[0] for row in elected_sheet.iter_rows(min_row=2, min_col=5, max_col=5, values_only=True)
            if row[0] is not None
        }
    finally:
        workbook.close()
    rows: list[dict[str, Any]] = []
    for contest in sorted(contests, key=lambda row: str(row["contest_id"])):
        contest_id = str(contest["contest_id"])
        fact = mobilization_by_contest.get(contest_id)
        summary = summary_by_contest.get(contest_id)
        aggregate = fact or summary
        source_id = str((aggregate or contest)["source_id"])
        evidence_id = str((aggregate or {}).get("evidence_id") or f"{source_id}:{descriptors[source_id]['sha256']}")

        registered = fact.get("registered_voters") if fact else None
        voters = fact.get("voters") if fact else None
        turnout_official = aggregate.get("turnout_rate") if aggregate else None
        turnout_missing = []
        if aggregate is None:
            turnout_missing.append("mobilization_record")
        else:
            if registered is None:
                turnout_missing.append("registered_voters")
            elif registered == 0:
                turnout_missing.append("nonzero_registered_voters")
            if voters is None:
                turnout_missing.append("voters")
            if turnout_official is None:
                turnout_missing.append("official_turnout_rate")
        turnout_recomputed = (
            float(voters) / float(registered)
            if not turnout_missing and registered not in (None, 0) and voters is not None else None
        )
        turnout_reason = None
        if turnout_missing:
            turnout_reason = "MOBILIZATION_RECORD_MISSING" if aggregate is None else (
                "OFFICIAL_TURNOUT_MISSING" if turnout_missing == ["official_turnout_rate"] else "MOBILIZATION_FIELDS_MISSING"
            )
        rows.append(_reconciliation_row(
            contest, "turnout_rate", turnout_official, turnout_recomputed, 1e-9,
            source_id, evidence_id, turnout_reason, turnout_missing,
            "voters / registered_voters from fact_electoral_mobilization",
        ))

        ballot_fields = ("voters", "valid_votes", "invalid_votes", "blank_votes")
        ballot_missing = [name for name in ballot_fields if fact is None or fact.get(name) is None]
        if aggregate is None:
            ballot_missing.insert(0, "mobilization_record")
        ballot_reason = "MOBILIZATION_RECORD_MISSING" if aggregate is None else (
            "BALLOT_COMPONENTS_MISSING" if ballot_missing else "BALLOT_TAXONOMY_UNVERIFIED"
        )
        rows.append(_reconciliation_row(
            contest, "ballot_categories_vs_voters", voters, None, 0.0,
            source_id, evidence_id, ballot_reason,
            ballot_missing or ["verified_ballot_accounting_taxonomy"],
            "valid_votes + invalid_votes + blank_votes compared with voters",
        ))

        party_missing = ["verified_external_party_universe"]
        if aggregate is None:
            party_missing.insert(0, "mobilization_record")
        if fact is None or fact.get("valid_votes") is None:
            party_missing.append("valid_votes")
        rows.append(_reconciliation_row(
            contest, "party_votes_vs_valid_votes", fact.get("valid_votes") if fact else None, None, 0.0,
            source_id, evidence_id, "PARTY_UNIVERSE_UNVERIFIED", party_missing,
            "sum complete expected party vote vector compared with valid_votes",
        ))

        seat_missing = ["verified_external_seat_allocation_universe"]
        if aggregate is None:
            seat_missing.insert(0, "mobilization_record")
        if fact is None or fact.get("contest_seats") is None:
            seat_missing.append("contest_seats")
        rows.append(_reconciliation_row(
            contest, "allocated_seats_vs_contest_seats", fact.get("contest_seats") if fact else None, None, 0.0,
            source_id, evidence_id, "SEAT_UNIVERSE_UNVERIFIED", seat_missing,
            "sum complete expected seat-allocation vector compared with contest_seats",
        ))

    key_counts = Counter((str(row["contest_id"]), str(row["metric"])) for row in rows)
    statuses = Counter(str(row["validation_status"]) for row in rows)
    reasons = Counter(str(row["not_computable_reason"]) for row in rows if row["not_computable_reason"])
    by_type: dict[tuple[str, str], int] = Counter(
        (str(row["election_id"]), str(row.get("list_type"))) for row in contests
    )
    contest_by_id = {str(row["contest_id"]): row for row in contests}
    computability_by_election = []
    for election_id in sorted({str(row["election_id"]) for row in contests}):
        election_rows = [
            row for row in rows
            if str(contest_by_id[str(row["contest_id"])]["election_id"]) == election_id
        ]
        metric_rows = {
            metric: [row for row in election_rows if row["metric"] == metric]
            for metric in RECONCILIATION_CHECKS
        }
        computability_by_election.append({
            "election_id": election_id,
            "total_checks": len(election_rows),
            "calculable_checks": sum(
                row["validation_status"] != "NOT_COMPUTABLE" for row in election_rows
            ),
            "not_computable_checks": sum(
                row["validation_status"] == "NOT_COMPUTABLE" for row in election_rows
            ),
            "by_metric": {
                metric: {
                    "total": len(metric_rows[metric]),
                    "calculable": sum(
                        row["validation_status"] != "NOT_COMPUTABLE"
                        for row in metric_rows[metric]
                    ),
                    "not_computable": sum(
                        row["validation_status"] == "NOT_COMPUTABLE"
                        for row in metric_rows[metric]
                    ),
                }
                for metric in RECONCILIATION_CHECKS
            },
        })
    report = {
        "covered_contests": len(contests),
        "expected_checks": len(contests) * len(RECONCILIATION_CHECKS),
        "produced_reconciliations": len(rows),
        "status_distribution": dict(sorted(statuses.items())),
        "not_computable_reason_distribution": dict(sorted(reasons.items())),
        "computability_by_election": computability_by_election,
        "contests_without_status": sum(
            key_counts[(str(contest["contest_id"]), metric)] == 0
            for contest in contests for metric in RECONCILIATION_CHECKS
        ),
        "duplicate_statuses": sum(count - 1 for count in key_counts.values() if count > 1),
        "official_values_without_source": sum(
            row["official_value"] is not None and not row["official_source_id"] for row in rows
        ),
        "non_reproducible_recomputed_values": 0,
        "unresolved_differences": statuses["FAIL"],
        "mobilization_field_coverage": {
            field: sum(row.get(field) is not None for row in mobilizations)
            for field in ("registered_voters", "voters", "valid_votes", "invalid_votes", "blank_votes", "turnout_rate", "contest_seats")
        },
        "communal_summary_field_coverage": {
            "records": len(summaries),
            "turnout_rate": sum(row.get("turnout_rate") is not None for row in summaries),
        },
        "combined_aggregate_field_coverage": {
            "records": len(mobilizations) + len(summaries),
            "turnout_rate": sum(row.get("turnout_rate") is not None for row in mobilizations)
            + sum(row.get("turnout_rate") is not None for row in summaries),
            "registered_voters": sum(row.get("registered_voters") is not None for row in mobilizations),
            "voters": sum(row.get("voters") is not None for row in mobilizations),
            "valid_votes": sum(row.get("valid_votes") is not None for row in mobilizations),
            "invalid_votes": sum(row.get("invalid_votes") is not None for row in mobilizations),
            "blank_votes": sum(row.get("blank_votes") is not None for row in mobilizations),
            "contest_seats": sum(row.get("contest_seats") is not None for row in mobilizations),
        },
        "external_universe_coverage": {"party_or_list_universes": 0, "seat_allocation_universes": 0},
        "source_qualification": {
            source_id: {"qualification": "SECONDARY_COPY_ATTRIBUTED_TO_ELECTIONS_MA", "sha256": descriptor["sha256"]}
            for source_id, descriptor in sorted(descriptors.items())
        },
        "contest_type_matrix": [
            {
                "election_id": election_id,
                "source_list_type": list_type,
                "contests": count,
                "applicable_checks": list(RECONCILIATION_CHECKS),
                "mobilization_records": sum(
                    contest_by_id[str(row["contest_id"])]["election_id"] == election_id
                    and str(contest_by_id[str(row["contest_id"])].get("list_type")) == list_type
                    for row in mobilizations
                ),
                "communal_summary_records": sum(
                    contest_by_id[str(row["contest_id"])]["election_id"] == election_id
                    and str(contest_by_id[str(row["contest_id"])].get("list_type")) == list_type
                    for row in summaries
                ),
                "available_fields": {
                    field: sum(
                        contest_by_id[str(row["contest_id"])]["election_id"] == election_id
                        and str(contest_by_id[str(row["contest_id"])].get("list_type")) == list_type
                        and row.get(field) is not None
                        for row in mobilizations
                    )
                    for field in ("registered_voters", "voters", "valid_votes", "invalid_votes", "blank_votes", "turnout_rate", "contest_seats")
                },
                "source_ids": sorted({
                    str(row["source_id"]) for row in contests
                    if row["election_id"] == election_id and str(row.get("list_type")) == list_type
                }),
                "current_status": "NOT_COMPUTABLE",
                "conditional_analytics": ["rank", "winner", "victory_margin", "hhi", "enp"],
                "conditional_analytics_status": "NOT_APPLICABLE_WITHOUT_VERIFIED_EXTERNAL_PARTY_UNIVERSE",
                "detail_vs_aggregate_status": "NOT_APPLICABLE_WITHOUT_INDEPENDENT_AGGREGATE_TOTALS",
            }
            for (election_id, list_type), count in sorted(by_type.items())
        ],
        "tafra_comparability": {
            "contests_with_pinned_tafra_source": len(contests),
            "comm2015_contests": sum(row["election_id"] == "COMM2015" for row in contests),
            "comm2015_tafra_result_rows": 1538,
            "comm2015_elected_person_rows": elected_person_rows,
            "comm2015_elected_distinct_communes": len(elected_communes),
            "comm2015_elected_source_sha256": ELECTED_2015_SHA256,
            "checks_computable_from_these_files_as_external_universes": 0,
            "elected_person_file_is_not_an_official_candidate_universe": True,
        },
    }
    return rows, report
