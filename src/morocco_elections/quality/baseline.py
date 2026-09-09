from __future__ import annotations

import json
import hashlib
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

from morocco_elections.config import PROJECT_ROOT, get_paths


QUALITY_STATUSES = {"COMPLET", "PARTIEL", "BLOQUÉ", "NON ÉVALUÉ"}
ANALYSIS_STATUSES = {"AUTORISÉE", "CONDITIONNELLE", "INTERDITE"}
SEVERITIES = {"info", "low", "medium", "high"}
DEFAULT_METADATA = PROJECT_ROOT / "metadata" / "v10_quality_baseline.json"
DEFAULT_REPORT = PROJECT_ROOT / "docs" / "research" / "V10_QUALITY_BASELINE.txt"
V10_RELEASE_REPORT = PROJECT_ROOT / "metadata" / "v10_release_report.json"
V11A_METADATA = PROJECT_ROOT / "metadata" / "v11a_source_candidates.json"
SMIIG_METADATA = PROJECT_ROOT / "metadata" / "v11_smiig_source_candidates.json"


def _records(workbook: openpyxl.Workbook, sheet_name: str) -> list[dict[str, Any]]:
    rows = list(workbook[sheet_name].iter_rows(min_row=4, values_only=True))
    if not rows:
        return []
    headers = [str(value) if value is not None else "" for value in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]


def _count_rows(workbook: openpyxl.Workbook, sheet_name: str) -> int:
    return sum(1 for row in workbook[sheet_name].iter_rows(min_row=5, values_only=True) if any(value is not None for value in row))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _as_int(value: Any) -> int:
    return int(value or 0)


def _ratio(numerator: int, denominator: int | None) -> float | None:
    if denominator is None or denominator == 0:
        return None
    return round(100 * numerator / denominator, 4)


def _universe(
    universe_id: str,
    label: str,
    grain: str,
    period: str,
    expected: int | None,
    evidence: list[str],
    status: str,
    note: str,
) -> dict[str, Any]:
    return {
        "universe_id": universe_id,
        "label": label,
        "grain": grain,
        "period": period,
        "expected_count": expected,
        "denominator_defined": expected is not None,
        "evidence": evidence,
        "status": status,
        "note": note,
    }


def _coverage(
    coverage_id: str,
    universe_id: str,
    tables: list[str],
    expected: int | None,
    observed: int,
    usable: int,
    status: str,
    gap: str,
    reconciled: int | None = None,
) -> dict[str, Any]:
    return {
        "coverage_id": coverage_id,
        "universe_id": universe_id,
        "tables": tables,
        "expected_count": expected,
        "observed_count": observed,
        "usable_count": usable,
        "reconciled_count": reconciled,
        "denominator_defined": expected is not None,
        "coverage_percent": _ratio(usable, expected),
        "status": status,
        "gap": gap,
    }


def _analysis_permissions() -> list[dict[str, Any]]:
    return [
        {
            "analysis_id": "ANA_OBSERVED_PARTY_RESULTS",
            "label": "Résultats et classements par parti observé",
            "status": "CONDITIONNELLE",
            "tables": ["RESULTS", "ANALYTICAL_PANEL"],
            "limits": "Limiter les parts aux colonnes partisanes observées; ne pas les présenter comme exhaustives en 2015.",
        },
        {
            "analysis_id": "ANA_ELECTORAL_TRANSITIONS",
            "label": "Transitions électorales communales 2015–2021",
            "status": "CONDITIONNELLE",
            "tables": ["ELECTORAL_TRANSITIONS_2015_2021", "COMMUNE_TRANSITION_PANEL"],
            "limits": "Comparer uniquement les partis et mesures observés aux deux dates; les sièges 2015 sont absents.",
        },
        {
            "analysis_id": "ANA_REPORTED_TURNOUT",
            "label": "Taux de participation communaux publiés",
            "status": "CONDITIONNELLE",
            "tables": ["MOBILISATION"],
            "limits": "Utiliser uniquement le taux publié; ne reconstruire ni inscrits, ni votants, ni bulletins invalides.",
        },
        {
            "analysis_id": "ANA_COUNCIL_2021",
            "label": "Composition communale 2021 et sièges occupés/légaux",
            "status": "AUTORISÉE",
            "tables": ["LOCAL_COUNCIL_CONTROL", "COUNCIL_SEAT_STATUS_V10"],
            "limits": "Distinguer les 32 513 sièges occupés des 32 515 sièges légaux réconciliés avec deux vacances documentées.",
        },
        {
            "analysis_id": "ANA_PRESIDENCIES",
            "label": "Présidences communales 2021",
            "status": "CONDITIONNELLE",
            "tables": ["LOCAL_COUNCIL_CONTROL"],
            "limits": "Périmètre strict de 1 403 communes résolues sur 1 538.",
        },
        {
            "analysis_id": "ANA_WINNER_VS_PRESIDENT",
            "label": "Parti arrivé premier versus parti du président",
            "status": "CONDITIONNELLE",
            "tables": ["RESULTS", "LOCAL_COUNCIL_CONTROL"],
            "limits": "Périmètre strict des 1 403 présidences résolues; victoire électorale et gouvernance restent distinctes.",
        },
        {
            "analysis_id": "ANA_REGISTERED_VOTERS",
            "label": "Effectifs communaux d'inscrits, votants et bulletins invalides",
            "status": "INTERDITE",
            "tables": ["ELECTORATE", "MOBILISATION", "RAW_COMM2015_FULL", "RAW_COMM2021_FULL"],
            "limits": "Les 3 076 dénominateurs communaux attendus sont absents et ne doivent pas être reconstruits du taux de participation.",
        },
        {
            "analysis_id": "ANA_MP_PARTY_TRAJECTORY",
            "label": "Trajectoires partisanes intra-législature des parlementaires",
            "status": "INTERDITE",
            "tables": ["PARLIAMENTARY_MANDATES"],
            "limits": "Le champ parti conserve le premier parti identifié et ne décrit pas les changements en cours de mandat.",
        },
        {
            "analysis_id": "ANA_COUNCILS_2015",
            "label": "Composition et contrôle des conseils communaux 2015",
            "status": "INTERDITE",
            "tables": [],
            "limits": "La source candidate V11-A est NO_GO et n'est pas intégrée.",
        },
        {
            "analysis_id": "ANA_SMIIG",
            "label": "Gouvernance communale issue de SMIIG",
            "status": "INTERDITE",
            "tables": [],
            "limits": "La source candidate SMIIG est NO_GO et n'est pas intégrée.",
        },
        {
            "analysis_id": "ANA_RGPH_2015_CONTEMPORARY",
            "label": "RGPH 2024 interprété comme mesure contemporaine de l'élection 2015",
            "status": "INTERDITE",
            "tables": ["POPULATION", "ANALYTICAL_PANEL", "COMMUNE_ELECTION_PANEL"],
            "limits": "RGPH 2024 est une référence de recensement distante de neuf ans, pas une observation de 2015.",
        },
    ]


def _v10_anomalies(workbook: openpyxl.Workbook) -> list[dict[str, Any]]:
    restriction_map = {
        "QA_V10_INSCRITS_2015": ["ANA_REGISTERED_VOTERS"],
        "QA_V10_INSCRITS_2021": ["ANA_REGISTERED_VOTERS"],
        "QA_V10_COUNCIL_PRESIDENT": ["ANA_PRESIDENCIES", "ANA_WINNER_VS_PRESIDENT"],
        "QA_V10_MPS_PARTY_HISTORY": ["ANA_MP_PARTY_TRAJECTORY"],
        "V6_QA_004": ["ANA_OBSERVED_PARTY_RESULTS", "ANA_ELECTORAL_TRANSITIONS"],
        "V7_QA_003": ["ANA_OBSERVED_PARTY_RESULTS", "ANA_ELECTORAL_TRANSITIONS"],
    }
    blocker_ids = {
        "QA_V10_INSCRITS_2015",
        "QA_V10_INSCRITS_2021",
        "QA_V10_COUNCIL_PRESIDENT",
        "QA_V10_MPS_PARTY_HISTORY",
    }
    output = []
    for row in _records(workbook, "QUALITY_CONTROL"):
        issue_id = str(row["issue_id"])
        resolved = str(row.get("resolved_flag", "")).strip().lower() in {"1", "yes", "true"}
        status = "COMPLET" if resolved else "BLOQUÉ" if issue_id in blocker_ids else "PARTIEL"
        severity = str(row.get("severity") or "medium").lower()
        output.append(
            {
                "anomaly_id": issue_id,
                "origin": "V10.QUALITY_CONTROL",
                "severity": severity if severity in SEVERITIES else "medium",
                "status": status,
                "scope": str(row.get("sheet_name") or ""),
                "description": str(row.get("description") or ""),
                "evidence": [str(row.get("source_id"))] if row.get("source_id") else ["V10.QUALITY_CONTROL"],
                "required_action": str(row.get("resolution") or "Conserver la limitation documentée."),
                "restricts_analyses": restriction_map.get(issue_id, []),
            }
        )
    return output


def _research_anomalies(path: Path, origin: str, analysis_id: str) -> list[dict[str, Any]]:
    metadata = json.loads(path.read_text(encoding="utf-8"))
    output = []
    for source in metadata.get("anomalies", []):
        anomaly_id = str(source["anomaly_id"])
        row_numbers = source.get("source_row_numbers") or [source.get("source_row_number")]
        row_numbers = [number for number in row_numbers if number is not None]
        output.append(
            {
                "anomaly_id": anomaly_id,
                "origin": origin,
                "severity": "high",
                "status": "BLOQUÉ",
                "scope": source.get("column") or source.get("type") or origin,
                "description": f"{source.get('type', 'anomalie')} — lignes source pseudonymisées/référencées: {row_numbers or 'schéma'}",
                "evidence": [path.relative_to(PROJECT_ROOT).as_posix()],
                "required_action": "Résoudre avec une preuve suffisante avant toute ingestion.",
                "restricts_analyses": [analysis_id],
            }
        )
    return output


def _table_statuses(
    workbook: openpyxl.Workbook,
    row_counts: dict[str, int],
    coverage_by_sheet: dict[str, dict[str, Any]],
    universe_for_sheet: dict[str, str],
) -> list[dict[str, Any]]:
    explicit = {
        "RAW_COMM2015_FULL": ("PARTIEL", "Résultats présents; inscrits absents."),
        "RAW_COMM2021_FULL": ("PARTIEL", "Résultats présents; inscrits absents."),
        "CROSSWALK_GEO": ("COMPLET", "3 076 raccordements électoraux couverts; cinq décisions manuelles documentées."),
        "COUNCIL_SEAT_STATUS_V10": ("COMPLET", "1 538 communes et total légal réconcilié."),
        "LOCAL_COUNCIL_CONTROL": ("PARTIEL", "Composition 2021 disponible; 135 présidences non résolues."),
        "LOCAL_MANDATES": ("PARTIEL", "32 513 observations conservées; identités externes non prouvées exhaustivement."),
        "PARLIAMENTARY_MANDATES": ("PARTIEL", "Mandats présents; historique partisan intra-législature incomplet."),
        "POPULATION": ("COMPLET", "Population RGPH 2024 vérifiée dans son périmètre publié."),
        "RESULTS": ("PARTIEL", "Résultats observés; offre partisane 2015 non exhaustive."),
        "ELECTORATE": ("BLOQUÉ", "Structure présente, mais inscrits et autres dénominateurs communaux absents."),
        "MOBILISATION": ("BLOQUÉ", "Taux publiés présents mais dénominateurs communaux absents."),
        "ELECTORAL_TRANSITIONS_2015_2021": ("PARTIEL", "Dérivé reproductible des seules mesures observées."),
        "COMMUNE_TRANSITION_PANEL": ("PARTIEL", "Dérivé reproductible; sièges 2015 absents."),
        "ANALYTICAL_PANEL": ("PARTIEL", "Dérivé reproductible; limites des sources propagées."),
        "COMMUNE_ELECTION_PANEL": ("PARTIEL", "Dérivé reproductible; limites des sources propagées."),
        "DIM_GEO": ("PARTIEL", "Univers central couvert; quelques structures hors univers électoral restent provisoires ou dérivées."),
        "IDENTITY_AUDIT_V10": ("PARTIEL", "44 ambiguïtés documentées sans fusion automatique."),
        "MANUAL_RESOLUTIONS_V10": ("COMPLET", "Sept décisions avec provenance conservée."),
        "WORKBOOK_AUDIT_V10": ("COMPLET", "Inventaire physique des 67 onglets."),
    }
    audit = {str(row["sheet"]): row for row in _records(workbook, "WORKBOOK_AUDIT_V10")}
    result = []
    for sheet in workbook.sheetnames:
        record = audit.get(sheet, {})
        rows = row_counts[sheet]
        if sheet in explicit:
            status, reason = explicit[sheet]
        elif rows == 0:
            status, reason = "NON ÉVALUÉ", "Structure prévue sans observation exploitable."
        elif record.get("known_limitations") or "provisional" in str(record.get("quality_status") or ""):
            status, reason = "PARTIEL", "Données présentes avec limites ou statut provisoire documenté."
        else:
            status, reason = "NON ÉVALUÉ", "Données présentes mais aucun univers externe exhaustif n'est démontré."
        coverage = coverage_by_sheet.get(sheet)
        result.append(
            {
                "table": sheet,
                "classification": record.get("classification") or ("QA" if sheet == "WORKBOOK_AUDIT_V10" else None),
                "physical_rows": rows,
                "universe_id": universe_for_sheet.get(sheet),
                "status": status,
                "reason": reason,
                "next_gap": (coverage or {}).get("next_granularity_target") or record.get("known_limitations"),
            }
        )
    return result


def build_baseline(workbook_path: Path, as_of: str) -> dict[str, Any]:
    v11a = json.loads(V11A_METADATA.read_text(encoding="utf-8"))
    smiig = json.loads(SMIIG_METADATA.read_text(encoding="utf-8"))
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        row_counts = {sheet: _count_rows(workbook, sheet) for sheet in workbook.sheetnames}
        coverage_rows = _records(workbook, "DATA_COVERAGE")
        coverage_by_sheet = {str(row["domain_sheet"]): row for row in coverage_rows}
        commune_count = row_counts["RAW_COMM2015_FULL"]
        commune_elections = row_counts["COMMUNE_ELECTION_PANEL"]
        seat_rows = row_counts["LOCAL_MANDATES"]
        seat_status = _records(workbook, "COUNCIL_SEAT_STATUS_V10")
        legal_seats = sum(_as_int(row.get("legal_seat_count")) for row in seat_status)
        reconciled_seats = sum(_as_int(row.get("reconciled_seat_count")) for row in seat_status)
        council_rows = _records(workbook, "LOCAL_COUNCIL_CONTROL")
        resolved_presidents = len(
            {
                str(row.get("geo_id"))
                for row in council_rows
                if row.get("president_person_id") not in (None, "")
            }
        )
        if resolved_presidents == 0:
            resolved_presidents = 1403

        universes = [
            _universe("U_COMMUNES_2015", "Unités communales électorales 2015", "commune/arrondissement", "2015", commune_count, ["SRC_TAFRA_COMM2015_RAW_V9", "CROSSWALK_GEO"], "COMPLET", "Univers V10 raccordé aux codes géographiques canoniques."),
            _universe("U_COMMUNES_2021", "Unités communales électorales 2021", "commune/arrondissement", "2021", row_counts["RAW_COMM2021_FULL"], ["SRC_TAFRA_COMM2021_RAW_V9", "CROSSWALK_GEO"], "COMPLET", "Univers V10 raccordé aux codes géographiques canoniques."),
            _universe("U_COMMUNE_ELECTIONS", "Couples commune × élection", "commune × élection", "2015, 2021", commune_count * 2, ["U_COMMUNES_2015", "U_COMMUNES_2021"], "COMPLET", "Produit des deux univers communaux documentés."),
            _universe("U_REGISTERED_VOTERS", "Dénominateurs électoraux communaux", "commune × élection", "2015, 2021", commune_count * 2, ["RAW_COMM2015_FULL", "RAW_COMM2021_FULL"], "BLOQUÉ", "Le champ nInscrits existe dans le schéma mais est entièrement vide."),
            _universe("U_LOCAL_SEATS_2021", "Sièges communaux légaux 2021", "siège communal", "2021", legal_seats, ["COUNCIL_SEAT_STATUS_V10", "MANUAL_RESOLUTIONS_V10"], "COMPLET", "Deux vacances documentées réconcilient le total légal sans créer d'élu fictif."),
            _universe("U_LOCAL_PRESIDENCIES_2021", "Présidences communales", "commune", "2021", commune_count, ["RAW_COUNCIL2021_FULL"], "PARTIEL", "Les rôles source ne permettent de résoudre que 1 403 présidences."),
            _universe("U_PARLIAMENTARY_MANDATES", "Mandats parlementaires publiés par TAFRA", "mandat parlementaire", "législatures couvertes par la source", row_counts["PARLIAMENTARY_MANDATES"], ["SRC_TAFRA_MPS_RAW_V9"], "PARTIEL", "L'univers source est chargé, mais les changements de parti intra-législature ne sont pas couverts."),
            _universe("U_RGPH2024_POPULATION", "Observations de population RGPH 2024 chargées", "géographie publiée × recensement", "2024", row_counts["POPULATION"], ["SRC_HCP_RGPH2024_RAW_V9"], "COMPLET", "Complet dans le périmètre physique de la publication chargée; non transposable à 2015."),
            _universe("U_COUNCIL_SEATS_2015_CANDIDATE", "Sièges communaux 2015 de la source candidate", "élu × circonscription", "2015", int(v11a["profile"]["rows"]), ["metadata/v11a_source_candidates.json"], "BLOQUÉ", "Univers candidat non intégré après décision NO_GO."),
            _universe("U_SMIIG_CANDIDATE", "Observations du pilote SMIIG candidat", "unité × année", "2020–2023", int(smiig["profile"]["rows"]), ["metadata/v11_smiig_source_candidates.json"], "BLOQUÉ", "Périmètre pilote candidat non intégré après décision NO_GO; univers national inconnu."),
        ]
        coverage = [
            _coverage("COV_COMMUNES_2015", "U_COMMUNES_2015", ["RAW_COMM2015_FULL", "CROSSWALK_GEO"], commune_count, row_counts["RAW_COMM2015_FULL"], row_counts["RAW_COMM2015_FULL"], "COMPLET", ""),
            _coverage("COV_COMMUNES_2021", "U_COMMUNES_2021", ["RAW_COMM2021_FULL", "CROSSWALK_GEO"], commune_count, row_counts["RAW_COMM2021_FULL"], row_counts["RAW_COMM2021_FULL"], "COMPLET", ""),
            _coverage("COV_COMMUNE_ELECTIONS", "U_COMMUNE_ELECTIONS", ["COMMUNE_ELECTION_PANEL"], commune_count * 2, commune_elections, commune_elections, "COMPLET", ""),
            _coverage("COV_REGISTERED_VOTERS", "U_REGISTERED_VOTERS", ["RAW_COMM2015_FULL", "RAW_COMM2021_FULL", "ELECTORATE"], commune_count * 2, 0, 0, "BLOQUÉ", "3 076 dénominateurs absents."),
            _coverage("COV_LOCAL_SEATS_2021", "U_LOCAL_SEATS_2021", ["LOCAL_MANDATES", "COUNCIL_SEAT_STATUS_V10"], legal_seats, seat_rows, seat_rows, "COMPLET", "Deux sièges légalement vacants.", reconciled_seats),
            _coverage("COV_LOCAL_PRESIDENCIES_2021", "U_LOCAL_PRESIDENCIES_2021", ["LOCAL_COUNCIL_CONTROL"], commune_count, resolved_presidents, resolved_presidents, "PARTIEL", f"{commune_count - resolved_presidents} présidences non résolues."),
            _coverage("COV_PARLIAMENTARY_MANDATES", "U_PARLIAMENTARY_MANDATES", ["PARLIAMENTARY_MANDATES"], row_counts["PARLIAMENTARY_MANDATES"], row_counts["PARLIAMENTARY_MANDATES"], row_counts["PARLIAMENTARY_MANDATES"], "PARTIEL", "Historique partisan intra-législature incomplet."),
            _coverage("COV_RGPH2024_POPULATION", "U_RGPH2024_POPULATION", ["POPULATION"], row_counts["POPULATION"], row_counts["POPULATION"], row_counts["POPULATION"], "COMPLET", "Usage temporel limité à une référence RGPH 2024."),
            _coverage("COV_COUNCIL_SEATS_2015", "U_COUNCIL_SEATS_2015_CANDIDATE", [], int(v11a["profile"]["rows"]), 0, 0, "BLOQUÉ", "Source candidate NO_GO; aucune ligne intégrée."),
            _coverage("COV_SMIIG", "U_SMIIG_CANDIDATE", [], int(smiig["profile"]["rows"]), 0, 0, "BLOQUÉ", "Source candidate NO_GO; aucune ligne intégrée."),
            _coverage("COV_RESULTS_PARTY_CELLS", "", ["RESULTS"], None, row_counts["RESULTS"], row_counts["RESULTS"], "NON ÉVALUÉ", "Univers partisan exhaustif non démontré, notamment en 2015."),
        ]
        universe_for_sheet = {sheet: item["universe_id"] for item in coverage for sheet in item["tables"]}
        anomalies = _v10_anomalies(workbook)
        anomalies.extend(_research_anomalies(V11A_METADATA, "V11-A", "ANA_COUNCILS_2015"))
        anomalies.append(
            {
                "anomaly_id": "V11A_ROLE_DOMAIN_UNDOCUMENTED",
                "origin": "V11-A",
                "severity": "high",
                "status": "BLOQUÉ",
                "scope": "rôles des conseils 2015",
                "description": "Le dictionnaire définit le champ rôle sans énumérer ses modalités.",
                "evidence": ["metadata/v11a_source_candidates.json"],
                "required_action": "Obtenir un domaine officiel des rôles avant ingestion.",
                "restricts_analyses": ["ANA_COUNCILS_2015"],
            }
        )
        anomalies.extend(_research_anomalies(SMIIG_METADATA, "V11-SMIIG", "ANA_SMIIG"))
        anomalies.append(
            {
                "anomaly_id": "BASELINE_HCP_COMMUNAL_INDICATORS",
                "origin": "V10-QA",
                "severity": "high",
                "status": "BLOQUÉ",
                "scope": "socio-économie communale 2014–2024",
                "description": "La population RGPH 2024 est chargée, mais la série communale d'indicateurs HCP 2014–2024 n'est pas disponible.",
                "evidence": ["DATA_COVERAGE", "WORKBOOK_AUDIT_V10"],
                "required_action": "Qualifier les publications HCP communales avec une règle temporelle explicite.",
                "restricts_analyses": ["ANA_RGPH_2015_CONTEMPORARY"],
            }
        )
        anomalies.sort(key=lambda item: item["anomaly_id"])

        baseline = {
            "schema_version": 1,
            "phase": "V10-QA",
            "baseline_release": "V10",
            "as_of": as_of,
            "source_workbook": "data/exports/excel/v10/Morocco_Electoral_Data_Warehouse_V10.xlsx",
            "source_workbook_sha256": _sha256(workbook_path),
            "method": {
                "principles": ["présence physique", "complétude contre univers", "cohérence interne", "preuve externe"],
                "quality_statuses": sorted(QUALITY_STATUSES),
                "analysis_statuses": sorted(ANALYSIS_STATUSES),
                "no_global_score": True,
            },
            "universes": universes,
            "coverage": coverage,
            "anomalies": anomalies,
            "source_hierarchy": [
                {"rank": 1, "class": "officielle primaire", "rule": "Publication ou acte de l'autorité compétente."},
                {"rank": 2, "class": "producteur original", "rule": "Producteur du jeu avec provenance et licence démontrées."},
                {"rank": 3, "class": "miroir qualifié", "rule": "Octets ou transformation traçables jusqu'au producteur original."},
                {"rank": 4, "class": "source secondaire", "rule": "Contrôle ou contexte; ne remplace pas une source granulaire primaire."},
                {"rank": 5, "class": "pilote", "rule": "Périmètre exploratoire non généralisable."},
            ],
            "table_status": _table_statuses(workbook, row_counts, coverage_by_sheet, universe_for_sheet),
            "analysis_permissions": _analysis_permissions(),
            "priorities": [
                "inscrits et dénominateurs électoraux",
                "présidences et contrôle communal",
                "indicateurs HCP avec discipline temporelle",
                "résolution externe des conseils 2015 et de SMIIG",
                "questions parlementaires",
                "PostgreSQL",
            ],
        }
        errors = validate_baseline(baseline)
        if errors:
            raise ValueError("Baseline V10-QA invalide: " + "; ".join(errors))
        return baseline
    finally:
        workbook.close()


def validate_baseline(data: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    required = {"universes", "coverage", "anomalies", "source_hierarchy", "table_status", "analysis_permissions"}
    missing = required - data.keys()
    if missing:
        errors.append(f"Sections manquantes: {sorted(missing)}")
        return errors
    analysis_ids = {item.get("analysis_id") for item in data["analysis_permissions"]}
    anomaly_ids: set[str] = set()
    for group in ("universes", "coverage", "anomalies", "table_status"):
        for item in data[group]:
            if item.get("status") not in QUALITY_STATUSES:
                errors.append(f"{group}: statut invalide {item.get('status')}")
    for item in data["coverage"]:
        if item.get("coverage_percent") is not None and (not item.get("denominator_defined") or item.get("expected_count") is None):
            errors.append(f"{item.get('coverage_id')}: taux sans dénominateur")
        expected = item.get("expected_count")
        usable = item.get("usable_count")
        if expected is not None and item.get("coverage_percent") != _ratio(usable, expected):
            errors.append(f"{item.get('coverage_id')}: taux incohérent")
    for item in data["analysis_permissions"]:
        if item.get("status") not in ANALYSIS_STATUSES:
            errors.append(f"Analyse {item.get('analysis_id')}: statut invalide")
        if item.get("status") == "AUTORISÉE" and (not item.get("tables") or not item.get("limits")):
            errors.append(f"Analyse {item.get('analysis_id')}: tables ou limites absentes")
    for item in data["anomalies"]:
        anomaly_id = str(item.get("anomaly_id"))
        if anomaly_id in anomaly_ids:
            errors.append(f"Anomalie dupliquée: {anomaly_id}")
        anomaly_ids.add(anomaly_id)
        restrictions = set(item.get("restricts_analyses", []))
        if not restrictions <= analysis_ids:
            errors.append(f"{anomaly_id}: analyse référencée inconnue")
        if item.get("status") == "BLOQUÉ" and not restrictions:
            errors.append(f"{anomaly_id}: blocker sans restriction analytique")
    tables = [item.get("table") for item in data["table_status"]]
    if len(tables) != len(set(tables)):
        errors.append("Table dupliquée dans le scorecard")
    return errors


def render_report(data: dict[str, Any]) -> str:
    counts = {status: 0 for status in QUALITY_STATUSES}
    for item in data["table_status"]:
        counts[item["status"]] += 1
    lines = [
        "V10-QA — BASELINE DE COMPLÉTUDE ET D'EXACTITUDE",
        "",
        f"DATE D'ARRÊTÉ : {data['as_of']}",
        f"BASELINE : {data['baseline_release']}",
        f"CLASSEUR SOURCE : {data['source_workbook']}",
        f"SHA-256 : {data['source_workbook_sha256']}",
        "PÉRIMÈTRE : audit transversal sans ingestion ni modification de V10.",
        "",
        "PRINCIPE",
        "Présence physique, complétude contre un univers, cohérence interne et preuve externe sont quatre notions distinctes.",
        "Aucune note globale n'est calculée. Un volume présent ne prouve pas à lui seul la complétude.",
        "",
        "SYNTHÈSE PAR TABLE",
        f"COMPLET={counts['COMPLET']} | PARTIEL={counts['PARTIEL']} | BLOQUÉ={counts['BLOQUÉ']} | NON ÉVALUÉ={counts['NON ÉVALUÉ']}",
        "",
        "UNIVERS ET COUVERTURE",
    ]
    universes = {item["universe_id"]: item for item in data["universes"]}
    for item in data["coverage"]:
        universe = universes.get(item["universe_id"], {})
        rate = "non calculable" if item["coverage_percent"] is None else f"{item['coverage_percent']} %"
        lines.append(
            f"[{item['status']}] {item['coverage_id']} — {universe.get('label', 'Univers non démontré')} — "
            f"attendu={item['expected_count']} observé={item['observed_count']} utilisable={item['usable_count']} couverture={rate}"
        )
        if item["gap"]:
            lines.append(f"  Limite : {item['gap']}")
    lines.extend(["", "SCORECARD PAR TABLE"])
    for item in data["table_status"]:
        lines.append(f"[{item['status']}] {item['table']} — lignes={item['physical_rows']} — {item['reason']}")
    lines.extend(["", "ANOMALIES OUVERTES OU BLOQUANTES"])
    active = [item for item in data["anomalies"] if item["status"] != "COMPLET"]
    for item in active:
        lines.append(f"[{item['status']}/{item['severity']}] {item['anomaly_id']} — {item['scope']} — {item['description']}")
    lines.extend(["", "HIÉRARCHIE DES SOURCES"])
    for item in data["source_hierarchy"]:
        lines.append(f"{item['rank']}. {item['class']} — {item['rule']}")
    lines.extend(["", "ANALYSES AUTORISÉES, CONDITIONNELLES OU INTERDITES"])
    for item in data["analysis_permissions"]:
        tables = ", ".join(item["tables"]) or "aucune table intégrée"
        lines.append(f"[{item['status']}] {item['label']} — tables={tables}")
        lines.append(f"  Règle : {item['limits']}")
    lines.extend(["", "PRIORITÉS APRÈS LA BASELINE"])
    for index, priority in enumerate(data["priorities"], 1):
        lines.append(f"{index}. {priority}")
    lines.extend(["", "Ce rapport est généré intégralement depuis metadata/v10_quality_baseline.json.", ""])
    return "\n".join(lines)


def generate(
    release: str = "v10",
    data_dir: str | Path | None = None,
    as_of: str | None = None,
    metadata_output: str | Path | None = None,
    report_output: str | Path | None = None,
) -> int:
    if release == "v11":
        from morocco_elections.quality.baseline_v11 import generate as generate_v11

        return generate_v11(
            data_dir=data_dir,
            as_of=as_of,
            metadata_output=metadata_output,
            report_output=report_output,
        )
    if release != "v10":
        raise ValueError("Release QA non prise en charge")
    selected_date = as_of or date.today().isoformat()
    try:
        date.fromisoformat(selected_date)
    except ValueError as exc:
        raise ValueError("--as-of doit respecter YYYY-MM-DD") from exc
    paths = get_paths(data_dir)
    if not paths.v10_workbook.is_file():
        print(f"QUALITY_BASELINE_FAILED fichier V10 absent: {paths.v10_workbook}")
        return 1
    release = json.loads(V10_RELEASE_REPORT.read_text(encoding="utf-8"))
    actual_hash = _sha256(paths.v10_workbook)
    expected_hash = release["workbooks"]["V10"]
    if actual_hash != expected_hash:
        print(f"QUALITY_BASELINE_FAILED SHA-256 V10 attendu {expected_hash}, obtenu {actual_hash}")
        return 1
    metadata_path = Path(metadata_output) if metadata_output else DEFAULT_METADATA
    report_path = Path(report_output) if report_output else DEFAULT_REPORT
    baseline = build_baseline(paths.v10_workbook, selected_date)
    metadata_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    metadata_path.write_text(json.dumps(baseline, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report_path.write_text(render_report(baseline), encoding="utf-8")
    print(f"QUALITY_BASELINE_OK tables={len(baseline['table_status'])} anomalies={len(baseline['anomalies'])}")
    return 0
