from __future__ import annotations

import argparse
import ast
import hashlib
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Iterable

import openpyxl

from morocco_elections.config import PROJECT_ROOT, get_paths, resolve_manifest_path

ROOT = PROJECT_ROOT
MANIFEST_PATH = get_paths().source_manifest
BACKLOG_PATH = get_paths().github_backlog
V10_REPORT_PATH = PROJECT_ROOT / "metadata" / "v10_release_report.json"
DOCUMENTATION_DIRS = {"v9": get_paths().documentation_v9, "v10": get_paths().documentation_v10}
EXPECTED_DOCUMENTS = [
    "00_INDEX_ET_MODE_EMPLOI.txt",
    "01_ONTOLOGIE_ELECTORALE_GLOBALE.txt",
    "02_ARCHITECTURE_DES_CUBES.txt",
    "03_DIMENSIONS_CONFORMES.txt",
    "04_HIERARCHIE_GEOGRAPHIQUE.txt",
    "05_CROSSWALKS_ET_IDENTITES.txt",
    "06_TEMPS_CYCLES_ET_ELECTIONS.txt",
    "07_FAITS_ELECTORAUX_ET_MOBILISATION.txt",
    "08_PERSONNES_MANDATS_ET_REPRESENTATION.txt",
    "09_POUVOIR_LOCAL_ET_GOUVERNANCE.txt",
    "10_SOCIO_ECONOMIE_CAMPAGNE_INFORMATION.txt",
    "11_PIPELINE_RAW_NORMALIZED_FACT_ANALYTICAL.txt",
    "12_REGLES_D_INTEGRITE_ET_JOINTURE.txt",
    "13_METRIQUES_DERIVEES_ET_REGLES_TEMPORELLES.txt",
    "14_CATALOGUE_COMPLET_DES_ONGLETS.txt",
]
REQUIRED_DOCUMENT_METADATA = (
    "VERSION",
    "CLASSEUR SOURCE",
    "DATE DE GÉNÉRATION",
    "PÉRIMÈTRE",
    "RENVOIS",
)
REQUIRED_SOURCE_FIELDS = {
    "source_id",
    "local_path",
    "source_url",
    "producer",
    "license",
    "version",
    "sha256",
    "byte_size",
    "grain",
    "rows",
    "columns",
    "status",
}
REQUIRED_ARTIFACT_FIELDS = {
    "artifact_id",
    "local_path",
    "sha256",
    "byte_size",
    "sheet_count",
    "status",
}
REQUIRED_PHYSICAL_FILE_FIELDS = {
    "file_id",
    "previous_local_path",
    "local_path",
    "sha256",
    "byte_size",
    "kind",
}
FORBIDDEN_SUFFIXES = {
    ".xlsx",
    ".xls",
    ".xlsm",
    ".parquet",
    ".shp",
    ".shx",
    ".dbf",
    ".prj",
    ".cpg",
    ".zip",
    ".7z",
    ".dump",
    ".backup",
    ".pem",
    ".key",
}
FORBIDDEN_DIRECTORIES = {"raw_sources", "artifacts", "backups", "secrets", "pgdata", "postgres-data"}
FORBIDDEN_DATA_PREFIXES = {
    ("data", "raw"),
    ("data", "legacy"),
    ("data", "staging"),
    ("data", "processed"),
    ("data", "exports"),
}
MAX_TRACKED_BYTES = 5 * 1024 * 1024
EXPECTED_V9_VOLUMES = {
    "LOCAL_MANDATES": 32513,
    "PARLIAMENTARY_MANDATES": 1654,
    "COMMUNE_ELECTION_PANEL": 3076,
    "ELECTORAL_TRANSITIONS_2015_2021": 14555,
    "ANALYTICAL_PANEL": 22054,
}
EXPECTED_V10_VOLUMES = {
    **EXPECTED_V9_VOLUMES,
    "DIM_PERSON": 33697,
    "IDENTITY_AUDIT_V10": 44,
    "MANUAL_RESOLUTIONS_V10": 7,
    "COUNCIL_SEAT_STATUS_V10": 1538,
}
REQUIRED_EVIDENCE_FIELDS = {"evidence_id", "local_path", "source_url", "publisher", "sha256", "byte_size", "status"}


class ValidationError(RuntimeError):
    pass


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_manifest(path: Path | None = None) -> dict:
    path = path or MANIFEST_PATH
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValidationError(f"Manifeste illisible: {path}: {exc}") from exc


def validate_manifest(manifest: dict) -> list[str]:
    errors: list[str] = []
    if manifest.get("schema_version") != 3:
        errors.append("metadata/source_manifest.json: schema_version doit valoir 3")
    if manifest.get("warehouse_version") != "V10":
        errors.append("metadata/source_manifest.json: warehouse_version doit valoir V10")

    sources = manifest.get("sources")
    artifacts = manifest.get("artifacts")
    physical_files = manifest.get("physical_files")
    evidence = manifest.get("evidence")
    if not isinstance(sources, list) or not sources:
        errors.append("Le manifeste doit contenir une liste sources non vide")
        sources = []
    if not isinstance(artifacts, list) or not artifacts:
        errors.append("Le manifeste doit contenir une liste artifacts non vide")
        artifacts = []
    if not isinstance(physical_files, list) or len(physical_files) != 22:
        errors.append("Le manifeste doit inventorier exactement 22 fichiers physiques")
        physical_files = []
    if not isinstance(evidence, list) or len(evidence) != 2:
        errors.append("Le manifeste doit inventorier exactement deux preuves officielles V10")
        evidence = []

    seen_ids: set[str] = set()
    for kind, records, required, id_field in (
        ("source", sources, REQUIRED_SOURCE_FIELDS, "source_id"),
        ("artifact", artifacts, REQUIRED_ARTIFACT_FIELDS, "artifact_id"),
        ("physical_file", physical_files, REQUIRED_PHYSICAL_FILE_FIELDS, "file_id"),
        ("evidence", evidence, REQUIRED_EVIDENCE_FIELDS, "evidence_id"),
    ):
        seen_paths: set[str] = set()
        for index, record in enumerate(records):
            label = f"{kind}[{index}]"
            if not isinstance(record, dict):
                errors.append(f"{label}: objet JSON attendu")
                continue
            missing = sorted(required - record.keys())
            if missing:
                errors.append(f"{label}: champs absents: {', '.join(missing)}")
            record_id = str(record.get(id_field, ""))
            local_path = str(record.get("local_path", ""))
            if not record_id or record_id in seen_ids:
                errors.append(f"{label}: identifiant vide ou dupliqué: {record_id!r}")
            if not local_path or local_path in seen_paths or Path(local_path).is_absolute() or ".." in Path(local_path).parts:
                errors.append(f"{label}: local_path invalide ou dupliqué: {local_path!r}")
            seen_ids.add(record_id)
            seen_paths.add(local_path)
            if not re.fullmatch(r"[0-9a-f]{64}", str(record.get("sha256", ""))):
                errors.append(f"{label}: SHA-256 invalide")
            for field in ("byte_size", "rows", "columns", "sheet_count"):
                if field in record and (not isinstance(record[field], int) or record[field] <= 0):
                    errors.append(f"{label}: {field} doit être un entier positif")
    return errors


def validate_backlog() -> list[str]:
    errors: list[str] = []
    try:
        backlog = json.loads(BACKLOG_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"Backlog GitHub illisible: {exc}"]
    milestones = backlog.get("milestones", [])
    labels = backlog.get("labels", [])
    issues = backlog.get("issues", [])
    milestone_titles = {item.get("title") for item in milestones if isinstance(item, dict)}
    label_names = {item.get("name") for item in labels if isinstance(item, dict)}
    issue_titles = {item.get("title") for item in issues if isinstance(item, dict)}
    if len(milestones) != 4 or len(milestone_titles) != 4:
        errors.append("Le backlog doit définir exactement quatre jalons uniques")
    if len(issues) != 7 or len(issue_titles) != 7:
        errors.append("Le backlog doit définir exactement sept issues uniques")
    required_labels = {"data", "qa", "source", "identity", "governance", "parliament", "postgresql"}
    if label_names != required_labels:
        errors.append(f"Labels GitHub incorrects: {sorted(label_names)}")
    for issue in issues:
        if issue.get("milestone") not in milestone_titles:
            errors.append(f"Issue sans jalon valide: {issue.get('title')}")
        if not issue.get("acceptance"):
            errors.append(f"Issue sans critères d’acceptation: {issue.get('title')}")
        unknown_labels = set(issue.get("labels", [])) - label_names
        if unknown_labels:
            errors.append(f"Issue {issue.get('title')}: labels inconnus {sorted(unknown_labels)}")
        unknown_dependencies = set(issue.get("depends_on", [])) - issue_titles
        if unknown_dependencies:
            errors.append(f"Issue {issue.get('title')}: dépendances inconnues {sorted(unknown_dependencies)}")
    return errors


def validate_v10_release_report(manifest: dict) -> list[str]:
    try:
        report = json.loads(V10_REPORT_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"Rapport V10 illisible: {exc}"]
    errors = []
    if report.get("release") != "V10" or report.get("baseline") != "V9" or report.get("status") != "validated":
        errors.append("Le rapport de release doit décrire V10 validé contre V9")
    artifacts = {item["artifact_id"]: item for item in manifest.get("artifacts", [])}
    for version in ("V8", "V9", "V10"):
        artifact = artifacts.get(f"WAREHOUSE_{version}_{'INPUT' if version == 'V8' else 'EXPORT'}")
        if not artifact or report.get("workbooks", {}).get(version) != artifact.get("sha256"):
            errors.append(f"Empreinte {version} incohérente dans le rapport de release")
    for sheet, expected in EXPECTED_V10_VOLUMES.items():
        if report.get("volumes", {}).get(sheet) != expected:
            errors.append(f"Volume {sheet} incohérent dans le rapport de release")
    return errors


def validate_python_sources() -> list[str]:
    errors: list[str] = []
    for path in sorted(ROOT.rglob("*.py")):
        if any(part in {".venv", "venv", "__pycache__"} for part in path.parts):
            continue
        try:
            ast.parse(path.read_text(encoding="utf-8-sig"), filename=str(path))
        except (OSError, SyntaxError, UnicodeError) as exc:
            errors.append(f"Python invalide: {path.relative_to(ROOT)}: {exc}")
    return errors


def validate_documentation(release: str = "all") -> list[str]:
    errors: list[str] = []
    releases = DOCUMENTATION_DIRS if release == "all" else {release: DOCUMENTATION_DIRS[release]}
    for version, directory in releases.items():
        actual = sorted(path.name for path in directory.glob("*.txt"))
        if actual != EXPECTED_DOCUMENTS:
            errors.append(f"Corpus documentaire {version} incorrect: attendu {EXPECTED_DOCUMENTS}, obtenu {actual}")
        for name in EXPECTED_DOCUMENTS:
            path = directory / name
            if not path.is_file():
                continue
            try:
                content = path.read_text(encoding="utf-8-sig")
            except UnicodeError as exc:
                errors.append(f"Document non UTF-8: {version}/{name}: {exc}")
                continue
            if not content.strip():
                errors.append(f"Document vide: {version}/{name}")
            for token in REQUIRED_DOCUMENT_METADATA:
                if token not in content:
                    errors.append(f"Métadonnée absente de {version}/{name}: {token}")
    return errors


def tracked_files() -> list[Path]:
    result = subprocess.run(
        ["git", "ls-files", "-z"],
        cwd=ROOT,
        capture_output=True,
        check=False,
    )
    if result.returncode != 0:
        return []
    return [ROOT / item.decode("utf-8") for item in result.stdout.split(b"\0") if item]


def candidate_repository_files() -> list[Path]:
    tracked = tracked_files()
    if tracked:
        return tracked
    candidates: list[Path] = []
    for path in ROOT.rglob("*"):
        if not path.is_file() or ".git" in path.parts:
            continue
        relative = path.relative_to(ROOT)
        if any(part in FORBIDDEN_DIRECTORIES or part in {".venv", "venv", "__pycache__"} for part in relative.parts):
            continue
        if tuple(relative.parts[:2]) in FORBIDDEN_DATA_PREFIXES:
            continue
        if path.suffix.lower() in FORBIDDEN_SUFFIXES:
            continue
        candidates.append(path)
    return candidates


def validate_repository_files(files: Iterable[Path] | None = None) -> list[str]:
    errors: list[str] = []
    selected = list(files) if files is not None else candidate_repository_files()
    secret_patterns = {
        "clé privée": re.compile("BEGIN" + r" [A-Z ]*PRIVATE KEY"),
        "jeton GitHub classique": re.compile("gh" + r"p_[A-Za-z0-9]{20,}"),
        "jeton GitHub fin": re.compile("github" + r"_pat_[A-Za-z0-9_]{20,}"),
        "clé AWS": re.compile("AK" + r"IA[0-9A-Z]{16}"),
    }
    for path in selected:
        try:
            relative = path.resolve().relative_to(ROOT.resolve())
        except ValueError:
            errors.append(f"Fichier hors dépôt: {path}")
            continue
        if any(part in FORBIDDEN_DIRECTORIES for part in relative.parts):
            errors.append(f"Répertoire de données interdit dans Git: {relative.as_posix()}")
        if tuple(relative.parts[:2]) in FORBIDDEN_DATA_PREFIXES:
            errors.append(f"Zone de données interdite dans Git: {relative.as_posix()}")
        if path.suffix.lower() in FORBIDDEN_SUFFIXES:
            errors.append(f"Type de fichier interdit dans Git: {relative.as_posix()}")
        if path.exists() and path.stat().st_size > MAX_TRACKED_BYTES:
            errors.append(f"Fichier suivi supérieur à 5 MiB: {relative.as_posix()}")
        if not path.exists() or path.suffix.lower() in FORBIDDEN_SUFFIXES:
            continue
        try:
            content = path.read_text(encoding="utf-8-sig")
        except (UnicodeError, OSError):
            continue
        for label, pattern in secret_patterns.items():
            if pattern.search(content):
                errors.append(f"Secret potentiel ({label}) dans {relative.as_posix()}")
    return errors


def validate_physical_record(record: dict, data_dir: str | Path | None = None) -> list[str]:
    errors: list[str] = []
    path = resolve_manifest_path(record["local_path"], data_dir)
    label = record.get("source_id") or record.get("artifact_id") or record.get("evidence_id") or record.get("file_id")
    if not path.is_file():
        return [f"{label}: fichier local absent: {record['local_path']}"]
    if path.stat().st_size != record["byte_size"]:
        errors.append(f"{label}: taille attendue {record['byte_size']}, obtenue {path.stat().st_size}")
    actual_hash = sha256(path)
    if actual_hash != record["sha256"]:
        errors.append(f"{label}: SHA-256 attendu {record['sha256']}, obtenu {actual_hash}")
    return errors


def _populated_rows(sheet) -> int:
    return sum(1 for row in sheet.iter_rows(min_row=5, values_only=True) if any(value is not None for value in row))


def compare_v10_to_v9(data_dir: str | Path | None = None) -> list[str]:
    errors: list[str] = []
    paths = get_paths(data_dir)
    if not paths.v9_workbook.is_file() or not paths.v10_workbook.is_file():
        return ["Comparaison V9/V10 impossible: un classeur est absent"]
    old = openpyxl.load_workbook(paths.v9_workbook, read_only=True, data_only=False)
    new = openpyxl.load_workbook(paths.v10_workbook, read_only=True, data_only=False)
    allowed = {"README", "SOURCES", "DIM_PERSON", "CROSSWALK_GEO", "LOCAL_MANDATES", "LOCAL_COUNCIL_CONTROL", "QUALITY_CONTROL", "DATA_COVERAGE", "WORKBOOK_AUDIT_V9", "WORKBOOK_AUDIT_V10", "IDENTITY_AUDIT_V10", "MANUAL_RESOLUTIONS_V10", "COUNCIL_SEAT_STATUS_V10"}
    for name in sorted((set(old.sheetnames) & set(new.sheetnames)) - allowed):
        left = list(old[name].iter_rows(min_row=4, values_only=True))
        right = list(new[name].iter_rows(min_row=4, values_only=True))
        if left != right:
            errors.append(f"Différence V9/V10 non autorisée dans {name}")
    expected_new = {"IDENTITY_AUDIT_V10", "MANUAL_RESOLUTIONS_V10", "COUNCIL_SEAT_STATUS_V10", "WORKBOOK_AUDIT_V10"}
    actual_new = set(new.sheetnames) - set(old.sheetnames)
    if actual_new != expected_new:
        errors.append(f"Onglets V10 nouveaux inattendus: {sorted(actual_new)}")
    old.close()
    new.close()
    return errors


def validate_full(manifest: dict, data_dir: str | Path | None = None, release: str = "all", baseline: str | None = None) -> list[str]:
    errors: list[str] = []
    for record in [*manifest["physical_files"], *manifest["evidence"]]:
        errors.extend(validate_physical_record(record, data_dir))
    for source in manifest["sources"]:
        errors.extend(validate_physical_record(source, data_dir))
        path = resolve_manifest_path(source["local_path"], data_dir)
        if path.suffix.lower() == ".xlsx" and path.is_file():
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            sheet = workbook.worksheets[source.get("sheet_index", 0)]
            actual_rows, actual_columns = sheet.max_row - 1, sheet.max_column
            workbook.close()
            if actual_rows != source["rows"] or actual_columns != source["columns"]:
                errors.append(f"{source['source_id']}: dimensions attendues {source['rows']}×{source['columns']}, obtenues {actual_rows}×{actual_columns}")

    for artifact in manifest["artifacts"]:
        errors.extend(validate_physical_record(artifact, data_dir))
        path = resolve_manifest_path(artifact["local_path"], data_dir)
        if not path.is_file():
            continue
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        if len(workbook.sheetnames) != artifact["sheet_count"]:
            errors.append(f"{artifact['artifact_id']}: {len(workbook.sheetnames)} onglets au lieu de {artifact['sheet_count']}")
        volumes = EXPECTED_V9_VOLUMES if artifact["artifact_id"] == "WAREHOUSE_V9_EXPORT" else EXPECTED_V10_VOLUMES if artifact["artifact_id"] == "WAREHOUSE_V10_EXPORT" else {}
        for sheet_name, expected in volumes.items():
            actual = _populated_rows(workbook[sheet_name])
            if actual != expected:
                errors.append(f"{artifact['artifact_id']}.{sheet_name}: {actual} lignes au lieu de {expected}")
        workbook.close()

    paths = get_paths(data_dir)
    if release in {"v9", "all"} and paths.v9_workbook.is_file():
        from morocco_elections.legacy.v9 import documentation as module
        module.configure_paths(data_dir)
        before_hash = sha256(paths.v9_workbook)
        model = module.load_model()
        generated, metric_rules = module.build_docs(model)
        try:
            module.validate(model, generated, metric_rules, before_hash)
        except RuntimeError as exc:
            errors.append(str(exc))
        for name, content in generated.items():
            expected = content.replace("\r\n", "\n").rstrip() + "\n"
            actual = (DOCUMENTATION_DIRS["v9"] / name).read_text(encoding="utf-8-sig").replace("\r\n", "\n")
            if actual != expected:
                errors.append(f"Documentation désynchronisée du V9: {name}")
        if sha256(paths.v9_workbook) != before_hash:
            errors.append("Le V9 a été modifié pendant la validation")

    if release in {"v10", "all"} and paths.v10_workbook.is_file():
        from morocco_elections.releases.v10 import documentation as module10
        module10.configure_paths(data_dir)
        before_hash = sha256(paths.v10_workbook)
        model = module10.base.load_model()
        generated, _ = module10.base.build_docs(model)
        generated = module10._adapt(generated, model)
        try:
            module10._validate(model, generated, before_hash)
        except RuntimeError as exc:
            errors.append(str(exc))
        for name, content in generated.items():
            expected = content.replace("\r\n", "\n").rstrip() + "\n"
            actual = (DOCUMENTATION_DIRS["v10"] / name).read_text(encoding="utf-8-sig").replace("\r\n", "\n")
            if actual != expected:
                errors.append(f"Documentation désynchronisée du V10: {name}")
        if sha256(paths.v10_workbook) != before_hash:
            errors.append("Le V10 a été modifié pendant la validation")
        errors.extend(compare_v10_to_v9(data_dir))
    return errors


def run(mode: str, data_dir: str | Path | None = None, release: str = "all", baseline: str | None = None) -> list[str]:
    manifest = load_manifest()
    errors = []
    errors.extend(validate_manifest(manifest))
    errors.extend(validate_backlog())
    errors.extend(validate_v10_release_report(manifest))
    errors.extend(validate_python_sources())
    errors.extend(validate_documentation(release))
    errors.extend(validate_repository_files())
    if mode == "full" and not errors:
        errors.extend(validate_full(manifest, data_dir, release, baseline))
    return errors


def report(mode: str, data_dir: str | Path | None = None, release: str = "all", baseline: str | None = None) -> int:
    errors = run(mode, data_dir, release, baseline)
    if errors:
        print("VALIDATION_FAILED")
        for error in errors:
            print(f"- {error}")
        return 1
    document_count = len(EXPECTED_DOCUMENTS) * (2 if release == "all" else 1)
    print(f"VALIDATION_OK mode={mode} release={release} documents={document_count}")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Valide le dépôt et les releases locales V9/V10.")
    parser.add_argument("--mode", choices=("ci", "full"), default="ci")
    parser.add_argument("--data-dir")
    parser.add_argument("--release", choices=("v9", "v10", "all"), default="all")
    parser.add_argument("--baseline", choices=("v9",))
    args = parser.parse_args(argv)
    return report(args.mode, args.data_dir, args.release, args.baseline)


if __name__ == "__main__":
    sys.exit(main())
