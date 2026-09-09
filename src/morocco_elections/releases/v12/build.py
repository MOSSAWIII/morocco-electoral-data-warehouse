from __future__ import annotations

import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from morocco_elections.config import PROJECT_ROOT, get_paths
from morocco_elections.domains.identity import person_name_match_key
from morocco_elections.legacy.v9 import build as legacy
from morocco_elections.provenance import sha256_file
from morocco_elections.releases.v11 import build as v11
from morocco_elections.research import parliamentary_questions as research


VERSION = "V12"
GENERATED_ON = "2026-09-09"
QUALIFICATION_REPORT = PROJECT_ROOT / "metadata" / "v12_parliament_qualification.json"
RELEASE_REPORT = PROJECT_ROOT / "metadata" / "v12_release_report.json"
DIFF_REPORT = PROJECT_ROOT / "docs" / "research" / "V12_VS_V11_DIFF.txt"
EXPECTED_QUESTIONS = 5_589
EXPECTED_LINKED_ROWS = 5_453
EXPECTED_ANSWERED = 3_015
ALLOWED_CHANGED_SHEETS = {
    "README",
    "SOURCES",
    "PARLIAMENTARY_QUESTIONS",
    "DATA_DICTIONARY",
    "DATA_COVERAGE",
    "QUALITY_CONTROL",
    "WORKBOOK_AUDIT_V12",
}


def canonical_source_inputs(data_dir: str | Path | None = None) -> tuple[Path, ...]:
    paths = get_paths(data_dir)
    return (*v11.canonical_source_inputs(data_dir), *paths.parliament_questions_2023_sources)


def _save_deterministic(workbook: openpyxl.Workbook, target: Path) -> None:
    fixed = datetime(2026, 9, 9)
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    workbook.properties.lastModifiedBy = "morocco_elections V12"
    workbook.save(target)
    canonical = target.with_suffix(".canonical.tmp")
    if canonical.exists():
        canonical.unlink()
    with zipfile.ZipFile(target, "r") as source, zipfile.ZipFile(
        canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as output:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 9, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(
                    rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
                    rb'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-09T00:00:00Z</dcterms:modified>',
                    payload,
                )
            output.writestr(info, payload)
    canonical.replace(target)


def _identity_context(workbook: openpyxl.Workbook) -> tuple[dict[str, set[str]], dict[str, str]]:
    headers, rows = legacy.rows_from_sheet(workbook["PARLIAMENTARY_MANDATES"])
    if "full_name_ar" not in headers:
        raise RuntimeError("PARLIAMENTARY_MANDATES sans full_name_ar")
    by_name: dict[str, set[str]] = {}
    party_by_person: dict[str, str] = {}
    for row in rows:
        if str(row.get("legislature")) != "2021-2026":
            continue
        key = person_name_match_key(row.get("full_name_ar"))
        if key:
            by_name.setdefault(key, set()).add(str(row["person_id"]))
        if row.get("party_id"):
            party_by_person[str(row["person_id"])] = str(row["party_id"])
    return by_name, party_by_person


def _question_rows(workbook: openpyxl.Workbook, source_paths: tuple[Path, ...]) -> list[dict[str, Any]]:
    by_name, party_by_person = _identity_context(workbook)
    output: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    for path in sorted(source_paths, key=lambda item: item.name):
        spec = research.SOURCE_SPECS[path.name]
        for row in research.load_question_rows(path):
            number = str(row["numero"]).strip()
            question_id = f"PQ_V12_{str(spec['resource_id']).replace('-', '').upper()}_{number}"
            if question_id in seen_ids:
                raise RuntimeError(f"Question dupliquée: {question_id}")
            seen_ids.add(question_id)
            matches = by_name.get(person_name_match_key(row["depute"]), set())
            if len(matches) > 1:
                raise RuntimeError(f"Identité parlementaire ambiguë à la ligne {row['source_row']} de {path.name}")
            person_id = next(iter(matches)) if matches else None
            response_date = research._iso_date(row["date_reponse"], allow_dash=True)
            output.append(
                {
                    "question_id": question_id,
                    "source_question_number": number,
                    "question_type": "written",
                    "legislature": "2021-2026",
                    "period_raw": row["periode"],
                    "deposit_date": research._iso_date(row["date_depot"]),
                    "person_id": person_id,
                    "party_id": party_by_person.get(person_id) if person_id else None,
                    "deputy_name_ar_raw": row["depute"],
                    "parliamentary_group_ar_raw": row["groupe"],
                    "ministry_ar_raw": row["ministere"],
                    "subject_ar": row["sujet"],
                    "question_text_ar": row["observation"],
                    "response_date": response_date,
                    "response_status": "answered" if response_date else "unanswered_as_published",
                    "source_id": spec["source_id"],
                    "source_url": spec["url"],
                    "source_row": row["source_row"],
                    "identity_match_method": "unicode_exact_name" if person_id else "unresolved_no_fuzzy_match",
                    "quality_status": "verified_linked" if person_id else "verified_unlinked_identity",
                    "notes": "Source officielle; nom RAW conservé; response_status dérivé uniquement de date_reponse publiée.",
                }
            )
    return output


def _add_questions(workbook: openpyxl.Workbook, source_paths: tuple[Path, ...]) -> list[dict[str, Any]]:
    rows = _question_rows(workbook, source_paths)
    if len(rows) != EXPECTED_QUESTIONS:
        raise RuntimeError(f"Volume parlementaire inattendu: {len(rows)}")
    linked = sum(row["person_id"] is not None for row in rows)
    answered = sum(row["response_date"] is not None for row in rows)
    if linked != EXPECTED_LINKED_ROWS or answered != EXPECTED_ANSWERED:
        raise RuntimeError(f"Contrôles parlementaires inattendus: linked={linked}, answered={answered}")
    headers = list(rows[0])
    worksheet = legacy.replace_standard_sheet(
        workbook,
        "PARLIAMENTARY_QUESTIONS",
        "V12 — questions parlementaires écrites",
        "Périmètre officiel partiel: quatre segments du cycle 2023–2024; identités non raccordées conservées sans fuzzy matching.",
        headers,
        [[row[header] for header in headers] for row in rows],
    )
    text_columns = {headers.index(name) + 1 for name in ("subject_ar", "question_text_ar")}
    for row_index in range(5, worksheet.max_row + 1):
        for column_index in text_columns:
            cell = worksheet.cell(row_index, column_index)
            if isinstance(cell.value, str):
                cell.data_type = "s"
    return rows


def _append_sources(workbook: openpyxl.Workbook, source_paths: tuple[Path, ...]) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["SOURCES"])
    existing = {str(row["source_id"]) for row in rows}
    for path in sorted(source_paths, key=lambda item: item.name):
        spec = research.SOURCE_SPECS[path.name]
        if spec["source_id"] in existing:
            raise RuntimeError(f"Source déjà présente: {spec['source_id']}")
        rows.append(
            {
                "source_id": spec["source_id"],
                "source_name": f"Questions parlementaires écrites — {spec['segment']}",
                "producer": "Parlement du Royaume du Maroc",
                "source_tier": "official primary",
                "domain": "parliamentary activity",
                "source_url": spec["url"],
                "coverage_start": "2023",
                "coverage_end": "2024",
                "geographic_level": "national institution",
                "grain": "parliamentary question",
                "format": "XLSX",
                "access_method": "direct download via data.gov.ma",
                "license_or_terms": "Open Data Commons Open Database License (ODbL)",
                "confidence_score": 100,
                "retrieval_date": GENERATED_ON,
                "ingestion_status": "ingested_partial_scope",
                "notes": f"SHA256={sha256_file(path)}; question écrite uniquement; noms publics des députés.",
            }
        )
    legacy.rewrite_existing(workbook["SOURCES"], headers, rows)
    workbook["SOURCES"]["A1"] = "V12 — registre des sources"


def _append_dictionary(workbook: openpyxl.Workbook) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["DATA_DICTIONARY"])
    rows.append(
        {
            "metric_id": "parliamentary_question_count",
            "domain": "PARLIAMENT",
            "metric_name": "Nombre de questions parlementaires écrites",
            "definition": "Comptage des question_id observés dans le périmètre source explicite.",
            "data_type": "integer",
            "unit": "questions",
            "preferred_geo_granularity": "none",
            "preferred_temporal_frequency": "date",
            "primary_fact_sheet": "PARLIAMENTARY_QUESTIONS",
            "required_keys": "question_id",
            "candidate_sources": "; ".join(spec["source_id"] for spec in research.SOURCE_SPECS.values()),
            "collection_status": "observed_partial",
            "priority": "P0_core",
            "quality_rule": "Count only source question records; never infer missing questions or identities.",
            "notes": "V12 covers written questions in four 2023–2024 segments, not all parliamentary activity.",
        }
    )
    legacy.rewrite_existing(workbook["DATA_DICTIONARY"], headers, rows)
    workbook["DATA_DICTIONARY"]["A1"] = "V12 — dictionnaire des métriques"


def _append_quality(workbook: openpyxl.Workbook) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["QUALITY_CONTROL"])
    additions = [
        ("QA_V12_QUESTION_GRAIN", "grain_uniqueness", "5,589 question_id et clés source uniques", "passed"),
        ("QA_V12_IDENTITY_LINK", "identity_linkage", "5,453/5,589 lignes liées exactement; 136 non liées; zéro ambiguïté", "partial"),
        ("QA_V12_RESPONSE", "response_semantics", "3,015 dates de réponse publiées; 2,574 tirets interprétés sans inventer de date", "passed"),
        ("QA_V12_SCOPE", "scope_control", "Questions écrites, quatre segments 2023–2024 seulement; aucune prétention d'exhaustivité générale", "partial"),
    ]
    source_ids = "; ".join(spec["source_id"] for spec in research.SOURCE_SPECS.values())
    for issue_id, issue_type, description, resolution in additions:
        rows.append(
            {
                "issue_id": issue_id,
                "detected_date": GENERATED_ON,
                "sheet_name": "PARLIAMENTARY_QUESTIONS",
                "record_or_range": "all rows",
                "issue_type": issue_type,
                "severity": "info" if resolution == "passed" else "medium",
                "description": description,
                "expected_rule": "Observed scope and unresolved links remain explicit.",
                "resolution": resolution,
                "resolved_flag": 1 if resolution == "passed" else 0,
                "resolved_date": GENERATED_ON if resolution == "passed" else None,
                "owner": "data engineering",
                "source_id": source_ids,
                "notes": "Validated by deterministic V12 build.",
            }
        )
    legacy.rewrite_existing(workbook["QUALITY_CONTROL"], headers, rows)
    workbook["QUALITY_CONTROL"]["A1"] = "V12 — contrôles qualité"


def _update_coverage(workbook: openpyxl.Workbook) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["DATA_COVERAGE"])
    rows.append(
        {
            "domain_sheet": "PARLIAMENTARY_QUESTIONS",
            "rows_loaded": EXPECTED_QUESTIONS,
            "verified_rows": EXPECTED_QUESTIONS,
            "provisional_rows": 0,
            "derived_rows": 0,
            "low_confidence_rows": EXPECTED_QUESTIONS - EXPECTED_LINKED_ROWS,
            "primary_key": "question_id",
            "granularity_now": "source × parliamentary question",
            "next_granularity_target": "additional official segments only after the same source gate",
        }
    )
    legacy.rewrite_existing(workbook["DATA_COVERAGE"], headers, rows)
    workbook["DATA_COVERAGE"]["A1"] = "V12 — couverture par table"


def _append_readme(workbook: openpyxl.Workbook) -> None:
    worksheet = workbook["README"]
    for row in (
        ("VERSION V12 — ACTIVITÉ PARLEMENTAIRE ÉCRITE", None),
        ("version", "V12"),
        ("date", GENERATED_ON),
        ("addition", "5,589 questions écrites officielles dans PARLIAMENTARY_QUESTIONS"),
        ("identity linkage", "5,453 exactes; 136 non raccordées conservées; zéro fuzzy matching"),
        ("scope", "quatre segments 2023–2024; ne représente pas toute l'activité parlementaire"),
    ):
        worksheet.append(row)


def _rebuild_audit(workbook: openpyxl.Workbook, baseline_sheets: set[str]) -> None:
    del workbook["WORKBOOK_AUDIT_V11"]
    legacy.audit_workbook(workbook, baseline_sheets - {"WORKBOOK_AUDIT_V11"})
    worksheet = workbook["WORKBOOK_AUDIT_V9"]
    worksheet.title = "WORKBOOK_AUDIT_V12"
    worksheet["A1"] = "V12 — audit exhaustif du classeur"
    headers, rows = legacy.rows_from_sheet(worksheet)
    changed = ALLOWED_CHANGED_SHEETS - {"WORKBOOK_AUDIT_V12"}
    for row in rows:
        if row["sheet"] == "PARLIAMENTARY_QUESTIONS":
            row["origin_version"] = "V12 added"
        elif row["sheet"] in changed:
            row["origin_version"] = "V12 updated"
        else:
            row["origin_version"] = "V11 preserved"
    legacy.rewrite_existing(worksheet, headers, rows)
    if len(workbook.sheetnames) != 68:
        raise RuntimeError(f"V12 doit contenir 68 onglets, obtenu {len(workbook.sheetnames)}")


def _sheet_values(workbook: openpyxl.Workbook, name: str) -> list[list[Any]]:
    return [list(row) for row in workbook[name].iter_rows(values_only=True)]


def compare_v12_to_v11(v11_path: Path, v12_path: Path) -> tuple[list[str], list[str]]:
    before = openpyxl.load_workbook(v11_path, read_only=True, data_only=False)
    after = openpyxl.load_workbook(v12_path, read_only=True, data_only=False)
    mapped = {("WORKBOOK_AUDIT_V12" if name == "WORKBOOK_AUDIT_V11" else name): name for name in before.sheetnames}
    expected = set(mapped) | {"PARLIAMENTARY_QUESTIONS"}
    if expected != set(after.sheetnames):
        raise RuntimeError("La structure des onglets diffère entre V11 et V12")
    changed: list[str] = ["PARLIAMENTARY_QUESTIONS"]
    unexpected: list[str] = []
    for target_name, source_name in mapped.items():
        if _sheet_values(before, source_name) != _sheet_values(after, target_name):
            changed.append(target_name)
            if target_name not in ALLOWED_CHANGED_SHEETS:
                unexpected.append(target_name)
    before.close()
    after.close()
    return sorted(changed), sorted(unexpected)


def render_diff_report(report: dict[str, Any]) -> str:
    return "\n".join(
        [
            "V12 — RAPPORT D'ÉCARTS AVEC V11",
            "",
            f"DATE : {report['generated_on']}",
            f"V11 SHA-256 : {report['workbooks']['V11']}",
            f"V12 SHA-256 : {report['workbooks']['V12']}",
            "",
            "CHANGEMENTS AUTORISÉS ET OBSERVÉS",
            *[f"- {name}" for name in report["actual_changed_sheets"]],
            "",
            "VOLUMES ET LIMITES",
            f"- Questions écrites observées : {report['volumes']['PARLIAMENTARY_QUESTIONS']}",
            f"- Identités raccordées exactement : {report['controls']['identity_rows_linked']}",
            f"- Identités non raccordées : {report['controls']['identity_rows_unlinked']}",
            f"- Questions avec date de réponse : {report['controls']['answered_rows']}",
            "- Couverture : quatre segments officiels du cycle 2023–2024; aucune extrapolation à toute l'activité parlementaire.",
            "",
            "V11 et tous les RAW antérieurs sont restés inchangés.",
            "",
        ]
    )


def build_from_sources(
    data_dir: str | Path | None = None,
    *,
    target: Path | None = None,
    report_output: Path | None = None,
) -> dict[str, Any]:
    paths = get_paths(data_dir)
    target = target or paths.v12_workbook
    required = canonical_source_inputs(data_dir)
    for path in required:
        if not path.is_file():
            raise FileNotFoundError(path)
    qualification = json.loads(QUALIFICATION_REPORT.read_text(encoding="utf-8"))
    if qualification.get("decision") != "GO":
        raise RuntimeError("La qualification parlementaire versionnée n'est pas GO")
    expected_hashes = {item["file_name"]: item["sha256"] for item in qualification["resources"]}
    for path in paths.parliament_questions_2023_sources:
        if expected_hashes.get(path.name) != sha256_file(path):
            raise RuntimeError(f"Source parlementaire différente de la qualification: {path.name}")

    input_hashes = {str(path): sha256_file(path) for path in required}
    target.parent.mkdir(parents=True, exist_ok=True)
    baseline_temp = target.parent / ".v12-v11-reconstructed.tmp.xlsx"
    if baseline_temp.exists():
        baseline_temp.unlink()
    workbook = None
    baseline_report: dict[str, Any] | None = None
    rows: list[dict[str, Any]] = []
    try:
        baseline_report = v11.build_from_sources(data_dir, target=baseline_temp)
        if baseline_report["workbooks"]["V11"] != qualification["baseline_sha256"]:
            raise RuntimeError("La reconstruction directe de V11 diffère de la qualification")
        workbook = openpyxl.load_workbook(baseline_temp, data_only=False)
        baseline_sheets = set(workbook.sheetnames)
        rows = _add_questions(workbook, paths.parliament_questions_2023_sources)
        _append_sources(workbook, paths.parliament_questions_2023_sources)
        _append_dictionary(workbook)
        _append_quality(workbook)
        _update_coverage(workbook)
        _append_readme(workbook)
        _rebuild_audit(workbook, baseline_sheets)
        _save_deterministic(workbook, target)
        if any(sha256_file(Path(path)) != digest for path, digest in input_hashes.items()):
            raise RuntimeError("Un fichier source a changé pendant la construction")
        changed, unexpected = compare_v12_to_v11(baseline_temp, target)
        if unexpected or set(changed) != ALLOWED_CHANGED_SHEETS:
            raise RuntimeError(f"Différences V12 non conformes: changed={changed}; unexpected={unexpected}")
    finally:
        if workbook is not None:
            workbook.close()
        if baseline_temp.exists():
            baseline_temp.unlink()
    if baseline_report is None:
        raise AssertionError("Rapport V11 non construit")

    report = {
        "release": "V12",
        "baseline": "V11",
        "generated_on": GENERATED_ON,
        "status": "validated",
        "scope_status": "PARTIAL",
        "workbooks": {**baseline_report["workbooks"], "V12": sha256_file(target)},
        "source_hashes": {
            research.SOURCE_SPECS[path.name]["source_id"]: sha256_file(path)
            for path in paths.parliament_questions_2023_sources
        },
        "volumes": {"sheets": 68, "PARLIAMENTARY_QUESTIONS": len(rows)},
        "controls": {
            "unique_question_ids": len({row["question_id"] for row in rows}),
            "identity_rows_linked": sum(row["person_id"] is not None for row in rows),
            "identity_rows_unlinked": sum(row["person_id"] is None for row in rows),
            "answered_rows": sum(row["response_date"] is not None for row in rows),
            "unanswered_rows": sum(row["response_date"] is None for row in rows),
            "fuzzy_matches": 0,
        },
        "allowed_changed_sheets": sorted(ALLOWED_CHANGED_SHEETS),
        "actual_changed_sheets": changed,
        "integrity": {
            "v11_reconstructed_from_sources": True,
            "source_hashes_unchanged": True,
            "deterministic_export": True,
        },
    }
    if report_output is not None:
        report_output.parent.mkdir(parents=True, exist_ok=True)
        report_output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


def main(data_dir: str | Path | None = None) -> None:
    paths = get_paths(data_dir)
    report = build_from_sources(data_dir, target=paths.v12_workbook, report_output=RELEASE_REPORT)
    DIFF_REPORT.write_text(render_diff_report(report), encoding="utf-8")
    print(
        f"V12_BUILD_OK path={paths.v12_workbook} questions={report['volumes']['PARLIAMENTARY_QUESTIONS']} "
        f"sha256={report['workbooks']['V12']}"
    )


if __name__ == "__main__":
    main()
