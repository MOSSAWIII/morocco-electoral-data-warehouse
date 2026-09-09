from __future__ import annotations

import hashlib
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from morocco_elections.config import PROJECT_ROOT, get_paths
from morocco_elections.legacy.v9 import build as legacy
from morocco_elections.research import hcp_indicators as hcp


VERSION = "V11"
GENERATED_ON = "2026-09-08"
V10_REPORT = PROJECT_ROOT / "metadata" / "v10_release_report.json"
RELEASE_REPORT = PROJECT_ROOT / "metadata" / "v11_release_report.json"
DIFF_REPORT = PROJECT_ROOT / "docs" / "research" / "V11_VS_V10_DIFF.txt"
SOURCE_2014 = "SRC_HCP_RGPH2014_INDIVIDUALS_V11"
SOURCE_2024 = "SRC_HCP_RGPH2024_INDICATORS_V11"
EXPECTED_FACT_BASELINE = 127
EXPECTED_FACT_V11 = 3_203
ALLOWED_CHANGED_SHEETS = {
    "README",
    "SOURCES",
    "FACT_OBSERVATION",
    "DATA_DICTIONARY",
    "DATA_COVERAGE",
    "QUALITY_CONTROL",
    "WORKBOOK_AUDIT_V11",
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _save_deterministic(workbook: openpyxl.Workbook, target: Path) -> None:
    fixed = datetime(2026, 9, 8)
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    workbook.properties.lastModifiedBy = "morocco_elections V11"
    workbook.save(target)
    canonical = target.with_suffix(".canonical.tmp")
    if canonical.exists():
        canonical.unlink()
    with zipfile.ZipFile(target, "r") as source, zipfile.ZipFile(
        canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as output:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 8, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(
                    rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
                    rb'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-08T00:00:00Z</dcterms:modified>',
                    payload,
                )
            output.writestr(info, payload)
    canonical.replace(target)


def _baseline_context(workbook_path: Path) -> dict[str, Any]:
    return hcp._baseline_context(workbook_path)


def _extract_2014(path: Path, context: dict[str, Any]) -> dict[str, tuple[int, str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["Indic.Ensemble"]
    output: dict[str, tuple[int, str]] = {}
    for row in worksheet.iter_rows(min_row=4, values_only=True):
        try:
            if not (all(row[index] not in (None, "") for index in (0, 1, 2, 4)) and row[5] in (None, "")):
                continue
            source_geo = hcp._canonical_2014(row)
            parent = f"MA-{int(row[0]):02d}-{int(row[1]):03d}"
        except (TypeError, ValueError):
            continue
        geo_id, method = hcp._match_geo(source_geo, row[7], parent, context, 2014)
        if geo_id is None:
            continue
        value = hcp._numeric(row[8])
        if value is None or value < 0 or not value.is_integer() or geo_id in output:
            raise RuntimeError(f"Valeur 2014 invalide ou dupliquée: {geo_id}")
        output[geo_id] = (int(value), method)
    workbook.close()
    return output


def _extract_2024(path: Path, context: dict[str, Any]) -> tuple[dict[str, tuple[int, str]], dict[str, int]]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    worksheet = workbook["Population"]
    municipal: dict[str, tuple[int, str]] = {}
    legal: dict[str, int] = {}
    for row_index in range(4, worksheet.max_row + 1):
        code_cell = worksheet.cell(row_index, 1)
        name = worksheet.cell(row_index, 2).value
        if code_cell.value is None or code_cell.number_format != '00"."000"."00"."00' or hcp._is_prefecture_aggregate(name):
            continue
        source_geo = hcp._canonical_2024(code_cell.value)
        parent = "-".join(source_geo.split("-")[:3])
        geo_id, method = hcp._match_geo(source_geo, name, parent, context, 2024)
        if geo_id is None:
            continue
        legal_value = hcp._numeric(worksheet.cell(row_index, 3).value)
        municipal_value = hcp._numeric(worksheet.cell(row_index, 4).value)
        if (
            legal_value is None
            or municipal_value is None
            or legal_value < 0
            or municipal_value < 0
            or not legal_value.is_integer()
            or not municipal_value.is_integer()
            or geo_id in municipal
        ):
            raise RuntimeError(f"Valeur 2024 invalide ou dupliquée: {geo_id}")
        legal[geo_id] = int(legal_value)
        municipal[geo_id] = (int(municipal_value), method)
    workbook.close()
    return municipal, legal


def _population_control(workbook: openpyxl.Workbook, universe: set[str], candidate: dict[str, int]) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["POPULATION"])
    existing = {
        str(row["geo_id"]): int(row["population_total"])
        for row in rows
        if str(row.get("geo_id")) in universe and row.get("population_total") is not None
    }
    if len(existing) != 1_538 or existing != candidate or sum(existing.values()) != 36_828_330:
        raise RuntimeError("La population légale 2024 ne concorde pas exactement avec V10")
    if headers[:4] != ["record_id", "geo_id", "time_id", "population_total"]:
        raise RuntimeError("Schéma POPULATION V10 inattendu")


def _append_fact_observations(
    workbook: openpyxl.Workbook,
    population_2014: dict[str, tuple[int, str]],
    municipal_2024: dict[str, tuple[int, str]],
) -> None:
    worksheet = workbook["FACT_OBSERVATION"]
    headers, rows = legacy.rows_from_sheet(worksheet)
    if len(rows) != EXPECTED_FACT_BASELINE:
        raise RuntimeError(f"FACT_OBSERVATION V10 inattendu: {len(rows)}")
    existing_ids = {str(row["observation_id"]) for row in rows}
    existing_grain = {(str(row.get("geo_id")), str(row.get("time_id")), str(row.get("metric_id"))) for row in rows}

    def add_rows(values: dict[str, tuple[int, str]], year: int, metric_id: str, source_id: str, url: str, publication: str) -> None:
        for geo_id, (value, match_method) in sorted(values.items()):
            observation_id = f"OBS_HCP_{year}_{metric_id.upper()}_{geo_id.replace('-', '_')}"
            grain = (geo_id, str(year), metric_id)
            if observation_id in existing_ids or grain in existing_grain:
                raise RuntimeError(f"Observation HCP dupliquée: {grain}")
            rows.append(
                {
                    "observation_id": observation_id,
                    "metric_id": metric_id,
                    "value_numeric": value,
                    "value_text": None,
                    "unit": "persons",
                    "geo_id": geo_id,
                    "time_id": str(year),
                    "election_id": None,
                    "party_id": None,
                    "person_id": None,
                    "date_start": None,
                    "date_end": None,
                    "geo_granularity": "commune_or_arrondissement",
                    "temporal_granularity": "census_year",
                    "method": f"observed_official; geo_match={match_method}",
                    "source_id": source_id,
                    "source_url": url,
                    "publication_date": publication,
                    "retrieval_date": GENERATED_ON,
                    "quality_status": "verified",
                    "notes": "Observation censitaire; aucune interpolation et aucune jointure électorale automatique.",
                }
            )
            existing_ids.add(observation_id)
            existing_grain.add(grain)

    add_rows(population_2014, 2014, "population_legal", SOURCE_2014, hcp.SOURCE_SPECS["RGPH2014_INDIVIDUALS"]["url"], "2017-05-30")
    add_rows(municipal_2024, 2024, "population_municipal", SOURCE_2024, hcp.SOURCE_SPECS["RGPH2024_INDICATORS"]["url"], "2024-12-17")
    if len(rows) != EXPECTED_FACT_V11:
        raise RuntimeError(f"FACT_OBSERVATION V11 inattendu: {len(rows)}")
    legacy.replace_standard_sheet(
        workbook,
        "FACT_OBSERVATION",
        "V11 — observations factuelles longues",
        "3 076 observations HCP communales validées ajoutées aux 127 observations V10; aucun indicateur NO_GO.",
        headers,
        [[row.get(header) for header in headers] for row in rows],
    )


def _append_sources(workbook: openpyxl.Workbook, paths) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["SOURCES"])
    existing = {str(row["source_id"]) for row in rows}
    specs = [
        (
            SOURCE_2014,
            "RGPH 2014 — Indicateurs communaux Individus",
            "population",
            hcp.SOURCE_SPECS["RGPH2014_INDIVIDUALS"]["url"],
            "2014",
            paths.hcp_individuals_2014,
            "Population légale uniquement; 1 538 unités; autres indicateurs exclus par V10-QA-3.",
        ),
        (
            SOURCE_2024,
            "RGPH 2024 — Indicateurs démographiques et socio-économiques",
            "population",
            hcp.SOURCE_SPECS["RGPH2024_INDICATORS"]["url"],
            "2024",
            paths.hcp_indicators_2024,
            "Population municipale uniquement; population légale utilisée comme contrôle; autres indicateurs exclus par V10-QA-3.",
        ),
    ]
    for source_id, name, domain, url, year, path, notes in specs:
        if source_id in existing:
            raise RuntimeError(f"Source V11 déjà présente: {source_id}")
        rows.append(
            dict(
                zip(
                    headers,
                    [
                        source_id,
                        name,
                        "Haut-Commissariat au Plan",
                        "official primary",
                        domain,
                        url,
                        year,
                        year,
                        "commune/arrondissement",
                        "census",
                        "XLSX",
                        "direct download",
                        "official public; attribution retained; source file not redistributed",
                        100,
                        GENERATED_ON,
                        "ingested_partial_go_only",
                        f"{notes} SHA256={sha256_file(path)}",
                    ],
                )
            )
        )
    legacy.rewrite_existing(workbook["SOURCES"], headers, rows)
    workbook["SOURCES"]["A1"] = "V11 — registre des sources"


def _append_dictionary(workbook: openpyxl.Workbook) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["DATA_DICTIONARY"])
    existing = {str(row["metric_id"]) for row in rows}
    additions = [
        ("population_legal", "Population légale publiée par le RGPH", "2014", SOURCE_2014),
        ("population_municipal", "Population municipale publiée par le RGPH", "2024", SOURCE_2024),
    ]
    for metric_id, definition, year, source_id in additions:
        if metric_id in existing:
            raise RuntimeError(f"Métrique déjà définie: {metric_id}")
        record = {
            "metric_id": metric_id,
            "domain": "POPULATION",
            "metric_name": metric_id,
            "definition": definition,
            "data_type": "integer",
            "unit": "persons",
            "preferred_geo_granularity": "commune/arrondissement",
            "preferred_temporal_frequency": "census",
            "primary_fact_sheet": "FACT_OBSERVATION",
            "required_keys": "geo_id,time_id,metric_id",
            "candidate_sources": source_id,
            "collection_status": "observed_verified",
            "priority": "P0_core",
            "quality_rule": "Exactly 1,538 positive integer values; no interpolation; source year mandatory.",
            "notes": f"V11 GO for observation year {year}; temporal permissions documented in V11-QA.",
        }
        rows.append(record)
    legacy.rewrite_existing(workbook["DATA_DICTIONARY"], headers, rows)
    workbook["DATA_DICTIONARY"]["A1"] = "V11 — dictionnaire des métriques"


def _append_quality(workbook: openpyxl.Workbook) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["QUALITY_CONTROL"])
    source_ids = f"{SOURCE_2014}; {SOURCE_2024}"
    additions = [
        ("QA_V11_HCP_GRAIN", "FACT_OBSERVATION", "3,076 rows", "grain_uniqueness", "Two complete metric-year universes; zero duplicate geo_id×time_id×metric_id.", "exactly 1,538 unique units per metric-year", source_ids),
        ("QA_V11_HCP_TOTALS", "FACT_OBSERVATION", "HCP V11 rows", "aggregate_reconciliation", "2014 legal population=33,848,242; 2024 municipal population=36,490,591.", "national sums equal qualified official exports", source_ids),
        ("QA_V11_HCP_2024_CONTROL", "POPULATION", "1,538 analytical units", "cross_source_reconciliation", "New official 2024 legal population equals V10 cell by cell; total=36,828,330.", "zero value difference", SOURCE_2024),
        ("QA_V11_HCP_NO_GO_EXCLUSION", "FACT_OBSERVATION", "25 V10-QA-3 decisions", "scope_control", "No NO_GO indicator was ingested; no derived value or interpolation was created.", "only population_legal 2014 and population_municipal 2024 added", source_ids),
    ]
    for issue_id, sheet, record_range, issue_type, description, rule, source_id in additions:
        rows.append(
            {
                "issue_id": issue_id,
                "detected_date": GENERATED_ON,
                "sheet_name": sheet,
                "record_or_range": record_range,
                "issue_type": issue_type,
                "severity": "info",
                "description": description,
                "expected_rule": rule,
                "resolution": "passed",
                "resolved_flag": 1,
                "resolved_date": GENERATED_ON,
                "owner": "data engineering",
                "source_id": source_id,
                "notes": "Validated by deterministic V11 build.",
            }
        )
    legacy.rewrite_existing(workbook["QUALITY_CONTROL"], headers, rows)
    workbook["QUALITY_CONTROL"]["A1"] = "V11 — contrôles qualité"


def _update_coverage(workbook: openpyxl.Workbook) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["DATA_COVERAGE"])
    if any(row.get("domain_sheet") == "FACT_OBSERVATION" for row in rows):
        raise RuntimeError("Couverture FACT_OBSERVATION déjà présente")
    rows.append(
        {
            "domain_sheet": "FACT_OBSERVATION",
            "rows_loaded": EXPECTED_FACT_V11,
            "verified_rows": 3_203,
            "provisional_rows": 0,
            "derived_rows": 0,
            "low_confidence_rows": 0,
            "primary_key": "observation_id",
            "granularity_now": "geography × time × metric",
            "next_granularity_target": "only indicator-year pairs independently qualified GO",
        }
    )
    legacy.rewrite_existing(workbook["DATA_COVERAGE"], headers, rows)
    workbook["DATA_COVERAGE"]["A1"] = "V11 — couverture par table"


def _append_readme(workbook: openpyxl.Workbook) -> None:
    worksheet = workbook["README"]
    for row in (
        ("VERSION V11 — HCP MINIMAL", None),
        ("version", "V11"),
        ("date", GENERATED_ON),
        ("HCP additions", "1,538 population_legal 2014 + 1,538 population_municipal 2024 in FACT_OBSERVATION"),
        ("2024 control", "1,538/1,538 population_legal values equal V10; no duplicate fact created"),
        ("temporal rule", "no interpolation; source observation year and distance to election must remain explicit"),
        ("excluded", "25 V10-QA-3 NO_GO decisions; electoral panels unchanged"),
    ):
        worksheet.append(row)


def _rebuild_audit(workbook: openpyxl.Workbook, baseline_sheets: set[str]) -> None:
    del workbook["WORKBOOK_AUDIT_V10"]
    legacy.audit_workbook(workbook, baseline_sheets - {"WORKBOOK_AUDIT_V10"})
    worksheet = workbook["WORKBOOK_AUDIT_V9"]
    worksheet.title = "WORKBOOK_AUDIT_V11"
    worksheet["A1"] = "V11 — audit exhaustif du classeur"
    headers, rows = legacy.rows_from_sheet(worksheet)
    changed = ALLOWED_CHANGED_SHEETS - {"WORKBOOK_AUDIT_V11"}
    for row in rows:
        row["origin_version"] = "V11 updated" if row["sheet"] in changed else "V10 preserved"
        if row["sheet"] in {"IDENTITY_AUDIT_V10", "MANUAL_RESOLUTIONS_V10", "COUNCIL_SEAT_STATUS_V10"}:
            row["classification"] = "QA"
    legacy.rewrite_existing(worksheet, headers, rows)
    if len(workbook.sheetnames) != 67:
        raise RuntimeError(f"V11 doit contenir 67 onglets, obtenu {len(workbook.sheetnames)}")


def _sheet_values(workbook: openpyxl.Workbook, name: str) -> list[list[Any]]:
    worksheet = workbook[name]
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def compare_v11_to_v10(v10_path: Path, v11_path: Path) -> tuple[list[str], list[str]]:
    before = openpyxl.load_workbook(v10_path, read_only=True, data_only=False)
    after = openpyxl.load_workbook(v11_path, read_only=True, data_only=False)
    mapped_before = {("WORKBOOK_AUDIT_V11" if name == "WORKBOOK_AUDIT_V10" else name): name for name in before.sheetnames}
    if set(mapped_before) != set(after.sheetnames):
        raise RuntimeError("La structure des 67 onglets diffère entre V10 et V11")
    changed: list[str] = []
    unexpected: list[str] = []
    for target_name, source_name in mapped_before.items():
        if _sheet_values(before, source_name) != _sheet_values(after, target_name):
            changed.append(target_name)
            if target_name not in ALLOWED_CHANGED_SHEETS:
                unexpected.append(target_name)
    before.close()
    after.close()
    return sorted(changed), sorted(unexpected)


def render_diff_report(report: dict[str, Any]) -> str:
    return "\n".join(
        [
            "V11 — RAPPORT D'ÉCARTS AVEC V10",
            "",
            f"DATE : {report['generated_on']}",
            f"V10 SHA-256 : {report['workbooks']['V10']}",
            f"V11 SHA-256 : {report['workbooks']['V11']}",
            "",
            "CHANGEMENTS AUTORISÉS ET OBSERVÉS",
            *[f"- {name}" for name in report["actual_changed_sheets"]],
            "",
            "VOLUMES",
            f"- FACT_OBSERVATION : {report['volumes']['FACT_OBSERVATION_V10']} → {report['volumes']['FACT_OBSERVATION_V11']}",
            f"- population_legal 2014 : {report['volumes']['population_legal_2014']}",
            f"- population_municipal 2024 : {report['volumes']['population_municipal_2024']}",
            "",
            "CONTRÔLES",
            f"- Population légale 2014 : {report['controls']['population_legal_2014_sum']}",
            f"- Population municipale 2024 : {report['controls']['population_municipal_2024_sum']}",
            f"- Population légale 2024 : {report['controls']['population_legal_2024_sum']} ; différences avec V10={report['controls']['population_legal_2024_differences']}",
            f"- Décisions NO_GO ingérées : {report['controls']['no_go_ingested']}",
            "",
            "Aucun panel électoral, RAW historique ou classeur V10 n'a été modifié.",
            "",
        ]
    )


def main(data_dir: str | Path | None = None) -> None:
    paths = get_paths(data_dir)
    required = [paths.v10_workbook, paths.hcp_individuals_2014, paths.hcp_indicators_2024]
    for path in required:
        if not path.is_file():
            raise FileNotFoundError(path)
    v10_release = json.loads(V10_REPORT.read_text(encoding="utf-8"))
    v10_hash = sha256_file(paths.v10_workbook)
    if v10_hash != v10_release["workbooks"]["V10"]:
        raise RuntimeError("Empreinte V10 différente du rapport de release")
    input_hashes = {str(path): sha256_file(path) for path in required}

    context = _baseline_context(paths.v10_workbook)
    population_2014 = _extract_2014(paths.hcp_individuals_2014, context)
    municipal_2024, legal_2024 = _extract_2024(paths.hcp_indicators_2024, context)
    if set(population_2014) != context["universe"] or set(municipal_2024) != context["universe"]:
        raise RuntimeError("Couverture HCP différente de 1 538 unités")
    if sum(value for value, _ in population_2014.values()) != 33_848_242:
        raise RuntimeError("Total population légale 2014 inattendu")
    if sum(value for value, _ in municipal_2024.values()) != 36_490_591:
        raise RuntimeError("Total population municipale 2024 inattendu")

    workbook = openpyxl.load_workbook(paths.v10_workbook, data_only=False)
    baseline_sheets = set(workbook.sheetnames)
    _population_control(workbook, context["universe"], legal_2024)
    _append_fact_observations(workbook, population_2014, municipal_2024)
    _append_sources(workbook, paths)
    _append_dictionary(workbook)
    _append_quality(workbook)
    _update_coverage(workbook)
    _append_readme(workbook)
    _rebuild_audit(workbook, baseline_sheets)
    paths.v11_workbook.parent.mkdir(parents=True, exist_ok=True)
    _save_deterministic(workbook, paths.v11_workbook)
    workbook.close()

    if sha256_file(paths.v10_workbook) != v10_hash or any(sha256_file(Path(path)) != digest for path, digest in input_hashes.items()):
        raise RuntimeError("Un fichier source ou V10 a changé pendant la construction")
    changed, unexpected = compare_v11_to_v10(paths.v10_workbook, paths.v11_workbook)
    if unexpected or set(changed) != ALLOWED_CHANGED_SHEETS:
        raise RuntimeError(f"Différences V11 non conformes: changed={changed}; unexpected={unexpected}")

    report = {
        "release": "V11",
        "baseline": "V10",
        "generated_on": GENERATED_ON,
        "status": "validated",
        "workbooks": {**v10_release["workbooks"], "V11": sha256_file(paths.v11_workbook)},
        "source_hashes": {
            SOURCE_2014: sha256_file(paths.hcp_individuals_2014),
            SOURCE_2024: sha256_file(paths.hcp_indicators_2024),
        },
        "volumes": {
            "sheets": 67,
            "FACT_OBSERVATION_V10": EXPECTED_FACT_BASELINE,
            "FACT_OBSERVATION_V11": EXPECTED_FACT_V11,
            "population_legal_2014": len(population_2014),
            "population_municipal_2024": len(municipal_2024),
        },
        "controls": {
            "population_legal_2014_sum": sum(value for value, _ in population_2014.values()),
            "population_municipal_2024_sum": sum(value for value, _ in municipal_2024.values()),
            "population_legal_2024_sum": sum(legal_2024.values()),
            "population_legal_2024_differences": 0,
            "no_go_ingested": 0,
            "derived_values_ingested": 0,
            "electoral_panels_changed": False,
        },
        "allowed_changed_sheets": sorted(ALLOWED_CHANGED_SHEETS),
        "actual_changed_sheets": changed,
        "integrity": {"v10_unchanged": True, "source_hashes_unchanged": True, "deterministic_export": True},
    }
    RELEASE_REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    DIFF_REPORT.write_text(render_diff_report(report), encoding="utf-8")
    print(f"V11_BUILD_OK path={paths.v11_workbook} facts={EXPECTED_FACT_V11} sha256={report['workbooks']['V11']}")


if __name__ == "__main__":
    main()
