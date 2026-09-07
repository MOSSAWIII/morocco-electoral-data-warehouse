from __future__ import annotations

import hashlib
import re
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl

from morocco_elections.config import get_paths

VERSION = "V9"
AUDIT_SHEET = "WORKBOOK_AUDIT_V9"
GENERATED_ON = date.today().isoformat()


def configure_paths(data_dir: str | Path | None = None) -> None:
    paths = get_paths(data_dir)
    global ROOT, WORKBOOK, OUTPUT_DIR
    ROOT = paths.project_root
    WORKBOOK = paths.v9_workbook
    OUTPUT_DIR = paths.documentation_v9


configure_paths()

DOCUMENTS = [
    "00_INDEX_ET_MODE_EMPLOI.txt",
    "01_ONTOLOGIE_ELECTORALE_GLOBALE.txt",
    "02_ARCHITECTURE_DES_CUBES.txt",
    "03_DIMENSIONS_CONFORMES.txt",
    "04_HIERARCHIE_GEOGRAPHIQUE.txt",
    "05_CROSSWALKS_ET_IDENTITES.txt",
    "06_TEMPS_CYCLES_ET_ELECTIONS.txt",
    "07_FAITS_ELECTORAUX_ET_MOBILISATION.txt",
    "08_PERSONNES_MANDATS_ET_REPRESENTATION.txt",
    "09_POUVOIR_LOCAL_ET_GOUVERNANCE.txt",
    "10_SOCIO_ECONOMIE_CAMPAGNE_INFORMATION.txt",
    "11_PIPELINE_RAW_NORMALIZED_FACT_ANALYTICAL.txt",
    "12_REGLES_D_INTEGRITE_ET_JOINTURE.txt",
    "13_METRIQUES_DERIVEES_ET_REGLES_TEMPORELLES.txt",
    "14_CATALOGUE_COMPLET_DES_ONGLETS.txt",
]

STATUS_LEGEND = """STATUTS DOCUMENTAIRES
OBSERVÉ       : donnée effectivement présente dans V9.
DÉRIVÉ        : valeur calculée à partir de données V9 traçables.
STRUCTURE VIDE: schéma existant, mais aucune ligne de données chargée.
PILOTE        : ancien échantillon conservé pour mémoire ou validation.
BLOQUÉ        : donnée ou relation empêchée par une limitation QA explicite.
"""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_text(value) -> str:
    return "" if value is None else str(value).strip()


def standard_header(ws):
    if ws.max_row < 4:
        return 1, []
    row4 = list(next(ws.iter_rows(min_row=4, max_row=4, max_col=ws.max_column, values_only=True)))
    if any(value is not None for value in row4):
        return 4, [value for value in row4 if value is not None]
    row1 = list(next(ws.iter_rows(min_row=1, max_row=1, max_col=ws.max_column, values_only=True)))
    return 1, [value for value in row1 if value is not None]


def populated_row_count(ws, header_row: int) -> int:
    return sum(
        1
        for row in ws.iter_rows(
            min_row=header_row + 1,
            max_row=ws.max_row,
            max_col=ws.max_column,
            values_only=True,
        )
        if any(value is not None for value in row)
    )


def table_records(wb, sheet_name, header_row=4):
    ws = wb[sheet_name]
    headers = list(next(ws.iter_rows(min_row=header_row, max_row=header_row, max_col=ws.max_column, values_only=True)))
    records = []
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, max_col=ws.max_column, values_only=True):
        if any(value is not None for value in row):
            records.append(dict(zip(headers, row)))
    return [header for header in headers if header is not None], records


def load_model():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=False)
    audit_headers, audit_rows = table_records(wb, AUDIT_SHEET)
    audit = {row["sheet"]: row for row in audit_rows}
    coverage_headers, coverage_rows = table_records(wb, "DATA_COVERAGE")
    coverage = {row["domain_sheet"]: row for row in coverage_rows}
    qa_headers, qa_rows = table_records(wb, "QUALITY_CONTROL")
    dictionary_headers, dictionary_rows = table_records(wb, "DATA_DICTIONARY")
    crosswalk_headers, crosswalk_rows = table_records(wb, "CROSSWALK_GEO")
    local_headers, local_rows = table_records(wb, "LOCAL_MANDATES")

    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        header_row, headers = standard_header(ws)
        count = populated_row_count(ws, header_row)
        sheet_audit = audit.get(name, {})
        sheets[name] = {
            "name": name,
            "header_row": header_row,
            "headers": headers,
            "rows": count,
            "classification": sheet_audit.get("classification") or "NON CLASSÉ",
            "grain": sheet_audit.get("grain") or "non tabulaire / voir titre",
            "primary_key": sheet_audit.get("primary_key"),
            "foreign_keys": sheet_audit.get("foreign_keys"),
            "sources": sheet_audit.get("sources"),
            "coverage_period": sheet_audit.get("coverage_period"),
            "geographic_level": sheet_audit.get("geographic_level"),
            "quality_status": sheet_audit.get("quality_status"),
            "limitations": sheet_audit.get("known_limitations"),
            "origin_version": sheet_audit.get("origin_version"),
        }

    manual_geo = [row for row in crosswalk_rows if row.get("manual_review_flag") == 1]
    unique_geo_reviews = sorted(
        {
            (
                row.get("source_geo_id"),
                row.get("source_geo_name"),
                row.get("source_province"),
                row.get("canonical_geo_id"),
                row.get("canonical_name"),
                row.get("confidence_score"),
            )
            for row in manual_geo
        }
    )
    unresolved_qa = [row for row in qa_rows if row.get("resolved_flag") in (0, None, "0")]
    local_person_ids = {row.get("person_id") for row in local_rows if row.get("person_id")}
    local_communes = {row.get("source_idcommune") for row in local_rows if row.get("source_idcommune") is not None}
    local_names = {row.get("person_name_ar") for row in local_rows if row.get("person_name_ar")}
    return {
        "wb": wb,
        "sheets": sheets,
        "audit": audit,
        "coverage": coverage,
        "qa": qa_rows,
        "unresolved_qa": unresolved_qa,
        "dictionary": dictionary_rows,
        "geo_reviews": unique_geo_reviews,
        "local_identity": {
            "rows": len(local_rows),
            "unique_person_ids": len(local_person_ids),
            "unique_communes": len(local_communes),
            "unique_names": len(local_names),
        },
        "workbook_hash": sha256(WORKBOOK),
    }


def header(title: str, scope: str, references: list[str]) -> str:
    refs = ", ".join(references) if references else "00_INDEX_ET_MODE_EMPLOI.txt"
    return "\n".join(
        [
            title,
            "=" * len(title),
            "",
            f"VERSION           : {VERSION}",
            f"CLASSEUR SOURCE   : {WORKBOOK.name}",
            f"DATE DE GÉNÉRATION: {GENERATED_ON}",
            f"PÉRIMÈTRE         : {scope}",
            f"RENVOIS           : {refs}",
            "",
            STATUS_LEGEND.strip(),
            "",
            "",
        ]
    )


def section(title: str, body: str) -> str:
    return f"\n{title}\n{'-' * len(title)}\n{body.strip()}\n"


def format_rows(rows, columns):
    lines = []
    for row in rows:
        lines.append(" | ".join(f"{column}={normalize_text(row.get(column)) or 'NULL'}" for column in columns))
    return "\n".join(lines) if lines else "Aucune ligne."


def sheet_count(model, name):
    return model["sheets"][name]["rows"]


def build_docs(model):
    counts = {name: data["rows"] for name, data in model["sheets"].items()}
    empty_sheets = sorted(name for name, data in model["sheets"].items() if data["rows"] == 0)
    pilot_sheets = sorted(name for name in model["sheets"] if "PILOT" in name or "SAMPLE" in name or name.startswith("MICRO_"))

    docs = {}

    body = """
Ce corpus décrit l'ontologie, les cubes, les hiérarchies, les règles de jointure,
les métriques et la gouvernance des données du warehouse V9. Il ne remplace pas
le classeur : le classeur reste l'unique source de vérité exécutable.

ORDRE DE LECTURE RECOMMANDÉ
1. Ontologie globale.
2. Architecture des cubes.
3. Dimensions conformes et hiérarchie géographique.
4. Faits électoraux, personnes et pouvoir local.
5. Pipeline, intégrité, métriques et règles temporelles.
6. Catalogue exhaustif des 64 onglets.

CONVENTIONS
- Les identifiants techniques sont reproduits sans traduction.
- NULL signifie absent ou non calculable, jamais zéro implicite.
- Une part est exprimée sur l'échelle 0–100 sauf mention explicite.
- Une ligne RAW n'est jamais corrigée silencieusement.
- Une relation temporelle doit conserver observation_year, election_year,
  temporal_distance et temporal_join_method lorsqu'ils existent.
- Les anciens onglets SAMPLE/PILOT restent historiques et ne remplacent jamais
  les tables FULL ou les panels V9.
- BLOQUÉ: le person_id local V9 est un proxy communal, pas une identité d'élu;
  les analyses individuelles doivent utiliser les lignes RAW en attendant sa correction.

LISTE DES DOCUMENTS
""" + "\n".join(f"{index + 1:02d}. {name}" for index, name in enumerate(DOCUMENTS))
    body += section("EMPREINTE DE LA SOURCE", f"SHA-256 V9: {model['workbook_hash']}\nNombre d'onglets: {len(model['sheets'])}\nNombre de documents: {len(DOCUMENTS)}")
    body += section("STRUCTURES VIDES", "STRUCTURE VIDE: " + (", ".join(empty_sheets) if empty_sheets else "aucune"))
    body += section("PILOTES CONSERVÉS", "PILOTE: " + ", ".join(pilot_sheets))
    docs[DOCUMENTS[0]] = header("INDEX ET MODE D'EMPLOI", "Corpus documentaire complet de V9.", DOCUMENTS[1:4]) + body

    body = """
L'ontologie modélise l'élection marocaine comme une chaîne de faits reliés par
des dimensions conformes :

TERRITOIRE
  → POPULATION ET CONDITIONS SOCIO-ÉCONOMIQUES
  → ÉLECTORAT
  → OFFRE POLITIQUE
  → CAMPAGNE ET INFORMATION
  → MOBILISATION
  → VOTE
  → SIÈGES ET REPRÉSENTATION
  → CONTRÔLE POLITIQUE
  → GOUVERNANCE ET ACTION PUBLIQUE
  → ÉLECTION SUIVANTE

Les nœuds ne sont pas tous remplis au même niveau. POPULATION, RESULTS,
MOBILISATION, LOCAL_MANDATES et PARLIAMENTARY_MANDATES portent des observations.
POLITICAL_ACTIVITY, LOCAL_FINANCE et SHOCKS sont des STRUCTURES VIDES dans V9.

UNITÉS ONTOLOGIQUES FONDAMENTALES
- Lieu : geo_id, hiérarchisé et versionné.
- Temps : time_id, année, date, cycle et législature.
- Organisation politique : party_id, incluant partis, alliances et codes source.
- Personne : person_id, stable pour les députés TAFRA; BLOQUÉ au niveau local car
  la clé V9 actuelle se réduit à un proxy communal.
- Scrutin : election_id.
- Provenance : source_id.
- Observation : record_id, observation_id ou identifiant de mandat.

RELATIONS CAUSALES NON AUTORISÉES
Le warehouse permet des associations longitudinales, mais aucune table V9 ne
prouve qu'une condition socio-économique cause un vote, qu'une campagne cause
une mobilisation ou qu'une gouvernance cause une réélection. Toute analyse de
ce type doit rester descriptive tant qu'un protocole causal séparé n'existe pas.
"""
    body += section("VOLUMES CENTRAUX OBSERVÉS", f"OBSERVÉ RESULTS: {counts['RESULTS']} lignes totales, dont 22 054 lignes communales V9.\nOBSERVÉ LOCAL_MANDATES: {counts['LOCAL_MANDATES']} lignes.\nOBSERVÉ PARLIAMENTARY_MANDATES: {counts['PARLIAMENTARY_MANDATES']} lignes.\nDÉRIVÉ ELECTORAL_TRANSITIONS_2015_2021: {counts['ELECTORAL_TRANSITIONS_2015_2021']} lignes.\nDÉRIVÉ COMMUNE_ELECTION_PANEL: {counts['COMMUNE_ELECTION_PANEL']} lignes.")
    docs[DOCUMENTS[1]] = header("ONTOLOGIE ÉLECTORALE GLOBALE", "Entités, événements et relations conceptuelles de V9.", [DOCUMENTS[2], DOCUMENTS[11], DOCUMENTS[14]]) + body

    body = f"""
CUBE CENTRAL A — COMMUNE × PARTI × ÉLECTION
Table de restitution : ANALYTICAL_PANEL.
Grain : une ligne par geo_id × party_id × election_id.
Volume : {counts['ANALYTICAL_PANEL']} lignes.
Dimensions conformes : DIM_GEO, DIM_PARTY, DIM_ELECTION, DIM_TIME.
Mesures : party_votes, vote_share, rank, winner, seats, seat_share,
previous_vote_share, swing_pp, vote_change et seat_change.
Population : population_reference est une référence RGPH 2024, jamais une
observation contemporaine de 2015 ou 2021.

CUBE CENTRAL B — COMMUNE × ÉLECTION
Table de restitution : COMMUNE_ELECTION_PANEL.
Grain : une ligne par geo_id × election_id.
Volume : {counts['COMMUNE_ELECTION_PANEL']} lignes = 1 538 unités × 2 scrutins.
Mesures : turnout, winning_share, margins, hhi, enp et volatility.
Attributs de pouvoir 2021 : largest_party et president_party, sans les confondre.

CUBE DES TRANSITIONS
ELECTORAL_TRANSITIONS_2015_2021 : commune × parti, {counts['ELECTORAL_TRANSITIONS_2015_2021']} lignes.
COMMUNE_TRANSITION_PANEL : commune, {counts['COMMUNE_TRANSITION_PANEL']} lignes.
Le cube compare 2015 à 2021 mais laisse les sièges 2015 et control_change à NULL.

CUBE DES MANDATS LOCAUX
LOCAL_MANDATES : personne locale prudente × commune × circonscription × mandat.
LOCAL_COUNCIL_CONTROL : commune × parti × élection.
Les effectifs d'élus peuvent être additionnés par parti puis commune, sous réserve
des deux écarts de siège documentés en QA.

CUBE PARLEMENTAIRE
PARLIAMENTARY_MANDATES : personne × législature × siège × période de mandat.
idPerson est stable; idSiege peut être partagé par des remplaçants successifs.

CUBE SOCIO-ÉCONOMIQUE
Grain cible : geo_id × time_id. Les faits résident notamment dans POPULATION,
SOCIAL_HUMAN, LIVING_CONDITIONS, DIGITAL_ACCESS et FACT_OBSERVATION.
La couverture communale exhaustive actuelle concerne surtout POPULATION 2024.

ADDITIVITÉ
- Additif : votes et effectifs d'élus, seulement sur des partitions sans doublon.
- Semi-additif : population, sièges, inscrits et stocks; ne pas sommer dans le temps.
- Non additif : taux, parts, HHI, ENP, rang, marge en points et volatilité.
- Recalcul obligatoire après agrégation : toute part ou indice dérivé.

CHEMINS DE FORAGE
Région → province/préfecture → commune/arrondissement.
Cycle → élection → parti → résultat.
Personne → mandat → législature/circonscription.
Commune → conseil → parti → président/exécutif.
"""
    docs[DOCUMENTS[2]] = header("ARCHITECTURE DES CUBES", "Grains, dimensions, mesures et règles d'agrégation.", [DOCUMENTS[3], DOCUMENTS[7], DOCUMENTS[13]]) + body

    dimensions = ["DIM_GEO", "DIM_TIME", "DIM_PARTY", "DIM_PERSON", "DIM_ELECTION"]
    body = "PRINCIPE\nLes dimensions conformes donnent le même sens aux clés utilisées dans plusieurs faits. Une clé absente ne doit jamais être remplacée par un libellé libre.\n"
    for name in dimensions:
        sheet = model["sheets"][name]
        body += section(name, f"STATUT: OBSERVÉ\nLIGNES: {sheet['rows']}\nGRAIN: {sheet['grain']}\nCLÉ PRIMAIRE: {sheet['primary_key']}\nCOLONNES: " + ", ".join(map(str, sheet["headers"])))
    body += section("RÈGLES COMMUNES", """
1. geo_id, time_id, party_id, person_id et election_id sont les identifiants canoniques.
2. Les libellés servent à l'affichage, jamais à une jointure autonome.
3. valid_from/valid_to encadrent la validité; NULL valid_to signifie fin ouverte, pas éternité prouvée.
4. source_id et quality_status accompagnent toute valeur de référence importante.
5. Une clé source reste distincte d'une clé canonique.
6. Une relation parti-personne dans DIM_PERSON est descriptive du dernier ou principal état disponible; elle ne remplace pas PERSON_PARTY_HISTORY.
""")
    docs[DOCUMENTS[3]] = header("DIMENSIONS CONFORMES", "Contrats des cinq dimensions maîtres.", [DOCUMENTS[4], DOCUMENTS[5], DOCUMENTS[12]]) + body

    body = f"""
HIÉRARCHIE CANONIQUE
MAR
  → MA-xx                         région
    → MA-xx-xxx                   province ou préfecture
      → MA-xx-xxx-xxxx            commune ou arrondissement
        → circonscription communale, lorsque disponible
        → fraction, cible future
          → douar, cible future

OBSERVÉ
- DIM_GEO contient {counts['DIM_GEO']} lignes, incluant les niveaux historiques et canoniques.
- Le référentiel électoral canonique couvre 1 538 unités : 1 497 communes et
  41 arrondissements selon la structure HCP exploitée par V9.
- CROSSWALK_GEO contient {counts['CROSSWALK_GEO']} lignes : deux raccords
  élection-spécifiques pour chacune des 1 538 unités.

RÈGLES DE PARENTÉ
- parent_geo_id doit exister dans DIM_GEO, sauf racine explicitement admise.
- Une commune/arrondissement a exactement une province/préfecture pour une période donnée.
- Une province/préfecture a exactement une région pour une période donnée.
- Un changement de frontière impose une nouvelle période de validité ou un statut de compatibilité temporelle.
- Une circonscription n'est pas assimilée à une commune, même si leurs noms coïncident.
- Un douar ne devient jamais grain électoral sans résultat électoral réellement publié à ce niveau.

JOINTURES AUTORISÉES
Fait canonique.geo_id → DIM_GEO.geo_id.
Clé TAFRA idCommune → CROSSWALK_GEO.source_geo_id → canonical_geo_id.
canonical_geo_id → DIM_GEO.geo_id.

JOINTURES INTERDITES
- commune seule;
- commune sans province/préfecture;
- correspondance fuzzy sans parent géographique;
- code administratif historique interprété comme code courant sans validité;
- arrondissement agrégé avec sa commune mère sans règle d'agrégation explicite.
"""
    docs[DOCUMENTS[4]] = header("HIÉRARCHIE GÉOGRAPHIQUE", "Niveaux territoriaux, parentés et règles de forage.", [DOCUMENTS[3], DOCUMENTS[5], DOCUMENTS[12]]) + body

    review_lines = [
        f"{source_id} | {source_name} ({province}) → {canonical_id} | {canonical_name} | confiance={confidence}"
        for source_id, source_name, province, canonical_id, canonical_name, confidence in model["geo_reviews"]
    ]
    body = f"""
CROSSWALK_GEO
Grain : source_system × source_geo_id × élection.
Clé logique : source_system + source_geo_id + contexte électoral.
Résultat : canonical_geo_id, canonical_region_id et canonical_province_id.

ORDRE DE MATCHING
1. Nom normalisé + province/préfecture exacte.
2. Variantes d'accents, apostrophes, espaces, traits d'union et suffixes administratifs.
3. Fuzzy matching uniquement à l'intérieur du même parent.
4. Codes administratifs/historiques lorsque disponibles.
5. Revue manuelle si le score ou l'unicité reste insuffisant.

OBSERVÉ: 3 052 raccords normalized_name+province et 24 raccords
fuzzy_name+province. Dix lignes élection-spécifiques correspondent à cinq unités
uniques encore signalées pour revue explicite.

UNITÉS À REVOIR
""" + "\n".join(review_lines) + """

CROSSWALK_PARTY
Grain : source_system × raw_source_code × période.
Le code brut, le nom brut, le canonical_party_id et la validité sont conservés.
AGD et AFG sont des coalitions temporelles, pas des synonymes intemporels.
SAP, SAP2 et SAP3 sont des emplacements de listes indépendantes, pas des
identités partisanes stables entre communes ou élections.

RÈGLES D'IDENTITÉ
- Ne jamais développer un acronyme ambigu sans preuve source.
- Ne jamais fusionner deux codes uniquement parce que leurs noms se ressemblent.
- Une alliance conserve alliance_flag et sa période.
- predecessor/successor restent NULL sans relation documentée.
- BLOQUÉ: l'identité locale V9 n'est pas commune + personne en pratique. Les
  32 513 mandats portent seulement 1 538 person_id, soit un identifiant par commune.
- L'identité parlementaire TAFRA_MP_* repose sur idPerson et peut être longitudinale.
"""
    docs[DOCUMENTS[5]] = header("CROSSWALKS ET IDENTITÉS", "Résolution géographique, partis et personnes.", [DOCUMENTS[3], DOCUMENTS[4], DOCUMENTS[8]]) + body

    body = f"""
DIM_TIME contient {counts['DIM_TIME']} lignes et DIM_ELECTION {counts['DIM_ELECTION']} lignes.

NIVEAUX TEMPORELS
- date exacte : time_id au format date lorsque disponible;
- année : référence censitaire, budgétaire ou statistique;
- campagne : campaign_start, campaign_end et campaign_phase;
- scrutin : election_date et election_id;
- cycle : electoral_cycle;
- législature : parlement/legislature, par exemple 2007-2011 à 2021-2026;
- validité : valid_from et valid_to.

DISCIPLINE
1. Une variable doit conserver son année d'observation réelle.
2. Une jointure à une élection renseigne election_year et temporal_distance.
3. temporal_join_method appartient à : exact_year, previous_observation,
   next_observation, interpolated ou census_reference.
4. Une valeur RGPH 2024 jointe à 2015 reste une référence 2024, distance 9.
5. La même valeur jointe à 2021 reste une référence 2024, distance 3.
6. Une législature est une période; elle ne doit pas être réduite à son année de début.
7. Un mandat de remplacement conserve dateEntree et dateSortie propres.

FRONTIÈRES
Le découpage utilisé est compatible avec le cadre administratif en vigueur depuis
2015, mais temporal_boundary_status doit rester consulté. Une homonymie ou une
continuité de nom ne prouve pas à elle seule une continuité de périmètre.
"""
    docs[DOCUMENTS[6]] = header("TEMPS, CYCLES ET ÉLECTIONS", "Hiérarchie temporelle et validité longitudinale.", [DOCUMENTS[3], DOCUMENTS[4], DOCUMENTS[13]]) + body

    body = f"""
RESULTS
OBSERVÉ: {counts['RESULTS']} lignes totales; 22 054 lignes communales V9.
Grain cible: geo_id × party_id × election_id.
Clé physique: record_id. Clé analytique candidate: geo_id + party_id + election_id.
Les 10 571 lignes 2015 utilisent 31 colonnes partisanes publiées; les 11 483
lignes 2021 utilisent 33 colonnes partisanes publiées.

MOBILISATION
OBSERVÉ: {counts['MOBILISATION']} lignes.
Grain: geo_id × election_id.
turnout_rate est publié; valid_votes est la somme des colonnes partisanes.
BLOQUÉ: nInscrits/registered_voters est entièrement absent dans les deux sources
communales. voters, invalid_votes et blank_votes ne doivent donc pas être reconstruits.

ELECTORATE
OBSERVÉ: {counts['ELECTORATE']} lignes à des niveaux non exhaustifs.
Ne pas extrapoler ces lignes à toutes les communes.

POLITICAL_SUPPLY
OBSERVÉ: {counts['POLITICAL_SUPPLY']} lignes, principalement agrégées ou pilotes.
Un élu observé n'est pas automatiquement un candidat exhaustivement observé.

COMPETITION
OBSERVÉ: {counts['COMPETITION']} lignes historiques agrégées.
Les métriques communales exhaustives se trouvent surtout dans les panels V9.

DÉNOMINATEURS AUTORISÉS
- vote_share = votes du parti / somme de toutes les colonnes partisanes publiées.
- turnout_rate reste une mesure source car registered_voters est absent.
- seat_share = sièges du parti / somme des élus observés du conseil, avec QA.
- winning_share et margin_pp utilisent valid_votes communal.

INTERDICTIONS
- Ne pas calculer voters = valid_votes / turnout sans invalides et inscrits fiables.
- Ne pas qualifier valid_votes d'officiel si le champ est seulement dérivé.
- Ne pas sommer des pourcentages entre communes.
- Ne pas confondre résultat en voix, élus/sièges et contrôle exécutif.
"""
    docs[DOCUMENTS[7]] = header("FAITS ÉLECTORAUX ET MOBILISATION", "Voix, participation, offre et compétition.", [DOCUMENTS[2], DOCUMENTS[12], DOCUMENTS[13]]) + body

    body = f"""
DIM_PERSON
OBSERVÉ: {counts['DIM_PERSON']} lignes : 1 185 identités parlementaires stables et
1 538 proxies locaux communaux. Les proxies locaux ne sont pas des personnes uniques.

LOCAL_MANDATES
OBSERVÉ: {counts['LOCAL_MANDATES']} observations d'élus locaux.
Clé: local_mandate_id.
Grain: personne prudente × commune × circonscription × mandat.
person_id local a la forme TAFRA_LOCAL_2021_<idCommune>_<hash>.
BLOQUÉ: la normalisation utilisée dans V9 élimine les caractères arabes avant le
hash. Résultat observé: {model['local_identity']['rows']} mandats, {model['local_identity']['unique_names']}
noms distincts, mais seulement {model['local_identity']['unique_person_ids']} person_id,
exactement autant que les {model['local_identity']['unique_communes']} communes.
Le person_id local est donc un proxy communal partagé par tous les élus de la
commune. Il est interdit pour dédupliquer des personnes, compter des individus,
mesurer des trajectoires ou joindre des élus locaux à DIM_PERSON comme identités.
Pour une analyse individuelle provisoire, utiliser la ligne RAW et la combinaison
idCommune + idCirconscription + prenomNom + parti, en conservant le risque d'homonymie.

PARLIAMENTARY_MANDATES
OBSERVÉ: {counts['PARLIAMENTARY_MANDATES']} périodes de siège pour 1 185 idPerson.
Clé: mandate_id; personne stable: TAFRA_MP_<idPerson>.
idSiege/source_seat_id peut apparaître plusieurs fois lorsque des députés se
succèdent sur un siège. La clé ne doit donc jamais être legislature + idSiege seule.

RÈGLES DE MANDAT
- dateEntree/start_date est obligatoire pour ordonner les périodes.
- dateSortie/end_date NULL signifie actif ou sortie non publiée selon le contexte.
- entry_reason, exit_reason et replacement_procedure restent des attributs de période.
- female/male est disponible pour les députés; le sexe n'est pas fourni pour les élus locaux.
- role_uncertain_flag et head_of_list_uncertain_flag interdisent une certitude silencieuse.

BLOQUÉ
Le champ parti TAFRA des députés représente le premier parti observé dans la
législature. Il ne constitue pas un historique complet du nomadisme politique.
Une future PERSON_PARTY_HISTORY devra utiliser personne × parti × période × source.
"""
    docs[DOCUMENTS[8]] = header("PERSONNES, MANDATS ET REPRÉSENTATION", "Identités prudentes et trajectoires de mandat.", [DOCUMENTS[3], DOCUMENTS[5], DOCUMENTS[9]]) + body

    body = f"""
LOCAL_COUNCIL_CONTROL
OBSERVÉ/DÉRIVÉ: {counts['LOCAL_COUNCIL_CONTROL']} lignes commune × parti pour 2021.
party_seats correspond au nombre de lignes d'élus du parti dans la commune.
seat_rank ordonne ces effectifs. largest_party_flag identifie le ou les maxima.

DISTINCTIONS OBLIGATOIRES
- winner dans RESULTS : parti ayant le plus de voix publiées.
- largest_party : parti ayant le plus d'élus observés au conseil.
- president_party : parti du président identifié avec un rôle non incertain.
- executive parties : partis possédant des membres aux rôles exécutifs observés.
Ces quatre notions peuvent diverger et ne doivent jamais être fusionnées.

PRÉSIDENCE
OBSERVÉ: 1 403 communes ont un président identifié avec flagRole=0.
BLOQUÉ: 135 présidences restent non résolues; president_party reste NULL.
Il est interdit de remplacer president_party par largest_party.

RÉCONCILIATION DES SIÈGES
BLOQUÉ: deux communes ont 15 lignes d'élus pour 16 sièges attendus :
- idCommune 770;
- idCommune 817.
Les lignes sources sont conservées et l'écart national est -2.

GOUVERNANCE
Les champs governance_score et control_change restent NULL lorsqu'aucune source
ne permet leur calcul. Les données SMIIG sont une priorité future; leur arrivée
devra produire une observation commune × année, jointe au parti de contrôle sans
interprétation causale prématurée.
"""
    docs[DOCUMENTS[9]] = header("POUVOIR LOCAL ET GOUVERNANCE", "Sièges, présidences, exécutifs et contrôle communal.", [DOCUMENTS[7], DOCUMENTS[8], DOCUMENTS[12]]) + body

    socio_sheets = ["POPULATION", "SOCIAL_HUMAN", "LIVING_CONDITIONS", "DIGITAL_ACCESS", "ECONOMY", "PUBLIC_CAPITAL"]
    campaign_sheets = ["CAMPAIGN_FINANCE", "MEDIA_OFFLINE", "DIGITAL_CAMPAIGN", "PUBLIC_OPINION", "SHOCKS", "LOCAL_FINANCE", "POLITICAL_ACTIVITY"]
    body = "DOMAINES SOCIO-ÉCONOMIQUES\n"
    for name in socio_sheets:
        status = "STRUCTURE VIDE" if counts[name] == 0 else "OBSERVÉ"
        body += f"{name}: {status}, {counts[name]} lignes, grain déclaré={model['sheets'][name]['grain']}.\n"
    body += "\nCAMPAGNE, INFORMATION ET ACTION PUBLIQUE\n"
    for name in campaign_sheets:
        status = "STRUCTURE VIDE" if counts[name] == 0 else "OBSERVÉ"
        body += f"{name}: {status}, {counts[name]} lignes.\n"
    body += """

RÈGLES D'UTILISATION
- POPULATION 2024 est exhaustive au niveau canonique communal/arrondissement.
- Les autres domaines ont surtout une couverture nationale, régionale ou pilote.
- geo_id et time_id doivent tous deux correspondre au niveau réel de l'observation.
- Une mesure nationale ne peut pas être répétée sur chaque commune comme si elle était locale.
- Une enquête conserve survey_name, wave, sample_size et pondération lorsqu'ils existent.
- Une dépense de campagne doit préciser parti/personne, scrutin, devise et périmètre.
- Les métriques numériques génériques peuvent être stockées dans FACT_OBSERVATION,
  mais metric_id doit exister dans DATA_DICTIONARY.

STRUCTURES VIDES V9
POLITICAL_ACTIVITY, LOCAL_FINANCE et SHOCKS n'ont aucune ligne. Leur schéma est
documenté, mais aucune mesure ne doit être présentée comme disponible.
"""
    docs[DOCUMENTS[10]] = header("SOCIO-ÉCONOMIE, CAMPAGNE ET INFORMATION", "Domaines explicatifs et couverture réellement disponible.", [DOCUMENTS[1], DOCUMENTS[6], DOCUMENTS[11]]) + body

    body = f"""
PIPELINE OBLIGATOIRE
SOURCE EXTERNE
  → RAW immuable
  → profil et dictionnaire source
  → dimensions/crosswalks
  → faits normalisés
  → panels analytiques
  → QA et couverture

RAW OBSERVÉ
RAW_COMM2015_FULL: {counts['RAW_COMM2015_FULL']} lignes, 45 colonnes.
RAW_COMM2021_FULL: {counts['RAW_COMM2021_FULL']} lignes, 56 colonnes.
RAW_COUNCIL2021_FULL: {counts['RAW_COUNCIL2021_FULL']} lignes, 18 colonnes.
RAW_MPS_FULL: {counts['RAW_MPS_FULL']} lignes, 24 colonnes.
RAW_HCP_RGPH2024_FULL: disposition littérale du classeur officiel HCP.

RÈGLES RAW
- Ne pas renommer, supprimer, remplir ou corriger une colonne brute.
- Toute correction appartient à une couche normalisée et doit être traçable.
- Les fichiers sources restent conservés hors du classeur et leurs SHA-256 sont dans SOURCES.
- Un sample ou miroir nettoyé ne remplace pas un fichier original complet.

NORMALISATION
- Les identifiants TAFRA deviennent des clés source, pas des geo_id canoniques.
- Les codes de parti sont résolus par CROSSWALK_PARTY avec période.
- Les noms de personne locale produisent des identités prudentes limitées à la commune.
- Les valeurs nulles restent nulles; zéro n'est créé que si sa signification source est prouvée.

LINEAGE MINIMALE
source_id, source_url, publication_date, retrieval_date, quality_status et notes.
Les états recommandés sont verified, provisional, derived et low_confidence.

CONTRÔLES DE GOUVERNANCE
SOURCES décrit la provenance. SOURCE_DOWNLOADS décrit les ressources physiques.
RAW_IMPORT_QUEUE décrit le backlog. DATA_COVERAGE décrit la couverture.
QUALITY_CONTROL conserve toute anomalie. WORKBOOK_AUDIT_V9 inventorie le classeur.
"""
    docs[DOCUMENTS[11]] = header("PIPELINE RAW → NORMALIZED → FACT → ANALYTICAL", "Flux d'ingestion, lineage et non-altération.", [DOCUMENTS[0], DOCUMENTS[12], DOCUMENTS[14]]) + body

    body = """
CLÉS ET CARDINALITÉS
DIM_GEO.geo_id                 1 → N faits territoriaux.
DIM_TIME.time_id               1 → N faits datés.
DIM_PARTY.party_id             1 → N résultats/mandats/compositions.
DIM_PERSON.person_id           1 → N périodes de mandat lorsque l'identité est stable.
DIM_ELECTION.election_id       1 → N résultats/mobilisations/mandats électoraux.
SOURCES.source_id              1 → N lignes de faits et dimensions.

CLÉS ANALYTIQUES
RESULTS: geo_id + election_id + party_id.
MOBILISATION: geo_id + election_id.
ANALYTICAL_PANEL: geo_id + election_id + party_id.
COMMUNE_ELECTION_PANEL: geo_id + election_id.
ELECTORAL_TRANSITIONS_2015_2021: geo_id + party_id.
LOCAL_COUNCIL_CONTROL: geo_id + election_id + party_id.
PARLIAMENTARY_MANDATES: mandate_id; ne pas utiliser idSiege seul.
LOCAL_MANDATES: local_mandate_id uniquement. person_id local V9 n'est pas une clé
d'individu et ne doit pas servir de clé analytique.

RÈGLES DE NULLITÉ
- NULL = inconnu, absent, inapplicable ou non calculable selon notes/QA.
- NULL n'est jamais converti automatiquement en zéro.
- Un dénominateur NULL ou nul produit une métrique dérivée NULL.
- Un champ BLOQUÉ reste NULL jusqu'à ingestion d'une source défendable.

RÈGLES NUMÉRIQUES
- Comptages non négatifs.
- turnout_rate, vote_share et seat_share dans [0,100].
- Somme vote_share ≈ 100 au grain commune-élection complet.
- Somme party_seats = sièges du conseil, sinon anomalie QA.
- winner doit avoir rank=1; les ex æquo doivent rester détectables.
- Les valeurs dérivées conservent la source de leurs intrants.

JOINTURES INTERDITES
- noms de commune seuls;
- noms de personne seuls entre communes;
- person_id de LOCAL_MANDATES utilisé comme identité d'élu;
- acronymes de parti ambigus;
- année seule lorsque plusieurs observations existent;
- agrégation de mesures appartenant à des niveaux géographiques différents;
- 2024 présenté comme observation 2015/2021;
- largest_party utilisé comme president_party;
- parti parlementaire utilisé comme historique complet de changement de parti.

GESTION DES DOUBLONS
Une duplication de clé primaire est bloquante. Une duplication de clé naturelle
candidate est conservée jusqu'à comparaison de toutes les colonnes. V9 contient
deux lignes candidates sur commune+circonscription+nom+parti qui ne doivent pas
être supprimées silencieusement.
"""
    docs[DOCUMENTS[12]] = header("RÈGLES D'INTÉGRITÉ ET DE JOINTURE", "Contraintes référentielles, cardinalités et interdictions.", [DOCUMENTS[3], DOCUMENTS[5], DOCUMENTS[13]]) + body

    metric_rules = [
        ("valid_votes", "Σ_p votes_p", "toutes les colonnes partisanes publiées de la commune-élection", "NULL si aucune voix publiée"),
        ("vote_share", "100 × votes_p / valid_votes", "valid_votes", "NULL si valid_votes est NULL ou 0"),
        ("rank", "1 + nombre de partis ayant strictement plus de voix", "ensemble des partis observés dans la commune-élection", "NULL si votes est NULL"),
        ("winner", "1 si rank=1, sinon 0", "rang calculable", "NULL si rang non calculable"),
        ("victory_margin_votes", "votes_top1 - votes_top2", "deux partis classables", "NULL pour les non-vainqueurs ou moins de deux partis"),
        ("victory_margin_pp", "100 × victory_margin_votes / valid_votes", "valid_votes", "NULL si valid_votes est NULL ou 0"),
        ("abstention_rate", "100 - turnout_rate", "turnout_rate source", "NULL si turnout_rate est NULL"),
        ("party_seat_share", "100 × party_seats / Σ_p party_seats", "élus observés du conseil", "NULL si total observé est 0"),
        ("hhi", "Σ_p vote_share_p²", "parts complètes sur échelle 0–100", "NULL si parts incomplètes"),
        ("enp", "10000 / hhi", "HHI sur échelle 0–100", "NULL si HHI est NULL ou 0"),
        ("swing_pp", "share_2021 - share_2015", "même geo_id et même party_id temporel", "NULL si une part est absente"),
        ("vote_change", "votes_2021 - votes_2015", "même geo_id et party_id", "NULL si une valeur est absente"),
        ("vote_change_pct", "100 × (votes_2021 - votes_2015) / votes_2015", "votes_2015", "NULL si votes_2015 est NULL ou 0"),
        ("turnout_change_pp", "turnout_2021 - turnout_2015", "deux taux source comparables", "NULL si un taux est absent"),
        ("pedersen_volatility", "0.5 × Σ_p |share_p,2021 - share_p,2015|", "union des codes politiques temporels", "NULL si un dénominateur communal est incomplet"),
        ("temporal_distance", "observation_year - election_year", "deux années explicites", "NULL si une année est absente"),
    ]
    body = "RÈGLES DES MÉTRIQUES DÉRIVÉES\n"
    for name, formula, denominator, null_rule in metric_rules:
        body += f"\nMÉTRIQUE: {name}\nFORMULE: {formula}\nDÉNOMINATEUR/PÉRIMÈTRE: {denominator}\nRÈGLE NULL: {null_rule}\nSTATUT: DÉRIVÉ\n"
    body += """

RÈGLES TEMPORELLES
- previous_vote_share exige la même commune canonique et le même code politique défendable.
- entered_market/exited_market décrivent la présence d'un code source, pas nécessairement la naissance ou dissolution juridique d'un parti.
- AGD 2015 et AFG 2021 restent distincts sans crosswalk politique documenté.
- SAP/SAP2/SAP3 ne portent pas une identité longitudinale stable.
- seats_2015, seat_change et control_change restent NULL dans V9.
- population_reference utilise temporal_join_method=census_reference,
  observation_year=2024 et temporal_distance explicite.

PRÉCAUTIONS D'AGRÉGATION
Recalculer HHI, ENP, parts et marges à partir des voix agrégées. Ne jamais faire
la moyenne simple des indices communaux pour obtenir un indice régional sans
définir une pondération et un objet analytique distinct.
"""
    docs[DOCUMENTS[13]] = header("MÉTRIQUES DÉRIVÉES ET RÈGLES TEMPORELLES", "Formules, dénominateurs, nullité et comparabilité.", [DOCUMENTS[6], DOCUMENTS[7], DOCUMENTS[12]]) + body

    catalog = "CONTRAT DU CATALOGUE\nChaque entrée reproduit le volume calculé, la classification d'audit, le grain, les clés, les colonnes physiques, la couverture, la qualité et les limites.\n"
    for name in model["wb"].sheetnames:
        data = model["sheets"][name]
        catalog += "\n" + "=" * 100 + "\n"
        catalog += f"ONGLET: {name}\n"
        catalog += f"CLASSIFICATION: {normalize_text(data['classification']) or 'NON CLASSÉ'}\n"
        catalog += f"STATUT DOCUMENTAIRE: {'STRUCTURE VIDE' if data['rows'] == 0 else ('PILOTE' if name in pilot_sheets else 'OBSERVÉ/DÉRIVÉ SELON quality_status')}\n"
        catalog += f"LIGNES_DOCUMENTEES: {data['rows']}\n"
        catalog += f"GRAIN: {normalize_text(data['grain']) or 'non déclaré'}\n"
        catalog += f"CLÉ PRIMAIRE: {normalize_text(data['primary_key']) or 'non déclarée'}\n"
        catalog += f"CLÉS ÉTRANGÈRES: {normalize_text(data['foreign_keys']) or 'aucune déclarée'}\n"
        if name == "RAW_HCP_RGPH2024_FULL":
            columns = "Disposition source A:G; libellés sur les lignes 6–7: Collectivités territoriales, Marocains, Étrangers, Population, Ménages, libellé arabe, Code géographique"
        else:
            columns = ", ".join(map(str, data["headers"])) or "onglet non tabulaire"
        catalog += f"COLONNES: {columns}\n"
        catalog += f"SOURCES: {normalize_text(data['sources']) or 'voir cellules/onglet SOURCES'}\n"
        catalog += f"COUVERTURE TEMPORELLE: {normalize_text(data['coverage_period']) or 'non calculée'}\n"
        catalog += f"NIVEAU GÉOGRAPHIQUE: {normalize_text(data['geographic_level']) or 'non déclaré'}\n"
        catalog += f"QUALITÉ: {normalize_text(data['quality_status']) or 'voir données/QA'}\n"
        catalog += f"LIMITATIONS: {normalize_text(data['limitations']) or 'aucune limitation synthétique déclarée'}\n"
        if name == "LOCAL_MANDATES":
            catalog += "ALERTE IDENTITÉ: BLOQUÉ — 32 513 lignes mais 1 538 person_id, un par commune; utiliser local_mandate_id pour l'unicité de ligne.\n"
        if name == "DIM_PERSON":
            catalog += "ALERTE IDENTITÉ: les 1 538 entrées TAFRA_LOCAL_2021_* sont des proxies communaux, pas des élus distincts.\n"
        catalog += f"ORIGINE: {normalize_text(data['origin_version']) or 'V9'}\n"

    glossary = """
GLOSSAIRE
geo_id: identifiant géographique canonique.
source_geo_id: identifiant géographique dans une source.
time_id: identifiant de date ou période.
party_id: identifiant canonique ou code politique conservé.
person_id: identifiant de personne avec portée documentée.
election_id: identifiant du scrutin.
source_id: identifiant de provenance.
record_id: identifiant physique d'une observation de fait.
grain: combinaison minimale d'entités décrite par une ligne.
dimension conforme: référentiel partagé entre plusieurs tables de faits.
crosswalk: correspondance versionnée entre clé source et clé canonique.
lineage: chaîne de provenance et de transformation.
valid_votes: somme dérivée des voix partisanes publiées.
turnout: taux de participation publié par la source communale.
HHI: somme des carrés des parts de voix.
ENP: nombre effectif de partis, 10000/HHI sur parts 0–100.
swing: différence de part entre deux élections.
volatilité de Pedersen: demi-somme des variations absolues de parts.
largest_party: parti ayant le plus d'élus observés.
president_party: parti du président du conseil identifié.
quality_status: niveau de confiance ou nature dérivée d'une ligne.
"""
    docs[DOCUMENTS[14]] = header("CATALOGUE COMPLET DES ONGLETS", "Inventaire exhaustif des 64 onglets et glossaire.", [DOCUMENTS[0], DOCUMENTS[11], DOCUMENTS[12]]) + catalog + glossary
    return docs, metric_rules


FACT_RELATIONS = {
    "POPULATION": ["geo_id", "time_id"],
    "SOCIAL_HUMAN": ["geo_id", "time_id"],
    "LIVING_CONDITIONS": ["geo_id", "time_id"],
    "ELECTORATE": ["geo_id", "time_id", "election_id"],
    "POLITICAL_SUPPLY": ["geo_id", "time_id", "election_id", "party_id", "person_id"],
    "COMPETITION": ["geo_id", "time_id", "election_id"],
    "CAMPAIGN_FINANCE": ["geo_id", "time_id", "election_id", "party_id", "person_id"],
    "MEDIA_OFFLINE": ["geo_id", "time_id", "election_id", "party_id", "person_id"],
    "DIGITAL_CAMPAIGN": ["geo_id", "time_id", "election_id", "party_id", "person_id"],
    "DIGITAL_ACCESS": ["geo_id", "time_id"],
    "MOBILISATION": ["geo_id", "time_id", "election_id"],
    "RESULTS": ["geo_id", "time_id", "election_id", "party_id", "person_id"],
    "POWER_REP": ["geo_id", "time_id", "election_id", "party_id", "person_id"],
    "POLITICAL_ACTIVITY": ["geo_id", "time_id", "election_id", "party_id", "person_id"],
    "LOCAL_FINANCE": ["geo_id", "time_id"],
    "ECONOMY": ["geo_id", "time_id"],
    "PUBLIC_CAPITAL": ["geo_id", "time_id"],
    "SHOCKS": ["geo_id", "time_id"],
    "INSTITUTIONAL_LEGAL": ["geo_id", "time_id", "election_id"],
    "FACT_OBSERVATION": ["metric_id", "geo_id", "time_id", "election_id", "party_id", "person_id", "source_id"],
}


def validate(model, docs, metric_rules, before_hash):
    errors = []
    if len(model["sheets"]) != 64:
        errors.append(f"Le V9 contient {len(model['sheets'])} onglets au lieu de 64.")
    if set(model["audit"]) != set(model["sheets"]) - {"WORKBOOK_AUDIT_V9"}:
        missing = sorted((set(model["sheets"]) - {"WORKBOOK_AUDIT_V9"}) - set(model["audit"]))
        extra = sorted(set(model["audit"]) - set(model["sheets"]))
        errors.append(f"Audit non exhaustif. Manquants={missing}; extras={extra}")

    # Workbook audit and coverage must agree with physical populated-row counts.
    for name, data in model["sheets"].items():
        if name in model["audit"]:
            expected = int(model["audit"][name]["data_rows"])
            if expected != data["rows"]:
                errors.append(f"Volume {name}: audit={expected}, physique={data['rows']}")
        if name in model["coverage"]:
            expected = int(model["coverage"][name]["rows_loaded"])
            if expected != data["rows"]:
                errors.append(f"Couverture {name}: DATA_COVERAGE={expected}, physique={data['rows']}")

    # Every declared fact relation must physically expose its foreign-key columns.
    for sheet, required in FACT_RELATIONS.items():
        headers = set(model["sheets"][sheet]["headers"])
        missing = [column for column in required if column not in headers]
        if missing:
            errors.append(f"Relations {sheet}: colonnes absentes {missing}")

    catalog = docs["14_CATALOGUE_COMPLET_DES_ONGLETS.txt"]
    for name, data in model["sheets"].items():
        if f"ONGLET: {name}" not in catalog:
            errors.append(f"Onglet absent du catalogue: {name}")
        if f"LIGNES_DOCUMENTEES: {data['rows']}" not in catalog:
            errors.append(f"Volume absent du catalogue: {name}")
        if name not in {"README", "RAW_HCP_RGPH2024_FULL"}:
            for column in data["headers"]:
                if str(column) not in catalog:
                    errors.append(f"Colonne absente du catalogue: {name}.{column}")

    if len(metric_rules) < 15:
        errors.append("Registre de métriques dérivées incomplet.")
    for rule in metric_rules:
        if len(rule) != 4 or not all(normalize_text(value) for value in rule):
            errors.append(f"Règle métrique incomplète: {rule}")

    corpus = "\n".join(docs.values())
    required_blockers = [
        "nInscrits", "idCommune 770", "idCommune 817", "135 présidences",
        "cinq unités", "sièges 2015", "historique complet du nomadisme politique",
        "32 513 mandats portent seulement 1 538 person_id",
    ]
    for phrase in required_blockers:
        if phrase not in corpus:
            errors.append(f"Blocker absent du corpus: {phrase}")

    for name, content in docs.items():
        for token in ["VERSION", "CLASSEUR SOURCE", "DATE DE GÉNÉRATION", "PÉRIMÈTRE", "RENVOIS"]:
            if token not in content:
                errors.append(f"Métadonnée {token} absente de {name}")
        if len(content.strip()) < 500:
            errors.append(f"Document trop court: {name}")
        content.encode("utf-8")

    if before_hash != sha256(WORKBOOK):
        errors.append("Le classeur V9 a été modifié pendant la génération.")
    if errors:
        raise RuntimeError("ÉCHEC DE VALIDATION\n- " + "\n- ".join(errors))


def write_documents(docs):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    unexpected = sorted(path.name for path in OUTPUT_DIR.glob("*.txt") if path.name not in DOCUMENTS)
    if unexpected:
        raise RuntimeError(f"Fichiers .txt inattendus dans {OUTPUT_DIR}: {unexpected}")
    for name in DOCUMENTS:
        content = docs[name].replace("\r\n", "\n").rstrip() + "\n"
        # UTF-8 with BOM remains UTF-8 and is detected correctly by Windows text tools.
        (OUTPUT_DIR / name).write_text(content, encoding="utf-8-sig", newline="\n")
    produced = sorted(path.name for path in OUTPUT_DIR.glob("*.txt"))
    if produced != sorted(DOCUMENTS):
        raise RuntimeError(f"Ensemble documentaire incorrect: {produced}")
    for path in OUTPUT_DIR.glob("*.txt"):
        if path.stat().st_size == 0:
            raise RuntimeError(f"Document vide: {path.name}")
        path.read_text(encoding="utf-8")


def main(data_dir: str | Path | None = None):
    configure_paths(data_dir)
    if not WORKBOOK.exists():
        raise FileNotFoundError(WORKBOOK)
    before_hash = sha256(WORKBOOK)
    model = load_model()
    docs, metric_rules = build_docs(model)
    if set(docs) != set(DOCUMENTS):
        raise RuntimeError("La liste des documents construits ne correspond pas au contrat.")
    validate(model, docs, metric_rules, before_hash)
    write_documents(docs)
    # Validate disk bytes and immutability after writes.
    if len(list(OUTPUT_DIR.glob("*.txt"))) != 15:
        raise RuntimeError("Le répertoire final ne contient pas exactement 15 fichiers .txt.")
    after_hash = sha256(WORKBOOK)
    if before_hash != after_hash:
        raise RuntimeError("Le V9 a changé après écriture de la documentation.")
    total_bytes = sum(path.stat().st_size for path in OUTPUT_DIR.glob("*.txt"))
    print(f"DOCUMENTATION_OK files=15 bytes={total_bytes}")
    print(f"WORKBOOK_SHA256={after_hash}")
    print(f"OUTPUT={OUTPUT_DIR}")


if __name__ == "__main__":
    main()
