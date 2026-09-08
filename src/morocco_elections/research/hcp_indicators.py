from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

from morocco_elections.config import PROJECT_ROOT, get_paths


EXPECTED_UNITS = 1_538
DEFAULT_METADATA = PROJECT_ROOT / "metadata" / "v10qa3_hcp_indicators.json"
DEFAULT_REPORT = PROJECT_ROOT / "docs" / "research" / "V10QA3_HCP_INDICATORS.txt"
V10_REPORT = PROJECT_ROOT / "metadata" / "v10_release_report.json"
MISSING_SYMBOLS = {"", ".", "…", "-"}

SOURCE_SPECS = {
    "RGPH2014_INDIVIDUALS": {
        "url": "https://www.hcp.ma/file/230045/",
        "publisher": "Haut-Commissariat au Plan",
        "publication_date": "2017-05-30",
        "expected_name": "rgph2014_individus.xlsx",
        "expected_sheets": ["Indic.Ensemble", "Indic.Urbain", "Indic.Rural"],
    },
    "RGPH2014_HOUSEHOLDS": {
        "url": "https://www.hcp.ma/file/230042/",
        "publisher": "Haut-Commissariat au Plan",
        "publication_date": "2017-05-30",
        "expected_name": "rgph2014_menages.xlsx",
        "expected_sheets": ["Indic.Ensemble", "Indic.Urbain", "Indic.Rural"],
    },
    "RGPH2024_INDICATORS": {
        "url": "https://www.hcp.ma/file/242671/",
        "publisher": "Haut-Commissariat au Plan",
        "publication_date": "2024-12-17",
        "expected_name": "rgph2024_indicateurs.xlsx",
        "expected_sheets": [
            "Population",
            "Population_Urbaine",
            "Population_Rurale",
            "Ménages",
            "Ménages_Urbains",
            "Ménages_Ruraux",
            "Définitions_des_Concepts",
            "Signes_Conventionnels",
            "Avis_aux_Utilisateurs",
        ],
    },
}

# Deux changements de codification internes aux publications HCP 2024 sont
# raccordés à la géographie officielle déjà manifestée dans V10. Aucun fuzzy match.
GEO_ALIASES_2024 = {
    "MA-06-385-0105": "MA-06-385-0305",  # Ouled/Oulad Salah
    "MA-06-355-0107": "MA-06-355-0301",  # Almajjatia/Al Majjatia Oulad Taleb
}

INDICATORS = [
    {
        "indicator_id": "population_legal",
        "unit": "persons",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_INDIVIDUALS", "ensemble", 9, "Population légale"),
        "2024": ("RGPH2024_INDICATORS", "population", 3, "Population légale"),
        "domain": "count",
    },
    {
        "indicator_id": "population_municipal",
        "unit": "persons",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_INDIVIDUALS", "ensemble", 10, "Population municipale"),
        "2024": ("RGPH2024_INDICATORS", "population", 4, "Population municipale"),
        "domain": "count",
    },
    {
        "indicator_id": "population_urban",
        "unit": "persons",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_INDIVIDUALS", "urban", 10, "Population municipale, milieu urbain"),
        "2024": ("RGPH2024_INDICATORS", "population_urban", 4, "Population municipale, milieu urbain"),
        "domain": "count",
    },
    {
        "indicator_id": "urbanization_rate",
        "unit": "percent",
        "kind": "DERIVED",
        "formula": "100 * population_urban / population_municipal",
        "denominator": "population_municipal",
        "domain": "percent",
    },
    {
        "indicator_id": "illiteracy_rate_10_plus",
        "unit": "percent",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_INDIVIDUALS", "ensemble", 36, "Taux d'analphabétisme"),
        "2024": ("RGPH2024_INDICATORS", "population", 35, "Taux d'analphabétisme des 10 ans et plus"),
        "domain": "percent",
    },
    {
        "indicator_id": "literacy_rate_10_plus",
        "unit": "percent",
        "kind": "DERIVED",
        "formula": "100 - illiteracy_rate_10_plus",
        "denominator": "population_10_plus defined by source",
        "domain": "percent",
    },
    {
        "indicator_id": "school_enrollment_rate",
        "unit": "percent",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_INDIVIDUALS", "ensemble", 35, "Taux de scolarisation des enfants âgés de 7 à 12 ans"),
        "2024": ("RGPH2024_INDICATORS", "population", 33, "Taux de scolarisation des 6-11 ans en 2023/2024"),
        "domain": "percent",
    },
    {
        "indicator_id": "activity_rate",
        "unit": "percent",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_INDIVIDUALS", "ensemble", 54, "Taux net d'activité"),
        "2024": ("RGPH2024_INDICATORS", "population", 57, "Taux d'activité des 15 ans et plus"),
        "domain": "percent",
    },
    {
        "indicator_id": "unemployment_rate",
        "unit": "percent",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_INDIVIDUALS", "ensemble", 55, "Taux de chômage"),
        "2024": ("RGPH2024_INDICATORS", "population", 58, "Taux de chômage"),
        "domain": "percent",
    },
    {
        "indicator_id": "household_electricity_rate",
        "unit": "percent",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_HOUSEHOLDS", "ensemble", 30, "Logements équipés en électricité"),
        "2024": ("RGPH2024_INDICATORS", "households", 24, "Électricité"),
        "domain": "percent",
    },
    {
        "indicator_id": "household_piped_water_rate",
        "unit": "percent",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_HOUSEHOLDS", "ensemble", 31, "Logements équipés en eau courante"),
        "2024": ("RGPH2024_INDICATORS", "households", 25, "Eau courante"),
        "domain": "percent",
    },
    {
        "indicator_id": "household_public_sewer_rate",
        "unit": "percent",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_HOUSEHOLDS", "ensemble", 32, "Évacuation des eaux usées par réseau public"),
        "2024": ("RGPH2024_INDICATORS", "households", 26, "Réseau public d'assainissement"),
        "domain": "percent",
    },
    {
        "indicator_id": "household_internet_rate",
        "unit": "percent",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_HOUSEHOLDS", "ensemble", 47, "Ménages équipés d'Internet"),
        "2024": None,
        "domain": "percent",
    },
    {
        "indicator_id": "household_computer_rate",
        "unit": "percent",
        "kind": "OBSERVED",
        "2014": ("RGPH2014_HOUSEHOLDS", "ensemble", 48, "Ménages équipés d'un ordinateur"),
        "2024": None,
        "domain": "percent",
    },
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _normalized_name(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).casefold()
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("’", "'")
    text = re.sub(r"\s*\((?:mun\.|arrond\.)\)\s*$", "", text)
    text = re.sub(r"^(?:commune|arrondissement)\s+(?:d'|de |du |des )", "", text)
    text = re.sub(r"[^a-z0-9\u0600-\u06ff]+", " ", text)
    return " ".join(text.split())


def _canonical_2014(row: tuple[Any, ...]) -> str:
    return f"MA-{int(row[0]):02d}-{int(row[1]):03d}-{int(row[2]):02d}{int(row[4]):02d}"


def _canonical_2024(value: Any) -> str:
    code = f"{int(value):09d}"
    return f"MA-{code[:2]}-{code[2:5]}-{code[5:7]}{code[7:9]}"


def _baseline_context(workbook: Path) -> dict[str, Any]:
    geo = pd.read_excel(workbook, sheet_name="DIM_GEO", header=3)
    control = pd.read_excel(workbook, sheet_name="LOCAL_COUNCIL_CONTROL", header=3)
    universe_ids = set(control["geo_id"].dropna().astype(str))
    analytical = geo[geo["geo_id"].astype(str).isin(universe_ids)]
    by_name_parent: dict[tuple[str, str], list[str]] = {}
    for row in analytical.itertuples(index=False):
        by_name_parent.setdefault((str(row.parent_geo_id), _normalized_name(row.geo_name)), []).append(str(row.geo_id))
    if len(universe_ids) != EXPECTED_UNITS:
        raise RuntimeError(f"Univers V10 inattendu: {len(universe_ids)} unités")
    return {"universe": universe_ids, "by_name_parent": by_name_parent}


def _source_record(path: Path | None, source_id: str, retrieved_on: str) -> dict[str, Any]:
    spec = SOURCE_SPECS[source_id]
    if path is None:
        return {
            "source_id": source_id,
            "local_name": spec["expected_name"],
            "source_url": spec["url"],
            "final_url": spec["url"],
            "publisher": spec["publisher"],
            "publication_date": spec["publication_date"],
            "retrieved_on": retrieved_on,
            "reuse_terms": "Téléchargement public officiel HCP; attribution conservée; fichier non redistribué.",
            "etag": None,
            "last_modified": None,
            "byte_size": None,
            "sha256": None,
            "sheets": [],
            "status": "not_provided",
            "tracked": False,
        }
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = [{"name": ws.title, "rows": ws.max_row, "columns": ws.max_column} for ws in workbook.worksheets]
    workbook.close()
    expected = spec["expected_sheets"]
    status = "qualified_candidate" if [item["name"] for item in sheets] == expected else "schema_mismatch"
    return {
        "source_id": source_id,
        "local_name": path.name,
        "source_url": spec["url"],
        "final_url": spec["url"],
        "publisher": spec["publisher"],
        "publication_date": spec["publication_date"],
        "retrieved_on": retrieved_on,
        "reuse_terms": "Téléchargement public officiel HCP; attribution conservée; fichier non redistribué.",
        "etag": None,
        "last_modified": None,
        "byte_size": path.stat().st_size,
        "sha256": sha256_file(path),
        "sheets": sheets,
        "status": status,
        "tracked": False,
    }


def _match_geo(
    source_geo: str,
    source_name: Any,
    parent_geo: str,
    context: dict[str, Any],
    year: int,
) -> tuple[str | None, str]:
    if source_geo in context["universe"]:
        return source_geo, "official_code"
    if year == 2024 and source_geo in GEO_ALIASES_2024:
        target = GEO_ALIASES_2024[source_geo]
        return (target, "documented_hcp_2024_code_alias") if target in context["universe"] else (None, "unresolved")
    matches = context["by_name_parent"].get((parent_geo, _normalized_name(source_name)), [])
    if len(matches) == 1:
        return matches[0], "exact_normalized_name_and_parent"
    return None, "unresolved"


def _extract_2014(path: Path, sheet: str, context: dict[str, Any]) -> tuple[dict[str, list[Any]], dict[str, int]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    values: dict[str, list[Any]] = {}
    methods: dict[str, int] = {}
    for row in worksheet.iter_rows(min_row=4, values_only=True):
        try:
            is_unit = all(row[index] not in (None, "") for index in (0, 1, 2, 4)) and row[5] in (None, "")
            if not is_unit:
                continue
            source_geo = _canonical_2014(row)
            parent = f"MA-{int(row[0]):02d}-{int(row[1]):03d}"
        except (TypeError, ValueError):
            continue
        target, method = _match_geo(source_geo, row[7], parent, context, 2014)
        methods[method] = methods.get(method, 0) + 1
        if target:
            values.setdefault(target, []).append(row)
    workbook.close()
    return values, methods


def _is_prefecture_aggregate(name: Any) -> bool:
    normalized = _normalized_name(name)
    return normalized.startswith("prefecture d arrondissement") or normalized.startswith("prefecture d arrondissements")


def _extract_2024(path: Path, sheet: str, context: dict[str, Any], header_rows: int) -> tuple[dict[str, list[Any]], dict[str, int]]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    worksheet = workbook[sheet]
    values: dict[str, list[Any]] = {}
    methods: dict[str, int] = {}
    for row_index in range(header_rows + 1, worksheet.max_row + 1):
        code_cell = worksheet.cell(row_index, 1)
        name = worksheet.cell(row_index, 2).value
        if code_cell.value is None or code_cell.number_format != '00"."000"."00"."00' or _is_prefecture_aggregate(name):
            continue
        source_geo = _canonical_2024(code_cell.value)
        parent = "-".join(source_geo.split("-")[:3])
        target, method = _match_geo(source_geo, name, parent, context, 2024)
        methods[method] = methods.get(method, 0) + 1
        if target:
            values.setdefault(target, []).append(tuple(worksheet.cell(row_index, column).value for column in range(1, worksheet.max_column + 1)))
    workbook.close()
    return values, methods


def _numeric(value: Any) -> float | None:
    if value is None or str(value).strip() in MISSING_SYMBOLS:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _evaluate_observed(
    indicator: dict[str, Any],
    year: int,
    source_rows: dict[str, list[Any]] | None,
    source_id: str | None,
    column: int | None,
    definition: str,
    context: dict[str, Any],
) -> tuple[dict[str, Any], dict[str, float | None]]:
    values: dict[str, float | None] = {}
    duplicates = 0
    invalid = 0
    if source_rows is not None and column is not None:
        for geo_id in context["universe"]:
            rows = source_rows.get(geo_id, [])
            if len(rows) != 1:
                duplicates += max(0, len(rows) - 1)
                values[geo_id] = None
                continue
            value = _numeric(rows[0][column - 1])
            if value is not None:
                if indicator["domain"] == "count" and value < 0:
                    invalid += 1
                    value = None
                if indicator["domain"] == "percent" and not 0 <= value <= 100:
                    invalid += 1
                    value = None
            values[geo_id] = value
    usable = sum(value is not None for value in values.values())
    matched = len(source_rows or {})
    go = matched == EXPECTED_UNITS and usable == EXPECTED_UNITS and duplicates == 0 and invalid == 0
    return (
        {
            "indicator_id": indicator["indicator_id"],
            "observation_year": year,
            "decision": "GO" if go else "NO_GO",
            "kind": indicator["kind"],
            "source_id": source_id,
            "definition": definition,
            "unit": indicator["unit"],
            "expected_units": EXPECTED_UNITS,
            "matched_units": matched,
            "usable_units": usable,
            "coverage_rate": round(usable / EXPECTED_UNITS, 6),
            "duplicates": duplicates,
            "invalid_values": invalid,
            "formula": indicator.get("formula"),
            "denominator": indicator.get("denominator"),
            "null_rule": "NULL si la valeur publiée est absente, non applicable ou hors domaine.",
        },
        values,
    )


def _evaluate_derived(
    indicator: dict[str, Any],
    year: int,
    components: dict[str, dict[str, float | None]],
) -> tuple[dict[str, Any], dict[str, float | None]]:
    values: dict[str, float | None] = {}
    if indicator["indicator_id"] == "urbanization_rate":
        component_ids = ("population_urban", "population_municipal")
    else:
        component_ids = ("illiteracy_rate_10_plus",)
    matched_ids = set.intersection(*(set(components.get(component_id, {})) for component_id in component_ids))
    for geo_id in matched_ids:
        if indicator["indicator_id"] == "urbanization_rate":
            numerator = components["population_urban"].get(geo_id)
            denominator = components["population_municipal"].get(geo_id)
            values[geo_id] = None if numerator is None or denominator is None or denominator <= 0 else 100 * numerator / denominator
        else:
            source = components["illiteracy_rate_10_plus"].get(geo_id)
            values[geo_id] = None if source is None else 100 - source
    usable = sum(value is not None and 0 <= value <= 100 for value in values.values())
    go = usable == EXPECTED_UNITS
    return (
        {
            "indicator_id": indicator["indicator_id"],
            "observation_year": year,
            "decision": "GO" if go else "NO_GO",
            "kind": "DERIVED",
            "source_id": "DERIVED_FROM_QUALIFIED_HCP",
            "definition": indicator["formula"],
            "unit": indicator["unit"],
            "expected_units": EXPECTED_UNITS,
            "matched_units": len(matched_ids),
            "usable_units": usable,
            "coverage_rate": round(usable / EXPECTED_UNITS, 6),
            "duplicates": 0,
            "invalid_values": 0,
            "formula": indicator["formula"],
            "denominator": indicator["denominator"],
            "null_rule": "NULL si une composante manque ou si le dénominateur est nul ou négatif.",
        },
        values,
    )


def _comparability(decisions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_key = {(item["indicator_id"], item["observation_year"]): item for item in decisions}
    results = []
    for indicator in INDICATORS:
        indicator_id = indicator["indicator_id"]
        left = by_key[(indicator_id, 2014)]
        right = by_key[(indicator_id, 2024)]
        reason = "Même définition, unité et univers communal; aucune interpolation."
        comparable = left["decision"] == right["decision"] == "GO"
        if indicator_id == "school_enrollment_rate":
            comparable = False
            reason = "Classes d'âge différentes: 7-12 ans en 2014 et 6-11 ans en 2024."
        elif indicator_id == "activity_rate":
            comparable = False
            reason = "Le libellé 2014 ne démontre pas le même univers d'âge que le taux 15 ans et plus de 2024."
        elif left["decision"] != "GO" or right["decision"] != "GO":
            reason = "Au moins un millésime ne couvre pas les 1 538 unités avec une valeur utilisable."
        results.append(
            {
                "indicator_id": indicator_id,
                "period": "2014-2024",
                "decision": "GO" if comparable else "NO_GO",
                "reason": reason,
                "interpolation_allowed": False,
            }
        )
    return results


def build_metadata(
    workbook_path: Path,
    candidates: dict[str, Path | None],
    as_of: str,
) -> dict[str, Any]:
    baseline_hash = sha256_file(workbook_path)
    context = _baseline_context(workbook_path)
    sources = {source_id: _source_record(candidates.get(source_id), source_id, as_of) for source_id in SOURCE_SPECS}
    extracted: dict[tuple[int, str], dict[str, list[Any]]] = {}
    geo_methods: dict[str, dict[str, int]] = {}

    path = candidates.get("RGPH2014_INDIVIDUALS")
    if path and sources["RGPH2014_INDIVIDUALS"]["status"] == "qualified_candidate":
        for key, sheet in (("ensemble", "Indic.Ensemble"), ("urban", "Indic.Urbain")):
            extracted[(2014, key)], geo_methods[f"2014_individuals_{key}"] = _extract_2014(path, sheet, context)
    path = candidates.get("RGPH2014_HOUSEHOLDS")
    if path and sources["RGPH2014_HOUSEHOLDS"]["status"] == "qualified_candidate":
        extracted[(2014, "households")], geo_methods["2014_households"] = _extract_2014(path, "Indic.Ensemble", context)
    path = candidates.get("RGPH2024_INDICATORS")
    if path and sources["RGPH2024_INDICATORS"]["status"] == "qualified_candidate":
        for key, sheet, headers in (
            ("population", "Population", 3),
            ("population_urban", "Population_Urbaine", 3),
            ("households", "Ménages", 2),
        ):
            extracted[(2024, key)], geo_methods[f"2024_{key}"] = _extract_2024(path, sheet, context, headers)

    decisions: list[dict[str, Any]] = []
    values_by_year: dict[int, dict[str, dict[str, float | None]]] = {2014: {}, 2024: {}}
    for year in (2014, 2024):
        for indicator in [item for item in INDICATORS if item["kind"] == "OBSERVED"]:
            spec = indicator.get(str(year))
            if spec is None:
                decision, values = _evaluate_observed(indicator, year, None, None, None, "Non publié", context)
            else:
                source_id, key, column, definition = spec
                decision, values = _evaluate_observed(
                    indicator,
                    year,
                    extracted.get((year, "households" if key == "ensemble" and source_id.endswith("HOUSEHOLDS") else key)),
                    source_id,
                    column,
                    definition,
                    context,
                )
            decisions.append(decision)
            values_by_year[year][indicator["indicator_id"]] = values
        for indicator in [item for item in INDICATORS if item["kind"] == "DERIVED"]:
            decision, values = _evaluate_derived(indicator, year, values_by_year[year])
            decisions.append(decision)
            values_by_year[year][indicator["indicator_id"]] = values

    decisions.sort(key=lambda item: (item["indicator_id"], item["observation_year"]))
    go_count = sum(item["decision"] == "GO" for item in decisions)
    if sha256_file(workbook_path) != baseline_hash:
        raise RuntimeError("La baseline V10 a changé pendant la qualification")
    return {
        "schema_version": 1,
        "phase": "V10-QA-3",
        "generated_on": as_of,
        "scope": "Qualification HCP communale 2014-2024; aucune ingestion.",
        "decision_model": "GO/NO_GO par indicateur canonique et millésime",
        "baseline": {
            "release": "V10",
            "sha256_before": baseline_hash,
            "sha256_after": sha256_file(workbook_path),
            "modified": False,
        },
        "gate": {
            "expected_units": EXPECTED_UNITS,
            "full_national_coverage_required": True,
            "fuzzy_matching_allowed": False,
            "screen_scraping_allowed": False,
            "interpolation_allowed": False,
            "missing_symbols_are_zero": False,
        },
        "sources": list(sources.values()),
        "geographic_matching": geo_methods,
        "geographic_resolutions": [
            {"source_geo_id": source, "canonical_geo_id": target, "method": "documented_hcp_2024_code_alias"}
            for source, target in sorted(GEO_ALIASES_2024.items())
        ],
        "indicator_decisions": decisions,
        "comparability": _comparability(decisions),
        "intercensal": {
            "decision": "NO_GO",
            "reason": "Aucune publication officielle nationale communale intercensitaire couvrant ces indicateurs n'a été identifiée.",
            "values_reconstructed": False,
        },
        "temporal_permissions": [
            {"election_year": 2015, "observation_year": 2014, "permission": "CONDITIONAL", "offset_years": -1},
            {"election_year": 2015, "observation_year": 2024, "permission": "FORBIDDEN", "offset_years": 9},
            {"election_year": 2021, "observation_year": 2014, "permission": "CONDITIONAL", "offset_years": -7},
            {"election_year": 2021, "observation_year": 2024, "permission": "CONDITIONAL", "offset_years": 3},
        ],
        "summary": {
            "indicator_year_decisions": len(decisions),
            "go": go_count,
            "no_go": len(decisions) - go_count,
            "comparability_go": sum(item["decision"] == "GO" for item in _comparability(decisions)),
        },
        "next_action": "Ouvrir feat/v11-hcp-indicators uniquement pour les indicateurs GO; conserver les autres comme lacunes explicites.",
    }


def validate_metadata(metadata: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    if metadata.get("schema_version") != 1 or metadata.get("phase") != "V10-QA-3":
        errors.append("Métadonnées racine invalides")
    gate = metadata.get("gate", {})
    if gate.get("expected_units") != EXPECTED_UNITS or gate.get("fuzzy_matching_allowed") is not False:
        errors.append("Gate géographique invalide")
    sources = metadata.get("sources", [])
    source_ids = {item.get("source_id") for item in sources}
    if len(sources) != len(SOURCE_SPECS) or source_ids != set(SOURCE_SPECS):
        errors.append("Registre des sources incomplet ou dupliqué")
    required_source_fields = {
        "source_url",
        "final_url",
        "publisher",
        "publication_date",
        "retrieved_on",
        "reuse_terms",
        "etag",
        "byte_size",
        "sha256",
        "sheets",
        "status",
    }
    for source in sources:
        if not required_source_fields <= source.keys():
            errors.append(f"Provenance incomplète: {source.get('source_id')}")
        if source.get("publisher") != "Haut-Commissariat au Plan" or not str(source.get("source_url", "")).startswith(
            "https://www.hcp.ma/"
        ):
            errors.append(f"Source non officielle: {source.get('source_id')}")
        if source.get("status") == "qualified_candidate":
            if not re.fullmatch(r"[0-9a-f]{64}", str(source.get("sha256", ""))) or not source.get("byte_size"):
                errors.append(f"Empreinte candidat invalide: {source.get('source_id')}")
    decisions = metadata.get("indicator_decisions", [])
    if len(decisions) != len(INDICATORS) * 2:
        errors.append("Décisions indicateur × millésime incomplètes")
    keys = {(item.get("indicator_id"), item.get("observation_year")) for item in decisions}
    expected = {(item["indicator_id"], year) for item in INDICATORS for year in (2014, 2024)}
    if keys != expected:
        errors.append("Clés de décision incorrectes ou dupliquées")
    for item in decisions:
        if item.get("decision") not in {"GO", "NO_GO"}:
            errors.append(f"Décision invalide: {item.get('indicator_id')}")
        if not item.get("definition") or not item.get("unit") or not item.get("null_rule"):
            errors.append(f"Contrat incomplet: {item.get('indicator_id')}")
        expected_decision = "GO" if item.get("usable_units") == EXPECTED_UNITS and item.get("matched_units") == EXPECTED_UNITS and not item.get("duplicates") and not item.get("invalid_values") else "NO_GO"
        if item.get("decision") != expected_decision:
            errors.append(f"Gate incohérent: {item.get('indicator_id')} {item.get('observation_year')}")
        if item.get("kind") == "DERIVED" and (not item.get("formula") or not item.get("denominator")):
            errors.append(f"Métrique dérivée sans formule: {item.get('indicator_id')}")
    comparability = metadata.get("comparability", [])
    if len(comparability) != len(INDICATORS) or {item.get("indicator_id") for item in comparability} != {
        item["indicator_id"] for item in INDICATORS
    }:
        errors.append("Décisions de comparabilité incomplètes")
    if any(item.get("decision") not in {"GO", "NO_GO"} for item in comparability):
        errors.append("Décision de comparabilité invalide")
    baseline = metadata.get("baseline", {})
    if baseline.get("modified") is not False or baseline.get("sha256_before") != baseline.get("sha256_after"):
        errors.append("Immutabilité V10 non démontrée")
    summary = metadata.get("summary", {})
    go_count = sum(item.get("decision") == "GO" for item in decisions)
    if summary.get("go") != go_count or summary.get("no_go") != len(decisions) - go_count:
        errors.append("Résumé incohérent")
    return errors


def render_report(metadata: dict[str, Any]) -> str:
    summary = metadata["summary"]
    lines = [
        "V10-QA-3 — QUALIFICATION DES INDICATEURS COMMUNAUX HCP 2014–2024",
        "",
        "VERSION : V10-QA-3 (recherche, sans export warehouse)",
        f"DATE DE GÉNÉRATION : {metadata['generated_on']}",
        "PÉRIMÈTRE : indicateurs HCP au grain des 1 538 unités analytiques V10",
        "RENVOIS : metadata/v10qa3_hcp_indicators.json ; V10-QA ; V10 immuable",
        "",
        f"BILAN : {summary['go']} GO et {summary['no_go']} NO_GO sur {summary['indicator_year_decisions']} décisions indicateur × millésime.",
        f"COMPARABILITÉ 2014–2024 : {summary['comparability_go']} indicateur(s) GO.",
        "",
        "Aucune valeur n'a été ingérée. Aucun classeur source n'est suivi par Git.",
        "Les symboles ., … et - restent manquants ou non applicables et ne deviennent jamais zéro.",
        "",
        "SOURCES",
    ]
    for source in metadata["sources"]:
        lines.append(
            f"[{source['status']}] {source['source_id']} — {source['source_url']} — octets={source['byte_size']} — sha256={source['sha256']}"
        )
    lines.extend(["", "DÉCISIONS PAR INDICATEUR"])
    for item in metadata["indicator_decisions"]:
        lines.append(
            f"[{item['decision']}] {item['indicator_id']} {item['observation_year']} — {item['kind']} — "
            f"couverture={item['usable_units']}/{item['expected_units']} — unité={item['unit']} — source={item['source_id']}"
        )
        if item.get("formula"):
            lines.append(f"  Formule : {item['formula']} ; dénominateur : {item['denominator']} ; {item['null_rule']}")
    lines.extend(["", "COMPARABILITÉ 2014–2024"])
    for item in metadata["comparability"]:
        lines.append(f"[{item['decision']}] {item['indicator_id']} — {item['reason']}")
    lines.extend(["", "DISCIPLINE TEMPORELLE"])
    for item in metadata["temporal_permissions"]:
        lines.append(
            f"[{item['permission']}] élection {item['election_year']} × observation {item['observation_year']} — écart signé={item['offset_years']} ans"
        )
    lines.extend(
        [
            "",
            f"INTERCENSITAIRE : {metadata['intercensal']['decision']} — {metadata['intercensal']['reason']}",
            "",
            "COMMANDE DE REPRODUCTION",
            "python -m morocco_elections qualify hcp-indicators --candidate-2014-individuals data/staging/v10qa3/source_candidates/rgph2014_individus.xlsx --candidate-2014-households data/staging/v10qa3/source_candidates/rgph2014_menages.xlsx --candidate-2024-indicators data/staging/v10qa3/source_candidates/rgph2024_indicateurs.xlsx --baseline v10 --as-of "
            + metadata["generated_on"],
            "",
            "SUITE",
            metadata["next_action"],
            "",
            "Ce rapport est généré intégralement depuis metadata/v10qa3_hcp_indicators.json.",
            "",
        ]
    )
    return "\n".join(lines)


def qualify(
    *,
    candidate_2014_individuals: str | Path | None = None,
    candidate_2014_households: str | Path | None = None,
    candidate_2024_indicators: str | Path | None = None,
    baseline: str = "v10",
    as_of: str | None = None,
    metadata_output: str | Path | None = None,
    decision_output: str | Path | None = None,
    data_dir: str | Path | None = None,
) -> int:
    if baseline != "v10":
        raise ValueError("Seule la baseline v10 est prise en charge")
    try:
        generated_on = as_of or date.today().isoformat()
        date.fromisoformat(generated_on)
        paths = get_paths(data_dir)
        if not paths.v10_workbook.is_file():
            raise RuntimeError(f"Baseline V10 absente: {paths.v10_workbook}")
        release = json.loads(V10_REPORT.read_text(encoding="utf-8"))
        if sha256_file(paths.v10_workbook) != release["workbooks"]["V10"]:
            raise RuntimeError("Empreinte V10 différente du rapport de release")
        raw_candidates = {
            "RGPH2014_INDIVIDUALS": candidate_2014_individuals,
            "RGPH2014_HOUSEHOLDS": candidate_2014_households,
            "RGPH2024_INDICATORS": candidate_2024_indicators,
        }
        candidates: dict[str, Path | None] = {}
        for source_id, value in raw_candidates.items():
            path = Path(value).resolve() if value else None
            if path is not None and not path.is_file():
                raise RuntimeError(f"Candidat absent: {path}")
            candidates[source_id] = path
        metadata = build_metadata(paths.v10_workbook, candidates, generated_on)
        errors = validate_metadata(metadata)
        if errors:
            raise RuntimeError("; ".join(errors))
        metadata_path = Path(metadata_output).resolve() if metadata_output else DEFAULT_METADATA
        report_path = Path(decision_output).resolve() if decision_output else DEFAULT_REPORT
        metadata_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        report_path.write_text(render_report(metadata), encoding="utf-8")
        print(f"HCP_INDICATORS_QUALIFIED go={metadata['summary']['go']} no_go={metadata['summary']['no_go']}")
        return 0
    except (OSError, KeyError, RuntimeError, ValueError) as exc:
        print(f"HCP_INDICATORS_FAILED: {exc}")
        return 2
