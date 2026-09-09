from __future__ import annotations

import json
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import openpyxl

from morocco_elections.config import PROJECT_ROOT, get_paths
from morocco_elections.domains.identity import person_name_match_key
from morocco_elections.provenance import sha256_file


EXPECTED_HEADERS = (
    "numero",
    "periode",
    "date_depot",
    "depute",
    "groupe",
    "ministere",
    "sujet",
    "observation",
    "date_reponse",
)
MIN_IDENTITY_ROW_COVERAGE = 0.95
DEFAULT_OUTPUT = PROJECT_ROOT / "metadata" / "v12_parliament_qualification.json"

SOURCE_SPECS = {
    "questions_written_april_2023.xlsx": {
        "source_id": "SRC_PARLIAMENT_WRITTEN_APRIL_2023_V12",
        "resource_id": "aa72201b-e51e-4b3a-99c5-19d60fb241a2",
        "url": "https://data.gov.ma/data/dataset/05e801e6-2e35-4fd2-b49f-225a68a23ca8/resource/aa72201b-e51e-4b3a-99c5-19d60fb241a2/download/-2023.xlsx",
        "segment": "april_session_2023",
    },
    "questions_written_apr_oct_2023.xlsx": {
        "source_id": "SRC_PARLIAMENT_WRITTEN_APR_OCT_2023_V12",
        "resource_id": "bda08153-7071-42e8-beb1-2861cb644110",
        "url": "https://data.gov.ma/data/dataset/f6646b5c-00a0-4447-ae34-8324c79c446a/resource/bda08153-7071-42e8-beb1-2861cb644110/download/-2023.xlsx",
        "segment": "between_april_october_2023",
    },
    "questions_written_october_2023.xlsx": {
        "source_id": "SRC_PARLIAMENT_WRITTEN_OCTOBER_2023_V12",
        "resource_id": "2e5441bd-ce3b-41e5-b1ad-9b76a33dc8a4",
        "url": "https://data.gov.ma/data/dataset/ba99bbb6-cd38-4d1f-8ea3-d80f2d4178c7/resource/2e5441bd-ce3b-41e5-b1ad-9b76a33dc8a4/download/-2023.xlsx",
        "segment": "october_session_2023",
    },
    "questions_written_2023_2024.xlsx": {
        "source_id": "SRC_PARLIAMENT_WRITTEN_OCT_APR_2023_2024_V12",
        "resource_id": "c892cfd6-4ced-4f68-812e-30783d76f7ec",
        "url": "https://data.gov.ma/data/fr/dataset/184f26ab-b762-444f-93d0-90fa2ae820c2/resource/c892cfd6-4ced-4f68-812e-30783d76f7ec/download/-2023-2024.xlsx",
        "segment": "between_october_2023_april_2024",
    },
}


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _iso_date(value: Any, *, allow_dash: bool = False) -> str | None:
    if allow_dash and str(value).strip() in {"-", "–", "—"}:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raise ValueError(f"Date source invalide: {type(value).__name__}")


def load_question_rows(path: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if len(workbook.sheetnames) != 1:
        workbook.close()
        raise ValueError(f"Une feuille attendue dans {path.name}")
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=False)
    header_cells = next(iterator)
    headers = tuple(str(cell.value).strip() if cell.value is not None else "" for cell in header_cells)
    if headers != EXPECTED_HEADERS:
        workbook.close()
        raise ValueError(f"Schéma inattendu dans {path.name}: {headers}")
    rows: list[dict[str, Any]] = []
    for source_row, cells in enumerate(iterator, start=2):
        if all(_is_blank(cell.value) for cell in cells):
            continue
        formula_columns = [headers[index] for index, cell in enumerate(cells) if cell.data_type == "f"]
        unsafe_formula_columns = set(formula_columns) - {"sujet", "observation"}
        if unsafe_formula_columns:
            workbook.close()
            raise ValueError(
                f"Formule interdite dans une clé ou dimension de {path.name}, ligne {source_row}: "
                f"{sorted(unsafe_formula_columns)}"
            )
        record = dict(zip(headers, (cell.value for cell in cells)))
        record["source_row"] = source_row
        record["formula_text_columns"] = formula_columns
        rows.append(record)
    workbook.close()
    return rows


def load_v11_identity_index(workbook_path: Path) -> dict[str, set[str]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook["PARLIAMENTARY_MANDATES"]
    iterator = worksheet.iter_rows(min_row=4, values_only=True)
    headers = tuple(str(value).strip() if value is not None else "" for value in next(iterator))
    index: dict[str, set[str]] = defaultdict(set)
    for values in iterator:
        row = dict(zip(headers, values))
        if _is_blank(row.get("mandate_id")) or str(row.get("legislature")) != "2021-2026":
            continue
        key = person_name_match_key(row.get("full_name_ar"))
        if key:
            index[key].add(str(row["person_id"]))
    workbook.close()
    return dict(index)


def qualify_sources(
    candidates: Iterable[str | Path],
    *,
    baseline_workbook: Path,
    as_of: str,
) -> dict[str, Any]:
    paths = [Path(value).resolve() for value in candidates]
    identity_index = load_v11_identity_index(baseline_workbook)
    resources: list[dict[str, Any]] = []
    all_keys: set[tuple[str, str]] = set()
    total_rows = matched_rows = ambiguous_rows = answered_rows = 0
    matched_people: set[str] = set()
    source_people: set[str] = set()
    unmatched_people: set[str] = set()
    failures: list[str] = []

    expected_names = set(SOURCE_SPECS)
    actual_names = {path.name for path in paths}
    if actual_names != expected_names or len(paths) != len(expected_names):
        failures.append("Les quatre ressources officielles attendues ne sont pas toutes présentes une seule fois.")

    for path in sorted(paths, key=lambda item: item.name):
        spec = SOURCE_SPECS.get(path.name)
        if spec is None:
            continue
        try:
            rows = load_question_rows(path)
        except (OSError, ValueError, KeyError) as exc:
            failures.append(f"{path.name}: {exc}")
            continue
        missing = {header: sum(_is_blank(row[header]) for row in rows) for header in EXPECTED_HEADERS}
        if any(missing.values()):
            failures.append(f"{path.name}: valeurs obligatoires absentes.")
        local_keys: set[str] = set()
        local_matched = local_ambiguous = local_unmatched = local_answered = local_formula_text = 0
        deposit_dates: list[str] = []
        for row in rows:
            local_formula_text += int(bool(row["formula_text_columns"]))
            number = str(row["numero"]).strip()
            composite = (str(spec["source_id"]), number)
            if number in local_keys or composite in all_keys:
                failures.append(f"{path.name}: numéro de question dupliqué.")
            local_keys.add(number)
            all_keys.add(composite)
            try:
                deposit_dates.append(_iso_date(row["date_depot"]) or "")
                response_date = _iso_date(row["date_reponse"], allow_dash=True)
            except ValueError as exc:
                failures.append(f"{path.name}: {exc}")
                response_date = None
            if response_date:
                local_answered += 1
            name_key = person_name_match_key(row["depute"])
            source_people.add(name_key)
            matches = identity_index.get(name_key, set())
            if len(matches) == 1:
                local_matched += 1
                matched_people.update(matches)
            elif len(matches) > 1:
                local_ambiguous += 1
            else:
                local_unmatched += 1
                unmatched_people.add(name_key)
        resources.append(
            {
                "source_id": spec["source_id"],
                "resource_id": spec["resource_id"],
                "segment": spec["segment"],
                "file_name": path.name,
                "url": spec["url"],
                "producer": "Parlement du Royaume du Maroc via data.gov.ma",
                "license": "Open Data Commons Open Database License (ODbL)",
                "sha256": sha256_file(path),
                "byte_size": path.stat().st_size,
                "rows": len(rows),
                "columns": len(EXPECTED_HEADERS),
                "missing_required_values": sum(missing.values()),
                "unique_source_question_numbers": len(local_keys),
                "deposit_date_min": min(deposit_dates) if deposit_dates else None,
                "deposit_date_max": max(deposit_dates) if deposit_dates else None,
                "answered_rows": local_answered,
                "identity_rows_exact_unique": local_matched,
                "identity_rows_ambiguous": local_ambiguous,
                "identity_rows_unmatched": local_unmatched,
                "formula_like_text_rows": local_formula_text,
            }
        )
        total_rows += len(rows)
        matched_rows += local_matched
        ambiguous_rows += local_ambiguous
        answered_rows += local_answered

    identity_coverage = matched_rows / total_rows if total_rows else 0.0
    if total_rows == 0:
        failures.append("Aucune question exploitable.")
    if ambiguous_rows:
        failures.append("Au moins un nom correspond à plusieurs person_id V11.")
    if identity_coverage < MIN_IDENTITY_ROW_COVERAGE:
        failures.append("Moins de 95 % des lignes sont raccordables exactement aux identités V11.")

    return {
        "schema_version": 1,
        "phase": "V12-PARLIAMENT-QUALIFICATION",
        "as_of": as_of,
        "baseline_release": "V11",
        "baseline_sha256": sha256_file(baseline_workbook),
        "decision": "GO" if not failures else "NO_GO",
        "scope": "Questions parlementaires écrites; quatre segments officiels du cycle 2023-2024.",
        "gate": {
            "accessible": len(resources) == 4,
            "provenance_identified": len(resources) == 4,
            "reuse_authorized": len(resources) == 4,
            "grain_usable": total_rows > 0 and len(all_keys) == total_rows,
            "identities_linkable": ambiguous_rows == 0 and identity_coverage >= MIN_IDENTITY_ROW_COVERAGE,
            "minimum_identity_row_coverage": MIN_IDENTITY_ROW_COVERAGE,
        },
        "resources": resources,
        "profile": {
            "rows": total_rows,
            "columns": len(EXPECTED_HEADERS),
            "unique_question_keys": len(all_keys),
            "answered_rows": answered_rows,
            "unanswered_rows": total_rows - answered_rows,
            "distinct_source_people": len(source_people),
            "matched_v11_people": len(matched_people),
            "unmatched_source_people": len(unmatched_people),
            "identity_rows_exact_unique": matched_rows,
            "identity_rows_ambiguous": ambiguous_rows,
            "identity_rows_unmatched": total_rows - matched_rows - ambiguous_rows,
            "identity_row_coverage": round(identity_coverage, 6),
            "formula_like_text_rows": sum(resource["formula_like_text_rows"] for resource in resources),
        },
        "rules": [
            "Le grain canonique est source_id × source_question_number.",
            "Les noms sont rapprochés uniquement par clé Unicode exacte; aucun rapprochement flou.",
            "Une identité non raccordée reste NULL et la valeur RAW demeure conservée.",
            "Les textes commençant par '=' sont exportés comme chaînes inertes, jamais comme formules Excel.",
            "Le périmètre est partiel: questions écrites seulement, cycle 2023-2024 seulement.",
        ],
        "failures": failures,
    }


def qualify(
    candidates: Iterable[str | Path] | None = None,
    *,
    baseline: str = "v11",
    as_of: str = "2026-09-09",
    output: str | Path | None = None,
    data_dir: str | Path | None = None,
) -> int:
    if baseline != "v11":
        raise ValueError("Seule la baseline V11 est acceptée")
    paths = get_paths(data_dir)
    selected = list(candidates or paths.parliament_questions_2023_sources)
    report = qualify_sources(selected, baseline_workbook=paths.v11_workbook, as_of=as_of)
    destination = Path(output) if output else DEFAULT_OUTPUT
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        f"PARLIAMENT_QUALIFICATION_{report['decision']} rows={report['profile']['rows']} "
        f"identity_coverage={report['profile']['identity_row_coverage']:.4f} output={destination}"
    )
    return 0 if report["decision"] == "GO" else 1
