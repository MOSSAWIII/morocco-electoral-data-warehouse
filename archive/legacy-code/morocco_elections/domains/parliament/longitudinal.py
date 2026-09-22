from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from morocco_elections.config import PROJECT_ROOT, get_paths
from morocco_elections.domains.identity import person_name_match_key


INVENTORY_PATH = PROJECT_ROOT / "metadata" / "acquisition_inventory.json"
QUESTION_SOURCE_TOKEN = "QUESTIONS"
LEGISLATURE_BOUNDARY = "2021-10-01"
DISPLAY_FIELDS = {
    "period_label_ar",
    "institution_name_ar",
    "subject_ar",
    "author_name_ar_raw",
    "group_name_ar_raw",
}


def _text(value: Any) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).strip().split())


def _key(value: Any) -> str:
    return _text(value).casefold()


def _number(value: Any) -> str:
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _iso_date(value: Any, *, required: bool) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None or _text(value) in {"", "-", "–", "—"}:
        if required:
            raise ValueError("date obligatoire absente")
        return None
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        return value.strip()
    raise ValueError(f"date source non reconnue: {value!r}")


def _stable_id(prefix: str, *values: str) -> str:
    payload = json.dumps(values, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return f"{prefix}_{hashlib.sha256(payload).hexdigest()[:32].upper()}"


def _resource_source_id(record: dict[str, Any]) -> str:
    return f"SRC_PARLIAMENT_{record['sha256'][:16].upper()}"


def load_source_records(inventory_path: Path = INVENTORY_PATH) -> list[dict[str, Any]]:
    inventory = json.loads(inventory_path.read_text(encoding="utf-8"))
    records = [
        record
        for record in inventory["records"]
        if record["domain"] == "parliament" and QUESTION_SOURCE_TOKEN in record["source_id"]
    ]
    if len(records) != 56:
        raise RuntimeError(f"56 ressources de questions attendues, obtenu {len(records)}")
    return sorted(records, key=lambda row: (row["sha256"], row["source_id"]))


def _resource_path(record: dict[str, Any], data_dir: str | Path | None) -> Path:
    relative = Path(record["local_path"])
    return get_paths(data_dir).data_root.joinpath(*relative.parts[1:])


def _headers_and_rows(path: Path) -> tuple[list[str], list[tuple[Any, ...]], int]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        workbook.close()
        raise ValueError(f"une feuille attendue: {path.name}")
    raw = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    while raw and not any(value not in (None, "") for value in raw[0]):
        raw.pop(0)
    if not raw:
        raise ValueError(f"ressource vide: {path.name}")
    first = [_key(value) for value in raw[0]]
    if "numero" in first:
        return first, raw[1:], 2
    if len(first) not in {9, 10} or _number(raw[0][0]) in {"", "numero"}:
        raise ValueError(f"en-tête non récupérable: {path.name}")
    recovered = [
        "numero", "type", "date_depot", "periode", "depute", "groupe", "ministere",
        "sujet", "question", "date_reponse",
    ][: len(first)]
    return recovered, raw, 1


def parse_resource(record: dict[str, Any], path: Path) -> list[dict[str, Any]]:
    headers, values, first_source_row = _headers_and_rows(path)
    question_type = "written" if "observation" in headers else "oral"
    result: list[dict[str, Any]] = []
    for source_row, cells in enumerate(values, start=first_source_row):
        if not any(value not in (None, "") for value in cells):
            continue
        row = {headers[index]: cells[index] if index < len(cells) else None for index in range(len(headers))}
        number = _number(row.get("numero"))
        if not number or _key(number) == "numero":
            continue
        deposit_date = _iso_date(row.get("date_depot"), required=True)
        result.append(
            {
                "question_type": question_type,
                "source_question_number": number,
                "deposit_date": deposit_date,
                "period_raw": _text(row.get("periode")),
                "deputy_name_ar_raw": _text(row.get("depute")),
                "parliamentary_group_ar_raw": _text(row.get("groupe")),
                "ministry_ar_raw": _text(row.get("ministere")),
                "subject_ar": _text(row.get("sujet")),
                "question_text_ar": _text(row.get("question") or row.get("observation")),
                "response_date": _iso_date(row.get("date_reponse"), required=False),
                "source_family_id": record["source_id"],
                "source_id": _resource_source_id(record),
                "source_sha256": record["sha256"],
                "source_row": source_row,
            }
        )
    return result


def load_acquired_questions(
    data_dir: str | Path | None = None,
    inventory_path: Path = INVENTORY_PATH,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    occurrences: list[dict[str, Any]] = []
    unique_sources: list[dict[str, Any]] = []
    seen_payloads: set[str] = set()
    for record in load_source_records(inventory_path):
        if record["sha256"] in seen_payloads:
            continue
        seen_payloads.add(record["sha256"])
        path = _resource_path(record, data_dir)
        if not path.is_file():
            raise FileNotFoundError(path)
        occurrences.extend(parse_resource(record, path))
        unique_sources.append(record)
    if len(unique_sources) != 55 or len(occurrences) != 61_385:
        raise RuntimeError(
            f"profil parlementaire inattendu: sources={len(unique_sources)} lignes={len(occurrences)}"
        )
    return occurrences, unique_sources


def _legislature(deposit_date: str) -> str:
    return "2021-2026" if deposit_date >= LEGISLATURE_BOUNDARY else "2016-2021"


def _identity_context(mandates: list[dict[str, Any]]) -> tuple[dict[tuple[str, str], set[str]], dict[tuple[str, str], set[str]]]:
    people: dict[tuple[str, str], set[str]] = defaultdict(set)
    parties: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in mandates:
        legislature = str(row.get("legislature"))
        person_id = str(row.get("person_id"))
        name_key = person_name_match_key(row.get("full_name_ar"))
        if name_key:
            people[(legislature, name_key)].add(person_id)
        if row.get("party_id"):
            parties[(legislature, person_id)].add(str(row["party_id"]))
    return people, parties


def _register_dimension(
    target: dict[str, dict[str, Any]],
    identifier: str,
    record: dict[str, Any],
) -> None:
    current = target.get(identifier)
    if current is None:
        target[identifier] = record
        return
    structural_fields = set(record) - DISPLAY_FIELDS
    if any(current.get(field) != record.get(field) for field in structural_fields):
        raise RuntimeError(f"collision d'identifiant canonique: {identifier}")
    for field in DISPLAY_FIELDS & set(record):
        candidates = [value for value in (current.get(field), record.get(field)) if value not in (None, "")]
        current[field] = min(candidates) if candidates else None


def build_longitudinal_model(
    occurrences: list[dict[str, Any]],
    mandates: list[dict[str, Any]],
    source_records: list[dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    people_index, party_index = _identity_context(mandates)
    by_question: dict[tuple[str, str, str], dict[str, Any]] = {}
    signatures: dict[tuple[str, str, str], tuple[str, ...]] = {}
    question_sources: list[dict[str, Any]] = []
    periods: dict[str, dict[str, Any]] = {}
    institutions: dict[str, dict[str, Any]] = {}
    subjects: dict[str, dict[str, Any]] = {}
    authors: dict[str, dict[str, Any]] = {}
    groups: dict[str, dict[str, Any]] = {}

    for row in occurrences:
        natural_key = (row["question_type"], _key(row["source_question_number"]), row["deposit_date"])
        signature = tuple(
            _key(row[field])
            for field in (
                "period_raw", "deputy_name_ar_raw", "parliamentary_group_ar_raw",
                "ministry_ar_raw", "subject_ar", "question_text_ar", "response_date",
            )
        )
        if natural_key in signatures and signatures[natural_key] != signature:
            raise RuntimeError(f"question contradictoire pour la clé {natural_key}")
        signatures[natural_key] = signature
        question_id = _stable_id("QUESTION_V14", *natural_key)
        by_question.setdefault(natural_key, {**row, "question_id": question_id})
        question_sources.append(
            {
                "question_source_id": _stable_id(
                    "QSOURCE_V14", question_id, row["source_id"], str(row["source_row"])
                ),
                "question_id": question_id,
                "source_id": row["source_id"],
                "source_row": row["source_row"],
                "source_sha256": row["source_sha256"],
            }
        )

    questions: list[dict[str, Any]] = []
    responses: list[dict[str, Any]] = []
    question_authors: list[dict[str, Any]] = []
    for row in sorted(by_question.values(), key=lambda item: item["question_id"]):
        legislature = _legislature(row["deposit_date"])
        period_id = _stable_id("PERIOD_V14", _key(row["period_raw"]))
        institution_id = _stable_id("INSTITUTION_V14", _key(row["ministry_ar_raw"]))
        subject_id = _stable_id("SUBJECT_V14", _key(row["subject_ar"]))
        name_key = person_name_match_key(row["deputy_name_ar_raw"])
        author_source_key = _key(row["deputy_name_ar_raw"])
        author_id = _stable_id("AUTHOR_V14", legislature, author_source_key)
        group_id = _stable_id("GROUP_V14", legislature, _key(row["parliamentary_group_ar_raw"]))
        matches = people_index.get((legislature, name_key), set())
        person_id = next(iter(matches)) if len(matches) == 1 else None
        party_matches = party_index.get((legislature, person_id), set()) if person_id else set()
        party_id = next(iter(party_matches)) if len(party_matches) == 1 else None
        identity_match_method = (
            "unicode_exact_legislature"
            if person_id
            else "ambiguous_unicode_exact_no_link"
            if len(matches) > 1
            else "unresolved_no_exact_match"
        )

        _register_dimension(
            periods,
            period_id,
            {"period_id": period_id, "period_label_ar": row["period_raw"], "period_key": _key(row["period_raw"])},
        )
        _register_dimension(
            institutions,
            institution_id,
            {
                "institution_id": institution_id,
                "institution_type": "ministry_or_government_department",
                "institution_name_ar": row["ministry_ar_raw"],
                "institution_key": _key(row["ministry_ar_raw"]),
            },
        )
        _register_dimension(
            subjects,
            subject_id,
            {"subject_id": subject_id, "subject_ar": row["subject_ar"], "subject_key": _key(row["subject_ar"])},
        )
        _register_dimension(
            authors,
            author_id,
            {
                "author_id": author_id,
                "legislature": legislature,
                "author_name_ar_raw": row["deputy_name_ar_raw"],
                "author_source_key": author_source_key,
                "person_name_match_key": name_key,
                "person_id": person_id,
                "identity_match_method": identity_match_method,
            },
        )
        _register_dimension(
            groups,
            group_id,
            {
                "group_id": group_id,
                "legislature": legislature,
                "group_name_ar_raw": row["parliamentary_group_ar_raw"],
                "group_key": _key(row["parliamentary_group_ar_raw"]),
            },
        )
        questions.append(
            {
                "question_id": row["question_id"],
                "question_type": row["question_type"],
                "source_question_number": row["source_question_number"],
                "deposit_date": row["deposit_date"],
                "legislature": legislature,
                "period_id": period_id,
                "institution_id": institution_id,
                "subject_id": subject_id,
                "question_text_ar": row["question_text_ar"],
                "response_status": "answered" if row["response_date"] else "unanswered_as_published",
                "quality_status": "verified_source_deduplicated",
            }
        )
        question_authors.append(
            {
                "question_author_id": _stable_id("QAUTHOR_V14", row["question_id"], author_id),
                "question_id": row["question_id"],
                "author_id": author_id,
                "person_id": person_id,
                "party_id": party_id,
                "group_id": group_id,
                "affiliation_method": "mandate_party_same_legislature" if party_id else "unresolved",
            }
        )
        if row["response_date"]:
            responses.append(
                {
                    "response_id": _stable_id("RESPONSE_V14", row["question_id"], row["response_date"]),
                    "question_id": row["question_id"],
                    "response_date": row["response_date"],
                    "response_status": "answered_date_published",
                }
            )

    questions_by_id = {row["question_id"]: row for row in questions}
    trajectory_groups: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    trajectory_parties: dict[tuple[str, str, str, str], set[str]] = defaultdict(set)
    for link in question_authors:
        person_id = link["person_id"]
        if not person_id:
            continue
        question = questions_by_id[link["question_id"]]
        grain = (
            person_id,
            question["legislature"],
            question["period_id"],
            question["question_type"],
        )
        trajectory_groups[grain].append(question)
        if link["party_id"]:
            trajectory_parties[grain].add(link["party_id"])

    trajectories: list[dict[str, Any]] = []
    for grain, items in sorted(trajectory_groups.items()):
        person_id, legislature, period_id, question_type = grain
        parties = trajectory_parties[grain]
        if len(parties) > 1:
            raise RuntimeError(f"affiliation partisane contradictoire pour la trajectoire {grain}")
        question_count = len(items)
        answered_count = sum(item["response_status"] == "answered" for item in items)
        dates = sorted(item["deposit_date"] for item in items)
        trajectories.append(
            {
                "trajectory_id": _stable_id("TRAJECTORY_V14", *grain),
                "person_id": person_id,
                "party_id": next(iter(parties)) if parties else None,
                "legislature": legislature,
                "period_id": period_id,
                "question_type": question_type,
                "question_count": question_count,
                "answered_question_count": answered_count,
                "response_rate_pct": round(100 * answered_count / question_count, 2),
                "first_deposit_date": dates[0],
                "last_deposit_date": dates[-1],
                "derivation_status": "DERIVED_FROM_PUBLISHED_QUESTIONS",
            }
        )

    source_dimension = [
        {
            "source_id": _resource_source_id(record),
            "source_family_id": record["source_id"],
            "title": record["title"],
            "producer": record["producer"],
            "url": record["initial_url"],
            "license": record["reuse_status"],
            "retrieved_at": record["retrieved_at"],
            "sha256": record["sha256"],
            "byte_size": record["byte_size"],
        }
        for record in source_records
    ]
    return {
        "dim_parliamentary_period": sorted(periods.values(), key=lambda row: row["period_id"]),
        "dim_institution": sorted(institutions.values(), key=lambda row: row["institution_id"]),
        "dim_parliamentary_subject": sorted(subjects.values(), key=lambda row: row["subject_id"]),
        "dim_parliamentary_author": sorted(authors.values(), key=lambda row: row["author_id"]),
        "dim_parliamentary_group": sorted(groups.values(), key=lambda row: row["group_id"]),
        "dim_parliamentary_source": sorted(source_dimension, key=lambda row: row["source_id"]),
        "fact_parliamentary_question": questions,
        "fact_parliamentary_response": sorted(responses, key=lambda row: row["response_id"]),
        "bridge_question_author": sorted(question_authors, key=lambda row: row["question_author_id"]),
        "bridge_question_source": sorted(question_sources, key=lambda row: row["question_source_id"]),
        "analytical_parliamentary_trajectory": trajectories,
    }
