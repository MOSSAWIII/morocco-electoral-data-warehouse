from __future__ import annotations

import argparse
import ast
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Iterable

import openpyxl

from morocco_elections.config import PROJECT_ROOT, get_paths, resolve_manifest_path
from morocco_elections.provenance import sha256_file as sha256

ROOT = PROJECT_ROOT
MANIFEST_PATH = get_paths().source_manifest
BACKLOG_PATH = get_paths().github_backlog
ACQUISITION_CATALOG_PATH = PROJECT_ROOT / "metadata" / "acquisition_catalog.json"
ACQUISITION_INVENTORY_PATH = PROJECT_ROOT / "metadata" / "acquisition_inventory.json"
ONTOLOGY_PATH = PROJECT_ROOT / "metadata" / "ontology_v1.json"
V13_SOURCE_PROFILE_PATH = PROJECT_ROOT / "metadata" / "v13_electoral_sources_profile.json"
V13_SOURCE_PROFILE_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V13_ELECTORAL_SOURCES_PROFILE.txt"
V13_IDENTITY_REGISTRY_PATH = PROJECT_ROOT / "metadata" / "v13_identity_registry.json"
V13_IDENTITY_REGISTRY_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V13_IDENTITY_REGISTRY.txt"
V13_ELECTORAL_QUALIFICATION_PATH = PROJECT_ROOT / "metadata" / "v13_electoral_qualification.json"
V13_ELECTORAL_QUALIFICATION_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V13_ELECTORAL_QUALIFICATION.txt"
V10_REPORT_PATH = PROJECT_ROOT / "metadata" / "v10_release_report.json"
V11_REPORT_PATH = PROJECT_ROOT / "metadata" / "v11_release_report.json"
V12_QUALIFICATION_PATH = PROJECT_ROOT / "metadata" / "v12_parliament_qualification.json"
V12_REPORT_PATH = PROJECT_ROOT / "metadata" / "v12_release_report.json"
V12_DIFF_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V12_VS_V11_DIFF.txt"
V13_REPORT_PATH = PROJECT_ROOT / "metadata" / "v13_release_report.json"
V13_DIFF_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V13_VS_V12_DIFF.txt"
V13_OPEN_DISTRIBUTION_PATH = PROJECT_ROOT / "metadata" / "v13_open_distribution.json"
V13_OPEN_README_PATH = PROJECT_ROOT / "docs" / "publication" / "V13_OPEN_DATA_README.txt"
V14_OPEN_DISTRIBUTION_PATH = PROJECT_ROOT / "metadata" / "v14_open_distribution.json"
V14_OPEN_README_PATH = PROJECT_ROOT / "docs" / "publication" / "V14_OPEN_DATA_README.txt"
V14_1_OPEN_DISTRIBUTION_PATH = PROJECT_ROOT / "metadata" / "v14_1_open_distribution.json"
V14_1_OPEN_README_PATH = PROJECT_ROOT / "docs" / "publication" / "V14_1_OPEN_DATA_README.txt"
V11A_METADATA_PATH = PROJECT_ROOT / "metadata" / "v11a_source_candidates.json"
V11A_DECISION_PATH = PROJECT_ROOT / "docs" / "research" / "V11A_DECISION_2015_COUNCILS.txt"
SMIIG_METADATA_PATH = PROJECT_ROOT / "metadata" / "v11_smiig_source_candidates.json"
SMIIG_DECISION_PATH = PROJECT_ROOT / "docs" / "research" / "V11_SMIIG_QUALIFICATION.txt"
QUALITY_BASELINE_PATH = PROJECT_ROOT / "metadata" / "v10_quality_baseline.json"
QUALITY_BASELINE_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V10_QUALITY_BASELINE.txt"
DENOMINATORS_METADATA_PATH = PROJECT_ROOT / "metadata" / "v10qa1_electoral_denominators.json"
DENOMINATORS_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V10QA1_ELECTORAL_DENOMINATORS.txt"
PRESIDENCIES_METADATA_PATH = PROJECT_ROOT / "metadata" / "v10qa2_local_presidencies.json"
PRESIDENCIES_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V10QA2_LOCAL_PRESIDENCIES.txt"
HCP_INDICATORS_METADATA_PATH = PROJECT_ROOT / "metadata" / "v10qa3_hcp_indicators.json"
HCP_INDICATORS_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V10QA3_HCP_INDICATORS.txt"
V11_QUALITY_BASELINE_PATH = PROJECT_ROOT / "metadata" / "v11_quality_baseline.json"
V11_QUALITY_BASELINE_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V11_QUALITY_BASELINE.txt"
V11_DIFF_REPORT_PATH = PROJECT_ROOT / "docs" / "research" / "V11_VS_V10_DIFF.txt"
DOCUMENTATION_DIRS = {
    "v9": get_paths().documentation_v9,
    "v10": get_paths().documentation_v10,
    "v11": get_paths().documentation_v11,
    "v12": get_paths().documentation_v12,
    "v13": get_paths().documentation_v13,
}
EXPECTED_DOCUMENTS = [
    "00_INDEX_ET_MODE_EMPLOI.txt",
    "01_ONTOLOGIE_ELECTORALE_GLOBALE.txt",
    "02_ARCHITECTURE_DES_CUBES.txt",
    "03_DIMENSIONS_CONFORMES.txt",
    "04_HIERARCHIE_GEOGRAPHIQUE.txt",
    "05_CROSSWALKS_ET_IDENTITES.txt",
    "06_TEMPS_CYCLES_ET_ELECTIONS.txt",
    "07_FAITS_ELECTORAUX_ET_MOBILISATION.txt",
    "08_PERSONNES_MANDATS_ET_REPRESENTATION.txt",
    "09_POUVOIR_LOCAL_ET_GOUVERNANCE.txt",
    "10_SOCIO_ECONOMIE_CAMPAGNE_INFORMATION.txt",
    "11_PIPELINE_RAW_NORMALIZED_FACT_ANALYTICAL.txt",
    "12_REGLES_D_INTEGRITE_ET_JOINTURE.txt",
    "13_METRIQUES_DERIVEES_ET_REGLES_TEMPORELLES.txt",
    "14_CATALOGUE_COMPLET_DES_ONGLETS.txt",
]
REQUIRED_DOCUMENT_METADATA = (
    "VERSION",
    "CLASSEUR SOURCE",
    "DATE DE GÉNÉRATION",
    "PÉRIMÈTRE",
    "RENVOIS",
)
REQUIRED_SOURCE_FIELDS = {
    "source_id",
    "local_path",
    "source_url",
    "producer",
    "license",
    "version",
    "sha256",
    "byte_size",
    "grain",
    "rows",
    "columns",
    "status",
}
REQUIRED_ARTIFACT_FIELDS = {
    "artifact_id",
    "local_path",
    "sha256",
    "byte_size",
    "sheet_count",
    "status",
}
REQUIRED_PHYSICAL_FILE_FIELDS = {
    "file_id",
    "previous_local_path",
    "local_path",
    "sha256",
    "byte_size",
    "kind",
}
FORBIDDEN_SUFFIXES = {
    ".xlsx",
    ".xls",
    ".xlsm",
    ".parquet",
    ".shp",
    ".shx",
    ".dbf",
    ".prj",
    ".cpg",
    ".zip",
    ".7z",
    ".dump",
    ".backup",
    ".pem",
    ".key",
}
FORBIDDEN_DIRECTORIES = {"raw_sources", "artifacts", "backups", "secrets", "pgdata", "postgres-data"}
FORBIDDEN_DATA_PREFIXES = {
    ("data", "raw"),
    ("data", "legacy"),
    ("data", "staging"),
    ("data", "processed"),
    ("data", "exports"),
}
MAX_TRACKED_BYTES = 5 * 1024 * 1024
EXPECTED_V9_VOLUMES = {
    "LOCAL_MANDATES": 32513,
    "PARLIAMENTARY_MANDATES": 1654,
    "COMMUNE_ELECTION_PANEL": 3076,
    "ELECTORAL_TRANSITIONS_2015_2021": 14555,
    "ANALYTICAL_PANEL": 22054,
}
EXPECTED_V10_VOLUMES = {
    **EXPECTED_V9_VOLUMES,
    "DIM_PERSON": 33697,
    "IDENTITY_AUDIT_V10": 44,
    "MANUAL_RESOLUTIONS_V10": 7,
    "COUNCIL_SEAT_STATUS_V10": 1538,
}
EXPECTED_V11_VOLUMES = {**EXPECTED_V10_VOLUMES, "FACT_OBSERVATION": 3_203}
EXPECTED_V12_VOLUMES = {**EXPECTED_V11_VOLUMES, "PARLIAMENTARY_QUESTIONS": 5_589}
EXPECTED_V13_VOLUMES = {
    **EXPECTED_V12_VOLUMES,
    "DIM_ELECTORAL_CONTEST": 639,
    "FACT_ELECTION_RESULT": 10_883,
    "FACT_ELECTORAL_MOBILIZATION": 639,
}
REQUIRED_EVIDENCE_FIELDS = {"evidence_id", "local_path", "source_url", "publisher", "sha256", "byte_size", "status"}


class ValidationError(RuntimeError):
    pass


def load_manifest(path: Path | None = None) -> dict:
    path = path or MANIFEST_PATH
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValidationError(f"Manifeste illisible: {path}: {exc}") from exc


def validate_manifest(manifest: dict) -> list[str]:
    errors: list[str] = []
    if manifest.get("schema_version") != 5:
        errors.append("metadata/source_manifest.json: schema_version doit valoir 5")
    if manifest.get("warehouse_version") != "V13":
        errors.append("metadata/source_manifest.json: warehouse_version doit valoir V13")

    sources = manifest.get("sources")
    artifacts = manifest.get("artifacts")
    physical_files = manifest.get("physical_files")
    evidence = manifest.get("evidence")
    if not isinstance(sources, list) or not sources:
        errors.append("Le manifeste doit contenir une liste sources non vide")
        sources = []
    if not isinstance(artifacts, list) or not artifacts:
        errors.append("Le manifeste doit contenir une liste artifacts non vide")
        artifacts = []
    if not isinstance(physical_files, list) or len(physical_files) != 37:
        errors.append("Le manifeste doit inventorier exactement 37 fichiers physiques")
        physical_files = []
    if not isinstance(evidence, list) or len(evidence) != 2:
        errors.append("Le manifeste doit inventorier exactement deux preuves officielles V10")
        evidence = []

    seen_ids: set[str] = set()
    for kind, records, required, id_field in (
        ("source", sources, REQUIRED_SOURCE_FIELDS, "source_id"),
        ("artifact", artifacts, REQUIRED_ARTIFACT_FIELDS, "artifact_id"),
        ("physical_file", physical_files, REQUIRED_PHYSICAL_FILE_FIELDS, "file_id"),
        ("evidence", evidence, REQUIRED_EVIDENCE_FIELDS, "evidence_id"),
    ):
        seen_paths: set[str] = set()
        for index, record in enumerate(records):
            label = f"{kind}[{index}]"
            if not isinstance(record, dict):
                errors.append(f"{label}: objet JSON attendu")
                continue
            missing = sorted(required - record.keys())
            if missing:
                errors.append(f"{label}: champs absents: {', '.join(missing)}")
            record_id = str(record.get(id_field, ""))
            local_path = str(record.get("local_path", ""))
            if not record_id or record_id in seen_ids:
                errors.append(f"{label}: identifiant vide ou dupliqué: {record_id!r}")
            if not local_path or local_path in seen_paths or Path(local_path).is_absolute() or ".." in Path(local_path).parts:
                errors.append(f"{label}: local_path invalide ou dupliqué: {local_path!r}")
            seen_ids.add(record_id)
            seen_paths.add(local_path)
            if not re.fullmatch(r"[0-9a-f]{64}", str(record.get("sha256", ""))):
                errors.append(f"{label}: SHA-256 invalide")
            for field in ("byte_size", "rows", "columns", "sheet_count"):
                if field in record and (not isinstance(record[field], int) or record[field] <= 0):
                    errors.append(f"{label}: {field} doit être un entier positif")
    return errors


def validate_backlog() -> list[str]:
    errors: list[str] = []
    try:
        backlog = json.loads(BACKLOG_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"Backlog GitHub illisible: {exc}"]
    milestones = backlog.get("milestones", [])
    labels = backlog.get("labels", [])
    issues = backlog.get("issues", [])
    milestone_titles = {item.get("title") for item in milestones if isinstance(item, dict)}
    label_names = {item.get("name") for item in labels if isinstance(item, dict)}
    issue_titles = {item.get("title") for item in issues if isinstance(item, dict)}
    if backlog.get("schema_version") != 2:
        errors.append("Le backlog GitHub doit utiliser schema_version=2")
    if len(milestones) != 4 or len(milestone_titles) != 4:
        errors.append("Le backlog doit définir exactement quatre jalons uniques")
    if len(issues) != 8 or len(issue_titles) != 8:
        errors.append("Le backlog doit définir exactement huit issues uniques")
    required_labels = {"data", "qa", "source", "identity", "governance", "parliament", "postgresql", "blocked"}
    if label_names != required_labels:
        errors.append(f"Labels GitHub incorrects: {sorted(label_names)}")
    for milestone in milestones:
        if milestone.get("state") not in {"open", "closed"}:
            errors.append(f"Jalon sans état valide: {milestone.get('title')}")
    for issue in issues:
        if issue.get("milestone") not in milestone_titles:
            errors.append(f"Issue sans jalon valide: {issue.get('title')}")
        if not issue.get("acceptance"):
            errors.append(f"Issue sans critères d’acceptation: {issue.get('title')}")
        if issue.get("state") not in {"open", "closed"}:
            errors.append(f"Issue sans état valide: {issue.get('title')}")
        if issue.get("state") == "closed" and issue.get("close_reason") not in {"completed", "not planned"}:
            errors.append(f"Issue fermée sans motif valide: {issue.get('title')}")
        if issue.get("state") == "open" and issue.get("close_reason") is not None:
            errors.append(f"Issue ouverte avec un motif de fermeture: {issue.get('title')}")
        unknown_labels = set(issue.get("labels", [])) - label_names
        if unknown_labels:
            errors.append(f"Issue {issue.get('title')}: labels inconnus {sorted(unknown_labels)}")
        unknown_dependencies = set(issue.get("depends_on", [])) - issue_titles
        if unknown_dependencies:
            errors.append(f"Issue {issue.get('title')}: dépendances inconnues {sorted(unknown_dependencies)}")
        unknown_blocks = set(issue.get("blocks", [])) - issue_titles
        if unknown_blocks:
            errors.append(f"Issue {issue.get('title')}: blocages inconnus {sorted(unknown_blocks)}")
    return errors


def validate_acquisition_catalog() -> list[str]:
    from morocco_elections.sources.acquisition import load_catalog, validate_catalog

    try:
        catalog = load_catalog(ACQUISITION_CATALOG_PATH)
    except RuntimeError as exc:
        return [str(exc)]
    return [f"Catalogue d'acquisition: {error}" for error in validate_catalog(catalog)]


def validate_acquisition_inventory(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.sources.acquisition import build_inventory, validate_inventory

    try:
        inventory = json.loads(ACQUISITION_INVENTORY_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"Inventaire d'acquisition illisible: {exc}"]
    errors = [f"Inventaire d'acquisition: {error}" for error in validate_inventory(inventory)]
    if full and not errors:
        try:
            rebuilt = build_inventory(data_dir, str(inventory["generated_on"]))
        except RuntimeError as exc:
            return [f"Inventaire d'acquisition: {exc}"]
        if rebuilt != inventory:
            errors.append("Inventaire d'acquisition non reproductible depuis les RAW locaux")
    return errors


def validate_ontology_contract() -> list[str]:
    from morocco_elections.ontology import load_ontology, validate_ontology

    try:
        ontology = load_ontology(ONTOLOGY_PATH)
    except (OSError, json.JSONDecodeError) as exc:
        return [f"Ontologie V1 illisible: {exc}"]
    return [f"Ontologie V1: {error}" for error in validate_ontology(ontology)]


def validate_v13_source_profile(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.research import electoral_archives

    try:
        profile = json.loads(V13_SOURCE_PROFILE_PATH.read_text(encoding="utf-8"))
        report = V13_SOURCE_PROFILE_REPORT_PATH.read_text(encoding="utf-8")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"Profil des archives électorales V13 illisible: {exc}"]
    errors = [f"Profil V13: {error}" for error in electoral_archives.validate_profile(profile)]
    if report != electoral_archives.render_report(profile):
        errors.append("Profil V13: rapport TXT désynchronisé du JSON")
    if full and not errors:
        try:
            rebuilt = electoral_archives.build_profile(data_dir, str(profile["as_of"]), "v12")
        except (OSError, RuntimeError, ValueError) as exc:
            errors.append(f"Profil V13: reconstruction impossible: {exc}")
        else:
            if rebuilt != profile:
                errors.append("Profil V13 non reproductible depuis les RAW et V12 locaux")
    return errors


def validate_v13_identity_registry(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.domains.identity import registry as identity_registry

    try:
        registry = json.loads(V13_IDENTITY_REGISTRY_PATH.read_text(encoding="utf-8"))
        report = V13_IDENTITY_REGISTRY_REPORT_PATH.read_text(encoding="utf-8")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"Registre d'identités V13 illisible: {exc}"]
    errors = [f"Registre V13: {error}" for error in identity_registry.validate_registry(registry)]
    try:
        inventory = json.loads(ACQUISITION_INVENTORY_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        errors.append(f"Registre V13: inventaire d'acquisition illisible: {exc}")
    else:
        acquisition_ids = {item.get("acquisition_id") for item in inventory.get("records", [])}
        evidence_ids = {item.get("evidence_id") for item in registry.get("crosswalks", [])}
        unknown_evidence = evidence_ids - acquisition_ids
        if unknown_evidence:
            errors.append(f"Registre V13: preuves absentes de l'inventaire: {sorted(unknown_evidence)}")
    if report != identity_registry.render_report(registry):
        errors.append("Registre V13: rapport TXT désynchronisé du JSON")
    if full and not errors:
        try:
            rebuilt = identity_registry.build_registry(data_dir, str(registry["as_of"]), "v12")
        except (OSError, RuntimeError, ValueError) as exc:
            errors.append(f"Registre V13: reconstruction impossible: {exc}")
        else:
            if rebuilt != registry:
                errors.append("Registre V13 non reproductible depuis les RAW et V12 locaux")
    return errors


def validate_v13_electoral_qualification() -> list[str]:
    from morocco_elections.research import electoral_qualification

    try:
        qualification = json.loads(V13_ELECTORAL_QUALIFICATION_PATH.read_text(encoding="utf-8"))
        report = V13_ELECTORAL_QUALIFICATION_REPORT_PATH.read_text(encoding="utf-8")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"Qualification électorale V13 illisible: {exc}"]
    errors = [
        f"Qualification électorale V13: {error}"
        for error in electoral_qualification.validate_qualification(qualification)
    ]
    if report != electoral_qualification.render_report(qualification):
        errors.append("Qualification électorale V13: rapport TXT désynchronisé du JSON")
    if not errors:
        try:
            rebuilt = electoral_qualification.build_qualification(
                str(qualification["as_of"]), str(qualification["baseline_release"])
            )
        except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
            errors.append(f"Qualification électorale V13: reconstruction impossible: {exc}")
        else:
            if rebuilt != qualification:
                errors.append("Qualification électorale V13 non reproductible depuis le profil et le registre")
    return errors


def validate_v10_release_report(manifest: dict) -> list[str]:
    try:
        report = json.loads(V10_REPORT_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"Rapport V10 illisible: {exc}"]
    errors = []
    if report.get("release") != "V10" or report.get("baseline") != "V9" or report.get("status") != "validated":
        errors.append("Le rapport de release doit décrire V10 validé contre V9")
    artifacts = {item["artifact_id"]: item for item in manifest.get("artifacts", [])}
    for version in ("V8", "V9", "V10"):
        artifact = artifacts.get(f"WAREHOUSE_{version}_{'INPUT' if version == 'V8' else 'EXPORT'}")
        if not artifact or report.get("workbooks", {}).get(version) != artifact.get("sha256"):
            errors.append(f"Empreinte {version} incohérente dans le rapport de release")
    for sheet, expected in EXPECTED_V10_VOLUMES.items():
        if report.get("volumes", {}).get(sheet) != expected:
            errors.append(f"Volume {sheet} incohérent dans le rapport de release")
    return errors


def validate_v11_release_report(manifest: dict) -> list[str]:
    from morocco_elections.releases.v11 import build as v11

    try:
        report = json.loads(V11_REPORT_PATH.read_text(encoding="utf-8"))
        diff_report = V11_DIFF_REPORT_PATH.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"Rapport V11 illisible: {exc}"]
    errors: list[str] = []
    if report.get("release") != "V11" or report.get("baseline") != "V10" or report.get("status") != "validated":
        errors.append("Le rapport de release doit décrire V11 validé contre V10")
    artifacts = {item["artifact_id"]: item for item in manifest.get("artifacts", [])}
    for version in ("V8", "V9", "V10", "V11"):
        artifact = artifacts.get(f"WAREHOUSE_{version}_{'INPUT' if version == 'V8' else 'EXPORT'}")
        if not artifact or report.get("workbooks", {}).get(version) != artifact.get("sha256"):
            errors.append(f"Empreinte {version} incohérente dans le rapport V11")
    volumes = report.get("volumes", {})
    if volumes.get("sheets") != 67 or volumes.get("FACT_OBSERVATION_V11") != 3_203:
        errors.append("Volumes principaux V11 incohérents")
    controls = report.get("controls", {})
    expected_controls = {
        "population_legal_2014_sum": 33_848_242,
        "population_municipal_2024_sum": 36_490_591,
        "population_legal_2024_sum": 36_828_330,
        "population_legal_2024_differences": 0,
        "no_go_ingested": 0,
        "derived_values_ingested": 0,
        "electoral_panels_changed": False,
    }
    for key, expected in expected_controls.items():
        if controls.get(key) != expected:
            errors.append(f"Contrôle V11 incohérent: {key}")
    if set(report.get("actual_changed_sheets", [])) != v11.ALLOWED_CHANGED_SHEETS:
        errors.append("Liste des onglets V11 modifiés incohérente")
    if diff_report != v11.render_diff_report(report):
        errors.append("Rapport TXT V11/V10 désynchronisé du JSON")
    return errors


def validate_v12_qualification(manifest: dict) -> list[str]:
    try:
        report = json.loads(V12_QUALIFICATION_PATH.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"Qualification parlementaire V12 illisible: {exc}"]
    errors: list[str] = []
    if report.get("schema_version") != 1 or report.get("phase") != "V12-PARLIAMENT-QUALIFICATION":
        errors.append("Qualification parlementaire V12: métadonnées racine invalides")
    if report.get("decision") != "GO" or report.get("failures"):
        errors.append("Qualification parlementaire V12: décision GO sans échec attendue")
    gate = report.get("gate", {})
    for key in ("accessible", "provenance_identified", "reuse_authorized", "grain_usable", "identities_linkable"):
        if gate.get(key) is not True:
            errors.append(f"Qualification parlementaire V12: gate non validé {key}")
    profile = report.get("profile", {})
    expected = {
        "rows": 5_589,
        "columns": 9,
        "unique_question_keys": 5_589,
        "answered_rows": 3_015,
        "unanswered_rows": 2_574,
        "identity_rows_exact_unique": 5_453,
        "identity_rows_ambiguous": 0,
        "identity_rows_unmatched": 136,
    }
    for key, value in expected.items():
        if profile.get(key) != value:
            errors.append(f"Qualification parlementaire V12: profil incohérent {key}")
    resources = report.get("resources", [])
    manifest_sources = {item["source_id"]: item for item in manifest.get("sources", [])}
    if len(resources) != 4:
        errors.append("Qualification parlementaire V12: quatre ressources attendues")
    for resource in resources:
        source = manifest_sources.get(resource.get("source_id"))
        if not source or source.get("sha256") != resource.get("sha256") or source.get("byte_size") != resource.get("byte_size"):
            errors.append(f"Qualification parlementaire V12: manifeste incohérent pour {resource.get('source_id')}")
    return errors


def validate_v12_release_report(manifest: dict) -> list[str]:
    from morocco_elections.releases.v12 import build as v12

    try:
        report = json.loads(V12_REPORT_PATH.read_text(encoding="utf-8"))
        diff_report = V12_DIFF_REPORT_PATH.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"Rapport V12 illisible: {exc}"]
    errors: list[str] = []
    if report.get("release") != "V12" or report.get("baseline") != "V11" or report.get("status") != "validated":
        errors.append("Le rapport de release doit décrire V12 validé contre V11")
    artifacts = {item["artifact_id"]: item for item in manifest.get("artifacts", [])}
    for version in ("V8", "V9", "V10", "V11", "V12"):
        artifact = artifacts.get(f"WAREHOUSE_{version}_{'INPUT' if version == 'V8' else 'EXPORT'}")
        if not artifact or report.get("workbooks", {}).get(version) != artifact.get("sha256"):
            errors.append(f"Empreinte {version} incohérente dans le rapport V12")
    if report.get("volumes", {}).get("sheets") != 68 or report.get("volumes", {}).get("PARLIAMENTARY_QUESTIONS") != 5_589:
        errors.append("Volumes principaux V12 incohérents")
    expected_controls = {
        "unique_question_ids": 5_589,
        "identity_rows_linked": 5_453,
        "identity_rows_unlinked": 136,
        "answered_rows": 3_015,
        "unanswered_rows": 2_574,
        "fuzzy_matches": 0,
    }
    for key, expected in expected_controls.items():
        if report.get("controls", {}).get(key) != expected:
            errors.append(f"Contrôle V12 incohérent: {key}")
    if set(report.get("actual_changed_sheets", [])) != v12.ALLOWED_CHANGED_SHEETS:
        errors.append("Liste des onglets V12 modifiés incohérente")
    if diff_report != v12.render_diff_report(report):
        errors.append("Rapport TXT V12/V11 désynchronisé du JSON")
    return errors


def validate_v13_release_report(manifest: dict) -> list[str]:
    from morocco_elections.releases.v13 import build as v13

    try:
        report = json.loads(V13_REPORT_PATH.read_text(encoding="utf-8"))
        diff_report = V13_DIFF_REPORT_PATH.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"Rapport V13 illisible: {exc}"]
    errors: list[str] = []
    if report.get("release") != "V13" or report.get("baseline") != "V12" or report.get("status") != "validated":
        errors.append("Le rapport de release doit décrire V13 validé contre V12")
    artifacts = {item["artifact_id"]: item for item in manifest.get("artifacts", [])}
    for version in ("V8", "V9", "V10", "V11", "V12", "V13"):
        artifact = artifacts.get(f"WAREHOUSE_{version}_{'INPUT' if version == 'V8' else 'EXPORT'}")
        if not artifact or report.get("workbooks", {}).get(version) != artifact.get("sha256"):
            errors.append(f"Empreinte {version} incohérente dans le rapport V13")
    manifest_sources = {item["source_id"]: item for item in manifest.get("sources", [])}
    if len(report.get("source_hashes", {})) != 6:
        errors.append("Six sources électorales V13 sont attendues")
    for source_id, digest in report.get("source_hashes", {}).items():
        if manifest_sources.get(source_id, {}).get("sha256") != digest:
            errors.append(f"Source V13 absente ou incohérente dans le manifeste: {source_id}")
    expected_volumes = {
        "sheets": 71,
        "DIM_ELECTORAL_CONTEST": 639,
        "FACT_ELECTION_RESULT": 10_883,
        "FACT_ELECTORAL_MOBILIZATION": 639,
    }
    for key, expected in expected_volumes.items():
        if report.get("volumes", {}).get(key) != expected:
            errors.append(f"Volume V13 incohérent: {key}")
    expected_controls = {
        "new_geo_entities": 159,
        "new_party_entities": 4,
        "unique_result_ids": 10_883,
        "unique_contest_ids": 639,
        "technical_zero_imputations": 0,
        "archive_2002_rows_ingested": 0,
    }
    for key, expected in expected_controls.items():
        if report.get("controls", {}).get(key) != expected:
            errors.append(f"Contrôle V13 incohérent: {key}")
    if set(report.get("actual_changed_sheets", [])) != v13.ALLOWED_CHANGED_SHEETS:
        errors.append("Liste des onglets V13 modifiés incohérente")
    if report.get("decision_hashes", {}).get("identity_registry") != sha256(V13_IDENTITY_REGISTRY_PATH):
        errors.append("Empreinte du registre d'identités incohérente dans V13")
    if report.get("decision_hashes", {}).get("electoral_qualification") != sha256(V13_ELECTORAL_QUALIFICATION_PATH):
        errors.append("Empreinte de la qualification électorale incohérente dans V13")
    if diff_report != v13.render_diff_report(report):
        errors.append("Rapport TXT V13/V12 désynchronisé du JSON")
    return errors


def validate_v13_open_distribution(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.exports import open_v13

    errors: list[str] = []
    if not V13_OPEN_DISTRIBUTION_PATH.is_file() or not V13_OPEN_README_PATH.is_file():
        return ["V13 open: manifeste ou README suivi absent"]
    try:
        metadata = json.loads(V13_OPEN_DISTRIBUTION_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"V13 open: manifeste illisible: {exc}"]
    errors.extend(f"V13 open: {error}" for error in open_v13.validate_manifest(metadata))
    release = json.loads(V13_REPORT_PATH.read_text(encoding="utf-8"))
    if metadata.get("source_workbook_sha256") != release.get("workbooks", {}).get("V13"):
        errors.append("V13 open: empreinte du classeur source incohérente")
    if V13_OPEN_README_PATH.read_text(encoding="utf-8") != open_v13.render_readme(metadata):
        errors.append("V13 open: README désynchronisé du manifeste")
    if not full or errors:
        return errors

    root = get_paths(data_dir).data_root / "exports" / "open" / "v13"
    local_manifest = root / "manifest.json"
    local_readme = root / "README.txt"
    checksum_path = root / "checksums.sha256"
    database_path = root / "morocco_elections_v13.duckdb"
    for required in (local_manifest, local_readme, checksum_path, database_path):
        if not required.is_file():
            errors.append(f"V13 open: fichier local absent: {required}")
    if errors:
        return errors
    if json.loads(local_manifest.read_text(encoding="utf-8")) != metadata:
        errors.append("V13 open: manifeste local désynchronisé")
    if local_readme.read_text(encoding="utf-8") != open_v13.render_readme(metadata):
        errors.append("V13 open: README local désynchronisé")
    expected_checksums: list[str] = []
    for item in metadata["files"]:
        path = root / item["path"]
        if not path.is_file():
            errors.append(f"V13 open: fichier absent: {item['path']}")
            continue
        digest = sha256(path)
        digest_mismatch = "sha256" in item and digest != item["sha256"]
        size_mismatch = "byte_size" in item and path.stat().st_size != item["byte_size"]
        if size_mismatch or digest_mismatch:
            errors.append(f"V13 open: empreinte ou taille divergente: {item['path']}")
        expected_checksums.append(f"{digest}  {item['path']}\n")
    expected_checksums.append(f"{sha256(local_manifest)}  manifest.json\n")
    if checksum_path.read_text(encoding="utf-8") != "".join(expected_checksums):
        errors.append("V13 open: checksums.sha256 désynchronisé")

    import duckdb

    connection = duckdb.connect(str(database_path), read_only=True)
    try:
        for table in metadata["tables"]:
            count = connection.execute(f'SELECT COUNT(*) FROM "{table["table_name"]}"').fetchone()[0]
            if count != table["rows"]:
                errors.append(f"V13 open: volume DuckDB divergent: {table['table_name']}")
        orphan_results = connection.execute(
            "SELECT COUNT(*) FROM fact_election_result r "
            "LEFT JOIN dim_electoral_contest c USING (contest_id) WHERE c.contest_id IS NULL"
        ).fetchone()[0]
        if orphan_results:
            errors.append("V13 open: résultats sans contest_id dans DuckDB")
    finally:
        connection.close()
    return errors


def validate_v14_open_distribution(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.domains.parliament.longitudinal import load_source_records
    from morocco_elections.exports import open_v14

    errors: list[str] = []
    if not V14_OPEN_DISTRIBUTION_PATH.is_file() or not V14_OPEN_README_PATH.is_file():
        return ["V14 open: manifeste ou README suivi absent"]
    try:
        metadata = json.loads(V14_OPEN_DISTRIBUTION_PATH.read_text(encoding="utf-8"))
        release = json.loads(V13_REPORT_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"V14 open: artefact illisible: {exc}"]
    errors.extend(f"V14 open: {error}" for error in open_v14.validate_manifest(metadata))
    if metadata.get("baseline_workbook_sha256") != release.get("workbooks", {}).get("V13"):
        errors.append("V14 open: empreinte du classeur V13 de base incohérente")
    if V14_OPEN_README_PATH.read_text(encoding="utf-8") != open_v14.render_readme(metadata):
        errors.append("V14 open: README désynchronisé du manifeste")

    expected_payloads: set[tuple[str, str, int]] = set()
    seen_payloads: set[str] = set()
    for record in load_source_records():
        if record["sha256"] in seen_payloads:
            continue
        seen_payloads.add(record["sha256"])
        expected_payloads.add((record["source_id"], record["sha256"], record["byte_size"]))
    source_manifest = load_manifest()
    for record in source_manifest["sources"]:
        family_id = open_v14.V13_SOURCE_FAMILIES.get(record["source_id"])
        if family_id:
            expected_payloads.add((family_id, record["sha256"], record["byte_size"]))
    actual_payloads = {
        (item.get("source_family_id"), item.get("sha256"), item.get("byte_size"))
        for item in metadata.get("source_payloads", [])
    }
    if actual_payloads != expected_payloads:
        errors.append("V14 open: provenance différente de l'inventaire d'acquisition")
    if not full or errors:
        return errors

    root = get_paths(data_dir).data_root / "exports" / "open" / "v14"
    local_manifest = root / "manifest.json"
    local_readme = root / "README.txt"
    checksum_path = root / "checksums.sha256"
    database_path = root / "morocco_elections_v14.duckdb"
    for required in (local_manifest, local_readme, checksum_path, database_path):
        if not required.is_file():
            errors.append(f"V14 open: fichier local absent: {required}")
    if errors:
        return errors
    if json.loads(local_manifest.read_text(encoding="utf-8")) != metadata:
        errors.append("V14 open: manifeste local désynchronisé")
    if local_readme.read_text(encoding="utf-8") != open_v14.render_readme(metadata):
        errors.append("V14 open: README local désynchronisé")
    v13_root = get_paths(data_dir).data_root / "exports" / "open" / "v13"
    for table_name in set(open_v14.open_v13.TABLES) - {"fact_parliamentary_activity"}:
        baseline_table = v13_root / "parquet" / f"{table_name}.parquet"
        current_table = root / "parquet" / f"{table_name}.parquet"
        if not baseline_table.is_file() or not current_table.is_file():
            errors.append(f"V14 open: table de comparaison absente: {table_name}")
        elif sha256(baseline_table) != sha256(current_table):
            errors.append(f"V14 open: table V13 modifiée: {table_name}")
    expected_checksums: list[str] = []
    for item in metadata["files"]:
        path = root / item["path"]
        if not path.is_file():
            errors.append(f"V14 open: fichier absent: {item['path']}")
            continue
        digest = sha256(path)
        if ("sha256" in item and digest != item["sha256"]) or (
            "byte_size" in item and path.stat().st_size != item["byte_size"]
        ):
            errors.append(f"V14 open: empreinte ou taille divergente: {item['path']}")
        expected_checksums.append(f"{digest}  {item['path']}\n")
    expected_checksums.append(f"{sha256(local_manifest)}  manifest.json\n")
    if checksum_path.read_text(encoding="utf-8") != "".join(expected_checksums):
        errors.append("V14 open: checksums.sha256 désynchronisé")

    import duckdb

    connection = duckdb.connect(str(database_path), read_only=True)
    try:
        for table in metadata["tables"]:
            table_name = table["table_name"]
            count = connection.execute(f'SELECT COUNT(*) FROM "{table_name}"').fetchone()[0]
            if count != table["rows"]:
                errors.append(f"V14 open: volume DuckDB divergent: {table_name}")
        for relationship in metadata["relationships"]:
            child = relationship["child_table"]
            child_column = relationship["child_column"]
            parent = relationship["parent_table"]
            parent_column = relationship["parent_column"]
            if not relationship["nullable"]:
                nulls = connection.execute(
                    f'SELECT COUNT(*) FROM "{child}" WHERE "{child_column}" IS NULL'
                ).fetchone()[0]
                if nulls:
                    errors.append(f"V14 open: clé obligatoire vide: {child}.{child_column}")
            orphans = connection.execute(
                f'SELECT COUNT(*) FROM "{child}" c LEFT JOIN "{parent}" p '
                f'ON c."{child_column}" = p."{parent_column}" '
                f'WHERE c."{child_column}" IS NOT NULL AND p."{parent_column}" IS NULL'
            ).fetchone()[0]
            if orphans:
                errors.append(f"V14 open: clé étrangère orpheline: {child}.{child_column}")
        trajectory_total = connection.execute(
            "SELECT SUM(question_count) FROM analytical_parliamentary_trajectory"
        ).fetchone()[0]
        if trajectory_total != metadata["controls"]["identity_linked_questions"]:
            errors.append("V14 open: trajectoires non réconciliées avec les questions raccordées")
        invalid_rates = connection.execute(
            "SELECT COUNT(*) FROM analytical_parliamentary_trajectory "
            "WHERE question_count <= 0 OR answered_question_count < 0 "
            "OR answered_question_count > question_count "
            "OR ABS(response_rate_pct - 100.0 * answered_question_count / question_count) > 0.011"
        ).fetchone()[0]
        if invalid_rates:
            errors.append("V14 open: taux de réponse de trajectoire invalide")
        v13_database = (v13_root / "morocco_elections_v13.duckdb").as_posix().replace("'", "''")
        connection.execute(f"ATTACH '{v13_database}' AS baseline_v13 (READ_ONLY)")
        missing_baseline = connection.execute(
            "SELECT COUNT(*) FROM baseline_v13.fact_parliamentary_activity old "
            "LEFT JOIN fact_parliamentary_question new "
            "ON lower(trim(cast(old.question_type AS varchar))) = lower(trim(cast(new.question_type AS varchar))) "
            "AND trim(cast(old.source_question_number AS varchar)) = trim(cast(new.source_question_number AS varchar)) "
            "AND cast(old.deposit_date AS date) = cast(new.deposit_date AS date) "
            "WHERE new.question_id IS NULL"
        ).fetchone()[0]
        if missing_baseline:
            errors.append("V14 open: des questions V13 ont disparu de la série longitudinale")
    finally:
        connection.close()
    return errors


def validate_v14_1_open_distribution(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.exports import open_v14_1

    errors: list[str] = []
    try:
        metadata = json.loads(V14_1_OPEN_DISTRIBUTION_PATH.read_text(encoding="utf-8"))
        baseline = json.loads(V14_OPEN_DISTRIBUTION_PATH.read_text(encoding="utf-8"))
        readme = V14_1_OPEN_README_PATH.read_text(encoding="utf-8")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"V14.1 open: artefact illisible: {exc}"]
    errors.extend(f"V14.1 open: {error}" for error in open_v14_1.validate_manifest(metadata))
    if metadata.get("source_payloads") != baseline.get("source_payloads"):
        errors.append("V14.1 open: les sources V14 ont été modifiées")
    if readme != open_v14_1.render_readme(metadata):
        errors.append("V14.1 open: README désynchronisé")
    if not full or errors:
        return errors

    root = get_paths(data_dir).data_root / "exports" / "open" / "v14.1"
    required = [
        root / "manifest.json", root / "README.txt", root / "checksums.sha256",
        root / "morocco_elections_v14_1.duckdb",
    ]
    for path in required:
        if not path.is_file():
            errors.append(f"V14.1 open: fichier local absent: {path}")
    if errors:
        return errors
    if json.loads((root / "manifest.json").read_text(encoding="utf-8")) != metadata:
        errors.append("V14.1 open: manifeste local désynchronisé")
    expected_checksums = []
    for item in metadata["files"]:
        path = root / item["path"]
        if not path.is_file():
            errors.append(f"V14.1 open: fichier absent: {item['path']}")
            continue
        digest = sha256(path)
        if ("sha256" in item and digest != item["sha256"]) or (
            "byte_size" in item and path.stat().st_size != item["byte_size"]
        ):
            errors.append(f"V14.1 open: empreinte ou taille divergente: {item['path']}")
        expected_checksums.append(f"{digest}  {item['path']}\n")
    expected_checksums.append(f"{sha256(root / 'manifest.json')}  manifest.json\n")
    if (root / "checksums.sha256").read_text(encoding="utf-8") != "".join(expected_checksums):
        errors.append("V14.1 open: checksums.sha256 désynchronisé")

    import duckdb

    connection = duckdb.connect(str(root / "morocco_elections_v14_1.duckdb"), read_only=True)
    try:
        for table in metadata["tables"]:
            count = connection.execute(
                f'SELECT COUNT(*) FROM "{table["table_name"]}"'
            ).fetchone()[0]
            if count != table["rows"]:
                errors.append(f"V14.1 open: volume DuckDB divergent: {table['table_name']}")
        invalid_coverage = connection.execute(
            "SELECT COUNT(*) FROM analytical_parliamentary_coverage "
            "WHERE expected_file_count IS NULL AND coverage_status <> 'UNKNOWN'"
        ).fetchone()[0]
        if invalid_coverage:
            errors.append("V14.1 open: couverture affirmée sans dénominateur")
        invalid_exposure = connection.execute(
            "SELECT COUNT(*) FROM analytical_person_period_exposure "
            "WHERE observed_mandate_days <= 0 OR published_question_count < 0"
        ).fetchone()[0]
        if invalid_exposure:
            errors.append("V14.1 open: exposition invalide")
    finally:
        connection.close()
    return errors


def validate_v11a_artifacts() -> list[str]:
    errors: list[str] = []
    try:
        metadata = json.loads(V11A_METADATA_PATH.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"Métadonnées V11-A illisibles: {exc}"]
    if metadata.get("schema_version") != 1 or metadata.get("phase") != "V11-A":
        errors.append("Les métadonnées V11-A doivent utiliser schema_version=1 et phase=V11-A")
    if metadata.get("decision") not in {"GO", "NO_GO"}:
        errors.append("La décision V11-A doit être GO ou NO_GO")
    candidate = metadata.get("candidate", {})
    required_candidate = {
        "candidate_id",
        "expected_local_name",
        "dataset_url",
        "download_url",
        "final_url",
        "producer",
        "license",
        "version",
        "retrieved_on",
        "etag",
        "last_modified",
        "byte_size",
        "sha256",
        "tracked",
    }
    if not required_candidate <= set(candidate):
        errors.append(f"Métadonnées V11-A: champs candidat absents {sorted(required_candidate - set(candidate))}")
    if not re.fullmatch(r"[0-9a-f]{64}", str(candidate.get("sha256", ""))):
        errors.append("Métadonnées V11-A: SHA-256 candidat invalide")
    if candidate.get("tracked") is not False or candidate.get("expected_local_name") != "communes-elus-2015-1-0.xlsx":
        errors.append("Le candidat V11-A doit être déclaré non suivi sous son nom attendu")
    try:
        manifest = load_manifest()
        v10_artifact = next(item for item in manifest["artifacts"] if item["artifact_id"] == "WAREHOUSE_V10_EXPORT")
        baseline = metadata.get("baseline", {})
        if baseline.get("sha256_before") != v10_artifact["sha256"] or baseline.get("sha256_after") != v10_artifact["sha256"]:
            errors.append("Métadonnées V11-A: empreinte de baseline différente du V10 manifesté")
        if baseline.get("modified") is not False:
            errors.append("Métadonnées V11-A: V10 doit être déclaré inchangé")
    except (KeyError, StopIteration, ValidationError) as exc:
        errors.append(f"Métadonnées V11-A: baseline V10 invérifiable: {exc}")
    checks = metadata.get("checks")
    if not isinstance(checks, list) or not checks:
        errors.append("Métadonnées V11-A: liste de contrôles obligatoire")
        checks = []
    check_ids = {item.get("check_id") for item in checks if isinstance(item, dict)}
    required_checks = {
        "xlsx_readable",
        "sheets",
        "data_shape",
        "commune_universe",
        "required_values",
        "communal_seat_reconciliation",
        "national_seat_total",
        "v10_geographic_crosswalk",
        "party_codes",
        "duplicate_candidates_explained",
        "documented_domains",
        "license",
        "file_dictionary_notes_consistency",
    }
    if check_ids != required_checks:
        errors.append(f"Métadonnées V11-A: contrôles incorrects {sorted(check_ids)}")
    failed = [item for item in checks if item.get("required") and item.get("status") != "PASS"]
    expected_decision = "NO_GO" if failed else "GO"
    if metadata.get("decision") != expected_decision:
        errors.append("La décision V11-A ne correspond pas au résultat des contrôles obligatoires")
    for anomaly in metadata.get("anomalies", []):
        if not re.fullmatch(r"DUP2015_[0-9a-f]{20}", str(anomaly.get("anomaly_id", ""))):
            errors.append("Métadonnées V11-A: identifiant d'anomalie non pseudonymisé")
        if not anomaly.get("source_row_numbers") or anomaly.get("decision") != "no_deletion_no_ingestion":
            errors.append("Métadonnées V11-A: anomalie incomplète")
    try:
        decision = V11A_DECISION_PATH.read_text(encoding="utf-8")
    except (OSError, UnicodeError) as exc:
        errors.append(f"Rapport V11-A illisible: {exc}")
        return errors
    for token in (
        "VERSION : V11-A",
        "DATE DE GÉNÉRATION",
        "PÉRIMÈTRE",
        "RENVOIS",
        f"DÉCISION BINAIRE : {metadata.get('decision')}",
        "Aucune ligne n'a été ingérée",
        "COMMANDE DE REPRODUCTION",
    ):
        if token not in decision:
            errors.append(f"Rapport V11-A: élément obligatoire absent: {token}")
    return errors


def validate_smiig_artifacts() -> list[str]:
    errors: list[str] = []
    try:
        metadata = json.loads(SMIIG_METADATA_PATH.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"Métadonnées SMIIG illisibles: {exc}"]
    if metadata.get("schema_version") != 1 or metadata.get("phase") != "V11-SMIIG-QUALIFICATION":
        errors.append("Les métadonnées SMIIG doivent utiliser schema_version=1 et la phase attendue")
    if metadata.get("decision") not in {"GO", "NO_GO"} or metadata.get("status") != "PILOTE":
        errors.append("La qualification SMIIG doit porter une décision binaire et le statut PILOTE")
    candidate = metadata.get("candidate", {})
    required_candidate = {
        "candidate_id", "expected_local_name", "dataset_url", "resource_url", "final_url", "producer", "license",
        "version", "retrieved_on", "etag", "last_modified", "byte_size", "sha256", "tracked",
    }
    if not required_candidate <= set(candidate):
        errors.append(f"Métadonnées SMIIG: champs candidat absents {sorted(required_candidate - set(candidate))}")
    if not re.fullmatch(r"[0-9a-f]{64}", str(candidate.get("sha256", ""))):
        errors.append("Métadonnées SMIIG: SHA-256 candidat invalide")
    if candidate.get("tracked") is not False:
        errors.append("Le candidat SMIIG doit être déclaré non suivi")
    try:
        manifest = load_manifest()
        v10_artifact = next(item for item in manifest["artifacts"] if item["artifact_id"] == "WAREHOUSE_V10_EXPORT")
        baseline = metadata.get("baseline", {})
        if baseline.get("sha256_before") != v10_artifact["sha256"] or baseline.get("sha256_after") != v10_artifact["sha256"]:
            errors.append("Métadonnées SMIIG: empreinte de baseline différente du V10 manifesté")
        if baseline.get("modified") is not False:
            errors.append("Métadonnées SMIIG: V10 doit être déclaré inchangé")
    except (KeyError, StopIteration, ValidationError) as exc:
        errors.append(f"Métadonnées SMIIG: baseline V10 invérifiable: {exc}")
    checks = metadata.get("checks")
    if not isinstance(checks, list):
        errors.append("Métadonnées SMIIG: liste de contrôles obligatoire")
        checks = []
    expected_checks = {
        "xlsx_and_sheets", "data_shape", "required_values", "temporal_coverage", "grain_uniqueness",
        "v10_universe_subset", "v10_geographic_crosswalk", "dictionary_consistency", "license", "indicator_domains",
        "component_formula", "normalized_score_formula", "arrondissement_replication_lineage", "identity_and_role_scope",
    }
    actual_checks = {item.get("check_id") for item in checks if isinstance(item, dict)}
    if actual_checks != expected_checks:
        errors.append(f"Métadonnées SMIIG: contrôles incorrects {sorted(actual_checks)}")
    failed = [item for item in checks if item.get("required") and item.get("status") != "PASS"]
    if metadata.get("decision") != ("NO_GO" if failed else "GO"):
        errors.append("La décision SMIIG ne correspond pas aux contrôles obligatoires")
    for anomaly in metadata.get("anomalies", []):
        anomaly_id = str(anomaly.get("anomaly_id", ""))
        if not re.fullmatch(r"SMIIG_(?:SCHEMA_)?[0-9a-f]{20}", anomaly_id):
            errors.append("Métadonnées SMIIG: identifiant d'anomalie non pseudonymisé")
        if anomaly.get("status") != "unresolved":
            errors.append("Métadonnées SMIIG: anomalie sans statut unresolved")
    try:
        report = SMIIG_DECISION_PATH.read_text(encoding="utf-8")
    except (OSError, UnicodeError) as exc:
        errors.append(f"Rapport SMIIG illisible: {exc}")
        return errors
    for token in (
        "VERSION : V11-SMIIG-QUALIFICATION", "DATE DE GÉNÉRATION", "PÉRIMÈTRE", "RENVOIS",
        f"DÉCISION BINAIRE : {metadata.get('decision')}", "STATUT DE COUVERTURE : PILOTE",
        "Aucune ligne SMIIG n'a été ingérée", "COMMANDE DE REPRODUCTION",
    ):
        if token not in report:
            errors.append(f"Rapport SMIIG: élément obligatoire absent: {token}")
    return errors


def validate_python_sources() -> list[str]:
    errors: list[str] = []
    for path in sorted(ROOT.rglob("*.py")):
        if any(part in {".venv", "venv", "__pycache__"} for part in path.parts):
            continue
        try:
            ast.parse(path.read_text(encoding="utf-8-sig"), filename=str(path))
        except (OSError, SyntaxError, UnicodeError) as exc:
            errors.append(f"Python invalide: {path.relative_to(ROOT)}: {exc}")
    return errors


def validate_documentation(release: str = "v13") -> list[str]:
    errors: list[str] = []
    releases = DOCUMENTATION_DIRS if release == "all" else {release: DOCUMENTATION_DIRS[release]}
    for version, directory in releases.items():
        actual = sorted(path.name for path in directory.glob("*.txt"))
        if actual != EXPECTED_DOCUMENTS:
            errors.append(f"Corpus documentaire {version} incorrect: attendu {EXPECTED_DOCUMENTS}, obtenu {actual}")
        for name in EXPECTED_DOCUMENTS:
            path = directory / name
            if not path.is_file():
                continue
            try:
                content = path.read_text(encoding="utf-8-sig")
            except UnicodeError as exc:
                errors.append(f"Document non UTF-8: {version}/{name}: {exc}")
                continue
            if not content.strip():
                errors.append(f"Document vide: {version}/{name}")
            for token in REQUIRED_DOCUMENT_METADATA:
                if token not in content:
                    errors.append(f"Métadonnée absente de {version}/{name}: {token}")
    return errors


def tracked_files() -> list[Path]:
    result = subprocess.run(
        ["git", "ls-files", "-z"],
        cwd=ROOT,
        capture_output=True,
        check=False,
    )
    if result.returncode != 0:
        return []
    return [ROOT / item.decode("utf-8") for item in result.stdout.split(b"\0") if item]


def candidate_repository_files() -> list[Path]:
    tracked = tracked_files()
    if tracked:
        return tracked
    candidates: list[Path] = []
    for path in ROOT.rglob("*"):
        if not path.is_file() or ".git" in path.parts:
            continue
        relative = path.relative_to(ROOT)
        if any(part in FORBIDDEN_DIRECTORIES or part in {".venv", "venv", "__pycache__"} for part in relative.parts):
            continue
        if tuple(relative.parts[:2]) in FORBIDDEN_DATA_PREFIXES:
            continue
        if path.suffix.lower() in FORBIDDEN_SUFFIXES:
            continue
        candidates.append(path)
    return candidates


def validate_repository_files(files: Iterable[Path] | None = None) -> list[str]:
    errors: list[str] = []
    selected = list(files) if files is not None else candidate_repository_files()
    secret_patterns = {
        "clé privée": re.compile("BEGIN" + r" [A-Z ]*PRIVATE KEY"),
        "jeton GitHub classique": re.compile("gh" + r"p_[A-Za-z0-9]{20,}"),
        "jeton GitHub fin": re.compile("github" + r"_pat_[A-Za-z0-9_]{20,}"),
        "clé AWS": re.compile("AK" + r"IA[0-9A-Z]{16}"),
    }
    for path in selected:
        try:
            relative = path.resolve().relative_to(ROOT.resolve())
        except ValueError:
            errors.append(f"Fichier hors dépôt: {path}")
            continue
        if any(part in FORBIDDEN_DIRECTORIES for part in relative.parts):
            errors.append(f"Répertoire de données interdit dans Git: {relative.as_posix()}")
        if tuple(relative.parts[:2]) in FORBIDDEN_DATA_PREFIXES:
            errors.append(f"Zone de données interdite dans Git: {relative.as_posix()}")
        if path.suffix.lower() in FORBIDDEN_SUFFIXES:
            errors.append(f"Type de fichier interdit dans Git: {relative.as_posix()}")
        if path.exists() and path.stat().st_size > MAX_TRACKED_BYTES:
            errors.append(f"Fichier suivi supérieur à 5 MiB: {relative.as_posix()}")
        if not path.exists() or path.suffix.lower() in FORBIDDEN_SUFFIXES:
            continue
        try:
            content = path.read_text(encoding="utf-8-sig")
        except (UnicodeError, OSError):
            continue
        for label, pattern in secret_patterns.items():
            if pattern.search(content):
                errors.append(f"Secret potentiel ({label}) dans {relative.as_posix()}")
    return errors


def validate_physical_record(record: dict, data_dir: str | Path | None = None) -> list[str]:
    errors: list[str] = []
    path = resolve_manifest_path(record["local_path"], data_dir)
    label = record.get("source_id") or record.get("artifact_id") or record.get("evidence_id") or record.get("file_id")
    if not path.is_file():
        return [f"{label}: fichier local absent: {record['local_path']}"]
    if path.stat().st_size != record["byte_size"]:
        errors.append(f"{label}: taille attendue {record['byte_size']}, obtenue {path.stat().st_size}")
    actual_hash = sha256(path)
    if actual_hash != record["sha256"]:
        errors.append(f"{label}: SHA-256 attendu {record['sha256']}, obtenu {actual_hash}")
    return errors


def _populated_rows(sheet) -> int:
    return sum(1 for row in sheet.iter_rows(min_row=5, values_only=True) if any(value is not None for value in row))


def _recorded_documentation_date(directory: Path) -> str:
    index = directory / "00_INDEX_ET_MODE_EMPLOI.txt"
    content = index.read_text(encoding="utf-8-sig")
    match = re.search(r"DATE DE GÉNÉRATION\s*:\s*(\d{4}-\d{2}-\d{2})", content)
    if not match:
        raise ValidationError(f"Date de génération documentaire introuvable: {index}")
    return match.group(1)


def compare_v10_to_v9(data_dir: str | Path | None = None) -> list[str]:
    errors: list[str] = []
    paths = get_paths(data_dir)
    if not paths.v9_workbook.is_file() or not paths.v10_workbook.is_file():
        return ["Comparaison V9/V10 impossible: un classeur est absent"]
    old = openpyxl.load_workbook(paths.v9_workbook, read_only=True, data_only=False)
    new = openpyxl.load_workbook(paths.v10_workbook, read_only=True, data_only=False)
    allowed = {"README", "SOURCES", "DIM_PERSON", "CROSSWALK_GEO", "LOCAL_MANDATES", "LOCAL_COUNCIL_CONTROL", "QUALITY_CONTROL", "DATA_COVERAGE", "WORKBOOK_AUDIT_V9", "WORKBOOK_AUDIT_V10", "IDENTITY_AUDIT_V10", "MANUAL_RESOLUTIONS_V10", "COUNCIL_SEAT_STATUS_V10"}
    for name in sorted((set(old.sheetnames) & set(new.sheetnames)) - allowed):
        left = list(old[name].iter_rows(min_row=4, values_only=True))
        right = list(new[name].iter_rows(min_row=4, values_only=True))
        if left != right:
            errors.append(f"Différence V9/V10 non autorisée dans {name}")
    expected_new = {"IDENTITY_AUDIT_V10", "MANUAL_RESOLUTIONS_V10", "COUNCIL_SEAT_STATUS_V10", "WORKBOOK_AUDIT_V10"}
    actual_new = set(new.sheetnames) - set(old.sheetnames)
    if actual_new != expected_new:
        errors.append(f"Onglets V10 nouveaux inattendus: {sorted(actual_new)}")
    old.close()
    new.close()
    return errors


def compare_v11_to_v10(data_dir: str | Path | None = None) -> list[str]:
    from morocco_elections.releases.v11.build import ALLOWED_CHANGED_SHEETS, compare_v11_to_v10 as compare

    paths = get_paths(data_dir)
    if not paths.v10_workbook.is_file() or not paths.v11_workbook.is_file():
        return ["Comparaison V10/V11 impossible: un classeur est absent"]
    try:
        changed, unexpected = compare(paths.v10_workbook, paths.v11_workbook)
    except RuntimeError as exc:
        return [str(exc)]
    errors = []
    if unexpected:
        errors.append(f"Différences V10/V11 non autorisées: {unexpected}")
    if set(changed) != ALLOWED_CHANGED_SHEETS:
        errors.append(f"Liste réelle des changements V11 incorrecte: {changed}")
    return errors


def compare_v12_to_v11(data_dir: str | Path | None = None) -> list[str]:
    from morocco_elections.releases.v12.build import ALLOWED_CHANGED_SHEETS, compare_v12_to_v11 as compare

    paths = get_paths(data_dir)
    if not paths.v11_workbook.is_file() or not paths.v12_workbook.is_file():
        return ["Comparaison V11/V12 impossible: un classeur est absent"]
    try:
        changed, unexpected = compare(paths.v11_workbook, paths.v12_workbook)
    except RuntimeError as exc:
        return [str(exc)]
    errors = []
    if unexpected:
        errors.append(f"Différences V11/V12 non autorisées: {unexpected}")
    if set(changed) != ALLOWED_CHANGED_SHEETS:
        errors.append(f"Liste réelle des changements V12 incorrecte: {changed}")
    return errors


def compare_v13_to_v12(data_dir: str | Path | None = None) -> list[str]:
    from morocco_elections.releases.v13.build import ALLOWED_CHANGED_SHEETS, compare_v13_to_v12 as compare

    paths = get_paths(data_dir)
    if not paths.v12_workbook.is_file() or not paths.v13_workbook.is_file():
        return ["Comparaison V12/V13 impossible: un classeur est absent"]
    try:
        changed, unexpected = compare(paths.v12_workbook, paths.v13_workbook)
    except RuntimeError as exc:
        return [str(exc)]
    errors = []
    if unexpected:
        errors.append(f"Différences V12/V13 non autorisées: {unexpected}")
    if set(changed) != ALLOWED_CHANGED_SHEETS:
        errors.append(f"Liste réelle des changements V13 incorrecte: {changed}")
    return errors


def validate_full(manifest: dict, data_dir: str | Path | None = None, release: str = "v13", baseline: str | None = None) -> list[str]:
    errors: list[str] = []
    for record in [*manifest["physical_files"], *manifest["evidence"]]:
        errors.extend(validate_physical_record(record, data_dir))
    for source in manifest["sources"]:
        errors.extend(validate_physical_record(source, data_dir))
        path = resolve_manifest_path(source["local_path"], data_dir)
        if path.suffix.lower() == ".xlsx" and path.is_file():
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            sheet = workbook.worksheets[source.get("sheet_index", 0)]
            actual_rows, actual_columns = sheet.max_row - 1, sheet.max_column
            workbook.close()
            if actual_rows != source["rows"] or actual_columns != source["columns"]:
                errors.append(f"{source['source_id']}: dimensions attendues {source['rows']}×{source['columns']}, obtenues {actual_rows}×{actual_columns}")

    selected_artifacts = {
        "v9": {"WAREHOUSE_V8_INPUT", "WAREHOUSE_V9_EXPORT"},
        "v10": {"WAREHOUSE_V9_EXPORT", "WAREHOUSE_V10_EXPORT"},
        "v11": {"WAREHOUSE_V10_EXPORT", "WAREHOUSE_V11_EXPORT"},
        "v12": {"WAREHOUSE_V8_INPUT", "WAREHOUSE_V11_EXPORT", "WAREHOUSE_V12_EXPORT"},
        "v13": {"WAREHOUSE_V8_INPUT", "WAREHOUSE_V12_EXPORT", "WAREHOUSE_V13_EXPORT"},
        "v14": {"WAREHOUSE_V8_INPUT", "WAREHOUSE_V12_EXPORT", "WAREHOUSE_V13_EXPORT"},
        "all": {item["artifact_id"] for item in manifest["artifacts"]},
    }[release]
    for artifact in manifest["artifacts"]:
        errors.extend(validate_physical_record(artifact, data_dir))
        if artifact["artifact_id"] not in selected_artifacts:
            continue
        path = resolve_manifest_path(artifact["local_path"], data_dir)
        if not path.is_file():
            continue
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        if len(workbook.sheetnames) != artifact["sheet_count"]:
            errors.append(f"{artifact['artifact_id']}: {len(workbook.sheetnames)} onglets au lieu de {artifact['sheet_count']}")
        volumes = (
            EXPECTED_V9_VOLUMES
            if artifact["artifact_id"] == "WAREHOUSE_V9_EXPORT"
            else EXPECTED_V10_VOLUMES
            if artifact["artifact_id"] == "WAREHOUSE_V10_EXPORT"
            else EXPECTED_V11_VOLUMES
            if artifact["artifact_id"] == "WAREHOUSE_V11_EXPORT"
            else EXPECTED_V12_VOLUMES
            if artifact["artifact_id"] == "WAREHOUSE_V12_EXPORT"
            else EXPECTED_V13_VOLUMES
            if artifact["artifact_id"] == "WAREHOUSE_V13_EXPORT"
            else {}
        )
        for sheet_name, expected in volumes.items():
            actual = _populated_rows(workbook[sheet_name])
            if actual != expected:
                errors.append(f"{artifact['artifact_id']}.{sheet_name}: {actual} lignes au lieu de {expected}")
        workbook.close()

    paths = get_paths(data_dir)
    if release == "all":
        v11a_candidate = paths.data_root / "staging" / "v11a" / "source_candidates" / "communes-elus-2015-1-0.xlsx"
        if not v11a_candidate.is_file():
            errors.append(f"V11-A: candidat local absent: {v11a_candidate}")
        else:
            v11a_metadata = json.loads(V11A_METADATA_PATH.read_text(encoding="utf-8"))
            expected_size = v11a_metadata["candidate"]["byte_size"]
            expected_hash = v11a_metadata["candidate"]["sha256"]
            if v11a_candidate.stat().st_size != expected_size:
                errors.append(f"V11-A: taille candidat attendue {expected_size}, obtenue {v11a_candidate.stat().st_size}")
            actual_hash = sha256(v11a_candidate)
            if actual_hash != expected_hash:
                errors.append(f"V11-A: SHA-256 candidat attendu {expected_hash}, obtenu {actual_hash}")
        smiig_candidate = paths.data_root / "staging" / "v11_smiig" / "source_candidates" / "2024-01-08-dataset-smiig-v2023-communes.xlsx"
        if not smiig_candidate.is_file():
            errors.append(f"SMIIG: candidat local absent: {smiig_candidate}")
        else:
            smiig_metadata = json.loads(SMIIG_METADATA_PATH.read_text(encoding="utf-8"))
            expected_size = smiig_metadata["candidate"]["byte_size"]
            expected_hash = smiig_metadata["candidate"]["sha256"]
            if smiig_candidate.stat().st_size != expected_size:
                errors.append(f"SMIIG: taille candidat attendue {expected_size}, obtenue {smiig_candidate.stat().st_size}")
            actual_hash = sha256(smiig_candidate)
            if actual_hash != expected_hash:
                errors.append(f"SMIIG: SHA-256 candidat attendu {expected_hash}, obtenu {actual_hash}")
    if release in {"v9", "all"} and paths.v9_workbook.is_file():
        from morocco_elections.legacy.v9 import documentation as module
        module.configure_paths(data_dir)
        module.GENERATED_ON = _recorded_documentation_date(DOCUMENTATION_DIRS["v9"])
        before_hash = sha256(paths.v9_workbook)
        model = module.load_model()
        generated, metric_rules = module.build_docs(model)
        try:
            module.validate(model, generated, metric_rules, before_hash)
        except RuntimeError as exc:
            errors.append(str(exc))
        for name, content in generated.items():
            expected = content.replace("\r\n", "\n").rstrip() + "\n"
            actual = (DOCUMENTATION_DIRS["v9"] / name).read_text(encoding="utf-8-sig").replace("\r\n", "\n")
            if actual != expected:
                errors.append(f"Documentation désynchronisée du V9: {name}")
        if sha256(paths.v9_workbook) != before_hash:
            errors.append("Le V9 a été modifié pendant la validation")

    if release in {"v10", "all"} and paths.v10_workbook.is_file():
        from morocco_elections.releases.v10 import documentation as module10
        module10.configure_paths(data_dir)
        module10.base.GENERATED_ON = _recorded_documentation_date(DOCUMENTATION_DIRS["v10"])
        before_hash = sha256(paths.v10_workbook)
        model = module10.base.load_model()
        generated, _ = module10.base.build_docs(model)
        generated = module10._adapt(generated, model)
        try:
            module10._validate(model, generated, before_hash)
        except RuntimeError as exc:
            errors.append(str(exc))
        for name, content in generated.items():
            expected = content.replace("\r\n", "\n").rstrip() + "\n"
            actual = (DOCUMENTATION_DIRS["v10"] / name).read_text(encoding="utf-8-sig").replace("\r\n", "\n")
            if actual != expected:
                errors.append(f"Documentation désynchronisée du V10: {name}")
        if sha256(paths.v10_workbook) != before_hash:
            errors.append("Le V10 a été modifié pendant la validation")
        errors.extend(compare_v10_to_v9(data_dir))
    if release in {"v11", "all"} and paths.v11_workbook.is_file():
        from morocco_elections.releases.v11 import documentation as module11

        module11.configure_paths(data_dir)
        module11.base.GENERATED_ON = _recorded_documentation_date(DOCUMENTATION_DIRS["v11"])
        before_hash = sha256(paths.v11_workbook)
        model = module11.base.load_model()
        generated, _ = module11.base.build_docs(model)
        generated = module11._adapt(generated, model)
        try:
            module11._validate(model, generated, before_hash)
        except RuntimeError as exc:
            errors.append(str(exc))
        for name, content in generated.items():
            expected = content.replace("\r\n", "\n").rstrip() + "\n"
            actual = (DOCUMENTATION_DIRS["v11"] / name).read_text(encoding="utf-8-sig").replace("\r\n", "\n")
            if actual != expected:
                errors.append(f"Documentation désynchronisée du V11: {name}")
        if sha256(paths.v11_workbook) != before_hash:
            errors.append("Le V11 a été modifié pendant la validation")
        errors.extend(compare_v11_to_v10(data_dir))
    if release in {"v12", "all"} and paths.v12_workbook.is_file():
        from morocco_elections.releases.v12 import documentation as module12

        module12.configure_paths(data_dir)
        module12.base.GENERATED_ON = _recorded_documentation_date(DOCUMENTATION_DIRS["v12"])
        before_hash = sha256(paths.v12_workbook)
        model = module12.base.load_model()
        generated, _ = module12.base.build_docs(model)
        generated = module12._adapt(generated, model)
        try:
            module12._validate(model, generated, before_hash)
        except RuntimeError as exc:
            errors.append(str(exc))
        for name, content in generated.items():
            expected = content.replace("\r\n", "\n").rstrip() + "\n"
            actual = (DOCUMENTATION_DIRS["v12"] / name).read_text(encoding="utf-8-sig").replace("\r\n", "\n")
            if actual != expected:
                errors.append(f"Documentation désynchronisée du V12: {name}")
        if sha256(paths.v12_workbook) != before_hash:
            errors.append("Le V12 a été modifié pendant la validation")
        errors.extend(compare_v12_to_v11(data_dir))
    if release in {"v13", "all"} and paths.v13_workbook.is_file():
        from morocco_elections.releases.v13 import documentation as module13

        module13.configure_paths(data_dir)
        module13.base.GENERATED_ON = _recorded_documentation_date(DOCUMENTATION_DIRS["v13"])
        before_hash = sha256(paths.v13_workbook)
        model = module13.base.load_model()
        generated, _ = module13.base.build_docs(model)
        generated = module13._adapt(generated, model)
        try:
            module13._validate(model, generated, before_hash)
        except RuntimeError as exc:
            errors.append(str(exc))
        for name, content in generated.items():
            expected = content.replace("\r\n", "\n").rstrip() + "\n"
            actual = (DOCUMENTATION_DIRS["v13"] / name).read_text(encoding="utf-8-sig").replace("\r\n", "\n")
            if actual != expected:
                errors.append(f"Documentation désynchronisée du V13: {name}")
        if sha256(paths.v13_workbook) != before_hash:
            errors.append("Le V13 a été modifié pendant la validation")
        errors.extend(compare_v13_to_v12(data_dir))
    return errors


def validate_quality_baseline_artifacts(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.quality import baseline as quality_baseline

    errors: list[str] = []
    if not QUALITY_BASELINE_PATH.is_file() or not QUALITY_BASELINE_REPORT_PATH.is_file():
        return ["V10-QA: artefacts JSON/TXT absents"]
    try:
        data = json.loads(QUALITY_BASELINE_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"V10-QA: JSON illisible: {exc}"]
    if data.get("schema_version") != 1 or data.get("phase") != "V10-QA" or data.get("baseline_release") != "V10":
        errors.append("V10-QA: métadonnées racine invalides")
    release_report = json.loads(V10_REPORT_PATH.read_text(encoding="utf-8"))
    if data.get("source_workbook_sha256") != release_report["workbooks"]["V10"]:
        errors.append("V10-QA: empreinte V10 différente du rapport de release")
    errors.extend(f"V10-QA: {error}" for error in quality_baseline.validate_baseline(data))
    if len(data.get("table_status", [])) != 67:
        errors.append("V10-QA: le scorecard doit inventorier les 67 onglets V10")
    required_blockers = {
        "QA_V10_INSCRITS_2015",
        "QA_V10_INSCRITS_2021",
        "QA_V10_COUNCIL_PRESIDENT",
        "QA_V10_MPS_PARTY_HISTORY",
        "V11A_ROLE_DOMAIN_UNDOCUMENTED",
        "BASELINE_HCP_COMMUNAL_INDICATORS",
    }
    actual_blockers = {item.get("anomaly_id") for item in data.get("anomalies", []) if item.get("status") == "BLOQUÉ"}
    if not required_blockers <= actual_blockers:
        errors.append(f"V10-QA: blockers obligatoires absents: {sorted(required_blockers - actual_blockers)}")
    required_analyses = {
        "ANA_OBSERVED_PARTY_RESULTS",
        "ANA_ELECTORAL_TRANSITIONS",
        "ANA_REPORTED_TURNOUT",
        "ANA_COUNCIL_2021",
        "ANA_PRESIDENCIES",
        "ANA_WINNER_VS_PRESIDENT",
        "ANA_REGISTERED_VOTERS",
        "ANA_COUNCILS_2015",
        "ANA_SMIIG",
        "ANA_RGPH_2015_CONTEMPORARY",
    }
    actual_analyses = {item.get("analysis_id") for item in data.get("analysis_permissions", [])}
    if not required_analyses <= actual_analyses:
        errors.append(f"V10-QA: analyses obligatoires absentes: {sorted(required_analyses - actual_analyses)}")
    expected_report = quality_baseline.render_report(data)
    actual_report = QUALITY_BASELINE_REPORT_PATH.read_text(encoding="utf-8-sig")
    if actual_report != expected_report:
        errors.append("V10-QA: rapport TXT désynchronisé du JSON")
    if full and not errors:
        workbook = get_paths(data_dir).v10_workbook
        if not workbook.is_file():
            errors.append(f"V10-QA: classeur local absent: {workbook}")
        else:
            rebuilt = quality_baseline.build_baseline(workbook, str(data["as_of"]))
            if rebuilt != data:
                errors.append("V10-QA: JSON désynchronisé des sources locales et du V10")
    return errors


def validate_electoral_denominator_artifacts(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.research import electoral_denominators

    errors: list[str] = []
    try:
        metadata = json.loads(DENOMINATORS_METADATA_PATH.read_text(encoding="utf-8"))
        report = DENOMINATORS_REPORT_PATH.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"V10-QA-1: artefacts illisibles: {exc}"]
    errors.extend(f"V10-QA-1: {error}" for error in electoral_denominators.validate_metadata(metadata))
    try:
        manifest = load_manifest()
        v10_artifact = next(item for item in manifest["artifacts"] if item["artifact_id"] == "WAREHOUSE_V10_EXPORT")
        baseline = metadata.get("baseline", {})
        if baseline.get("sha256_before") != v10_artifact["sha256"] or baseline.get("sha256_after") != v10_artifact["sha256"]:
            errors.append("V10-QA-1: empreinte de baseline différente du V10 manifesté")
    except (KeyError, StopIteration, ValidationError) as exc:
        errors.append(f"V10-QA-1: baseline invérifiable: {exc}")
    source_ids = {item.get("source_id") for item in metadata.get("research_sources", [])}
    required_sources = {
        "OFFICIAL_ELECTIONS_MA_RESULTS",
        "OFFICIAL_ELECTORAL_LISTS_STATS",
        "WAYBACK_ELECTIONS_MA_2015",
        "CNDH_2015_CONTROL",
        "CNDH_2021_CONTROL",
        "TAFRA_COMMUNAL_RESULTS_2015_2021",
    }
    if source_ids != required_sources:
        errors.append("V10-QA-1: registre de recherche incomplet")
    if report != electoral_denominators.render_report(metadata):
        errors.append("V10-QA-1: rapport TXT désynchronisé du JSON")
    if full and not errors:
        paths = get_paths(data_dir)
        candidates: dict[int, Path | None] = {}
        for year in electoral_denominators.YEARS:
            record = metadata["candidates"][str(year)]
            if record is None:
                candidates[year] = None
                continue
            candidate = paths.data_root / "staging" / "v10qa1" / "source_candidates" / record["local_name"]
            if not candidate.is_file():
                errors.append(f"V10-QA-1: candidat local absent: {candidate}")
                continue
            if candidate.stat().st_size != record["byte_size"] or electoral_denominators.sha256_file(candidate) != record["sha256"]:
                errors.append(f"V10-QA-1: candidat {year} différent de son empreinte")
            candidates[year] = candidate
        if not errors:
            rebuilt = electoral_denominators.build_metadata(
                workbook_path=paths.v10_workbook,
                candidate_paths=candidates,
                as_of=str(metadata["generated_on"]),
            )
            if rebuilt != metadata:
                errors.append("V10-QA-1: décision désynchronisée des sources locales et de V10")
    return errors


def validate_local_presidency_artifacts(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.research import local_presidencies

    errors: list[str] = []
    try:
        metadata = json.loads(PRESIDENCIES_METADATA_PATH.read_text(encoding="utf-8"))
        report = PRESIDENCIES_REPORT_PATH.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"V10-QA-2: artefacts illisibles: {exc}"]
    errors.extend(f"V10-QA-2: {error}" for error in local_presidencies.validate_metadata(metadata))
    try:
        manifest = load_manifest()
        v10_artifact = next(item for item in manifest["artifacts"] if item["artifact_id"] == "WAREHOUSE_V10_EXPORT")
        baseline = metadata.get("baseline", {})
        if baseline.get("sha256_before") != v10_artifact["sha256"] or baseline.get("sha256_after") != v10_artifact["sha256"]:
            errors.append("V10-QA-2: empreinte de baseline différente du V10 manifesté")
    except (KeyError, StopIteration, ValidationError) as exc:
        errors.append(f"V10-QA-2: baseline invérifiable: {exc}")
    source_ids = {item.get("source_id") for item in metadata.get("research_sources", [])}
    if source_ids != {"DGCT_OPEN_DATA", "TAFRA_COUNCILS_2021", "SECONDARY_INSTALLATION_REPORTS"}:
        errors.append("V10-QA-2: registre de recherche incomplet")
    if report != local_presidencies.render_report(metadata):
        errors.append("V10-QA-2: rapport TXT désynchronisé du JSON")
    if full and not errors:
        paths = get_paths(data_dir)
        index_record = metadata.get("evidence_index", {})
        index_path: Path | None = None
        if index_record.get("provided"):
            local_path = str(index_record.get("local_path", ""))
            index_path = resolve_manifest_path(local_path, data_dir)
            if not index_path.is_file():
                errors.append(f"V10-QA-2: index local absent: {index_path}")
            elif local_presidencies.sha256_file(index_path) != index_record.get("sha256"):
                errors.append("V10-QA-2: index local différent de son empreinte")
        if not errors:
            rebuilt = local_presidencies.build_metadata(paths.v10_workbook, index_path, str(metadata["generated_on"]))
            if rebuilt != metadata:
                errors.append("V10-QA-2: décision désynchronisée des preuves locales et de V10")
    return errors


def validate_hcp_indicator_artifacts(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.research import hcp_indicators

    errors: list[str] = []
    try:
        metadata = json.loads(HCP_INDICATORS_METADATA_PATH.read_text(encoding="utf-8"))
        report = HCP_INDICATORS_REPORT_PATH.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"V10-QA-3: artefacts illisibles: {exc}"]
    errors.extend(f"V10-QA-3: {error}" for error in hcp_indicators.validate_metadata(metadata))
    try:
        manifest = load_manifest()
        v10_artifact = next(item for item in manifest["artifacts"] if item["artifact_id"] == "WAREHOUSE_V10_EXPORT")
        baseline = metadata.get("baseline", {})
        if baseline.get("sha256_before") != v10_artifact["sha256"] or baseline.get("sha256_after") != v10_artifact["sha256"]:
            errors.append("V10-QA-3: empreinte de baseline différente du V10 manifesté")
    except (KeyError, StopIteration, ValidationError) as exc:
        errors.append(f"V10-QA-3: baseline invérifiable: {exc}")
    source_ids = {item.get("source_id") for item in metadata.get("sources", [])}
    if source_ids != set(hcp_indicators.SOURCE_SPECS):
        errors.append("V10-QA-3: registre des trois sources HCP incomplet")
    if report != hcp_indicators.render_report(metadata):
        errors.append("V10-QA-3: rapport TXT désynchronisé du JSON")
    if full and not errors:
        paths = get_paths(data_dir)
        candidates: dict[str, Path | None] = {}
        for record in metadata["sources"]:
            source_id = str(record["source_id"])
            if record.get("status") == "not_provided":
                candidates[source_id] = None
                continue
            candidate = paths.data_root / "staging" / "v10qa3" / "source_candidates" / str(record["local_name"])
            promoted = {
                "RGPH2014_INDIVIDUALS": paths.hcp_individuals_2014,
                "RGPH2024_INDICATORS": paths.hcp_indicators_2024,
            }.get(source_id)
            if not candidate.is_file() and promoted is not None:
                candidate = promoted
            if not candidate.is_file():
                errors.append(f"V10-QA-3: candidat local absent: {candidate}")
                continue
            if candidate.stat().st_size != record.get("byte_size") or hcp_indicators.sha256_file(candidate) != record.get("sha256"):
                errors.append(f"V10-QA-3: candidat {source_id} différent de son empreinte")
            candidates[source_id] = candidate
        if not errors:
            rebuilt = hcp_indicators.build_metadata(paths.v10_workbook, candidates, str(metadata["generated_on"]))
            if rebuilt != metadata:
                errors.append("V10-QA-3: décision désynchronisée des sources locales et de V10")
    return errors


def validate_v11_quality_baseline_artifacts(data_dir: str | Path | None = None, full: bool = False) -> list[str]:
    from morocco_elections.quality import baseline as base
    from morocco_elections.quality import baseline_v11

    errors: list[str] = []
    try:
        metadata = json.loads(V11_QUALITY_BASELINE_PATH.read_text(encoding="utf-8"))
        report = V11_QUALITY_BASELINE_REPORT_PATH.read_text(encoding="utf-8-sig")
        release = json.loads(V11_REPORT_PATH.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return [f"V11-QA: artefacts illisibles: {exc}"]
    if metadata.get("phase") != "V11-QA" or metadata.get("baseline_release") != "V11":
        errors.append("V11-QA: métadonnées racine invalides")
    if metadata.get("source_workbook_sha256") != release.get("workbooks", {}).get("V11"):
        errors.append("V11-QA: empreinte différente du rapport de release")
    errors.extend(f"V11-QA: {error}" for error in base.validate_baseline(metadata))
    statuses = {item.get("anomaly_id"): item.get("status") for item in metadata.get("anomalies", [])}
    if statuses.get("BASELINE_HCP_COMMUNAL_INDICATORS") != "PARTIEL":
        errors.append("V11-QA: le chantier HCP doit être PARTIEL")
    permissions = {item.get("analysis_id"): item.get("status") for item in metadata.get("analysis_permissions", [])}
    expected_permissions = {
        "ANA_HCP_2014_FOR_ELECTION_2015": "CONDITIONNELLE",
        "ANA_HCP_2024_FOR_ELECTION_2015": "INTERDITE",
        "ANA_HCP_FOR_ELECTION_2021": "CONDITIONNELLE",
        "ANA_HCP_OTHER_INDICATORS": "INTERDITE",
    }
    if not all(permissions.get(key) == value for key, value in expected_permissions.items()):
        errors.append("V11-QA: permissions temporelles HCP incorrectes")
    if report != baseline_v11.render_report(metadata):
        errors.append("V11-QA: rapport TXT désynchronisé du JSON")
    if full and not errors:
        workbook = get_paths(data_dir).v11_workbook
        if not workbook.is_file():
            errors.append(f"V11-QA: classeur local absent: {workbook}")
        else:
            rebuilt = baseline_v11.build_baseline(workbook, str(metadata["as_of"]))
            if rebuilt != metadata:
                errors.append("V11-QA: JSON désynchronisé du V11 local")
    return errors


def run(mode: str, data_dir: str | Path | None = None, release: str = "v14", baseline: str | None = None) -> list[str]:
    manifest = load_manifest()
    errors = []
    errors.extend(validate_manifest(manifest))
    if release in {"v13", "v14", "v14.1", "all"}:
        errors.extend(validate_acquisition_catalog())
        errors.extend(validate_acquisition_inventory())
        errors.extend(validate_ontology_contract())
        errors.extend(validate_v13_source_profile())
        errors.extend(validate_v13_identity_registry())
        errors.extend(validate_v13_electoral_qualification())
        errors.extend(validate_v13_release_report(manifest))
        errors.extend(validate_v13_open_distribution())
    if release in {"v14", "v14.1", "all"}:
        errors.extend(validate_v14_open_distribution())
    if release in {"v14.1", "all"}:
        errors.extend(validate_v14_1_open_distribution())
    if release in {"v12", "v13", "v14", "v14.1", "all"}:
        errors.extend(validate_v12_qualification(manifest))
        errors.extend(validate_v12_release_report(manifest))
    if release in {"v10", "v11", "v12", "all"}:
        errors.extend(validate_v10_release_report(manifest))
    if release in {"v11", "v12", "all"}:
        errors.extend(validate_v11_release_report(manifest))
    if release in {"v11", "all"}:
        errors.extend(validate_v11_quality_baseline_artifacts())
    if release == "all":
        errors.extend(validate_backlog())
        errors.extend(validate_v11a_artifacts())
        errors.extend(validate_smiig_artifacts())
        errors.extend(validate_quality_baseline_artifacts())
        errors.extend(validate_electoral_denominator_artifacts())
        errors.extend(validate_local_presidency_artifacts())
        errors.extend(validate_hcp_indicator_artifacts())
    errors.extend(validate_python_sources())
    if release not in {"v14", "v14.1"}:
        errors.extend(validate_documentation(release))
    errors.extend(validate_repository_files())
    if mode == "full" and not errors:
        if release == "v14.1" and not get_paths(data_dir).v13_workbook.is_file():
            errors.append(f"V14.1 open: fichier local absent: {get_paths(data_dir).v13_workbook}")
        if release != "v14.1":
            errors.extend(validate_full(manifest, data_dir, release, baseline))
        if not errors and release == "all":
            errors.extend(validate_quality_baseline_artifacts(data_dir, full=True))
        if not errors and release == "all":
            errors.extend(validate_electoral_denominator_artifacts(data_dir, full=True))
        if not errors and release == "all":
            errors.extend(validate_local_presidency_artifacts(data_dir, full=True))
        if not errors and release == "all":
            errors.extend(validate_hcp_indicator_artifacts(data_dir, full=True))
        if not errors and release in {"v11", "all"}:
            errors.extend(validate_v11_quality_baseline_artifacts(data_dir, full=True))
        if not errors and release in {"v13", "v14", "v14.1", "all"}:
            errors.extend(validate_acquisition_inventory(data_dir, full=True))
        if not errors and release in {"v13", "v14", "v14.1", "all"}:
            errors.extend(validate_v13_source_profile(data_dir, full=True))
        if not errors and release in {"v13", "v14", "v14.1", "all"}:
            errors.extend(validate_v13_identity_registry(data_dir, full=True))
        if not errors and release in {"v13", "v14", "v14.1", "all"}:
            errors.extend(validate_v13_open_distribution(data_dir, full=True))
        if not errors and release in {"v14", "v14.1", "all"}:
            errors.extend(validate_v14_open_distribution(data_dir, full=True))
        if not errors and release in {"v14.1", "all"}:
            errors.extend(validate_v14_1_open_distribution(data_dir, full=True))
    return errors


def report(mode: str, data_dir: str | Path | None = None, release: str = "v14", baseline: str | None = None) -> int:
    errors = run(mode, data_dir, release, baseline)
    if errors:
        print("VALIDATION_FAILED")
        for error in errors:
            print(f"- {error}")
        return 1
    document_count = (
        len(EXPECTED_DOCUMENTS) * len(DOCUMENTATION_DIRS)
        if release == "all"
        else 0
        if release in {"v14", "v14.1"}
        else len(EXPECTED_DOCUMENTS)
    )
    print(f"VALIDATION_OK mode={mode} release={release} documents={document_count}")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Valide le dépôt et les releases locales V9 à V14.1.")
    parser.add_argument("--mode", choices=("ci", "full"), default="ci")
    parser.add_argument("--data-dir")
    parser.add_argument("--release", choices=("v9", "v10", "v11", "v12", "v13", "v14", "v14.1", "all"), default="v14.1")
    parser.add_argument("--baseline", choices=("v9", "v10", "v11", "v12", "v13", "v14"))
    args = parser.parse_args(argv)
    return report(args.mode, args.data_dir, args.release, args.baseline)


if __name__ == "__main__":
    sys.exit(main())
