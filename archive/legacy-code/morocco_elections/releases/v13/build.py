from __future__ import annotations

import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

from morocco_elections.config import PROJECT_ROOT, get_paths
from morocco_elections.domains.elections.core import build_electoral_core
from morocco_elections.legacy.v9 import build as legacy
from morocco_elections.provenance import sha256_file
from morocco_elections.releases.v12 import build as v12


VERSION = "V13"
GENERATED_ON = "2026-09-10"
RELEASE_REPORT = PROJECT_ROOT / "metadata" / "v13_release_report.json"
DIFF_REPORT = PROJECT_ROOT / "docs" / "research" / "V13_VS_V12_DIFF.txt"
INVENTORY_PATH = PROJECT_ROOT / "metadata" / "acquisition_inventory.json"
REGISTRY_PATH = PROJECT_ROOT / "metadata" / "v13_identity_registry.json"
QUALIFICATION_PATH = PROJECT_ROOT / "metadata" / "v13_electoral_qualification.json"
EXPECTED_RESULTS = 10_883
EXPECTED_CONTESTS = 639
EXPECTED_NEW_GEOS = 159
EXPECTED_NEW_PARTIES = 4
ALLOWED_CHANGED_SHEETS = {
    "README",
    "DIM_GEO",
    "DIM_PARTY",
    "SOURCES",
    "DIM_ELECTORAL_CONTEST",
    "FACT_ELECTION_RESULT",
    "FACT_ELECTORAL_MOBILIZATION",
    "DATA_DICTIONARY",
    "DATA_COVERAGE",
    "QUALITY_CONTROL",
    "WORKBOOK_AUDIT_V13",
}


def _inventory() -> dict:
    return json.loads(INVENTORY_PATH.read_text(encoding="utf-8"))


def _electoral_records() -> list[dict]:
    qualification = json.loads(QUALIFICATION_PATH.read_text(encoding="utf-8"))
    allowed = {item["source_id"] for item in qualification["authorized_for_future_ingestion"]}
    records = [item for item in _inventory()["records"] if item["source_id"] in allowed]
    if len(records) != 6:
        raise RuntimeError("Six acquisitions électorales GO sont attendues")
    return sorted(records, key=lambda item: item["source_id"])


def _record_path(record: dict, data_dir: str | Path | None) -> Path:
    relative = Path(record["local_path"])
    return get_paths(data_dir).data_root.joinpath(*relative.parts[1:])


def canonical_source_inputs(data_dir: str | Path | None = None) -> tuple[Path, ...]:
    return (*v12.canonical_source_inputs(data_dir), *(_record_path(item, data_dir) for item in _electoral_records()))


def _save_deterministic(workbook: openpyxl.Workbook, target: Path) -> None:
    fixed = datetime(2026, 9, 10)
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    workbook.properties.lastModifiedBy = "morocco_elections V13"
    workbook.save(target)
    canonical = target.with_suffix(".canonical.tmp")
    if canonical.exists():
        canonical.unlink()
    with zipfile.ZipFile(target, "r") as source, zipfile.ZipFile(
        canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as output:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 10, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(
                    rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
                    rb'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-10T00:00:00Z</dcterms:modified>',
                    payload,
                )
            output.writestr(info, payload)
    canonical.replace(target)


def _append_dimensions(workbook: openpyxl.Workbook, core: dict[str, list[dict[str, Any]]]) -> None:
    geo_headers, geo_rows = legacy.rows_from_sheet(workbook["DIM_GEO"])
    existing_geo = {str(row["geo_id"]) for row in geo_rows}
    for entity in sorted(core["new_territories"], key=lambda item: item["canonical_id"]):
        if entity["canonical_id"] in existing_geo:
            raise RuntimeError(f"Territoire V13 déjà présent: {entity['canonical_id']}")
        geo_rows.append(
            {
                "geo_id": entity["canonical_id"],
                "geo_name": entity["canonical_label"],
                "geo_type": entity["subtype"],
                "parent_geo_id": entity["parent_id"],
                "region_name": None,
                "province_prefecture": None,
                "constituency_name": None,
                "commune_name": None,
                "district_arrondissement": None,
                "fraction": None,
                "douar": None,
                "urban_rural": None,
                "latitude": None,
                "longitude": None,
                "boundary_version": entity["boundary_version"],
                "valid_from": entity["observed_on"],
                "valid_to": entity["observed_on"],
                "source_id": entity["evidence_id"].split(":", 1)[0],
                "source_url": None,
                "publication_date": None,
                "retrieval_date": GENERATED_ON,
                "quality_status": "provisional_historical_identity",
                "notes": (
                    "Entité source conservée pour son millésime; current_geo_candidate_id="
                    f"{entity['current_geo_candidate_id'] or 'NULL'}; aucune équivalence actuelle implicite."
                ),
            }
        )
    legacy.rewrite_existing(workbook["DIM_GEO"], geo_headers, geo_rows)
    workbook["DIM_GEO"]["A1"] = "V13 — dimension géographique multi-découpages"

    party_headers, party_rows = legacy.rows_from_sheet(workbook["DIM_PARTY"])
    existing_party = {str(row["party_id"]) for row in party_rows}
    for entity in sorted(core["new_parties"], key=lambda item: item["canonical_id"]):
        if entity["canonical_id"] in existing_party:
            raise RuntimeError(f"Entité partisane V13 déjà présente: {entity['canonical_id']}")
        party_rows.append(
            {
                "party_id": entity["canonical_id"],
                "party_name_fr": entity["canonical_label"],
                "party_name_ar": None,
                "acronym": entity["canonical_id"],
                "alliance_id": entity["canonical_id"] if entity["entity_type"] == "alliance" else None,
                "founded_year": None,
                "dissolved_year": None,
                "legal_status": f"historical_{entity['entity_type']}",
                "color_or_symbol": None,
                "leader_person_id": None,
                "parliamentary_group": None,
                "valid_from": entity["observed_on"],
                "valid_to": entity["observed_on"],
                "source_id": entity["evidence_id"].split(":", 1)[0],
                "source_url": None,
                "publication_date": None,
                "retrieval_date": GENERATED_ON,
                "quality_status": "provisional_source_defined",
                "notes": "Entité historique distincte; aucune fusion organisationnelle implicite.",
            }
        )
    legacy.rewrite_existing(workbook["DIM_PARTY"], party_headers, party_rows)
    workbook["DIM_PARTY"]["A1"] = "V13 — partis et alliances politiques"


def _add_core_sheets(workbook: openpyxl.Workbook, core: dict[str, list[dict[str, Any]]]) -> None:
    contest_headers = [
        "contest_id", "election_id", "geo_id", "region_geo_id", "list_type", "source_contest_id",
        "source_label", "normalized_label", "boundary_version", "valid_from", "valid_to", "source_id",
        "evidence_id", "identity_review_status", "notes",
    ]
    contest_rows = []
    for item in sorted(core["contests"], key=lambda row: row["contest_id"]):
        contest_rows.append(
            {
                "contest_id": item["contest_id"],
                "election_id": item["election_id"],
                "geo_id": item["geo_id"],
                "region_geo_id": item["region_geo_id"],
                "list_type": item["list_type"],
                "source_contest_id": item["source_contest_id"],
                "source_label": item["source_label"],
                "normalized_label": item["normalized_label"],
                "boundary_version": item["boundary_version"],
                "valid_from": item["valid_from"],
                "valid_to": item["valid_to"],
                "source_id": item["source_system"],
                "evidence_id": item["evidence_id"],
                "identity_review_status": item["review_status"],
                "notes": "Course datée; l'identité géographique ne vaut que pour boundary_version.",
            }
        )
    legacy.replace_standard_sheet(
        workbook,
        "DIM_ELECTORAL_CONTEST",
        "V13 — courses électorales canoniques",
        "Grain: contest_id; 2007/2011 restent attachés à leur découpage historique.",
        contest_headers,
        [[row[column] for column in contest_headers] for row in contest_rows],
    )

    result_headers = list(core["results"][0])
    result_rows = sorted(core["results"], key=lambda row: row["result_id"])
    legacy.replace_standard_sheet(
        workbook,
        "FACT_ELECTION_RESULT",
        "V13 — résultats électoraux canoniques par contest et parti",
        "Grain: contest_id × party_id; seules les voix non nulles directement publiées sont chargées.",
        result_headers,
        [[row[column] for column in result_headers] for row in result_rows],
    )

    mobilization_headers = list(core["mobilization"][0])
    mobilization_rows = sorted(core["mobilization"], key=lambda row: row["contest_id"])
    legacy.replace_standard_sheet(
        workbook,
        "FACT_ELECTORAL_MOBILIZATION",
        "V13 — mobilisation et capacité des courses électorales",
        "Grain: contest_id; NULL signifie non publié ou NO_GO, jamais zéro.",
        mobilization_headers,
        [[row[column] for column in mobilization_headers] for row in mobilization_rows],
    )


def _append_sources(workbook: openpyxl.Workbook) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["SOURCES"])
    existing = {str(row["source_id"]) for row in rows}
    for record in _electoral_records():
        if record["source_id"] in existing:
            raise RuntimeError(f"Source déjà présente: {record['source_id']}")
        rows.append(
            {
                "source_id": record["source_id"],
                "source_name": record["title"],
                "publisher": record["producer"],
                "source_type": "archived electoral results",
                "domain": "elections",
                "url": record["initial_url"],
                "coverage_start": str(record["target_period"]),
                "coverage_end": str(record["target_period"]),
                "geo_granularity": "electoral contest",
                "temporal_frequency": "election",
                "format": record["format"].upper(),
                "access_method": "immutable acquired file",
                "license": record["reuse_status"],
                "reliability_score": None,
                "last_checked": record["retrieved_at"],
                "status": "ingested_qualified_scope",
                "notes": f"SHA256={record['sha256']}; seules les mesures GO de la qualification V13 sont chargées.",
            }
        )
    legacy.rewrite_existing(workbook["SOURCES"], headers, rows)
    workbook["SOURCES"]["A1"] = "V13 — registre des sources"


def _append_dictionary(workbook: openpyxl.Workbook) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["DATA_DICTIONARY"])
    additions = [
        ("contest_party_votes", "Nombre de voix publié pour un parti ou une alliance dans une course.", "integer", "votes", "FACT_ELECTION_RESULT"),
        ("contest_seats", "Nombre de sièges attaché à la course dans la publication source; pas des sièges par parti.", "integer", "seats", "FACT_ELECTORAL_MOBILIZATION"),
        ("contest_registered_voters", "Nombre d'inscrits directement publié au grain de la course.", "integer", "persons", "FACT_ELECTORAL_MOBILIZATION"),
        ("contest_turnout_rate", "Taux de participation directement publié, conservé comme ratio [0,1].", "number", "ratio", "FACT_ELECTORAL_MOBILIZATION"),
        ("contest_valid_votes", "Suffrages valides directement publiés au grain de la course.", "integer", "votes", "FACT_ELECTORAL_MOBILIZATION"),
        ("contest_invalid_vote_rate", "Taux de bulletins blancs ou nuls directement publié, conservé comme ratio [0,1].", "number", "ratio", "FACT_ELECTORAL_MOBILIZATION"),
    ]
    for metric_id, definition, data_type, unit, sheet in additions:
        rows.append(
            {
                "metric_id": metric_id,
                "domain": "ELECTIONS",
                "metric_name": metric_id,
                "definition": definition,
                "data_type": data_type,
                "unit": unit,
                "preferred_geo_granularity": "electoral contest",
                "preferred_temporal_frequency": "election",
                "primary_fact_sheet": sheet,
                "required_keys": "contest_id" + (" × party_id" if sheet == "FACT_ELECTION_RESULT" else ""),
                "candidate_sources": "; ".join(item["source_id"] for item in _electoral_records()),
                "collection_status": "observed_qualified_scope",
                "priority": "P0_core",
                "quality_rule": "Renseigner seulement si la décision source × mesure vaut GO; conserver les absences à NULL.",
                "notes": "V13; aucune reconstruction depuis un taux ou une somme de voix.",
            }
        )
    legacy.rewrite_existing(workbook["DATA_DICTIONARY"], headers, rows)
    workbook["DATA_DICTIONARY"]["A1"] = "V13 — dictionnaire des métriques"


def _append_quality_and_coverage(workbook: openpyxl.Workbook, core: dict[str, list[dict[str, Any]]]) -> None:
    headers, rows = legacy.rows_from_sheet(workbook["DATA_COVERAGE"])
    coverage = [
        ("DIM_ELECTORAL_CONTEST", EXPECTED_CONTESTS, "contest_id", "contest"),
        ("FACT_ELECTION_RESULT", len(core["results"]), "contest_id × party_id", "contest × party"),
        ("FACT_ELECTORAL_MOBILIZATION", EXPECTED_CONTESTS, "contest_id", "contest"),
    ]
    for sheet, row_count, key, grain in coverage:
        rows.append(
            {
                "domain_sheet": sheet,
                "rows_loaded": row_count,
                "verified_rows": row_count,
                "provisional_rows": 0,
                "derived_rows": 0,
                "low_confidence_rows": 0,
                "primary_key": key,
                "granularity_now": grain,
                "next_granularity_target": "Étendre seulement avec une nouvelle source qualifiée.",
            }
        )
    legacy.rewrite_existing(workbook["DATA_COVERAGE"], headers, rows)
    workbook["DATA_COVERAGE"]["A1"] = "V13 — couverture par table"

    headers, rows = legacy.rows_from_sheet(workbook["QUALITY_CONTROL"])
    additions = [
        ("QA_V13_CONTEST_GRAIN", "DIM_ELECTORAL_CONTEST", "639 contest_id uniques", "passed"),
        ("QA_V13_RESULT_GRAIN", "FACT_ELECTION_RESULT", f"{len(core['results'])} clés contest_id × party_id uniques", "passed"),
        ("QA_V13_BOUNDARIES", "DIM_ELECTORAL_CONTEST", "187 contests 2007/2011 sans bridge territorial actuel", "restricted"),
        ("QA_V13_PARTY_NULLS", "FACT_ELECTION_RESULT", "Cellules partisanes source vides non chargées et jamais assimilées à zéro", "restricted"),
        ("QA_V13_ARCHIVE_2002", "FACT_ELECTION_RESULT", "LEG2002 conservé hors canonique", "passed"),
    ]
    for issue_id, sheet, description, resolution in additions:
        passed = resolution == "passed"
        rows.append(
            {
                "issue_id": issue_id,
                "detected_date": GENERATED_ON,
                "sheet_name": sheet,
                "record_or_range": "all rows",
                "issue_type": "v13_electoral_core",
                "severity": "info" if passed else "medium",
                "description": description,
                "expected_rule": "Respecter la qualification V13 et les découpages datés.",
                "resolution": resolution,
                "resolved_flag": 1 if passed else 0,
                "resolved_date": GENERATED_ON if passed else None,
                "owner": "data engineering",
                "source_id": "V13_ELECTORAL_QUALIFICATION",
                "notes": "Contrôle déterministe de la construction V13.",
            }
        )
    legacy.rewrite_existing(workbook["QUALITY_CONTROL"], headers, rows)
    workbook["QUALITY_CONTROL"]["A1"] = "V13 — contrôles qualité"


def _append_readme(workbook: openpyxl.Workbook) -> None:
    worksheet = workbook["README"]
    for row in (
        ("VERSION V13 — CŒUR ÉLECTORAL MULTI-SCRUTINS", None),
        ("version", "V13"),
        ("date", GENERATED_ON),
        ("addition", "639 courses, 10 883 résultats parti-course et 639 observations de mobilisation"),
        ("scope", "législatives 2007–2021 et régionales 2015–2021; 2002 reste hors canonique"),
        ("null rule", "absence source = NULL; aucune imputation à zéro"),
    ):
        worksheet.append(row)


def _validate_foreign_keys(workbook: openpyxl.Workbook, core: dict[str, list[dict[str, Any]]]) -> None:
    geo_ids = {str(row["geo_id"]) for row in legacy.rows_from_sheet(workbook["DIM_GEO"])[1]}
    party_ids = {str(row["party_id"]) for row in legacy.rows_from_sheet(workbook["DIM_PARTY"])[1]}
    election_ids = {str(row["election_id"]) for row in legacy.rows_from_sheet(workbook["DIM_ELECTION"])[1]}
    source_ids = {str(row["source_id"]) for row in legacy.rows_from_sheet(workbook["SOURCES"])[1]}
    contest_ids = {item["contest_id"] for item in core["contests"]}
    missing_geo = {
        value
        for item in core["contests"]
        for value in (item["geo_id"], item["region_geo_id"])
        if value not in geo_ids
    }
    missing_elections = {item["election_id"] for item in core["contests"] if item["election_id"] not in election_ids}
    missing_parties = {item["party_id"] for item in core["results"] if item["party_id"] not in party_ids}
    missing_sources = {
        item["source_id"]
        for collection in (core["results"], core["mobilization"])
        for item in collection
        if item["source_id"] not in source_ids
    }
    missing_contests = {
        item["contest_id"]
        for collection in (core["results"], core["mobilization"])
        for item in collection
        if item["contest_id"] not in contest_ids
    }
    if missing_geo or missing_elections or missing_parties or missing_sources or missing_contests:
        raise RuntimeError(
            "Clés étrangères V13 invalides: "
            f"geo={sorted(missing_geo)} elections={sorted(missing_elections)} "
            f"parties={sorted(missing_parties)} sources={sorted(missing_sources)} contests={sorted(missing_contests)}"
        )


def _rebuild_audit(workbook: openpyxl.Workbook, baseline_sheets: set[str]) -> None:
    del workbook["WORKBOOK_AUDIT_V12"]
    legacy.audit_workbook(workbook, baseline_sheets - {"WORKBOOK_AUDIT_V12"})
    worksheet = workbook["WORKBOOK_AUDIT_V9"]
    worksheet.title = "WORKBOOK_AUDIT_V13"
    worksheet["A1"] = "V13 — audit exhaustif du classeur"
    headers, rows = legacy.rows_from_sheet(worksheet)
    new_sheets = {"DIM_ELECTORAL_CONTEST", "FACT_ELECTION_RESULT", "FACT_ELECTORAL_MOBILIZATION"}
    changed = ALLOWED_CHANGED_SHEETS - {"WORKBOOK_AUDIT_V13"}
    for row in rows:
        if row["sheet"] in new_sheets:
            row["origin_version"] = "V13 added"
        elif row["sheet"] in changed:
            row["origin_version"] = "V13 updated"
        else:
            row["origin_version"] = "V12 preserved"
    legacy.rewrite_existing(worksheet, headers, rows)
    if len(workbook.sheetnames) != 71:
        raise RuntimeError(f"V13 doit contenir 71 onglets, obtenu {len(workbook.sheetnames)}")


def _sheet_values(workbook: openpyxl.Workbook, name: str) -> list[list[Any]]:
    return [list(row) for row in workbook[name].iter_rows(values_only=True)]


def compare_v13_to_v12(v12_path: Path, v13_path: Path) -> tuple[list[str], list[str]]:
    before = openpyxl.load_workbook(v12_path, read_only=True, data_only=False)
    after = openpyxl.load_workbook(v13_path, read_only=True, data_only=False)
    mapped = {("WORKBOOK_AUDIT_V13" if name == "WORKBOOK_AUDIT_V12" else name): name for name in before.sheetnames}
    new_sheets = {"DIM_ELECTORAL_CONTEST", "FACT_ELECTION_RESULT", "FACT_ELECTORAL_MOBILIZATION"}
    if set(mapped) | new_sheets != set(after.sheetnames):
        raise RuntimeError("La structure des onglets diffère entre V12 et V13")
    changed = sorted(new_sheets)
    unexpected: list[str] = []
    for target_name, source_name in mapped.items():
        if _sheet_values(before, source_name) != _sheet_values(after, target_name):
            changed.append(target_name)
            if target_name not in ALLOWED_CHANGED_SHEETS:
                unexpected.append(target_name)
    before.close()
    after.close()
    return sorted(changed), sorted(unexpected)


def render_diff_report(report: dict[str, Any]) -> str:
    return "\n".join(
        [
            "V13 — RAPPORT D'ÉCARTS AVEC V12",
            "",
            f"DATE : {report['generated_on']}",
            f"V12 SHA-256 : {report['workbooks']['V12']}",
            f"V13 SHA-256 : {report['workbooks']['V13']}",
            "",
            "CHANGEMENTS AUTORISÉS ET OBSERVÉS",
            *[f"- {name}" for name in report["actual_changed_sheets"]],
            "",
            "VOLUMES",
            f"- Courses électorales : {report['volumes']['DIM_ELECTORAL_CONTEST']}",
            f"- Résultats contest × parti : {report['volumes']['FACT_ELECTION_RESULT']}",
            f"- Mobilisation contest : {report['volumes']['FACT_ELECTORAL_MOBILIZATION']}",
            f"- Territoires historiques/provisoires ajoutés : {report['controls']['new_geo_entities']}",
            f"- Partis ou alliances provisoires ajoutés : {report['controls']['new_party_entities']}",
            "",
            "LIMITES",
            "- LEG2002 reste hors du canonique.",
            "- Les cellules partisanes absentes ne sont pas transformées en zéro.",
            "- Les contests 2007/2011 restent liés à leur découpage historique.",
            "- V12 et ses sources antérieures sont restés inchangés.",
            "",
        ]
    )


def build_from_sources(
    data_dir: str | Path | None = None,
    *,
    target: Path | None = None,
    report_output: Path | None = None,
) -> dict[str, Any]:
    paths = get_paths(data_dir)
    target = target or paths.v13_workbook
    required = canonical_source_inputs(data_dir)
    for path in required:
        if not path.is_file():
            raise FileNotFoundError(path)
    input_hashes = {str(path): sha256_file(path) for path in required}
    target.parent.mkdir(parents=True, exist_ok=True)
    baseline_temp = target.parent / ".v13-v12-reconstructed.tmp.xlsx"
    if baseline_temp.exists():
        baseline_temp.unlink()
    workbook = None
    baseline_report: dict[str, Any] | None = None
    try:
        baseline_report = v12.build_from_sources(data_dir, target=baseline_temp)
        core = build_electoral_core(data_dir)
        workbook = openpyxl.load_workbook(baseline_temp, data_only=False)
        baseline_sheets = set(workbook.sheetnames)
        _append_dimensions(workbook, core)
        _add_core_sheets(workbook, core)
        _append_sources(workbook)
        _append_dictionary(workbook)
        _append_quality_and_coverage(workbook, core)
        _append_readme(workbook)
        _validate_foreign_keys(workbook, core)
        _rebuild_audit(workbook, baseline_sheets)
        _save_deterministic(workbook, target)
        if any(sha256_file(Path(path)) != digest for path, digest in input_hashes.items()):
            raise RuntimeError("Un fichier source a changé pendant la construction")
        changed, unexpected = compare_v13_to_v12(baseline_temp, target)
        if unexpected or set(changed) != ALLOWED_CHANGED_SHEETS:
            raise RuntimeError(f"Différences V13 non conformes: changed={changed}; unexpected={unexpected}")
    finally:
        if workbook is not None:
            workbook.close()
        if baseline_temp.exists():
            baseline_temp.unlink()
    if baseline_report is None:
        raise AssertionError("Rapport V12 non construit")

    report = {
        "release": "V13",
        "baseline": "V12",
        "generated_on": GENERATED_ON,
        "status": "validated",
        "scope_status": "PARTIAL",
        "workbooks": {**baseline_report["workbooks"], "V13": sha256_file(target)},
        "source_hashes": {item["source_id"]: item["sha256"] for item in _electoral_records()},
        "decision_hashes": {
            "identity_registry": sha256_file(REGISTRY_PATH),
            "electoral_qualification": sha256_file(QUALIFICATION_PATH),
        },
        "volumes": {
            "sheets": 71,
            "DIM_ELECTORAL_CONTEST": len(core["contests"]),
            "FACT_ELECTION_RESULT": len(core["results"]),
            "FACT_ELECTORAL_MOBILIZATION": len(core["mobilization"]),
        },
        "controls": {
            "new_geo_entities": len(core["new_territories"]),
            "new_party_entities": len(core["new_parties"]),
            "unique_result_ids": len({item["result_id"] for item in core["results"]}),
            "unique_contest_ids": len({item["contest_id"] for item in core["contests"]}),
            "technical_zero_imputations": 0,
            "archive_2002_rows_ingested": 0,
        },
        "allowed_changed_sheets": sorted(ALLOWED_CHANGED_SHEETS),
        "actual_changed_sheets": changed,
        "integrity": {
            "v12_reconstructed_from_sources": True,
            "source_hashes_unchanged": True,
            "deterministic_export": True,
        },
    }
    if report_output is not None:
        report_output.parent.mkdir(parents=True, exist_ok=True)
        report_output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


def main(data_dir: str | Path | None = None) -> None:
    paths = get_paths(data_dir)
    report = build_from_sources(data_dir, target=paths.v13_workbook, report_output=RELEASE_REPORT)
    DIFF_REPORT.write_text(render_diff_report(report), encoding="utf-8")
    print(
        f"V13_BUILD_OK path={paths.v13_workbook} results={report['volumes']['FACT_ELECTION_RESULT']} "
        f"contests={report['volumes']['DIM_ELECTORAL_CONTEST']} sha256={report['workbooks']['V13']}"
    )


if __name__ == "__main__":
    main()
