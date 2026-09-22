from __future__ import annotations

import re
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

from morocco_elections.config import get_paths
from morocco_elections.domains.geography import GEO_DECISIONS, HCP_2014_SHA256, HCP_2014_URL, HCP_2024_SHA256, HCP_2024_URL
from morocco_elections.domains.governance import OFFICIAL_BULLETIN_SHA256, OFFICIAL_BULLETIN_URL, SEAT_EXCEPTIONS
from morocco_elections.domains.identity import audit_local_identities, local_person_id, normalize_person_name, person_name_match_key
from morocco_elections.legacy.v9 import build as v9

VERSION = "V10"
TODAY = "2026-09-07"


def _save_deterministic(wb, target: Path) -> None:
    """Save XLSX with fixed document and ZIP timestamps for reproducible hashes."""
    fixed = datetime(2026, 9, 7)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.lastModifiedBy = "morocco_elections V10"
    wb.save(target)
    canonical = target.with_suffix(".canonical.tmp")
    if canonical.exists():
        canonical.unlink()
    with zipfile.ZipFile(target, "r") as source, zipfile.ZipFile(canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as output:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 7, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(rb"<dcterms:modified[^>]*>.*?</dcterms:modified>", rb'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-07T00:00:00Z</dcterms:modified>', payload)
            output.writestr(info, payload)
    canonical.replace(target)


def _resolved_crosswalk(rows: list[list], headers: list[str]) -> list[list]:
    extra = ["review_status", "resolution_id", "evidence_url", "evidence_locator", "evidence_digest"]
    headers.extend(extra)
    output = []
    for row in rows:
        rec = dict(zip(headers[:-len(extra)], row))
        commune_id = int(str(rec["source_geo_id"]).rsplit("_", 1)[-1])
        if commune_id in GEO_DECISIONS:
            expected_id, expected_name, resolution_id = GEO_DECISIONS[commune_id]
            if rec["canonical_geo_id"] != expected_id:
                raise RuntimeError(f"Manual geography decision conflicts for {commune_id}")
            rec.update({"canonical_name": expected_name, "match_method": "manual_code+parent_validation", "confidence_score": 1.0, "manual_review_flag": 0,
                        "review_status": "validated_manual", "resolution_id": resolution_id,
                        "evidence_url": HCP_2014_URL if commune_id in {1042, 1453, 1474} else HCP_2024_URL,
                        "evidence_locator": f"official geographic code {expected_id.replace('MA-', '').replace('-', '.')}",
                        "notes": f"{rec['notes']} Historical source spelling retained as alias; official code and parent validated."})
            rec["evidence_digest"] = HCP_2014_SHA256 if commune_id in {1042, 1453, 1474} else HCP_2024_SHA256
        else:
            rec.update({"review_status": "automatic", "resolution_id": None, "evidence_url": None, "evidence_locator": None, "evidence_digest": None})
        output.append([rec.get(h) for h in headers])
    return output


def _replace_local_identities(wb, council: pd.DataFrame, mapping: dict[int, str]) -> tuple[list[dict], dict[int, tuple[str, str]], int]:
    source_id = v9.SOURCE_IDS["COUNCIL2021"]
    raw_records = council.to_dict("records")
    audit_rows, id_map = audit_local_identities(raw_records, source_id)
    if len(audit_rows) != 44 or len(id_map) != 32512:
        raise RuntimeError(f"Unexpected identity profile: audit={len(audit_rows)} persons={len(id_map)}")

    headers = ["local_mandate_id", "person_id", "person_name_ar", "person_name_normalized", "person_name_match_key", "identity_scope", "identity_status",
               "source_idcommune", "source_idcirconscription", "geo_id_source", "canonical_geo_id", "commune", "province", "party_id", "head_of_list", "role",
               "role_normalized", "executive_flag", "role_uncertain_flag", "head_of_list_uncertain_flag", "election_id", "source_id", "quality_status", "notes"]
    rows = []
    identity_counts = Counter()
    for idx, rec in enumerate(raw_records, 1):
        party = str(rec["parti"]).upper()
        pid = local_person_id(source_id, rec["idCommune"], rec["idCirconscription"], party, rec["prenomNom"])
        identity_counts[pid] += 1
        rows.append([f"LM2021_{idx:05d}", pid, rec["prenomNom"], normalize_person_name(rec["prenomNom"]), person_name_match_key(rec["prenomNom"]),
                     "source+commune+constituency+party+name", "repeated_source_identity" if identity_counts[pid] > 1 else "conservative",
                     int(rec["idCommune"]), int(rec["idCirconscription"]), f"TAFRA_COMM_{int(rec['idCommune'])}", mapping[int(rec["idCommune"])], rec["commune"], rec["prefProv"], party,
                     int(rec["teteDeListe"]), rec["role"], v9.role_normalized(rec["role"]), 0 if v9.role_normalized(rec["role"]) == "councillor" else 1,
                     int(rec["flagRole"]), int(rec["flagTeteDeListe"]), "COMM2021", source_id,
                     "low_confidence" if int(rec["flagRole"]) or int(rec["flagTeteDeListe"]) else "provisional",
                     "RAW name retained; normalized fields are derived; no cross-constituency or cross-party merge."])
    v9.replace_standard_sheet(wb, "LOCAL_MANDATES", "V10 — mandats locaux exhaustifs 2021", "32 513 observations conservées; identité Unicode déterministe et prudente.", headers, rows)

    dim_headers, dim_rows = v9.rows_from_sheet(wb["DIM_PERSON"])
    dim_rows = [r for r in dim_rows if not str(r.get("person_id") or "").startswith("TAFRA_LOCAL_2021_")]
    first_by_id = {}
    for rec in raw_records:
        party = str(rec["parti"]).upper()
        pid = local_person_id(source_id, rec["idCommune"], rec["idCirconscription"], party, rec["prenomNom"])
        first_by_id.setdefault(pid, {"person_id": pid, "full_name": rec["prenomNom"], "gender": None, "party_id": party, "role": "Élu communal",
            "incumbent_flag": "unknown", "first_elected_year": 2021, "valid_from": "2021-09-08", "source_id": source_id,
            "source_url": v9.URLS["COUNCIL2021"], "retrieval_date": TODAY, "quality_status": "provisional",
            "notes": "Deterministic V10 identity scoped to source+commune+constituency+party+Unicode name; not a national identity."})
    dim_rows.extend(first_by_id.values())
    v9.rewrite_existing(wb["DIM_PERSON"], dim_headers, dim_rows)

    audit_headers = ["audit_id", "audit_type", "source_idcommune", "normalized_name_key", "source_rows", "source_row_count", "distinct_raw_name_count", "constituency_ids", "party_ids", "assigned_person_count", "person_ids", "resolution", "quality_status", "notes"]
    v9.replace_standard_sheet(wb, "IDENTITY_AUDIT_V10", "V10 — audit des identités locales", "Groupes de noms répétés dans une commune; aucune fusion spéculative.", audit_headers, [[r.get(h) for h in audit_headers] for r in audit_rows])

    presidents = {}
    confident = council[(council["role"].map(v9.role_normalized) == "president") & (council["flagRole"] == 0)]
    for rec in confident.to_dict("records"):
        party = str(rec["parti"]).upper()
        presidents[int(rec["idCommune"])] = (local_person_id(source_id, rec["idCommune"], rec["idCirconscription"], party, rec["prenomNom"]), party)
    return audit_rows, presidents, len(first_by_id)


def _patch_control(wb, elections: dict[int, pd.DataFrame], council: pd.DataFrame, mapping: dict[int, str], presidents: dict[int, tuple[str, str]]) -> None:
    headers, records = v9.rows_from_sheet(wb["LOCAL_COUNCIL_CONTROL"])
    headers = headers[:5] + ["party_seat_share_legal"] + headers[5:]
    legal = elections[2021].set_index("idCommune")["nSieges"].astype(int).to_dict()
    reverse = {geo: source for source, geo in mapping.items()}
    patched = []
    for rec in records:
        source = reverse[rec["geo_id"]]
        pid, party = presidents.get(source, (None, None))
        rec["president_person_id"], rec["president_party_id"] = pid, party
        rec["party_seat_share_legal"] = rec["party_seats"] / legal[source] * 100
        rec["notes"] = "party_seat_share uses observed occupied seats; party_seat_share_legal uses published legal seats."
        patched.append([rec.get(h) for h in headers])
    v9.replace_standard_sheet(wb, "LOCAL_COUNCIL_CONTROL", "V10 — contrôle et composition des conseils communaux 2021", "Grain commune × parti; dénominateurs occupé et légal séparés.", headers, patched)

    observed = council.groupby("idCommune").size().to_dict()
    status_headers = ["source_idcommune", "geo_id", "commune", "legal_seat_count", "observed_elected_count", "documented_vacancy_count", "reconciled_seat_count", "observed_difference", "reconciled_difference", "snapshot_date", "resolution_id", "evidence_url", "quality_status", "notes"]
    names = elections[2021].set_index("idCommune")["commune"].to_dict()
    status_rows = []
    for source in sorted(legal):
        vacancy = int(SEAT_EXCEPTIONS.get(source, {}).get("vacant_seats", 0))
        resolution = SEAT_EXCEPTIONS.get(source, {}).get("resolution_id")
        actual = int(observed.get(source, 0))
        reconciled = actual + vacancy
        status_rows.append([source, mapping[source], names[source], legal[source], actual, vacancy, reconciled, actual - legal[source], reconciled - legal[source], "2021-09-08", resolution,
                            OFFICIAL_BULLETIN_URL if resolution else None, "verified" if reconciled == legal[source] else "blocked",
                            "Vacancy documented by arrêté 3267.21, BO 7037." if resolution else "Observed elected rows reconcile with legal seats."])
    v9.replace_standard_sheet(wb, "COUNCIL_SEAT_STATUS_V10", "V10 — statut des sièges communaux", "Une ligne par commune au snapshot du 8 septembre 2021.", status_headers, status_rows)


def _manual_resolutions(wb) -> None:
    headers = ["resolution_id", "resolution_type", "source_key", "canonical_key", "decision", "effective_date", "review_method", "review_status", "evidence_publisher", "evidence_url", "evidence_locator", "evidence_digest", "notes"]
    rows = []
    for source, (canonical, name, resolution) in GEO_DECISIONS.items():
        url = HCP_2014_URL if source in {1042, 1453, 1474} else HCP_2024_URL
        locator = f"official geographic code {canonical.replace('MA-', '').replace('-', '.')}"
        digest = HCP_2014_SHA256 if source in {1042, 1453, 1474} else HCP_2024_SHA256
        rows.append([resolution, "geographic_crosswalk", f"TAFRA_COMM_{source}", canonical, f"Validated canonical entity {name}; source spelling retained as alias.", "2015-02-20", "manual_documentary", "validated", "Haut-Commissariat au Plan", url, locator, digest, "Code and administrative parent are decisive."])
    for source, item in SEAT_EXCEPTIONS.items():
        locator = f"arrêté 3267.21; Taroudannt; {item['commune']}; circonscription {item['constituency']}"
        rows.append([item["resolution_id"], "council_seat_vacancy", str(source), str(item["constituency"]), "One legal seat vacant at the 2021-09-08 snapshot; no person imputed.", "2021-09-08", "manual_documentary", "validated", "Secrétariat général du gouvernement", OFFICIAL_BULLETIN_URL, locator, OFFICIAL_BULLETIN_SHA256, "A later partial election is outside V10 scope."])
    v9.replace_standard_sheet(wb, "MANUAL_RESOLUTIONS_V10", "V10 — décisions manuelles documentées", "Cinq crosswalks géographiques et deux vacances de siège.", headers, rows)


def _patch_quality_and_coverage(wb, identity_count: int) -> None:
    headers, records = v9.rows_from_sheet(wb["QUALITY_CONTROL"])
    for rec in records:
        issue = str(rec.get("issue_id") or "")
        if issue.startswith("QA_V9_"):
            rec["issue_id"] = issue.replace("QA_V9_", "QA_V10_", 1)
        if rec["issue_id"] == "QA_V10_GEO_MANUAL_REVIEW":
            rec.update(description="Five source variants (ten election rows) manually validated by official code and administrative parent.", resolution="validated with evidence in MANUAL_RESOLUTIONS_V10", resolved_flag=1, resolved_date=TODAY, severity="info", notes="Historical spellings retained as aliases.")
        if rec["issue_id"] == "QA_V10_COUNCIL_SEAT_SUM":
            rec.update(description="Observed elected-row difference remains -2; one documented vacancy in each of communes 770 and 817 reconciles the legal total.", expected_rule="observed + documented vacancies = legal seats", resolution="explained by arrêté 3267.21; no person imputed", resolved_flag=1, resolved_date=TODAY, severity="info", notes="Observed difference=-2; reconciled difference=0.")
        if rec["issue_id"] == "QA_V10_COUNCIL_NATURAL_DUP":
            rec.update(resolution="retained as two RAW mandate observations linked to one conservative person identity", resolved_flag=1, resolved_date=TODAY, notes="Bni Makada source rows differ in head-of-list value; no RAW deletion.")
    records.extend([
        {"issue_id": "QA_V10_IDENTITY_UNICODE", "detected_date": TODAY, "sheet_name": "LOCAL_MANDATES", "record_or_range": "all rows", "issue_type": "identity_integrity", "severity": "info", "description": f"Unicode-safe deterministic person IDs: {identity_count}/32513 distinct; no empty key and no technical collision.", "expected_rule": "stable ID for source+commune+constituency+party+normalized name", "resolution": "passed", "resolved_flag": 1, "resolved_date": TODAY, "owner": "data engineering", "source_id": v9.SOURCE_IDS["COUNCIL2021"], "notes": "Derived normalization never overwrites RAW names."},
        {"issue_id": "QA_V10_IDENTITY_AMBIGUITY", "detected_date": TODAY, "sheet_name": "IDENTITY_AUDIT_V10", "record_or_range": "44 groups", "issue_type": "identity_ambiguity", "severity": "medium", "description": "44 repeated name-within-commune groups are retained and explicitly scoped; no speculative merge.", "expected_rule": "name equality alone cannot join people", "resolution": "documented conservative separation", "resolved_flag": 1, "resolved_date": TODAY, "owner": "data engineering", "source_id": v9.SOURCE_IDS["COUNCIL2021"], "notes": "Future evidence may create cross-observation identity links without changing RAW."},
    ])
    v9.rewrite_existing(wb["QUALITY_CONTROL"], headers, records)

    headers, records = v9.rows_from_sheet(wb["DATA_COVERAGE"])
    targets = {"CROSSWALK_GEO": "incremental boundary-version review", "LOCAL_MANDATES": "external identity evidence and gender", "LOCAL_COUNCIL_CONTROL": "resolve 135 presidents"}
    for rec in records:
        if rec.get("domain_sheet") in targets:
            rec["next_granularity_target"] = targets[rec["domain_sheet"]]
    for name, pk, grain, target in [
        ("IDENTITY_AUDIT_V10", "audit_id", "ambiguous name group within commune", "external identity evidence"),
        ("MANUAL_RESOLUTIONS_V10", "resolution_id", "manual evidence decision", "append-only reviews"),
        ("COUNCIL_SEAT_STATUS_V10", "source_idcommune", "commune × snapshot", "later dated council states"),
    ]:
        _, rows = v9.rows_from_sheet(wb[name])
        q = Counter(str(r.get("quality_status")) for r in rows if r.get("quality_status"))
        records.append({"domain_sheet": name, "rows_loaded": len(rows), "verified_rows": q.get("verified", 0), "provisional_rows": q.get("provisional", 0), "derived_rows": q.get("derived", 0), "low_confidence_rows": q.get("low_confidence", 0), "primary_key": pk, "granularity_now": grain, "next_granularity_target": target})
    v9.replace_standard_sheet(wb, "DATA_COVERAGE", "V10 — couverture quantitative", "Couverture recalculée après correction d'identité et intégrité.", headers, [[r.get(h) for h in headers] for r in records])


def _retitle_and_audit(wb, v8_sheetnames: set[str]) -> None:
    for ws in wb.worksheets:
        if isinstance(ws["A1"].value, str) and "V9" in ws["A1"].value:
            ws["A1"] = ws["A1"].value.replace("V9", "V10")
    readme = wb["README"]
    readme["A1"] = "Morocco Electoral Quantitative Data Warehouse — V10"
    start = readme.max_row + 2
    additions = [
        ["VERSION V10 — IDENTITÉS ET INTÉGRITÉ", None], ["version", "V10"], ["date", TODAY],
        ["identity correction", "32,513 mandate observations; 32,512 deterministic local person IDs; 44 ambiguity groups audited"],
        ["manual resolutions", "five geographic variants validated; two legal-seat gaps explained as dated vacancies"],
        ["known blockers", "registered voters absent; 2015 party seats absent; 135 presidents unresolved; parliamentary party history incomplete"],
    ]
    for row in additions:
        readme.append(row)
    readme.cell(start, 1).font = v9.Font(bold=True, size=12)

    v9.audit_workbook(wb, v8_sheetnames)
    ws = wb["WORKBOOK_AUDIT_V9"]
    ws.title = "WORKBOOK_AUDIT_V10"
    ws["A1"] = "V10 — audit exhaustif du classeur"
    headers, records = v9.rows_from_sheet(ws)
    for rec in records:
        rec["origin_version"] = "V8 preserved" if rec["sheet"] in v8_sheetnames else "V10 rebuilt/new"
        if rec["sheet"] in {"IDENTITY_AUDIT_V10", "MANUAL_RESOLUTIONS_V10", "COUNCIL_SEAT_STATUS_V10"}:
            rec["classification"] = "QA"
    v9.rewrite_existing(ws, headers, records)
    if len(wb.sheetnames) != 67:
        raise RuntimeError(f"V10 must contain 67 sheets, got {len(wb.sheetnames)}")


def assemble_workbook(data_dir: str | Path | None = None) -> tuple[openpyxl.Workbook, dict[str, int]]:
    """Construire V10 en mémoire directement depuis V8 et les RAW immuables."""
    paths = get_paths(data_dir)
    v9.configure_paths(data_dir)
    for path in [paths.v8_workbook, *v9.FILES.values()]:
        if not path.exists():
            raise FileNotFoundError(path)
    wb = openpyxl.load_workbook(paths.v8_workbook, data_only=False)
    v8_sheetnames = set(wb.sheetnames)
    elections = {2015: v9.dataframe(v9.FILES["COMM2015"], 0), 2021: v9.dataframe(v9.FILES["COMM2021"], 0)}
    council = v9.dataframe(v9.FILES["COUNCIL2021"], 0)
    mps = v9.dataframe(v9.FILES["MPS"], 0)
    if elections[2015].shape != (1538, 45) or elections[2021].shape != (1538, 56) or council.shape != (32513, 18) or mps.shape != (1654, 24):
        raise RuntimeError("Unexpected immutable source shape")

    v9.copy_raw_excel_sheet(wb, "RAW_COMM2015_FULL", v9.FILES["COMM2015"], "données", f"Copie immuable; SHA256={v9.sha256(v9.FILES['COMM2015'])}")
    v9.copy_raw_excel_sheet(wb, "RAW_COMM2021_FULL", v9.FILES["COMM2021"], "données", f"Copie immuable; SHA256={v9.sha256(v9.FILES['COMM2021'])}")
    v9.copy_raw_excel_sheet(wb, "RAW_COUNCIL2021_FULL", v9.FILES["COUNCIL2021"], "données", f"Copie immuable; SHA256={v9.sha256(v9.FILES['COUNCIL2021'])}")
    v9.copy_raw_excel_sheet(wb, "RAW_MPS_FULL", v9.FILES["MPS"], "donnees", f"Copie immuable; SHA256={v9.sha256(v9.FILES['MPS'])}")
    v9.copy_raw_hcp(wb)
    regions, provinces, entities = v9.parse_hcp_geography()
    cw15, map15 = v9.build_crosswalk(elections[2015], entities, "COMM2015")
    cw21, map21 = v9.build_crosswalk(elections[2021], entities, "COMM2021")
    if map15 != map21:
        raise RuntimeError("2015/2021 geography mappings differ")
    cross_headers = ["source_system", "source_geo_id", "source_geo_name", "source_region", "source_province", "canonical_geo_id", "canonical_name", "canonical_region_id", "canonical_province_id", "valid_from", "valid_to", "match_method", "confidence_score", "manual_review_flag", "temporal_boundary_status", "source_id", "notes"]
    cross_rows = _resolved_crosswalk(cw15 + cw21, cross_headers)
    v9.replace_standard_sheet(wb, "CROSSWALK_GEO", "V10 — crosswalk géographique exhaustif TAFRA ↔ HCP", "3 076 raccords; cinq variantes validées manuellement par code et parent.", cross_headers, cross_rows)

    party_names = v9.source_party_names(v9.FILES["COMM2015"])
    party_names.update(v9.source_party_names(v9.FILES["COMM2021"]))
    v9.update_dimensions_and_population(wb, regions, provinces, entities, party_names, council, mps)
    v9.build_party_crosswalk(wb, party_names, [("TAFRA COMM2015", 2015, list(elections[2015].columns[elections[2015].columns.get_loc("AGD"):])), ("TAFRA COMM2021", 2021, list(elections[2021].columns[elections[2021].columns.get_loc("AFG"):])), ("TAFRA COUNCIL2021", 2021, list(council["parti"].dropna().unique())), ("TAFRA MPS", "2007-2026", list(mps["parti"].dropna().unique()))])
    normalized, result_rows, mob_rows = v9.build_results_and_mobilisation(wb, elections, {2015: map15, 2021: map21})
    _, _, council_counts = v9.build_local_and_mps(wb, council, mps, map21)
    identity_audit, presidents, identity_count = _replace_local_identities(wb, council, map21)
    _patch_control(wb, elections, council, map21, presidents)
    trans_rows, commune_rows, analytical_rows, commune_election_rows = v9.build_transitions_and_panels(wb, normalized, entities, council_counts, presidents)
    _manual_resolutions(wb)
    v9.append_sources(wb)
    v9.update_tracking(wb)
    v9.add_quality(wb, elections, {2015: map15, 2021: map21}, normalized, council, mps, [])
    v9.append_readme_version(wb, {"results": len(result_rows), "mob": len(mob_rows), "transitions": len(trans_rows), "communes": len(commune_rows), "analytical": len(analytical_rows), "commune_election": len(commune_election_rows)})
    v9.build_data_coverage(wb)
    _patch_quality_and_coverage(wb, identity_count)
    _retitle_and_audit(wb, v8_sheetnames)
    return wb, {
        "sheets": len(wb.sheetnames),
        "local_mandates": len(council),
        "local_persons": identity_count,
        "identity_audit": len(identity_audit),
    }


def main(data_dir: str | Path | None = None) -> None:
    paths = get_paths(data_dir)
    workbook, stats = assemble_workbook(data_dir)
    paths.v10_workbook.parent.mkdir(parents=True, exist_ok=True)
    _save_deterministic(workbook, paths.v10_workbook)
    workbook.close()
    print(f"Saved {paths.v10_workbook}")
    print(
        f"Sheets={stats['sheets']} LocalMandates={stats['local_mandates']} "
        f"LocalPersons={stats['local_persons']} IdentityAudit={stats['identity_audit']}"
    )


if __name__ == "__main__":
    main()
