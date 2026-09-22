from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from morocco_elections.domains.parliament import longitudinal


def _occurrence(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {
        "question_type": "written",
        "source_question_number": "42",
        "deposit_date": "2022-04-10",
        "period_raw": "Période synthétique",
        "deputy_name_ar_raw": "Personne Synthétique",
        "parliamentary_group_ar_raw": "Groupe Synthétique",
        "ministry_ar_raw": "Institution Synthétique",
        "subject_ar": "Sujet Synthétique",
        "question_text_ar": "Texte synthétique",
        "response_date": "2022-05-01",
        "source_family_id": "SYNTHETIC_FAMILY",
        "source_id": "SRC_PARLIAMENT_A",
        "source_sha256": "a" * 64,
        "source_row": 2,
    }
    row.update(overrides)
    return row


def _source(source_id: str, digest: str) -> dict[str, object]:
    return {
        "source_id": source_id,
        "title": "Source synthétique",
        "producer": "Producteur synthétique",
        "initial_url": "https://example.test/source",
        "reuse_status": "Synthétique",
        "retrieved_at": "2026-09-12",
        "sha256": digest,
        "byte_size": 10,
    }


def test_model_deduplicates_questions_but_preserves_provenance() -> None:
    first = _occurrence()
    second = _occurrence(
        source_id="SRC_PARLIAMENT_B", source_sha256="b" * 64, source_row=8
    )
    mandates = [
        {
            "legislature": "2021-2026",
            "person_id": "PERSON_SYNTHETIC",
            "full_name_ar": "Personne Synthétique",
            "party_id": "PARTY_SYNTHETIC",
        }
    ]

    model = longitudinal.build_longitudinal_model(
        [first, second],
        mandates,
        [_source("FAMILY_A", "a" * 64), _source("FAMILY_B", "b" * 64)],
    )

    assert len(model["fact_parliamentary_question"]) == 1
    assert len(model["bridge_question_source"]) == 2
    assert len(model["fact_parliamentary_response"]) == 1
    author = model["bridge_question_author"][0]
    assert author["person_id"] == "PERSON_SYNTHETIC"
    assert author["party_id"] == "PARTY_SYNTHETIC"
    assert model["analytical_parliamentary_trajectory"] == [
        {
            "trajectory_id": model["analytical_parliamentary_trajectory"][0]["trajectory_id"],
            "person_id": "PERSON_SYNTHETIC",
            "party_id": "PARTY_SYNTHETIC",
            "legislature": "2021-2026",
            "period_id": model["analytical_parliamentary_trajectory"][0]["period_id"],
            "question_type": "written",
            "question_count": 1,
            "answered_question_count": 1,
            "response_rate_pct": 100.0,
            "first_deposit_date": "2022-04-10",
            "last_deposit_date": "2022-04-10",
            "derivation_status": "DERIVED_FROM_PUBLISHED_QUESTIONS",
        }
    ]


def test_model_rejects_conflicting_reuse_of_natural_key() -> None:
    with pytest.raises(RuntimeError, match="contradictoire"):
        longitudinal.build_longitudinal_model(
            [_occurrence(), _occurrence(subject_ar="Autre sujet", source_row=3)],
            [],
            [_source("FAMILY_A", "a" * 64)],
        )


def test_model_keeps_unresolved_author_without_trajectory() -> None:
    model = longitudinal.build_longitudinal_model(
        [_occurrence()],
        [],
        [_source("FAMILY_A", "a" * 64)],
    )

    author = model["dim_parliamentary_author"][0]
    assert author["person_id"] is None
    assert author["identity_match_method"] == "unresolved_no_exact_match"
    assert model["analytical_parliamentary_trajectory"] == []


def test_parser_recovers_missing_header_without_dropping_first_row(tmp_path: Path) -> None:
    path = tmp_path / "candidate.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(
        [
            7, "oral", "unused", "unused", "unused", "unused", "unused", "unused", "unused"
        ]
    )
    sheet.cell(1, 3).value = __import__("datetime").datetime(2022, 1, 2)
    sheet.cell(1, 4).value = "Période"
    sheet.cell(1, 5).value = "Personne"
    sheet.cell(1, 6).value = "Groupe"
    sheet.cell(1, 7).value = "Institution"
    sheet.cell(1, 8).value = "Sujet"
    sheet.cell(1, 9).value = "Question"
    workbook.save(path)
    record = _source("FAMILY_A", "a" * 64)

    rows = longitudinal.parse_resource(record, path)

    assert len(rows) == 1
    assert rows[0]["source_question_number"] == "7"
    assert rows[0]["source_row"] == 1
